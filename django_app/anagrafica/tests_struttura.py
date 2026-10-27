"""Struttura del Mansionario: robustezza, export/import Excel, editor unico, ruoli in blocco.

Dati sintetici.
"""
from __future__ import annotations

from datetime import timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from assets.models import Asset
from core.legacy_models import AnagraficaDipendente

from .models import (
    AbilitazioneMacchina, AbilitazioneMacchinaStorico, AbilitazioneMacchinaVoce, AmbitoRuolo, CompetenzaSkm,
    DipendenteAnagraficaAziendale, DipendenteQualifica, DipendenteRuoloOperativo, Mansione, Reparto,
    RuoloOperativo, TipoQualifica, VoceMacchinaCatalogo,
)
from .services import mansionario_competenze as comp
from .services import mansionario_struttura as svc
from .services import ruoli_in_blocco as blocco
from .services import struttura_excel as xl
from .tests import _ensure_anagrafica_table

User = get_user_model()


def _xlsx(fogli: dict[str, list[list]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for nome, righe in fogli.items():
        ws = wb.create_sheet(nome)
        for r in righe:
            ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@override_settings(LEGACY_AUTH_ENABLED=False)
class _Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        _ensure_anagrafica_table()

    def setUp(self):
        self.admin = User.objects.create_superuser("str_admin", "str@example.invalid", "x")
        self.asset = Asset.objects.create(name="Fresa sintetica FR-09", reparto="FRESE T")
        self.fresa = CompetenzaSkm.objects.create(competenza_key="FR-09", display="Fresa FR-09",
                                                  tipo=CompetenzaSkm.TIPO_MACCHINA, asset=self.asset)
        self.processo = CompetenzaSkm.objects.create(competenza_key="P-SBA", display="Sbavatura",
                                                     tipo=CompetenzaSkm.TIPO_PROCESSO)
        self.ruolo = RuoloOperativo.objects.create(nome="Operatore fresa sintetico")
        self.capo = RuoloOperativo.objects.create(nome="Capo fresatura sintetico")
        self.rischio = Mansione.objects.create(nome="ADDETTO FRESE T")
        self.reparto = Reparto.objects.create(nome="SBAVATURA T")
        self.qualifica = TipoQualifica.objects.create(nome="Patentino sintetico")
        self.p1 = AnagraficaDipendente.objects.create(nome="UNO", cognome="STRUTTURA", aliasusername="str.uno.t")
        self.p2 = AnagraficaDipendente.objects.create(nome="DUE", cognome="STRUTTURA", aliasusername="str.due.t")
        self.client.force_login(self.admin)


# ═══════════════════════════════════════════════════════════════════════════
# Fase A: robustezza
# ═══════════════════════════════════════════════════════════════════════════

class RobustezzaTests(_Base):
    def test_competenza_tolta_si_riattiva_con_i_suoi_livelli(self):
        voce = comp.salva_competenza(self.processo, nome="Uso flessibile")
        ab = svc.assegna(self.processo, legacy_anagrafica_id=self.p1.pk, livello="L")
        AbilitazioneMacchinaVoce.objects.create(abilitazione=ab, voce=voce, livello="U")
        voce.attiva = False
        voce.save()
        di_nuovo = comp.salva_competenza(self.processo, nome="uso flessibile")
        self.assertEqual(di_nuovo.pk, voce.pk)
        self.assertTrue(di_nuovo.attiva)
        self.assertEqual(VoceMacchinaCatalogo.objects.filter(mansione=self.processo).count(), 1)

    def test_togli_due_volte_non_duplica_lo_storico(self):
        ab = svc.assegna(self.processo, legacy_anagrafica_id=self.p1.pk, livello="L")
        svc.togli(ab)
        svc.togli(ab)
        self.assertEqual(AbilitazioneMacchinaStorico.objects.filter(
            competenza=self.processo, note="Mansionario: tolto dalla lista").count(), 1)

    def test_assegna_due_volte_aggiorna(self):
        svc.assegna(self.processo, legacy_anagrafica_id=self.p1.pk, livello="L")
        svc.assegna(self.processo, legacy_anagrafica_id=self.p1.pk, livello="U")
        self.assertEqual(AbilitazioneMacchina.objects.get(competenza=self.processo).livello, "U")

    def test_reparto_non_numerico_non_manda_in_errore(self):
        r = self.client.post(reverse("anagrafica:mansionario_mansione", args=[self.processo.pk]),
                             {"reparto": "abc", "attiva": "on"})
        self.assertEqual(r.status_code, 302)
        self.processo.refresh_from_db()
        self.assertIsNone(self.processo.reparto_id)

    def test_requisiti_precaricati_danno_lo_stesso_esito(self):
        voce = comp.salva_competenza(self.processo, nome="Brasatura", qualifica_id=self.qualifica.pk)
        ab1 = svc.assegna(self.processo, legacy_anagrafica_id=self.p1.pk, livello="U")
        ab2 = svc.assegna(self.processo, legacy_anagrafica_id=self.p2.pk, livello="U")
        for ab in (ab1, ab2):
            AbilitazioneMacchinaVoce.objects.create(abilitazione=ab, voce=voce, livello="U")
        DipendenteQualifica.objects.create(legacy_anagrafica_id=self.p1.pk, tipo=self.qualifica,
                                           data_scadenza=timezone.localdate() + timedelta(days=30))
        DipendenteQualifica.objects.create(legacy_anagrafica_id=self.p2.pk, tipo=self.qualifica,
                                           data_scadenza=timezone.localdate() - timedelta(days=1))
        voci = list(comp.competenze(self.processo))
        req = comp.Requisiti([self.p1.pk, self.p2.pk], voci)
        for ab in (ab1, ab2):
            self.assertEqual(comp.valuta(ab, voci=voci).esito, comp.valuta(ab, voci=voci, requisiti=req).esito)
        self.assertEqual(comp.valuta(ab2, voci=voci, requisiti=req).esito, comp.NON_ABILITATO)
        self.assertEqual(comp.valuta(ab1, voci=voci, requisiti=req).esito, comp.ABILITATO)


# ═══════════════════════════════════════════════════════════════════════════
# Fase B: export / import Excel
# ═══════════════════════════════════════════════════════════════════════════

class ExcelTests(_Base):
    def test_export_e_reimport_non_cambiano_niente(self):
        self.processo.ruoli.set([self.ruolo])
        self.processo.mansioni_rischio.set([self.rischio])
        comp.salva_competenza(self.processo, nome="Lettura disegno", livello_minimo="L", bloccante=True)
        esito = xl.importa(xl.esporta_xlsx(), applica=True)
        self.assertEqual(esito.errori, 0, [r for r in esito.righe if r.azione == xl.ERRORE])
        self.assertEqual(esito.cambi, 0, [r for r in esito.righe if r.azione != xl.INVARIATO])

    def test_la_prova_non_scrive_e_l_applica_si(self):
        file = _xlsx({"Competenze": [
            ["Chiave mansione", "Mansione", "Competenza", "Livello minimo", "Bloccante", "Qualifica"],
            ["", "Sbavatura", "Uso mola", "U", "SI", "Patentino sintetico"],
            ["FR-09", "", "Set-up", "L", "", ""],
        ]})
        prova = xl.importa(file, applica=False)
        self.assertEqual(prova.cambi, 2)
        self.assertFalse(VoceMacchinaCatalogo.objects.exists())
        xl.importa(file, applica=True)
        mola = VoceMacchinaCatalogo.objects.get(mansione=self.processo, nome="Uso mola")
        self.assertTrue(mola.bloccante)
        self.assertEqual(mola.qualifica_id, self.qualifica.pk)
        self.assertEqual(VoceMacchinaCatalogo.objects.get(mansione=self.fresa).livello_minimo, "L")

    def test_cella_vuota_non_cambia_e_no_spegne(self):
        voce = comp.salva_competenza(self.processo, nome="Uso mola", livello_minimo="O", bloccante=True)
        xl.importa(_xlsx({"Competenze": [["Mansione", "Competenza", "Livello minimo", "Bloccante"],
                                         ["Sbavatura", "Uso mola", "", ""]]}), applica=True)
        voce.refresh_from_db()
        self.assertEqual((voce.livello_minimo, voce.bloccante, voce.attiva), ("O", True, True))
        xl.importa(_xlsx({"Competenze": [["Mansione", "Competenza", "Attiva"], ["Sbavatura", "Uso mola", "NO"]]}),
                   applica=True)
        voce.refresh_from_db()
        self.assertFalse(voce.attiva)

    def test_righe_sbagliate_si_scartano_le_altre_passano(self):
        esito = xl.importa(_xlsx({"Mansioni": [
            ["Mansione", "Ruoli", "Mansioni di rischio"],
            ["Sbavatura", "Operatore fresa sintetico", "RISCHIO CHE NON ESISTE"],
            ["Fresa FR-09", "Operatore fresa sintetico; Capo fresatura sintetico", "ADDETTO FRESE T"],
            ["Collaudo nuovo", "Operatore fresa sintetico", ""],
            ["Macchina nuova", "", ""],
        ]}), applica=True)
        per_riga = {r.numero: r for r in esito.righe}
        self.assertEqual(per_riga[2].azione, xl.ERRORE)
        self.assertIn("RISCHIO CHE NON ESISTE", per_riga[2].dettaglio)
        self.assertEqual(per_riga[3].azione, xl.MODIFICATO)
        self.assertEqual(set(self.fresa.ruoli.values_list("nome", flat=True)),
                         {"Operatore fresa sintetico", "Capo fresatura sintetico"})
        self.processo.refresh_from_db()
        self.assertFalse(self.processo.ruoli.exists())  # la riga sbagliata non ha scritto a metà
        nuova = CompetenzaSkm.objects.get(display="Collaudo nuovo")
        self.assertEqual(nuova.tipo, CompetenzaSkm.TIPO_PROCESSO)
        self.assertTrue(nuova.competenza_key.startswith("MANS-"))
        self.assertEqual(per_riga[5].azione, xl.NUOVO)

    def test_ruoli_creati_e_gerarchia_senza_giri(self):
        ambito = AmbitoRuolo.objects.create(nome="Produttivo T")
        esito = xl.importa(_xlsx({"Ruoli": [
            ["Ruolo", "Ambito", "Riporta a"],
            ["Addetto collaudo T", "Produttivo T", "Capo fresatura sintetico"],
            ["Capo fresatura sintetico", "", "Addetto collaudo T"],
        ]}), applica=True)
        nuovo = RuoloOperativo.objects.get(nome="Addetto collaudo T")
        self.assertEqual((nuovo.ambito_id, nuovo.riporta_a_id), (ambito.pk, self.capo.pk))
        self.capo.refresh_from_db()
        self.assertIsNone(self.capo.riporta_a_id)
        self.assertTrue(any("giro" in a for a in esito.avvisi))

    def test_file_non_excel(self):
        with self.assertRaises(ValueError):
            xl.importa(b"non sono un excel", applica=False)

    def test_view_prova_poi_applica(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        file = _xlsx({"Competenze": [["Mansione", "Competenza"], ["Sbavatura", "Pulizia bava"]]})
        url = reverse("anagrafica:mansionario_struttura_import")
        r = self.client.post(url, {"azione": "prova", "file": SimpleUploadedFile("s.xlsx", file)})
        self.assertContains(r, "Anteprima dell")
        self.assertFalse(VoceMacchinaCatalogo.objects.exists())
        r = self.client.post(url, {"azione": "applica", "sha": r.context["sha"]})
        self.assertContains(r, "Importazione registrata")
        self.assertTrue(VoceMacchinaCatalogo.objects.filter(nome="Pulizia bava").exists())
        r = self.client.post(url, {"azione": "applica", "sha": "x"})
        self.assertEqual(r.status_code, 302)

    def test_export_view(self):
        r = self.client.get(reverse("anagrafica:mansionario_struttura_export"))
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])


# ═══════════════════════════════════════════════════════════════════════════
# Fase C: editor unico
# ═══════════════════════════════════════════════════════════════════════════

class EditorTests(_Base):
    def _post(self, mansione, righe, **extra):
        dati = {"nome": mansione.nome, "attiva": "on", "n_righe": str(len(righe)), **extra}
        for i, r in enumerate(righe):
            for k, v in r.items():
                dati[f"r{i}_{k}"] = v
        return self.client.post(reverse("anagrafica:mansionario_struttura_mansione", args=[mansione.pk]), dati)

    def test_pagina_e_frammento(self):
        self.processo.ruoli.set([self.ruolo])
        r = self.client.get(reverse("anagrafica:mansionario_struttura"))
        self.assertContains(r, "Operatore fresa sintetico")
        self.assertContains(r, "Da collegare a un ruolo")  # la fresa non ha ruolo
        r = self.client.get(reverse("anagrafica:mansionario_struttura_mansione", args=[self.processo.pk]))
        self.assertContains(r, 'name="r0_nome"')

    def test_salva_configurazione_e_competenze_insieme(self):
        r = self._post(self.processo, [{"nome": "Uso mola", "livello_minimo": "L", "bloccante": "on"},
                                       {"nome": ""}], ruoli=[self.ruolo.pk], reparto=str(self.reparto.pk))
        self.assertContains(r, "Salvato.")
        self.processo.refresh_from_db()
        self.assertEqual(list(self.processo.ruoli.all()), [self.ruolo])
        self.assertEqual(self.processo.reparto_id, self.reparto.pk)
        voce = VoceMacchinaCatalogo.objects.get(mansione=self.processo)
        self.assertEqual((voce.nome, voce.livello_minimo, voce.bloccante), ("Uso mola", "L", True))

    def test_errore_annulla_tutto(self):
        r = self._post(self.processo, [{"nome": "Doppia"}, {"nome": "doppia"}], ruoli=[self.ruolo.pk])
        self.assertContains(r, "Niente salvato")
        self.assertFalse(VoceMacchinaCatalogo.objects.exists())
        self.assertFalse(self.processo.ruoli.exists())

    def test_togli_competenza(self):
        voce = comp.salva_competenza(self.processo, nome="Uso mola")
        self._post(self.processo, [{"id": voce.pk, "nome": "Uso mola", "togli": "on"}])
        voce.refresh_from_db()
        self.assertFalse(voce.attiva)

    def test_crea_processo(self):
        r = self.client.post(reverse("anagrafica:mansionario_struttura_nuova"),
                             {"nome": "Imballo T", "ruoli": self.ruolo.pk})
        m = CompetenzaSkm.objects.get(display="Imballo T")
        self.assertRedirects(r, reverse("anagrafica:mansionario_mansione", args=[m.pk]),
                             fetch_redirect_response=False)
        r = self.client.post(reverse("anagrafica:mansionario_struttura_nuova"), {"nome": "imballo t"})
        self.assertContains(r, "Esiste già una mansione")
        self.assertEqual(CompetenzaSkm.objects.filter(display__iexact="imballo t").count(), 1)

    def test_non_admin_non_entra(self):
        from unittest.mock import patch

        with patch("anagrafica.views._helpers._is_anagrafica_admin", return_value=False):
            r = self.client.get(reverse("anagrafica:mansionario_struttura"))
            self.assertEqual(r.status_code, 302)
            r = self.client.post(reverse("anagrafica:mansionario_struttura_mansione", args=[self.processo.pk]), {})
            self.assertEqual(r.status_code, 403)
            r = self.client.get(reverse("anagrafica:ruoli_in_blocco"))
            self.assertEqual(r.status_code, 302)


# ═══════════════════════════════════════════════════════════════════════════
# Fase D: cambio ruoli in blocco
# ═══════════════════════════════════════════════════════════════════════════

class RuoliInBloccoTests(_Base):
    def test_anteprima_non_scrive(self):
        esiti = blocco.esegui([self.p1.pk, self.p2.pk], ruolo=self.ruolo, modo=blocco.AGGIUNGI,
                              data=timezone.localdate(), applica=False)
        self.assertEqual([e.azione for e in esiti], ["cambia", "cambia"])
        self.assertFalse(DipendenteRuoloOperativo.objects.exists())

    def test_aggiungi_e_diventa_principale_se_manca(self):
        blocco.esegui([self.p1.pk], ruolo=self.ruolo, modo=blocco.AGGIUNGI, data=timezone.localdate(), applica=True)
        self.assertTrue(DipendenteRuoloOperativo.objects.filter(legacy_anagrafica_id=self.p1.pk, ruolo=self.ruolo).exists())
        self.assertEqual(DipendenteAnagraficaAziendale.objects.get(legacy_anagrafica_id=self.p1.pk).ruolo_aziendale,
                         self.ruolo.nome)
        esiti = blocco.esegui([self.p1.pk], ruolo=self.ruolo, modo=blocco.AGGIUNGI, data=timezone.localdate())
        self.assertEqual(esiti[0].azione, "salta")

    def test_sostituisci_chiude_il_vecchio(self):
        oggi = timezone.localdate()
        blocco.esegui([self.p1.pk], ruolo=self.ruolo, modo=blocco.AGGIUNGI, data=oggi - timedelta(days=30), applica=True)
        esiti = blocco.esegui([self.p1.pk], ruolo=self.capo, modo=blocco.SOSTITUISCI, data=oggi, applica=True)
        self.assertEqual(esiti[0].azione, "cambia", esiti[0].dettaglio)
        self.assertEqual(DipendenteAnagraficaAziendale.objects.get(legacy_anagrafica_id=self.p1.pk).ruolo_aziendale,
                         self.capo.nome)
        vecchio = DipendenteRuoloOperativo.objects.get(legacy_anagrafica_id=self.p1.pk, ruolo=self.ruolo)
        self.assertEqual(vecchio.data_fine, oggi - timedelta(days=1))

    def test_chiudi_il_principale_promuove_il_rimasto(self):
        oggi = timezone.localdate()
        blocco.esegui([self.p1.pk], ruolo=self.ruolo, modo=blocco.AGGIUNGI, data=oggi, applica=True)
        blocco.esegui([self.p1.pk], ruolo=self.capo, modo=blocco.AGGIUNGI, data=oggi, applica=True)
        blocco.esegui([self.p1.pk], ruolo=self.ruolo, modo=blocco.CHIUDI, data=oggi, applica=True)
        self.assertEqual(DipendenteRuoloOperativo.objects.get(legacy_anagrafica_id=self.p1.pk, ruolo=self.ruolo).data_fine, oggi)
        self.assertEqual(DipendenteAnagraficaAziendale.objects.get(legacy_anagrafica_id=self.p1.pk).ruolo_aziendale,
                         self.capo.nome)

    def test_controlli(self):
        sicurezza = AmbitoRuolo.objects.create(nome="SGSL T")
        preposto = RuoloOperativo.objects.create(nome="Preposto T", ambito=sicurezza)
        AmbitoRuolo.objects.create(nome="Produttivo T2", alimenta_scheda=True)
        with self.assertRaises(ValueError):
            blocco.esegui([self.p1.pk], ruolo=preposto, modo=blocco.SOSTITUISCI, data=timezone.localdate())
        with self.assertRaises(ValueError):
            blocco.esegui([self.p1.pk], ruolo=self.ruolo, modo=blocco.CHIUDI,
                          data=timezone.localdate() + timedelta(days=3))
        with self.assertRaises(ValueError):
            blocco.esegui([], ruolo=self.ruolo, modo=blocco.AGGIUNGI, data=timezone.localdate())

    def test_view_anteprima_e_registra(self):
        url = reverse("anagrafica:ruoli_in_blocco")
        self.assertContains(self.client.get(url), "STRUTTURA")
        dati = {"ruolo": self.ruolo.pk, "modo": "aggiungi", "data": timezone.localdate().isoformat(),
                "persone": [self.p1.pk]}
        r = self.client.post(url, {**dati, "azione": "prova"})
        self.assertContains(r, "Anteprima")
        self.assertFalse(DipendenteRuoloOperativo.objects.exists())
        self.client.post(url, {**dati, "azione": "applica"})
        self.assertTrue(DipendenteRuoloOperativo.objects.filter(ruolo=self.ruolo, legacy_anagrafica_id=self.p1.pk).exists())


# ═══════════════════════════════════════════════════════════════════════════
# Mansioni in blocco e import Excel delle assegnazioni (27/10/2026)
# ═══════════════════════════════════════════════════════════════════════════

class MansioniInBloccoTests(_Base):
    def test_assegna_cambia_livello_e_togli(self):
        from .services import mansioni_in_blocco as mb

        esiti = mb.esegui([self.p1.pk, self.p2.pk], mansione=self.processo, modo=mb.ASSEGNA, livello="L")
        self.assertEqual([e.azione for e in esiti], ["cambia", "cambia"])
        self.assertFalse(AbilitazioneMacchina.objects.exists())  # anteprima
        mb.esegui([self.p1.pk], mansione=self.processo, modo=mb.ASSEGNA, livello="L", applica=True)
        self.assertEqual(mb.esegui([self.p1.pk], mansione=self.processo, modo=mb.ASSEGNA, livello="L")[0].azione, "salta")
        mb.esegui([self.p1.pk], mansione=self.processo, modo=mb.ASSEGNA, livello="U", applica=True)
        self.assertEqual(AbilitazioneMacchina.objects.get(legacy_anagrafica_id=self.p1.pk, competenza=self.processo).livello, "U")
        mb.esegui([self.p1.pk], mansione=self.processo, modo=mb.TOGLI, applica=True)
        self.assertFalse(AbilitazioneMacchina.objects.get(legacy_anagrafica_id=self.p1.pk, competenza=self.processo).in_lista)
        with self.assertRaises(ValueError):
            mb.esegui([self.p1.pk], mansione=self.processo, modo=mb.ASSEGNA, livello="Z")

    def test_view_mansione_in_blocco(self):
        url = reverse("anagrafica:ruoli_in_blocco")
        self.client.post(url, {"cosa": "mansione", "mansione": self.processo.pk, "modo_mansione": "assegna",
                               "livello": "O", "persone": [self.p1.pk, self.p2.pk], "azione": "applica"})
        self.assertEqual(AbilitazioneMacchina.objects.filter(competenza=self.processo, in_lista=True, livello="O").count(), 2)
        r = self.client.get(url, {"filtro_mansione": self.processo.pk})
        self.assertEqual({p["legacy_id"] for p in r.context["persone"]}, {self.p1.pk, self.p2.pk})


class AssegnazioniExcelTests(_Base):
    def _file(self, ruoli=(), mansioni=()):
        from .services import assegnazioni_excel as ax

        return _xlsx({ax.FOGLIO_RUOLI: [ax.COL_RUOLI, *ruoli], ax.FOGLIO_MANSIONI: [ax.COL_MANSIONI, *mansioni]})

    def test_prova_non_scrive_e_conferma_scrive(self):
        from .services import assegnazioni_excel as ax

        f = self._file(ruoli=[[self.p1.pk, "", "", "", self.ruolo.nome, "", ""]],
                       mansioni=[["", "STRUTTURA UNO", "", "P-SBA", "", "U", "", "", ""],
                                 [self.p2.pk, "", "", "", "Sbavatura", "", "", "", ""]])
        esito = ax.importa(f, applica=False)
        self.assertEqual([r.azione for r in esito.righe], [xl.NUOVO, xl.NUOVO, xl.ERRORE])  # p2: livello mancante
        self.assertFalse(DipendenteRuoloOperativo.objects.exists())
        self.assertFalse(AbilitazioneMacchina.objects.exists())
        ax.importa(f, applica=True)
        self.assertTrue(DipendenteRuoloOperativo.objects.filter(legacy_anagrafica_id=self.p1.pk, ruolo=self.ruolo).exists())
        self.assertTrue(AbilitazioneMacchina.objects.filter(legacy_anagrafica_id=self.p1.pk, competenza=self.processo,
                                                            livello="U", in_lista=True).exists())

    def test_file_scaricato_e_ricaricato_non_cambia_niente(self):
        from openpyxl import load_workbook

        from .services import assegnazioni_excel as ax

        blocco.esegui([self.p1.pk], ruolo=self.ruolo, modo=blocco.AGGIUNGI, data=timezone.localdate(), applica=True)
        svc.assegna(self.processo, legacy_anagrafica_id=self.p1.pk, livello="L")
        contenuto = ax.esporta_xlsx()
        wb = load_workbook(BytesIO(contenuto))
        self.assertEqual(wb[ax.FOGLIO_MANSIONI].max_row, 2)
        esito = ax.importa(contenuto, applica=False)
        self.assertEqual(esito.cambi, 0)
        self.assertEqual(esito.errori, 0)

    def test_chiudi_e_togli(self):
        from .services import assegnazioni_excel as ax

        blocco.esegui([self.p1.pk], ruolo=self.ruolo, modo=blocco.AGGIUNGI, data=timezone.localdate(), applica=True)
        svc.assegna(self.processo, legacy_anagrafica_id=self.p1.pk, livello="L")
        oggi = timezone.localdate().strftime("%d/%m/%Y")
        esito = ax.importa(self._file(ruoli=[[self.p1.pk, "", "", "", self.ruolo.nome, oggi, "CHIUDI"]],
                                      mansioni=[[self.p1.pk, "", "", "P-SBA", "", "", "", "", "TOGLI"]]), applica=True)
        self.assertEqual([r.azione for r in esito.righe], [xl.MODIFICATO, xl.MODIFICATO])
        self.assertEqual(DipendenteRuoloOperativo.objects.get(legacy_anagrafica_id=self.p1.pk, ruolo=self.ruolo).data_fine,
                         timezone.localdate())
        self.assertFalse(AbilitazioneMacchina.objects.get(legacy_anagrafica_id=self.p1.pk, competenza=self.processo).in_lista)

    def test_errori_per_riga(self):
        from .services import assegnazioni_excel as ax

        esito = ax.importa(self._file(ruoli=[[999999, "", "", "", self.ruolo.nome, "", ""],
                                             ["", "NESSUNO", "", "", self.ruolo.nome, "", ""],
                                             [self.p1.pk, "", "", "", "Ruolo che non c'è", "", ""],
                                             [self.p1.pk, "", "", "", self.ruolo.nome, "", "BOH"]]), applica=False)
        self.assertEqual(esito.errori, 4)

    def test_view_export_e_import(self):
        r = self.client.get(reverse("anagrafica:assegnazioni_export"), {"vuoto": "1"})
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        f = self._file(ruoli=[[self.p1.pk, "", "", "", self.ruolo.nome, "", ""]])
        from django.core.files.uploadedfile import SimpleUploadedFile

        r = self.client.post(reverse("anagrafica:assegnazioni_import"),
                             {"azione": "prova", "file": SimpleUploadedFile("a.xlsx", f)})
        self.assertContains(r, "non è stato salvato niente")
        self.assertFalse(DipendenteRuoloOperativo.objects.exists())
        r = self.client.post(reverse("anagrafica:assegnazioni_import"), {"azione": "applica", "sha": r.context["sha"]})
        self.assertTrue(DipendenteRuoloOperativo.objects.filter(legacy_anagrafica_id=self.p1.pk, ruolo=self.ruolo).exists())
