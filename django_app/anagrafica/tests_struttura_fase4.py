"""Struttura Fase 4: accessi al portale dai ruoli d'organigramma, modificabili e tracciati.

Registro automatico delle variazioni per persona, modifica dalla scheda del ruolo,
foglio nel riesame accessi. Solo dati sintetici.
"""
from __future__ import annotations

from datetime import timedelta

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from core.legacy_models import AnagraficaDipendente
from core.models import AccessGroup, AuditLog
from core.permessi.variazioni import AZIONE

from .models import AmbitoRuolo, DipendenteRuoloOperativo, RuoloOperativo
from .tests import _ensure_anagrafica_table

User = get_user_model()


@override_settings(LEGACY_AUTH_ENABLED=False)
class AccessiDaRuoloTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        _ensure_anagrafica_table()

    def setUp(self):
        self.admin = User.objects.create_superuser("f4_admin", "f4@example.invalid", "x")
        ambito, _ = AmbitoRuolo.objects.get_or_create(nome="Produttivo", defaults={"alimenta_scheda": True})
        self.ruolo = RuoloOperativo.objects.create(nome="Capoturno sintetico F4", ambito=ambito)
        self.altro = RuoloOperativo.objects.create(nome="Preposto sintetico F4", ambito=ambito)
        self.g_turni = AccessGroup.objects.create(code="f4-turni", label="Turni sintetico F4")
        self.g_qualita = AccessGroup.objects.create(code="f4-qualita", label="Qualità sintetico F4")
        self.g_admin = AccessGroup.objects.create(code="f4-admin", label="Admin sintetico F4", amministratore=True)
        self.ruolo.profili_permessi.add(self.g_turni)
        self.persona = AnagraficaDipendente.objects.create(nome="Elena", cognome="PROVA", aliasusername="f4.elena")

    def _voci(self):
        return list(AuditLog.objects.filter(azione=AZIONE, oggetto_id=str(self.persona.pk)).order_by("id"))

    def test_assegnare_e_chiudere_il_ruolo_si_registra(self):
        a = DipendenteRuoloOperativo.objects.create(legacy_anagrafica_id=self.persona.pk, ruolo=self.ruolo)
        voci = self._voci()
        self.assertEqual(len(voci), 1)
        self.assertEqual(voci[0].dettaglio["ricevuti"], ["Turni sintetico F4"])
        self.assertIn("assegnato", voci[0].dettaglio["causa"])
        a.data_fine = timezone.localdate() - timedelta(days=1)
        a.save()
        voci = self._voci()
        self.assertEqual(voci[-1].dettaglio["tolti"], ["Turni sintetico F4"])
        self.assertIn("chiuso", voci[-1].dettaglio["causa"])

    def test_collegamento_profili_e_doppioni(self):
        DipendenteRuoloOperativo.objects.create(legacy_anagrafica_id=self.persona.pk, ruolo=self.ruolo)
        DipendenteRuoloOperativo.objects.create(legacy_anagrafica_id=self.persona.pk, ruolo=self.altro)
        prima = len(self._voci())
        # Lo stesso ruolo del portale già dato da un altro ruolo: nessuna variazione reale.
        self.altro.profili_permessi.add(self.g_turni)
        self.assertEqual(len(self._voci()), prima)
        self.altro.profili_permessi.add(self.g_qualita)
        self.assertEqual(self._voci()[-1].dettaglio["ricevuti"], ["Qualità sintetico F4"])
        # Un ruolo amministratore non arriva mai dall'organigramma.
        self.altro.profili_permessi.add(self.g_admin)
        self.assertEqual(self._voci()[-1].dettaglio["ricevuti"], ["Qualità sintetico F4"])

    def test_modifica_dalla_scheda_ruolo(self):
        DipendenteRuoloOperativo.objects.create(legacy_anagrafica_id=self.persona.pk, ruolo=self.ruolo)
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("anagrafica:struttura_scheda_ruolo", args=[self.ruolo.pk]))
        self.assertContains(resp, "Cambia gli accessi che concede")
        self.assertNotContains(resp, "Admin sintetico F4")
        url = reverse("anagrafica:struttura_ruolo_accessi", args=[self.ruolo.pk])
        resp = self.client.post(url, {"accessi": [self.g_qualita.pk, self.g_admin.pk]})
        self.assertRedirects(resp, reverse("anagrafica:struttura_org") + f"#scheda-ruolo-{self.ruolo.pk}",
                             fetch_redirect_response=False)
        self.assertEqual(set(self.ruolo.profili_permessi.all()), {self.g_qualita})
        # Aggiunta e rimozione: due voci, una per operazione.
        ultime = [v.dettaglio for v in self._voci()[-2:]]
        self.assertEqual(ultime[0]["ricevuti"], ["Qualità sintetico F4"])
        self.assertEqual(ultime[1]["tolti"], ["Turni sintetico F4"])
        self.assertTrue(AuditLog.objects.filter(azione="permessi_ruolo_organigramma").exists())

    def test_modifica_negata_a_chi_non_amministra_il_portale(self):
        utente = User.objects.create_user("f4_hr", "f4hr@example.invalid", "x")
        self.client.force_login(utente)
        self.client.post(reverse("anagrafica:struttura_ruolo_accessi", args=[self.ruolo.pk]),
                         {"accessi": [self.g_qualita.pk]})
        self.assertEqual(set(self.ruolo.profili_permessi.all()), {self.g_turni})

    def test_scheda_dipendente_e_riesame(self):
        DipendenteRuoloOperativo.objects.create(legacy_anagrafica_id=self.persona.pk, ruolo=self.ruolo)
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("anagrafica:dipendente_detail", args=[self.persona.pk]))
        self.assertContains(resp, "Variazioni degli accessi dal ruolo")
        from io import BytesIO

        from openpyxl import load_workbook

        from admin_portale.riesame_accessi import genera_riesame_xlsx

        wb = load_workbook(BytesIO(genera_riesame_xlsx(preparato_da="Test", sistema_attivo=True)))
        ws = wb["Variazioni da organigramma"]
        self.assertEqual(ws.cell(2, 2).value, "PROVA Elena")
        self.assertIn("Turni sintetico F4", ws.cell(2, 4).value)
