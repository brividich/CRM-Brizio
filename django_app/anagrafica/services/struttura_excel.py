"""Struttura ruoli → mansioni → competenze in un file Excel, andata e ritorno.

L'export scrive tre fogli (Ruoli, Mansioni, Competenze) più le istruzioni; lo
stesso file, corretto o completato, si reimporta. L'import:

- **non cancella mai**: aggiunge e modifica; per spegnere una voce si scrive
  ``NO`` nella colonna «Attiva»/«Attivo»;
- **cella vuota = non cambia**: un file che porta solo le competenze non tocca
  ruoli e impostazioni delle mansioni;
- prima gira **in prova** (tutto dentro una transazione annullata) e mostra
  riga per riga cosa farebbe; solo la conferma scrive;
- una riga sbagliata si scarta con il motivo e non ferma le altre.

Cataloghi che l'import **non crea**: mansioni di rischio (DVR, medico), reparti,
ambiti, corsi, qualifiche, processi MOD.128 e macchine (nascono dalla skill
matrix). Un nome che non esiste lì è un errore della riga, non un dato nuovo.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO

from django.db import transaction

from ..models import CompetenzaSkm, LivelloSkm, VoceMacchinaCatalogo
from .mansionario_struttura import TIPI_MANSIONE, chiave_nuova

FOGLIO_RUOLI = "Ruoli"
FOGLIO_MANSIONI = "Mansioni"
FOGLIO_COMPETENZE = "Competenze"
FOGLIO_ISTRUZIONI = "Istruzioni"

COL_RUOLI = ["Ruolo", "Ambito", "Riporta a", "Descrizione", "Attivo"]
COL_MANSIONI = ["Chiave", "Mansione", "Tipo", "Macchina", "Reparto", "Ruoli", "Mansioni di rischio", "Attiva"]
COL_COMPETENZE = ["Chiave mansione", "Mansione", "Competenza", "Livello minimo", "Bloccante", "Corso",
                  "Qualifica", "Processo MOD.128", "Ordine", "Attiva"]

SEP_LISTA = re.compile(r"\s*[;\n]\s*")
_SI = {"SI", "SÌ", "S", "X", "1", "TRUE", "VERO", "YES", "Y"}
_NO = {"NO", "N", "0", "FALSE", "FALSO"}
LIVELLI = {c for c, _ in LivelloSkm.choices}
MAX_RIGHE = 5000

ISTRUZIONI = [
    "Struttura ruoli → mansioni → competenze del Mansionario (NOVICROM HUB).",
    "",
    "Come si usa",
    "1. Correggi o completa i fogli Ruoli, Mansioni e Competenze. Non rinominare le intestazioni.",
    "2. Ricarica il file da Mansionario → Struttura → Importa: prima vedi l'anteprima, poi confermi.",
    "",
    "Regole",
    "• Cella vuota = il dato non cambia. L'import non cancella niente.",
    "• Per spegnere un ruolo, una mansione o una competenza scrivi NO nella colonna Attivo/Attiva.",
    "• Più valori nella stessa cella (Ruoli, Mansioni di rischio): separali con ; oppure vai a capo.",
    "  Quei valori sostituiscono l'elenco di oggi.",
    "• Livelli: I = in formazione, L = intermedio, U = autonomo, O = esperto.",
    "• Sì/No: SI, NO (vanno bene anche X, 1, 0).",
    "",
    "Mansioni",
    "• La colonna Chiave identifica la mansione: lasciala com'è. Una riga senza chiave si abbina per nome.",
    "• Una mansione nuova (nome mai visto) nasce come processo. Le macchine nascono dalla skill matrix:",
    "  qui si possono solo collegare a ruoli e rischi.",
    "• Il Reparto si sceglie solo per i processi: per le macchine vale quello della macchina.",
    "",
    "Competenze",
    "• Una riga per competenza. Si abbina alla mansione con la Chiave mansione o, se vuota, con il nome.",
    "• Corso: codice del corso (o titolo esatto). Qualifica e Processo MOD.128: nome esatto.",
    "",
    "Cataloghi che l'import non crea: mansioni di rischio, reparti, ambiti, corsi, qualifiche,",
    "processi MOD.128, macchine. Un nome che non esiste è segnalato come errore della riga.",
]


# ═══════════════════════════════════════════════════════════════════════════
# Export
# ═══════════════════════════════════════════════════════════════════════════

def _si_no(valore) -> str:
    return "SI" if valore else "NO"


def _codice_corso(corso) -> str:
    return f"{corso.codice} · {corso.titolo}" if corso else ""


def righe_export() -> dict[str, list[list]]:
    """Le righe dei tre fogli, nello stesso formato che l'import rilegge."""
    from ..models import RuoloOperativo

    ruoli = [
        [r.nome, r.ambito.nome if r.ambito_id else "", r.riporta_a.nome if r.riporta_a_id else "",
         r.descrizione, _si_no(r.is_active)]
        for r in RuoloOperativo.objects.select_related("ambito", "riporta_a").order_by("nome")
    ]
    mansioni = []
    for m in (CompetenzaSkm.objects.filter(tipo__in=TIPI_MANSIONE)
              .select_related("asset", "reparto").prefetch_related("ruoli", "mansioni_rischio")
              .order_by("tipo", "display", "competenza_key")):
        mansioni.append([
            m.competenza_key, m.nome, m.tipo, m.asset.name if m.asset_id else "",
            m.reparto_effettivo, "; ".join(sorted(r.nome for r in m.ruoli.all())),
            "; ".join(sorted(x.nome for x in m.mansioni_rischio.all())), _si_no(m.attiva),
        ])
    competenze = [
        [v.mansione.competenza_key, v.mansione.nome, v.nome, v.livello_minimo, _si_no(v.bloccante),
         _codice_corso(v.corso), v.qualifica.nome if v.qualifica_id else "",
         v.processo_mpq.nome if v.processo_mpq_id else "", v.ordine, _si_no(v.attiva)]
        for v in (VoceMacchinaCatalogo.objects.filter(mansione__isnull=False, mansione__tipo__in=TIPI_MANSIONE)
                  .select_related("mansione", "corso", "qualifica", "processo_mpq")
                  .order_by("mansione__display", "mansione__competenza_key", "ordine", "pk"))
    ]
    return {FOGLIO_RUOLI: ruoli, FOGLIO_MANSIONI: mansioni, FOGLIO_COMPETENZE: competenze}


def esporta_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    from core.excel_export import write_cell

    dati = righe_export()
    wb = Workbook()
    ws = wb.active
    ws.title = FOGLIO_ISTRUZIONI
    for i, riga in enumerate(ISTRUZIONI, 1):
        c = write_cell(ws, i, 1, riga)
        if i == 1 or riga in {"Come si usa", "Regole", "Mansioni", "Competenze"}:
            c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 110

    testata = PatternFill("solid", fgColor="1E3A5F")
    for nome, colonne in ((FOGLIO_RUOLI, COL_RUOLI), (FOGLIO_MANSIONI, COL_MANSIONI),
                          (FOGLIO_COMPETENZE, COL_COMPETENZE)):
        ws = wb.create_sheet(nome)
        for j, col in enumerate(colonne, 1):
            c = write_cell(ws, 1, j, col)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = testata
            c.alignment = Alignment(vertical="center")
        for i, riga in enumerate(dati[nome], 2):
            for j, valore in enumerate(riga, 1):
                write_cell(ws, i, j, valore)
        for j, col in enumerate(colonne, 1):
            lungh = max([len(str(col))] + [len(str(r[j - 1] or "")) for r in dati[nome][:500]])
            ws.column_dimensions[get_column_letter(j)].width = min(max(lungh + 2, 10), 60)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(colonne))}{max(len(dati[nome]) + 1, 2)}"
        # Tendine su Sì/No e livelli: aiutano chi compila, l'import resta tollerante.
        fine = MAX_RIGHE + 1
        for j, col in enumerate(colonne, 1):
            lettera = get_column_letter(j)
            if col in {"Attivo", "Attiva", "Bloccante"}:
                dv = DataValidation(type="list", formula1='"SI,NO"', allow_blank=True)
            elif col == "Livello minimo":
                dv = DataValidation(type="list", formula1='"I,L,U,O"', allow_blank=True)
            elif col == "Tipo":
                dv = DataValidation(type="list", formula1='"processo,macchina"', allow_blank=True)
            else:
                continue
            ws.add_data_validation(dv)
            dv.add(f"{lettera}2:{lettera}{fine}")
    wb.active = 1
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════════════
# Import
# ═══════════════════════════════════════════════════════════════════════════

NUOVO, MODIFICATO, INVARIATO, ERRORE = "nuovo", "modificato", "invariato", "errore"


@dataclass
class Riga:
    foglio: str
    numero: int
    chiave: str
    azione: str
    dettaglio: str = ""


@dataclass
class Esito:
    righe: list[Riga] = field(default_factory=list)
    applicato: bool = False
    avvisi: list[str] = field(default_factory=list)

    def aggiungi(self, foglio, numero, chiave, azione, dettaglio=""):
        self.righe.append(Riga(foglio, numero, chiave, azione, dettaglio))

    def conteggi(self) -> dict[str, dict[str, int]]:
        out: dict[str, dict[str, int]] = {}
        for r in self.righe:
            out.setdefault(r.foglio, {NUOVO: 0, MODIFICATO: 0, INVARIATO: 0, ERRORE: 0})[r.azione] += 1
        return out

    @property
    def errori(self) -> int:
        return sum(1 for r in self.righe if r.azione == ERRORE)

    @property
    def cambi(self) -> int:
        return sum(1 for r in self.righe if r.azione in (NUOVO, MODIFICATO))


class ErroreRiga(Exception):
    pass


class _Annulla(Exception):
    """Fa tornare indietro la transazione della prova."""


def _testo(valore) -> str:
    if valore is None:
        return ""
    if isinstance(valore, float) and valore.is_integer():
        valore = int(valore)
    return str(valore).strip()


def _bool(valore, colonna: str):
    t = _testo(valore).upper()
    if not t:
        return None
    if t in _SI:
        return True
    if t in _NO:
        return False
    raise ErroreRiga(f"«{colonna}»: scrivi SI o NO, non «{_testo(valore)}»")


def _lista(valore) -> list[str] | None:
    t = _testo(valore)
    if not t:
        return None
    return [x for x in SEP_LISTA.split(t) if x]


def leggi_fogli(contenuto: bytes) -> dict[str, list[tuple[int, dict]]]:
    """{foglio: [(numero riga Excel, {colonna: valore})]} dei fogli riconosciuti."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(BytesIO(contenuto), read_only=True, data_only=True)
    except Exception as exc:  # file non xlsx, corrotto, protetto
        raise ValueError("Il file non è un Excel .xlsx leggibile.") from exc
    attesi = {FOGLIO_RUOLI: COL_RUOLI, FOGLIO_MANSIONI: COL_MANSIONI, FOGLIO_COMPETENZE: COL_COMPETENZE}
    per_nome = {ws.title.strip().casefold(): ws for ws in wb.worksheets}
    out: dict[str, list[tuple[int, dict]]] = {}
    for nome, colonne in attesi.items():
        ws = per_nome.get(nome.casefold())
        if ws is None:
            continue
        righe = ws.iter_rows(values_only=True)
        intestazione = next(righe, None) or ()
        noti = {c.casefold(): c for c in colonne}
        mappa = {i: noti[_testo(h).casefold()] for i, h in enumerate(intestazione) if _testo(h).casefold() in noti}
        if not mappa:
            continue
        dati = []
        for numero, valori in enumerate(righe, 2):
            if numero > MAX_RIGHE + 1:
                raise ValueError(f"Il foglio {nome} supera le {MAX_RIGHE} righe.")
            riga = {mappa[i]: v for i, v in enumerate(valori or ()) if i in mappa}
            if any(_testo(v) for v in riga.values()):
                dati.append((numero, riga))
        out[nome] = dati
    wb.close()
    if not out:
        raise ValueError("Nel file non ci sono i fogli Ruoli, Mansioni o Competenze con le loro intestazioni.")
    return out


class _Cataloghi:
    """Nomi → oggetti, caricati una volta (case-insensitive)."""

    def __init__(self):
        from ..models import AmbitoRuolo, Mansione, Reparto, RuoloOperativo, TipoQualifica, TrainingCourse
        from ..models_mpq import ProcessoQualificato

        self.ambiti = {a.nome.casefold(): a for a in AmbitoRuolo.objects.all()}
        self.reparti = {r.nome.casefold(): r for r in Reparto.objects.filter(is_active=True)}
        self.rischi = {m.nome.casefold(): m for m in Mansione.objects.filter(is_active=True)}
        self.qualifiche = {q.nome.casefold(): q for q in TipoQualifica.objects.filter(is_active=True)}
        self.corsi_codice = {c.codice.casefold(): c for c in TrainingCourse.objects.filter(is_active=True)}
        self.corsi_titolo: dict[str, list] = {}
        for c in self.corsi_codice.values():
            self.corsi_titolo.setdefault(c.titolo.strip().casefold(), []).append(c)
        self.processi: dict[str, list] = {}
        for p in ProcessoQualificato.objects.all():
            self.processi.setdefault(p.nome.strip().casefold(), []).append(p)
            if p.codice:
                self.processi.setdefault(p.codice.strip().casefold(), []).append(p)
        self.ruoli = {r.nome.casefold(): r for r in RuoloOperativo.objects.all()}

    def uno(self, mappa, nome, cosa):
        ogg = mappa.get(nome.casefold())
        if ogg is None:
            raise ErroreRiga(f"{cosa} «{nome}» non esiste nel catalogo")
        return ogg

    def corso(self, testo):
        t = testo.strip()
        codice = t.split(" · ", 1)[0].strip()
        c = self.corsi_codice.get(codice.casefold())
        if c is not None:
            return c
        trovati = self.corsi_titolo.get(t.casefold(), [])
        if len(trovati) == 1:
            return trovati[0]
        raise ErroreRiga(f"corso «{t}» " + ("ambiguo: usa il codice" if trovati else "non esiste o non è attivo"))

    def processo(self, testo):
        trovati = {p.pk: p for p in self.processi.get(testo.strip().casefold(), [])}
        if len(trovati) == 1:
            return next(iter(trovati.values()))
        raise ErroreRiga(f"processo MOD.128 «{testo}» " + ("ambiguo" if trovati else "non esiste"))

    def ruolo(self, nome):
        return self.uno(self.ruoli, nome, "ruolo")


def _crea_ciclo(ruolo, capo) -> bool:
    visti, corrente = set(), capo
    while corrente is not None:
        if corrente.pk == ruolo.pk or corrente.pk in visti:
            return True
        visti.add(corrente.pk)
        corrente = corrente.riporta_a
    return False


def _importa_ruoli(righe, cat: _Cataloghi, esito: Esito):
    from ..models import RuoloOperativo

    riporti = []
    for numero, r in righe:
        nome = _testo(r.get("Ruolo"))[:100]
        if not nome:
            esito.aggiungi(FOGLIO_RUOLI, numero, "", ERRORE, "manca il nome del ruolo")
            continue
        try:
            with transaction.atomic():
                ruolo = cat.ruoli.get(nome.casefold())
                nuovo = ruolo is None
                if nuovo:
                    ruolo = RuoloOperativo(nome=nome)
                cambi = []
                ambito = _testo(r.get("Ambito"))
                if ambito:
                    a = cat.uno(cat.ambiti, ambito, "ambito")
                    if ruolo.ambito_id != a.pk:
                        ruolo.ambito = a
                        cambi.append("ambito")
                descr = _testo(r.get("Descrizione"))
                if descr and descr != ruolo.descrizione:
                    ruolo.descrizione = descr
                    cambi.append("descrizione")
                attivo = _bool(r.get("Attivo"), "Attivo")
                if attivo is not None and attivo != ruolo.is_active:
                    ruolo.is_active = attivo
                    cambi.append("attivo" if attivo else "spento")
                if nuovo or cambi:
                    ruolo.save()
                cat.ruoli[nome.casefold()] = ruolo
        except ErroreRiga as exc:
            esito.aggiungi(FOGLIO_RUOLI, numero, nome, ERRORE, str(exc))
            continue
        capo = _testo(r.get("Riporta a"))
        if capo:
            riporti.append((numero, ruolo, capo))
        esito.aggiungi(FOGLIO_RUOLI, numero, nome, NUOVO if nuovo else MODIFICATO if cambi else INVARIATO,
                       ", ".join(cambi))
    # Il «riporta a» si risolve dopo: il capo può essere una riga più in basso.
    for numero, ruolo, capo_nome in riporti:
        riga = next(x for x in reversed(esito.righe) if x.foglio == FOGLIO_RUOLI and x.numero == numero)
        try:
            capo = cat.ruolo(capo_nome)
            if ruolo.riporta_a_id == capo.pk:
                continue
            if _crea_ciclo(ruolo, capo):
                raise ErroreRiga(f"«riporta a {capo.nome}» creerebbe un giro nella gerarchia")
            ruolo.riporta_a = capo
            ruolo.save(update_fields=["riporta_a"])
            if riga.azione == INVARIATO:
                riga.azione = MODIFICATO
            riga.dettaglio = ", ".join(x for x in (riga.dettaglio, f"riporta a {capo.nome}") if x)
        except ErroreRiga as exc:
            riga.dettaglio = ", ".join(x for x in (riga.dettaglio, f"riporta a non applicato: {exc}") if x)
            esito.avvisi.append(f"Ruoli, riga {numero}: {exc}")


class _Mansioni:
    def __init__(self):
        self.per_chiave: dict[str, CompetenzaSkm] = {}
        self.per_nome: dict[str, list[CompetenzaSkm]] = {}
        self.per_alias: dict[str, list[CompetenzaSkm]] = {}
        for m in CompetenzaSkm.objects.filter(tipo__in=TIPI_MANSIONE).select_related("asset", "reparto"):
            self.registra(m)

    def registra(self, m):
        from .mansionario_struttura import nomi_precedenti

        self.per_chiave[m.competenza_key.casefold()] = m
        for nome in {m.nome.strip().casefold(), m.competenza_key.casefold()}:
            lista = self.per_nome.setdefault(nome, [])
            if m not in lista:
                lista.append(m)
        # Nomi precedenti a parte: un file preparato prima di una rinomina si abbina
        # ancora, ma un nome attuale vince sempre su un vecchio nome.
        for nome in {a.casefold() for a in nomi_precedenti(m)}:
            lista = self.per_alias.setdefault(nome, [])
            if m not in lista:
                lista.append(m)

    def trova(self, chiave: str, nome: str):
        if chiave:
            m = self.per_chiave.get(chiave.casefold())
            if m is None:
                raise ErroreRiga(f"nessuna mansione con chiave «{chiave}»")
            return m
        trovate = self.per_nome.get(nome.casefold()) or self.per_alias.get(nome.casefold(), [])
        if len(trovate) > 1:
            raise ErroreRiga(f"più mansioni si chiamano «{nome}»: indica la chiave")
        return trovate[0] if trovate else None


def _importa_mansioni(righe, cat: _Cataloghi, mans: _Mansioni, esito: Esito):
    for numero, r in righe:
        chiave, nome = _testo(r.get("Chiave")), _testo(r.get("Mansione"))[:255]
        etichetta = nome or chiave
        if not etichetta:
            esito.aggiungi(FOGLIO_MANSIONI, numero, "", ERRORE, "manca il nome della mansione")
            continue
        try:
            with transaction.atomic():
                m = mans.trova(chiave, nome)
                nuovo = m is None
                cambi = []
                tipo = _testo(r.get("Tipo")).casefold()
                if tipo and tipo not in TIPI_MANSIONE:
                    raise ErroreRiga(f"tipo «{tipo}»: scrivi processo o macchina")
                if nuovo:
                    if not nome:
                        raise ErroreRiga("manca il nome della mansione nuova")
                    if tipo == CompetenzaSkm.TIPO_MACCHINA:
                        raise ErroreRiga("le macchine nascono dalla skill matrix: qui si collegano soltanto")
                    m = CompetenzaSkm.objects.create(
                        competenza_key=chiave_nuova(nome), display=nome, tipo=CompetenzaSkm.TIPO_PROCESSO,
                    )
                    mans.registra(m)
                else:
                    if tipo and tipo != m.tipo:
                        raise ErroreRiga(f"la mansione è di tipo {m.tipo}: il tipo non si cambia da qui")
                    if nome and chiave and nome != m.nome:
                        from .mansionario_struttura import rinomina
                        try:
                            prima = rinomina(m, nome, salva=False)
                        except ValueError as exc:
                            raise ErroreRiga(str(exc)) from exc
                        if prima:
                            cambi.append(f"nome (prima «{prima}»)")
                            mans.registra(m)
                reparto = _testo(r.get("Reparto"))
                if reparto:
                    if m.asset_id:
                        if reparto.casefold() != (m.reparto_effettivo or "").casefold():
                            esito.avvisi.append(
                                f"Mansioni, riga {numero}: «{m.nome}» è una macchina, il reparto viene dalla "
                                f"macchina ({m.reparto_effettivo or 'non indicato'}): «{reparto}» ignorato.")
                    else:
                        rep = cat.uno(cat.reparti, reparto, "reparto")
                        if m.reparto_id != rep.pk:
                            m.reparto = rep
                            cambi.append("reparto")
                attiva = _bool(r.get("Attiva"), "Attiva")
                if attiva is not None and attiva != m.attiva:
                    m.attiva = attiva
                    cambi.append("attivata" if attiva else "spenta")
                if cambi:
                    m.save()
                ruoli = _lista(r.get("Ruoli"))
                if ruoli is not None:
                    trovati = [cat.ruolo(x) for x in ruoli]
                    spenti = [x.nome for x in trovati if not x.is_active]
                    if spenti:
                        raise ErroreRiga(f"ruolo spento: {', '.join(spenti)}")
                    nuovi = {x.pk for x in trovati}
                    if nuovi != set(m.ruoli.values_list("pk", flat=True)):
                        m.ruoli.set(nuovi)
                        cambi.append("ruoli")
                rischi = _lista(r.get("Mansioni di rischio"))
                if rischi is not None:
                    nuovi = {cat.uno(cat.rischi, x, "mansione di rischio").pk for x in rischi}
                    if nuovi != set(m.mansioni_rischio.values_list("pk", flat=True)):
                        m.mansioni_rischio.set(nuovi)
                        cambi.append("mansioni di rischio")
        except ErroreRiga as exc:
            esito.aggiungi(FOGLIO_MANSIONI, numero, etichetta, ERRORE, str(exc))
            continue
        esito.aggiungi(FOGLIO_MANSIONI, numero, m.nome, NUOVO if nuovo else MODIFICATO if cambi else INVARIATO,
                       f"chiave {m.competenza_key}" if nuovo else ", ".join(cambi))


def _importa_competenze(righe, cat: _Cataloghi, mans: _Mansioni, esito: Esito):
    from .mansionario_competenze import salva_competenza

    for numero, r in righe:
        chiave, nome_m = _testo(r.get("Chiave mansione")), _testo(r.get("Mansione"))
        nome = _testo(r.get("Competenza"))[:150]
        etichetta = f"{nome_m or chiave} › {nome}"
        if not nome or not (chiave or nome_m):
            esito.aggiungi(FOGLIO_COMPETENZE, numero, etichetta, ERRORE, "servono mansione e competenza")
            continue
        try:
            with transaction.atomic():
                m = mans.trova(chiave, nome_m)
                if m is None:
                    raise ErroreRiga(f"la mansione «{nome_m}» non esiste: aggiungila nel foglio Mansioni")
                voce = (VoceMacchinaCatalogo.objects.filter(mansione=m, nome__iexact=nome)
                        .order_by("-attiva", "pk").first())
                nuovo = voce is None
                attiva = _bool(r.get("Attiva"), "Attiva")
                if attiva is False:
                    if nuovo:
                        raise ErroreRiga("competenza nuova segnata come non attiva: niente da fare")
                    if voce.attiva:
                        voce.attiva = False
                        voce.save(update_fields=["attiva"])
                        esito.aggiungi(FOGLIO_COMPETENZE, numero, etichetta, MODIFICATO, "spenta")
                    else:
                        esito.aggiungi(FOGLIO_COMPETENZE, numero, etichetta, INVARIATO)
                    continue
                livello = _testo(r.get("Livello minimo")).upper()
                if livello and livello not in LIVELLI:
                    raise ErroreRiga(f"livello minimo «{livello}»: scrivi I, L, U oppure O")
                bloccante = _bool(r.get("Bloccante"), "Bloccante")
                corso = _testo(r.get("Corso"))
                qualifica = _testo(r.get("Qualifica"))
                processo = _testo(r.get("Processo MOD.128"))
                ordine_t = _testo(r.get("Ordine"))
                if ordine_t and not ordine_t.isdigit():
                    raise ErroreRiga(f"ordine «{ordine_t}»: serve un numero")
                prima = None if nuovo else (voce.livello_minimo, voce.bloccante, voce.corso_id, voce.qualifica_id,
                                            voce.processo_mpq_id, voce.ordine, voce.attiva, voce.nome)
                voce = salva_competenza(
                    m, voce=voce, nome=nome if nuovo or nome != voce.nome else voce.nome,
                    livello_minimo=livello or (voce.livello_minimo if voce else LivelloSkm.AUTONOMO),
                    bloccante=bloccante if bloccante is not None else (voce.bloccante if voce else False),
                    corso_id=cat.corso(corso).pk if corso else (voce.corso_id if voce else None),
                    qualifica_id=(cat.uno(cat.qualifiche, qualifica, "qualifica").pk if qualifica
                                  else (voce.qualifica_id if voce else None)),
                    processo_id=cat.processo(processo).pk if processo else (voce.processo_mpq_id if voce else None),
                    ordine=int(ordine_t) if ordine_t else (voce.ordine if voce else 0),
                )
                dopo = (voce.livello_minimo, voce.bloccante, voce.corso_id, voce.qualifica_id,
                        voce.processo_mpq_id, voce.ordine, voce.attiva, voce.nome)
        except ErroreRiga as exc:
            esito.aggiungi(FOGLIO_COMPETENZE, numero, etichetta, ERRORE, str(exc))
            continue
        except ValueError as exc:  # dalle regole di salva_competenza
            esito.aggiungi(FOGLIO_COMPETENZE, numero, etichetta, ERRORE, str(exc))
            continue
        if nuovo:
            esito.aggiungi(FOGLIO_COMPETENZE, numero, etichetta, NUOVO,
                           f"minimo {voce.livello_minimo}" + (", bloccante" if voce.bloccante else ""))
        else:
            esito.aggiungi(FOGLIO_COMPETENZE, numero, etichetta, MODIFICATO if dopo != prima else INVARIATO)


def importa(contenuto: bytes, *, applica: bool) -> Esito:
    """Legge il file e applica la struttura; con ``applica=False`` annulla tutto alla fine."""
    fogli = leggi_fogli(contenuto)
    esito = Esito()
    try:
        with transaction.atomic():
            cat = _Cataloghi()
            mans = _Mansioni()
            if FOGLIO_RUOLI in fogli:
                _importa_ruoli(fogli[FOGLIO_RUOLI], cat, esito)
            if FOGLIO_MANSIONI in fogli:
                _importa_mansioni(fogli[FOGLIO_MANSIONI], cat, mans, esito)
            if FOGLIO_COMPETENZE in fogli:
                _importa_competenze(fogli[FOGLIO_COMPETENZE], cat, mans, esito)
            if not applica:
                raise _Annulla
    except _Annulla:
        pass
    esito.applicato = applica
    return esito
