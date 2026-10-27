"""Ruoli e mansioni delle persone in un file Excel, andata e ritorno (27/10/2026).

Inserimento e aggiornamento massivo: si scarica il file con le assegnazioni di
oggi (o vuoto, come modello), si aggiungono o si cambiano le righe, si ricarica.

- Foglio «Ruoli persone»: una riga per persona e ruolo. «Azione» vuota o
  ``AGGIUNGI`` = la persona deve avere il ruolo (dal giorno «Dal»); ``CHIUDI``
  = il ruolo finisce il giorno «Dal».
- Foglio «Mansioni persone»: una riga per persona e mansione. «Azione» vuota o
  ``ASSEGNA`` = in lista col «Livello» indicato; ``TOGLI`` = fuori lista.

Come l'import della struttura: **non cancella niente** se non lo chiede una
riga, una riga uguale a com'è oggi non cambia nulla (il file scaricato e
ricaricato tale e quale non fa niente), prima gira **in prova** in una
transazione annullata e solo la conferma scrive; una riga sbagliata si scarta
con il motivo e non ferma le altre. La persona si riconosce dall'«ID persona»
o, se manca, dal nome (solo se c'è una persona in forza con quel nome); ruoli
e mansioni devono esistere: l'import non li crea.
"""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from django.db import DatabaseError, transaction
from django.utils import timezone

from .struttura_excel import ERRORE, INVARIATO, MODIFICATO, NUOVO, Esito, ErroreRiga, _testo

FOGLIO_RUOLI = "Ruoli persone"
FOGLIO_MANSIONI = "Mansioni persone"
FOGLIO_ISTRUZIONI = "Istruzioni"
COL_RUOLI = ["ID persona", "Persona", "Reparto", "Area", "Ruolo", "Dal", "Azione"]
COL_MANSIONI = ["ID persona", "Persona", "Reparto", "Chiave mansione", "Mansione", "Livello", "Classificazione",
                "Coefficiente", "Azione"]
MAX_RIGHE = 5000
AGGIUNGI, CHIUDI, ASSEGNA, TOGLI = "AGGIUNGI", "CHIUDI", "ASSEGNA", "TOGLI"

ISTRUZIONI = [
    "Ruoli e mansioni delle persone — NOVICROM HUB",
    "",
    "Come si usa",
    "1. Scarica il file (con le assegnazioni di oggi o vuoto) dalla pagina «Ruoli e mansioni in blocco».",
    "2. Aggiungi righe o cambia quelle che ci sono; puoi cancellare le righe che non ti servono.",
    "3. Ricarica il file: prima vedi la prova riga per riga, poi confermi.",
    "",
    "Regole",
    "- La persona si riconosce dall'«ID persona»; se è vuoto, dal nome (cognome e nome come nel portale).",
    "- Ruoli e mansioni devono esistere nel portale: l'import non li crea (per crearli c'è «Struttura e import»).",
    "- Una riga uguale a com'è oggi non cambia niente. Niente si cancella se una riga non lo chiede.",
    "",
    "Ruoli persone",
    "- Azione vuota o AGGIUNGI: la persona deve avere il ruolo, dal giorno «Dal» (vuoto = oggi).",
    "- CHIUDI: il ruolo finisce il giorno «Dal» (oggi o prima); resta nello storico.",
    "",
    "Mansioni persone",
    "- Azione vuota o ASSEGNA: in lista sulla mansione con il «Livello» (I in formazione, L intermedio, U autonomo, O esperto).",
    "- TOGLI: fuori lista (lo storico resta).",
    "- La mansione si riconosce dalla «Chiave mansione» o, se vuota, dal nome.",
    "- Classificazione: D1T Diretto 1 turno, D2T Diretto 2 turni, I1T Indiretto 1 turno, I2T Indiretto 2 turni.",
    "  Vuota = non si cambia (una mansione nuova prende la classificazione prevalente della persona).",
    "- Coefficiente: quanto pesa la mansione sul tempo della persona (es. 0,25). Le mansioni in lista di una",
    "  persona devono sommare 1: se scrivi dei coefficienti, la somma si controlla persona per persona e una",
    "  persona che non fa 1 si scarta (le altre righe restano valide). Vuoto = non si cambia.",
]


# ═══════════════════════════════════════════════════════════════════════════
# Export
# ═══════════════════════════════════════════════════════════════════════════

def _persone(reparto: str = "") -> list[dict]:
    from .ruoli_in_blocco import persone

    return persone(reparto=reparto)


def righe_export(*, reparto: str = "", vuoto: bool = False) -> dict[str, list[list]]:
    from django.db.models import Q

    from ..models import AbilitazioneMacchina, DipendenteAnagraficaAziendale, DipendenteRuoloOperativo

    if vuoto:
        return {FOGLIO_RUOLI: [], FOGLIO_MANSIONI: []}
    persone = {p["legacy_id"]: p for p in _persone(reparto)}
    ids = list(persone)
    aree = {a.legacy_anagrafica_id: a.area_aziendale.nome for a in DipendenteAnagraficaAziendale.objects.filter(
        legacy_anagrafica_id__in=ids, area_aziendale__isnull=False).select_related("area_aziendale")}
    oggi = timezone.localdate()
    ruoli = []
    for a in (DipendenteRuoloOperativo.objects.filter(Q(data_fine__isnull=True) | Q(data_fine__gte=oggi),
                                                      legacy_anagrafica_id__in=ids)
              .select_related("ruolo").order_by("ruolo__nome")):
        p = persone[a.legacy_anagrafica_id]
        ruoli.append([a.legacy_anagrafica_id, p["nome"], p["reparto"], aree.get(a.legacy_anagrafica_id, ""),
                      a.ruolo.nome, a.data_inizio, ""])
    mansioni = []
    for ab in (AbilitazioneMacchina.objects.filter(legacy_anagrafica_id__in=ids, in_lista=True,
                                                   competenza__isnull=False)
               .select_related("competenza").order_by("competenza__display")):
        p = persone[ab.legacy_anagrafica_id]
        mansioni.append([ab.legacy_anagrafica_id, p["nome"], p["reparto"], ab.competenza.competenza_key,
                         ab.competenza.nome, ab.livello, ab.classificazione,
                         float(ab.coefficiente) if ab.coefficiente is not None else None, ""])
    ruoli.sort(key=lambda r: (r[1].casefold(), r[4].casefold()))
    mansioni.sort(key=lambda r: (r[1].casefold(), r[4].casefold()))
    return {FOGLIO_RUOLI: ruoli, FOGLIO_MANSIONI: mansioni}


def esporta_xlsx(*, reparto: str = "", vuoto: bool = False) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    from core.excel_export import write_cell

    dati = righe_export(reparto=reparto, vuoto=vuoto)
    wb = Workbook()
    ws = wb.active
    ws.title = FOGLIO_ISTRUZIONI
    for i, riga in enumerate(ISTRUZIONI, 1):
        c = write_cell(ws, i, 1, riga)
        if i == 1 or riga in {"Come si usa", "Regole", "Ruoli persone", "Mansioni persone"}:
            c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 110
    testata = PatternFill("solid", fgColor="1E3A5F")
    for nome, colonne in ((FOGLIO_RUOLI, COL_RUOLI), (FOGLIO_MANSIONI, COL_MANSIONI)):
        ws = wb.create_sheet(nome)
        for j, col in enumerate(colonne, 1):
            c = write_cell(ws, 1, j, col)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = testata
            c.alignment = Alignment(vertical="center")
        for i, riga in enumerate(dati[nome], 2):
            for j, valore in enumerate(riga, 1):
                c = write_cell(ws, i, j, valore)
                if isinstance(valore, date):
                    c.number_format = "DD/MM/YYYY"
        for j, col in enumerate(colonne, 1):
            lungh = max([len(str(col))] + [len(str(r[j - 1] or "")) for r in dati[nome][:500]])
            ws.column_dimensions[get_column_letter(j)].width = min(max(lungh + 2, 10), 50)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(colonne))}{max(len(dati[nome]) + 1, 2)}"
        fine = MAX_RIGHE + 1
        lettera_azione = get_column_letter(colonne.index("Azione") + 1)
        scelte = f'"{AGGIUNGI},{CHIUDI}"' if nome == FOGLIO_RUOLI else f'"{ASSEGNA},{TOGLI}"'
        dv = DataValidation(type="list", formula1=scelte, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{lettera_azione}2:{lettera_azione}{fine}")
        if nome == FOGLIO_MANSIONI:
            lettera = get_column_letter(colonne.index("Livello") + 1)
            dv = DataValidation(type="list", formula1='"I,L,U,O"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{lettera}2:{lettera}{fine}")
            lettera = get_column_letter(colonne.index("Classificazione") + 1)
            dv = DataValidation(type="list", formula1='"D1T,D2T,I1T,I2T"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{lettera}2:{lettera}{fine}")
            lettera = get_column_letter(colonne.index("Coefficiente") + 1)
            for riga in range(2, len(dati[nome]) + 2):
                ws[f"{lettera}{riga}"].number_format = "0.0000"
    wb.active = 1
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════════════
# Import
# ═══════════════════════════════════════════════════════════════════════════

class _Annulla(Exception):
    """Fa tornare indietro la transazione della prova."""


def leggi_fogli(contenuto: bytes) -> dict[str, list[tuple[int, dict]]]:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        wb = load_workbook(BytesIO(contenuto), read_only=True, data_only=True)
    except (InvalidFileException, OSError, KeyError, ValueError) as exc:  # non xlsx, corrotto, protetto
        raise ValueError("Il file non è un Excel .xlsx leggibile.") from exc
    attesi = {FOGLIO_RUOLI: COL_RUOLI, FOGLIO_MANSIONI: COL_MANSIONI}
    per_nome = {ws.title.strip().casefold(): ws for ws in wb.worksheets}
    out: dict[str, list[tuple[int, dict]]] = {}
    for nome, colonne in attesi.items():
        ws = per_nome.get(nome.casefold())
        if ws is None:
            continue
        righe = ws.iter_rows(values_only=True)
        intestazione = next(righe, None) or ()
        noti = {c.casefold(): c for c in colonne}
        mappa = {i: noti[_testo(h).casefold()] for i, h in enumerate(intestazione) if _testo(h).casefold() in noti}
        if not mappa:
            continue
        dati = []
        for numero, valori in enumerate(righe, 2):
            if numero > MAX_RIGHE + 1:
                raise ValueError(f"Il foglio {nome} supera le {MAX_RIGHE} righe.")
            riga = {mappa[i]: v for i, v in enumerate(valori or ()) if i in mappa}
            if any(_testo(v) for v in riga.values()):
                dati.append((numero, riga))
        out[nome] = dati
    wb.close()
    if not out:
        raise ValueError(f"Nel file non ci sono i fogli «{FOGLIO_RUOLI}» o «{FOGLIO_MANSIONI}» con le loro intestazioni.")
    return out


def _data(valore, *, default: date) -> date:
    if valore in (None, ""):
        return default
    if isinstance(valore, datetime):
        return valore.date()
    if isinstance(valore, date):
        return valore
    t = _testo(valore)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            continue
    raise ErroreRiga(f"«Dal»: data non valida «{t}» (usa gg/mm/aaaa)")


class _Cataloghi:
    def __init__(self):
        from ..models import CompetenzaSkm, RuoloOperativo
        from .mansionario_struttura import TIPI_MANSIONE

        self.persone = {p["legacy_id"]: p for p in _persone()}
        per_nome: dict[str, list[int]] = {}
        for p in self.persone.values():
            per_nome.setdefault(p["nome"].casefold(), []).append(p["legacy_id"])
        self.per_nome = per_nome
        self.ruoli = {}
        for r in RuoloOperativo.objects.all().select_related("ambito"):
            self.ruoli.setdefault(r.nome.casefold(), r)
            if getattr(r, "sigla", None):
                self.ruoli.setdefault(str(r.sigla).casefold(), r)
        self.mansioni_chiave = {}
        self.mansioni_nome: dict[str, list] = {}
        for m in CompetenzaSkm.objects.filter(tipo__in=TIPI_MANSIONE):
            self.mansioni_chiave[m.competenza_key.casefold()] = m
            self.mansioni_nome.setdefault(m.nome.casefold(), []).append(m)

    def persona(self, riga: dict) -> dict:
        raw = _testo(riga.get("ID persona"))
        if raw:
            if not raw.isdigit() or int(raw) not in self.persone:
                raise ErroreRiga(f"«ID persona» {raw}: non è un dipendente in forza")
            return self.persone[int(raw)]
        nome = _testo(riga.get("Persona")).casefold()
        if not nome:
            raise ErroreRiga("manca la persona (ID persona o nome)")
        ids = self.per_nome.get(nome, [])
        if not ids:
            raise ErroreRiga(f"nessuna persona in forza si chiama «{_testo(riga.get('Persona'))}»")
        if len(ids) > 1:
            raise ErroreRiga(f"più persone si chiamano «{_testo(riga.get('Persona'))}»: scrivi l'ID persona")
        return self.persone[ids[0]]

    def ruolo(self, riga: dict):
        nome = _testo(riga.get("Ruolo"))
        if not nome:
            raise ErroreRiga("manca il ruolo")
        r = self.ruoli.get(nome.casefold())
        if r is None:
            raise ErroreRiga(f"il ruolo «{nome}» non esiste nel portale")
        return r

    def mansione(self, riga: dict):
        chiave = _testo(riga.get("Chiave mansione"))
        if chiave:
            m = self.mansioni_chiave.get(chiave.casefold())
            if m is None:
                raise ErroreRiga(f"nessuna mansione con chiave «{chiave}»")
            return m
        nome = _testo(riga.get("Mansione"))
        if not nome:
            raise ErroreRiga("manca la mansione (chiave o nome)")
        trovate = self.mansioni_nome.get(nome.casefold(), [])
        if not trovate:
            raise ErroreRiga(f"la mansione «{nome}» non esiste nel portale")
        if len(trovate) > 1:
            raise ErroreRiga(f"più mansioni si chiamano «{nome}»: scrivi la chiave")
        return trovate[0]


def _importa_ruoli(righe, cat: _Cataloghi, esito: Esito, user) -> None:
    from . import ruoli_in_blocco as rb

    oggi = timezone.localdate()
    for numero, riga in righe:
        chiave = f"{_testo(riga.get('Persona')) or _testo(riga.get('ID persona'))} · {_testo(riga.get('Ruolo'))}"
        try:
            p = cat.persona(riga)
            ruolo = cat.ruolo(riga)
            azione_xls = (_testo(riga.get("Azione")) or AGGIUNGI).upper()
            if azione_xls not in (AGGIUNGI, CHIUDI):
                raise ErroreRiga(f"«Azione»: scrivi {AGGIUNGI} o {CHIUDI}, non «{azione_xls}»")
            modo = rb.AGGIUNGI if azione_xls == AGGIUNGI else rb.CHIUDI
            data = _data(riga.get("Dal"), default=oggi)
            rb._controlla(ruolo, modo, data)
        except (ErroreRiga, ValueError) as exc:
            esito.aggiungi(FOGLIO_RUOLI, numero, chiave, ERRORE, str(exc))
            continue
        chiave = f"{p['nome']} · {ruolo.nome}"
        assegnati = rb.ruoli_in_essere([p["legacy_id"]]).get(p["legacy_id"], [])
        azione, dettaglio = rb._previsione(p, assegnati, ruolo, modo, data)
        if azione != "cambia":
            esito.aggiungi(FOGLIO_RUOLI, numero, chiave, INVARIATO, dettaglio)
            continue
        try:
            with transaction.atomic():
                dettaglio = rb._applica_una(p["legacy_id"], p, ruolo, modo, data, "Import Excel ruoli", user) or dettaglio
        except (ValueError, DatabaseError) as exc:
            esito.aggiungi(FOGLIO_RUOLI, numero, chiave, ERRORE, str(exc) or exc.__class__.__name__)
            continue
        esito.aggiungi(FOGLIO_RUOLI, numero, chiave, NUOVO if modo == rb.AGGIUNGI else MODIFICATO, dettaglio)


def _importa_mansioni(righe, cat: _Cataloghi, esito: Esito, user) -> None:
    from . import mansionario_struttura as mstr
    from . import mansioni_coefficienti as coeff
    from . import mansioni_in_blocco as mb
    from ..models import AbilitazioneMacchina

    coefficienti: dict[int, dict] = {}  # legacy_id → {"nome", "valori": {ab_id: Decimal}, "righe": [numeri]}
    for numero, riga in righe:
        chiave = f"{_testo(riga.get('Persona')) or _testo(riga.get('ID persona'))} · "                  f"{_testo(riga.get('Mansione')) or _testo(riga.get('Chiave mansione'))}"
        try:
            p = cat.persona(riga)
            m = cat.mansione(riga)
            azione_xls = (_testo(riga.get("Azione")) or ASSEGNA).upper()
            if azione_xls not in (ASSEGNA, TOGLI):
                raise ErroreRiga(f"«Azione»: scrivi {ASSEGNA} o {TOGLI}, non «{azione_xls}»")
            modo = mb.ASSEGNA if azione_xls == ASSEGNA else mb.TOGLI
            livello = _testo(riga.get("Livello")).upper()
            mb._controlla(m, modo, livello)
            classe = coeff.leggi_classificazione(riga.get("Classificazione")) if modo == mb.ASSEGNA else ""
            valore = coeff.leggi_coefficiente(riga.get("Coefficiente")) if modo == mb.ASSEGNA else None
        except (ErroreRiga, ValueError) as exc:
            esito.aggiungi(FOGLIO_MANSIONI, numero, chiave, ERRORE, str(exc))
            continue
        chiave = f"{p['nome']} · {m.nome}"
        ab = mb.in_lista(m, [p["legacy_id"]]).get(p["legacy_id"])
        azione, dettaglio = mb.previsione(ab, modo, livello, classe)
        ab_id = None
        if azione == "cambia":
            try:
                with transaction.atomic():
                    if modo == mb.ASSEGNA:
                        nuova = mstr.assegna(m, legacy_anagrafica_id=p["legacy_id"], livello=livello, user=user)
                        if classe and nuova.classificazione != classe:
                            AbilitazioneMacchina.objects.filter(pk=nuova.pk).update(classificazione=classe)
                        ab_id = nuova.pk
                    else:
                        mstr.togli(ab)
            except (ValueError, DatabaseError) as exc:
                esito.aggiungi(FOGLIO_MANSIONI, numero, chiave, ERRORE, str(exc) or exc.__class__.__name__)
                continue
        else:
            ab_id = ab.pk if ab is not None else None
        if valore is not None and ab_id is not None:
            dati = coefficienti.setdefault(p["legacy_id"], {"nome": p["nome"], "valori": {}, "righe": []})
            dati["valori"][ab_id] = valore
            dati["righe"].append((numero, chiave, azione, dettaglio, ab))
            continue
        if azione != "cambia":
            esito.aggiungi(FOGLIO_MANSIONI, numero, chiave, INVARIATO, dettaglio)
        else:
            esito.aggiungi(FOGLIO_MANSIONI, numero, chiave, NUOVO if ab is None else MODIFICATO, dettaglio)
    # Coefficienti: alla fine, persona per persona (la somma vale sull'insieme delle sue mansioni).
    for lid, dati in coefficienti.items():
        prima = {r.pk: r.coefficiente for r in coeff.conteggiate(lid)}
        errore = ""
        try:
            with transaction.atomic():
                coeff.imposta_coefficienti(lid, dati["valori"], user=user)
        except (ValueError, DatabaseError) as exc:
            errore = str(exc) or exc.__class__.__name__
        for numero, chiave, azione, dettaglio, ab in dati["righe"]:
            if errore:
                esito.aggiungi(FOGLIO_MANSIONI, numero, chiave, ERRORE,
                               f"coefficienti di {dati['nome']} non salvati: {errore}"
                               + (f" (la mansione è registrata lo stesso: {dettaglio})" if azione == "cambia" else ""))
                continue
            nuovo_coeff = dati["valori"].get(ab.pk) if ab is not None else None
            cambiato = ab is not None and prima.get(ab.pk) != nuovo_coeff
            if azione == "cambia":
                esito.aggiungi(FOGLIO_MANSIONI, numero, chiave, NUOVO if ab is None else MODIFICATO, dettaglio)
            elif cambiato:
                esito.aggiungi(FOGLIO_MANSIONI, numero, chiave, MODIFICATO,
                               f"coefficiente {prima.get(ab.pk) if prima.get(ab.pk) is not None else '—'} → "
                               f"{nuovo_coeff}")
            else:
                esito.aggiungi(FOGLIO_MANSIONI, numero, chiave, INVARIATO, dettaglio)


def importa(contenuto: bytes, *, applica: bool, user=None) -> Esito:
    """Legge il file e registra le assegnazioni; con ``applica=False`` annulla tutto alla fine."""
    fogli = leggi_fogli(contenuto)
    esito = Esito()
    try:
        with transaction.atomic():
            cat = _Cataloghi()
            if FOGLIO_RUOLI in fogli:
                _importa_ruoli(fogli[FOGLIO_RUOLI], cat, esito, user)
            if FOGLIO_MANSIONI in fogli:
                _importa_mansioni(fogli[FOGLIO_MANSIONI], cat, esito, user)
            if not applica:
                raise _Annulla
    except _Annulla:
        pass
    esito.applicato = applica
    return esito
