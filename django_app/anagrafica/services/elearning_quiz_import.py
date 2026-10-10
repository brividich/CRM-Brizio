"""Import delle domande del quiz da un foglio Excel (prompt 05, fase 2).

Formato (prima riga = intestazioni, l'ordine delle colonne è libero):

- ``Tipo``: singola, multipla, vero/falso (anche S, M, VF). Vuoto = singola,
  o multipla se le corrette sono più di una.
- ``Domanda``: il testo.
- ``Risposta 1`` … ``Risposta 8``: le opzioni (ignorate per vero/falso).
- ``Corrette``: i numeri delle risposte giuste separati da virgola (``1`` o
  ``1,3``); per vero/falso ``V``/``Vero`` oppure ``F``/``Falso``.

L'import è tutto o niente: se una riga non va, non si salva nulla e si dice quale
riga e perché. Le domande si accodano a quelle esistenti.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from django.db import transaction

MAX_RIGHE = 500
MAX_RISPOSTE = 8

_TIPI = {
    "": "", "singola": "SINGOLA", "s": "SINGOLA", "risposta singola": "SINGOLA",
    "multipla": "MULTIPLA", "m": "MULTIPLA", "risposta multipla": "MULTIPLA",
    "vero/falso": "VERO_FALSO", "vero falso": "VERO_FALSO", "vf": "VERO_FALSO", "v/f": "VERO_FALSO",
}


class ImportQuizError(Exception):
    def __init__(self, errori: list[str]):
        super().__init__("; ".join(errori))
        self.errori = errori


@dataclass
class _Riga:
    numero: int
    tipo: str
    testo: str
    risposte: list[str] = field(default_factory=list)
    corrette: set[int] = field(default_factory=set)  # indici 1-based; per VF {1}=vero {2}=falso


def _testo(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _intestazioni(riga) -> dict[str, int]:
    out = {}
    for i, v in enumerate(riga):
        nome = re.sub(r"\s+", " ", _testo(v).lower())
        if nome:
            out[nome] = i
    return out


def leggi(file_obj) -> list[_Riga]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(file_obj.read()), read_only=True, data_only=True)
    except Exception as exc:  # file corrotto o non xlsx
        raise ImportQuizError(["Il file non è un foglio Excel leggibile (.xlsx)."]) from exc
    ws = wb.worksheets[0]
    righe = ws.iter_rows(values_only=True)
    try:
        col = _intestazioni(next(righe))
    except StopIteration:
        raise ImportQuizError(["Il foglio è vuoto."])
    if "domanda" not in col:
        raise ImportQuizError(["Manca la colonna «Domanda» nella prima riga."])
    if "corrette" not in col:
        raise ImportQuizError(["Manca la colonna «Corrette» nella prima riga."])
    col_risposte = [col[f"risposta {n}"] for n in range(1, MAX_RISPOSTE + 1) if f"risposta {n}" in col]

    errori, out = [], []
    for numero, valori in enumerate(righe, start=2):
        valori = list(valori) + [None] * 64
        testo = _testo(valori[col["domanda"]])
        risposte = [_testo(valori[i]) for i in col_risposte]
        while risposte and not risposte[-1]:
            risposte.pop()
        corrette_txt = _testo(valori[col["corrette"]])
        tipo_txt = _testo(valori[col["tipo"]]).lower() if "tipo" in col else ""
        if not testo and not any(risposte) and not corrette_txt:
            continue  # riga vuota
        if len(out) >= MAX_RIGHE:
            errori.append(f"Troppe domande: massimo {MAX_RIGHE} per import.")
            break
        tipo = _TIPI.get(tipo_txt)
        if tipo is None:
            errori.append(f"Riga {numero}: tipo «{tipo_txt}» non riconosciuto (singola, multipla, vero/falso).")
            continue
        if not testo:
            errori.append(f"Riga {numero}: manca il testo della domanda.")
            continue
        riga = _Riga(numero, tipo, testo[:4000])
        if tipo == "VERO_FALSO":
            v = corrette_txt.lower()
            if v in ("v", "vero", "vera", "1"):
                riga.corrette = {1}
            elif v in ("f", "falso", "falsa", "0"):
                riga.corrette = {2}
            else:
                errori.append(f"Riga {numero}: per vero/falso «Corrette» vale V oppure F.")
                continue
            out.append(riga)
            continue
        if any(not r for r in risposte):
            errori.append(f"Riga {numero}: c'è una risposta vuota in mezzo alle altre.")
            continue
        if len(risposte) < 2:
            errori.append(f"Riga {numero}: servono almeno due risposte.")
            continue
        try:
            corrette = {int(x) for x in re.split(r"[,;\s]+", corrette_txt) if x}
        except ValueError:
            errori.append(f"Riga {numero}: «Corrette» deve contenere numeri di risposta (es. 1 oppure 1,3).")
            continue
        if not corrette or any(n < 1 or n > len(risposte) for n in corrette):
            errori.append(f"Riga {numero}: le corrette devono essere fra 1 e {len(risposte)}.")
            continue
        if not tipo:
            tipo = "MULTIPLA" if len(corrette) > 1 else "SINGOLA"
        if tipo == "SINGOLA" and len(corrette) != 1:
            errori.append(f"Riga {numero}: risposta singola con più corrette (usa «multipla»).")
            continue
        riga.tipo, riga.risposte, riga.corrette = tipo, [r[:500] for r in risposte], corrette
        out.append(riga)
    if errori:
        raise ImportQuizError(errori)
    if not out:
        raise ImportQuizError(["Nessuna domanda trovata nel foglio."])
    return out


def importa(corso, file_obj) -> int:
    """Crea le domande lette dal foglio, in coda alle esistenti. Ritorna quante."""
    from ..models_formazione import TrainingQuizOption, TrainingQuizQuestion
    from .elearning_quiz import FALSO, VERO

    righe = leggi(file_obj)
    with transaction.atomic():
        ultimo = corso.quiz_domande.order_by("-ordine").values_list("ordine", flat=True).first() or 0
        for n, riga in enumerate(righe, start=1):
            domanda = TrainingQuizQuestion.objects.create(
                corso=corso, ordine=ultimo + n, testo=riga.testo, tipo=riga.tipo)
            risposte = [VERO, FALSO] if riga.tipo == "VERO_FALSO" else riga.risposte
            TrainingQuizOption.objects.bulk_create([
                TrainingQuizOption(domanda=domanda, ordine=i, testo=testo, corretta=i in riga.corrette)
                for i, testo in enumerate(risposte, start=1)
            ])
    return len(righe)


def modello_xlsx() -> bytes:
    """Foglio di esempio da compilare (dati sintetici)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Domande"
    ws.append(["Tipo", "Domanda", "Risposta 1", "Risposta 2", "Risposta 3", "Risposta 4", "Corrette"])
    ws.append(["singola", "Chi deve indossare i DPI previsti per la mansione?", "Il lavoratore", "Solo il preposto",
               "Nessuno", "", "1"])
    ws.append(["multipla", "Quali sono misure di protezione collettiva?", "Parapetti", "Guanti", "Aspirazione localizzata",
               "", "1,3"])
    ws.append(["vero/falso", "Un'uscita di emergenza può restare chiusa a chiave durante il turno.", "", "", "", "", "F"])
    for colonna, larghezza in zip("ABCDEFG", (12, 60, 24, 24, 24, 24, 10)):
        ws.column_dimensions[colonna].width = larghezza
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
