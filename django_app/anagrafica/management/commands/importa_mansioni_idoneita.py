"""Assegna la mansione ai dipendenti dall'elenco idoneità del medico competente.

Sorgente: il foglio Excel con **Codice fiscale** e **MANSIONE DI RISCHIO** (es.
«ELENCO CONFRONTO PDF 2025» di ``elenco_idoneita_da_controllare.xlsx``). La
mansione del file è la mansione base del catalogo (``anagrafica.Mansione``),
talvolta con rischi aggiunti in coda: «ADDETTO MACCHINE UTENSILI, CARRELLO E
LAVORI IN ALTEZZA» → mansione «ADDETTO MACCHINE UTENSILI» + rischi aggiuntivi
«CARRELLO», «LAVORI IN ALTEZZA». I rischi aggiuntivi **non** vengono scritti:
sono riportati a parte, perché il loro posto sono le mansioni di rischio, non
la mansione lavorativa.

Scrittura come **spostamento organizzativo** (``DipendenteAssegnazione`` +
``attiva_assegnazione``): campo vivo, storico per-campo, notifica SDS in-app e
ricalcolo scadenze passano dallo stesso percorso della UI. Il piano di
adeguamento al cambio mansione **non** si genera di default: questo è un
caricamento iniziale, non un cambio reale (``--con-piano`` per generarlo).

Di default tocca solo chi ha la mansione **vuota**; ``--sovrascrivi`` corregge
anche chi ne ha una diversa. DRY-RUN di default, ``--commit`` per scrivere.

Privacy: il file contiene dati sanitari (art. 9 GDPR). Si leggono solo codice
fiscale e mansione; l'output e il report CSV non contengono nominativi né CF
(solo id anagrafica, esito e mansione).
"""
from __future__ import annotations

import csv
import re
from collections import Counter
from datetime import date

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone

FOGLIO_DEFAULT = "ELENCO CONFRONTO PDF 2025"

# Varianti di scrittura nel file → forma del catalogo.
ALIAS_PAROLE = {
    "MANUTENZIONI": "MANUTENZIONE",
}
# Code note: rischi che si sommano alla mansione base.
RISCHI_AGGIUNTIVI = {
    "CARRELLO": "CARRELLO",
    "LAVORI IN ALTEZZA": "LAVORI IN ALTEZZA",
    "SOSPESO CARRELLO": "CARRELLO (SOSPESO)",
}
# Suffissi di variante della stessa mansione (MAGAZZINO A / B).
VARIANTI = {"A", "B"}

ESITI = (
    "ASSEGNATA", "SOVRASCRITTA", "GIA_UGUALE", "DIVERSA_NON_TOCCATA",
    "NON_MAPPATA", "CF_NON_TROVATO", "CESSATO", "CF_MANCANTE", "CF_DUPLICATO_INCOERENTE",
)


def normalizza(testo: str) -> str:
    """Maiuscolo, «ADD.» → «ADDETTO», punteggiatura come separatore, spazi singoli."""
    s = str(testo or "").upper()
    s = re.sub(r"\bADD\s*\.\s*", "ADDETTO ", s)
    s = re.sub(r"[(),;/]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return " ".join(ALIAS_PAROLE.get(p, p) for p in s.split(" ")) if s else ""


def scomponi(valore: str, catalogo: dict[str, str]) -> tuple[str | None, list[str], list[str]]:
    """Mansione del file → (nome a catalogo, rischi aggiuntivi, residui non riconosciuti).

    ``catalogo`` mappa il nome normalizzato al nome esatto della ``Mansione``.
    Vince il nome di catalogo più lungo che apre la stringa.
    """
    s = normalizza(valore)
    base = next(
        (n for n in sorted(catalogo, key=len, reverse=True) if s == n or s.startswith(n + " ")),
        None,
    )
    if base is None:
        return None, [], [s] if s else []
    coda = s[len(base):].strip()
    rischi: list[str] = []
    residui: list[str] = []
    # «… E ADDETTO RITOCCHI», «… E RITOCCHI»: seconda mansione, riportata come rischio.
    for pezzo in re.split(r"\bE\b", coda):
        pezzo = pezzo.strip()
        if not pezzo or pezzo in VARIANTI:
            continue
        for chiave, etichetta in sorted(RISCHI_AGGIUNTIVI.items(), key=lambda kv: -len(kv[0])):
            if chiave in pezzo:
                rischi.append(etichetta)
                pezzo = pezzo.replace(chiave, " ").strip()
        if not pezzo:
            continue
        seconda = pezzo if pezzo.startswith("ADDETTO ") else f"ADDETTO {pezzo}"
        if seconda in catalogo:
            rischi.append(f"seconda mansione: {catalogo[seconda]}")
        else:
            residui.append(pezzo)
    return catalogo[base], rischi, residui


class Command(BaseCommand):
    help = "Assegna la mansione ai dipendenti dall'elenco idoneità (CF + MANSIONE DI RISCHIO). DRY-RUN di default."

    def add_arguments(self, parser):
        parser.add_argument("xlsx", help="Percorso del file Excel.")
        parser.add_argument("--foglio", default=FOGLIO_DEFAULT, help=f"Foglio (default «{FOGLIO_DEFAULT}»).")
        parser.add_argument("--dal", help="Decorrenza dell'assegnazione, AAAA-MM-GG (default: oggi).")
        parser.add_argument("--sovrascrivi", action="store_true",
                            help="Corregge anche chi ha già una mansione diversa (default: solo mansione vuota).")
        parser.add_argument("--con-piano", action="store_true",
                            help="Genera anche il piano di adeguamento al cambio mansione.")
        parser.add_argument("--report", help="CSV di esito per riga (id anagrafica, esito, mansioni; niente nominativi/CF).")
        parser.add_argument("--commit", action="store_true", help="Scrive davvero (senza: dry-run).")

    # ── lettura ──────────────────────────────────────────────────────────────
    def _righe(self, path: str, foglio: str) -> list[tuple[int, str, str]]:
        import openpyxl

        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise CommandError(f"File non leggibile: {exc}") from exc
        try:
            if foglio not in wb.sheetnames:
                raise CommandError(f"Foglio «{foglio}» assente. Fogli: {', '.join(wb.sheetnames)}")
            # In read_only il file resta aperto finché non si chiude il workbook:
            # su Windows lo bloccherebbe (anche sulla share) per tutto il comando.
            righe = iter(list(wb[foglio].iter_rows(values_only=True)))
        finally:
            wb.close()
        i_cf = i_mans = None
        n = 0
        for n, row in enumerate(righe, start=1):
            intest = [normalizza(v) for v in (row or ())]
            if "CODICE FISCALE" in intest and "MANSIONE DI RISCHIO" in intest:
                i_cf, i_mans = intest.index("CODICE FISCALE"), intest.index("MANSIONE DI RISCHIO")
                break
        if i_cf is None:
            raise CommandError("Intestazioni «Codice fiscale» e «MANSIONE DI RISCHIO» non trovate.")
        out = []
        for riga, row in enumerate(righe, start=n + 1):
            row = row or ()
            cf = str(row[i_cf] or "").strip().upper() if i_cf < len(row) else ""
            mans = str(row[i_mans] or "").strip() if i_mans < len(row) else ""
            if cf or mans:
                out.append((riga, cf, mans))
        return out

    # ── comando ──────────────────────────────────────────────────────────────
    def handle(self, *args, **opts):
        from anagrafica.models import (
            DipendenteAnagraficaAziendale, DipendenteAnagraficaCivile, DipendenteAssegnazione, Mansione,
        )

        dal = timezone.localdate()
        if opts.get("dal"):
            try:
                dal = date.fromisoformat(opts["dal"])
            except ValueError as exc:
                raise CommandError("--dal deve essere AAAA-MM-GG.") from exc

        righe = self._righe(opts["xlsx"], opts["foglio"])
        catalogo = {normalizza(n): n for n in Mansione.objects.filter(is_active=True).values_list("nome", flat=True)}
        cf_a_id = {
            (cf or "").strip().upper(): lid
            for cf, lid in DipendenteAnagraficaCivile.objects.exclude(codice_fiscale="")
            .values_list("codice_fiscale", "legacy_anagrafica_id")
        }
        cessati = set(
            DipendenteAnagraficaAziendale.objects.filter(data_cessazione__isnull=False)
            .values_list("legacy_anagrafica_id", flat=True)
        )

        # Un CF ripetuto con mansioni diverse non si decide da qui.
        per_cf: dict[str, set[str]] = {}
        for _, cf, mans in righe:
            if cf:
                per_cf.setdefault(cf, set()).add(normalizza(mans))

        from core.legacy_anagrafica import fetch_anagrafica_rows

        esiti: Counter = Counter()
        per_mansione: Counter = Counter()
        rischi_tot: Counter = Counter()
        non_mappate: Counter = Counter()
        report: list[dict] = []
        da_scrivere: list[tuple[int, str, str, list[str]]] = []
        visti: set[int] = set()

        for riga, cf, mans in righe:
            base, rischi, residui = scomponi(mans, catalogo)
            lid = cf_a_id.get(cf)
            if not cf:
                esito = "CF_MANCANTE"
            elif len(per_cf.get(cf, ())) > 1:
                esito = "CF_DUPLICATO_INCOERENTE"
            elif lid is None:
                esito = "CF_NON_TROVATO"
            elif lid in cessati:
                esito = "CESSATO"
            elif base is None:
                esito = "NON_MAPPATA"
                non_mappate[normalizza(mans) or "(vuota)"] += 1
            elif lid in visti:
                continue  # stessa persona, stessa mansione: già contata
            else:
                visti.add(lid)
                attuale = ((fetch_anagrafica_rows(ids=[lid]) or [{}])[0].get("mansione") or "").strip()
                programmata = DipendenteAssegnazione.objects.filter(
                    legacy_anagrafica_id=lid, attivata_il__isnull=True, mansione__iexact=base,
                ).exists()
                if attuale.casefold() == base.casefold() or programmata:
                    esito = "GIA_UGUALE"
                elif not attuale:
                    esito = "ASSEGNATA"
                elif opts["sovrascrivi"]:
                    esito = "SOVRASCRITTA"
                else:
                    esito = "DIVERSA_NON_TOCCATA"
                if esito in ("ASSEGNATA", "SOVRASCRITTA"):
                    da_scrivere.append((lid, base, mans, rischi))
                    per_mansione[base] += 1
                for r in rischi:
                    rischi_tot[r] += 1
                if residui:
                    non_mappate["(coda) " + " / ".join(residui)] += 1
            esiti[esito] += 1
            report.append({
                "riga_file": riga, "legacy_anagrafica_id": lid or "", "esito": esito,
                "mansione_file": normalizza(mans), "mansione_assegnata": base or "",
                "rischi_aggiuntivi": "; ".join(rischi),
            })

        scritte = errori = 0
        if opts["commit"]:
            for lid, base, mans, rischi in da_scrivere:
                try:
                    self._assegna(lid, base, mans, rischi, dal, con_piano=opts["con_piano"])
                    scritte += 1
                except Exception as exc:  # una persona non ferma le altre
                    errori += 1
                    self.stderr.write(f"  id {lid}: errore {type(exc).__name__}: {exc}")

        if opts.get("report"):
            with open(opts["report"], "w", newline="", encoding="utf-8-sig") as fh:
                w = csv.DictWriter(fh, fieldnames=list(report[0].keys()) if report else ["esito"], delimiter=";")
                w.writeheader()
                w.writerows(report)

        modo = "COMMIT" if opts["commit"] else "DRY-RUN (nessuna scrittura)"
        self.stdout.write(f"\n{modo} — righe lette: {len(righe)} — decorrenza {dal:%d/%m/%Y}")
        self.stdout.write("Esiti:")
        for e in ESITI:
            if esiti[e]:
                self.stdout.write(f"  {e:<26} {esiti[e]:>4}")
        self.stdout.write("Mansioni da assegnare:")
        for nome, n in per_mansione.most_common():
            self.stdout.write(f"  {n:>4}  {nome}")
        if rischi_tot:
            self.stdout.write("Rischi aggiuntivi trovati (NON scritti, vanno nelle mansioni di rischio):")
            for nome, n in rischi_tot.most_common():
                self.stdout.write(f"  {n:>4}  {nome}")
        if non_mappate:
            self.stdout.write("Valori non riconosciuti:")
            for nome, n in non_mappate.most_common():
                self.stdout.write(f"  {n:>4}  {nome}")
        if opts["commit"]:
            self.stdout.write(f"Scritte: {scritte} — errori: {errori}")

    # ── scrittura ────────────────────────────────────────────────────────────
    @transaction.atomic
    def _assegna(self, lid: int, base: str, mans: str, rischi: list[str], dal, *, con_piano: bool):
        from anagrafica.models import DipendenteAssegnazione
        from anagrafica.services import assegnazioni as svc

        note = f"Mansione da elenco idoneità del medico competente («{normalizza(mans)}»)"
        if rischi:
            note += f"; rischi aggiuntivi: {', '.join(rischi)}"
        if con_piano:
            svc.crea_assegnazione(lid, data_inizio=dal, mansione=base, note=note)
            return
        corrente = svc.assetto_corrente(lid)
        idoneita = svc.verifica_idoneita(lid, base)
        DipendenteAssegnazione.chiudi_aperta(lid, dal)
        assegnazione = DipendenteAssegnazione.objects.create(
            legacy_anagrafica_id=lid,
            data_inizio=dal,
            reparto=corrente["reparto"],
            area_aziendale_id=corrente["area_aziendale_id"],
            mansione=base,
            ruolo_aziendale=corrente["ruolo_aziendale"],
            note=note,
            idoneita_esito=idoneita["esito"],
            idoneita_mancanti=idoneita["mancanti"],
            idoneita_scaduti=idoneita["scaduti"],
            idoneita_verificata_il=timezone.now(),
        )
        if dal <= timezone.localdate():
            svc.attiva_assegnazione(assegnazione)
