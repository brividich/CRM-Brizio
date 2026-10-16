"""Test dei fix dell'analisi modulo HR (16/10/2026). Dati sintetici.

- deduplica anagrafiche: omonimi con CF o matricola diversi restano due persone;
- richiamo indicato dal medico: sopravvive a save() e al ricalcolo notturno;
- export scadenze visite: prescrizioni solo a chi apre i referti;
- retention: report in sola lettura, un attivo non perde mai documenti;
- ticket: il fallback del reparto legge l'anagrafica legacy.
"""
from datetime import date, timedelta
from io import BytesIO, StringIO
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from anagrafica.models import (
    DipendenteAnagraficaAziendale,
    DipendenteAnagraficaCivile,
    DocumentoDipendente,
    NOTA_SCADENZA_MEDICO,
    TipoVisitaMedica,
    VisitaMedica,
    _add_months,
)
from anagrafica.tests import _ensure_anagrafica_table
from core.legacy_anagrafica import merge_duplicate_anagrafica_rows
from core.legacy_models import AnagraficaDipendente

User = get_user_model()


class DeduplicaGuardiaCodiciTests(TestCase):
    def _row(self, pk, **extra):
        return {"id": pk, "nome": "Mario", "cognome": "Rossi", **extra}

    def test_omonimi_senza_codici_si_fondono_come_prima(self):
        rows = merge_duplicate_anagrafica_rows([self._row(1), self._row(2)])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["_merged_ids"], [1, 2])

    def test_omonimi_con_cf_diverso_restano_due_persone(self):
        DipendenteAnagraficaCivile.objects.create(legacy_anagrafica_id=1, codice_fiscale="RSSMRA80A01H501U")
        DipendenteAnagraficaCivile.objects.create(legacy_anagrafica_id=2, codice_fiscale="RSSMRA92C15F205X")
        rows = merge_duplicate_anagrafica_rows([self._row(1), self._row(2)])
        self.assertEqual(sorted(r["id"] for r in rows), [1, 2])

    def test_stesso_cf_si_fonde(self):
        DipendenteAnagraficaCivile.objects.create(legacy_anagrafica_id=1, codice_fiscale="RSSMRA80A01H501U")
        DipendenteAnagraficaCivile.objects.create(legacy_anagrafica_id=2, codice_fiscale="rssmra80a01h501u")
        rows = merge_duplicate_anagrafica_rows([self._row(1), self._row(2)])
        self.assertEqual(len(rows), 1)

    def test_matricola_diversa_separa_anche_senza_cf(self):
        rows = merge_duplicate_anagrafica_rows([self._row(1, matricola="100"), self._row(2, matricola="200")])
        self.assertEqual(len(rows), 2)

    def test_riga_senza_codici_si_aggancia_al_primo_compatibile(self):
        rows = merge_duplicate_anagrafica_rows([
            self._row(1, matricola="100"), self._row(2, matricola="200"), self._row(3),
        ])
        self.assertEqual(len(rows), 2)
        fusa = next(r for r in rows if 3 in r["_merged_ids"])
        self.assertIn(1, fusa["_merged_ids"])


class ScadenzaMedicoTests(TestCase):
    @classmethod
    def setUpClass(cls):
        _ensure_anagrafica_table()
        super().setUpClass()

    def setUp(self):
        self.dip = AnagraficaDipendente.objects.create(nome="Anna", cognome="Verdi", reparto="Produzione")
        self.tipo = TipoVisitaMedica.objects.create(nome="Periodica sintetica", durata_mesi=12, obbligatoria=True)
        self.svolta = timezone.localdate() - timedelta(days=30)

    def test_save_usa_il_richiamo_del_medico(self):
        richiamo = self.svolta + timedelta(days=180)
        v = VisitaMedica.objects.create(legacy_anagrafica_id=self.dip.id, tipo=self.tipo,
                                        data_svolgimento=self.svolta, scadenza_medico=richiamo)
        v.refresh_from_db()
        self.assertEqual(v.data_scadenza, richiamo)
        self.assertEqual(v.scadenza_nota, NOTA_SCADENZA_MEDICO)

    def test_senza_richiamo_resta_la_durata_del_tipo(self):
        v = VisitaMedica.objects.create(legacy_anagrafica_id=self.dip.id, tipo=self.tipo,
                                        data_svolgimento=self.svolta)
        v.refresh_from_db()
        self.assertEqual(v.data_scadenza, _add_months(self.svolta, 12))

    def test_ricalcolo_notturno_rispetta_il_richiamo(self):
        from anagrafica.services.scadenze import ricalcola_visite

        richiamo = self.svolta + timedelta(days=500)
        v = VisitaMedica.objects.create(legacy_anagrafica_id=self.dip.id, tipo=self.tipo,
                                        data_svolgimento=self.svolta, scadenza_medico=richiamo)
        ricalcola_visite()
        v.refresh_from_db()
        self.assertEqual(v.data_scadenza, richiamo)

    def test_form_rifiuta_richiamo_prima_della_visita(self):
        from anagrafica.forms import VisitaMedicaForm

        form = VisitaMedicaForm(data={
            "tipo": self.tipo.pk, "data_svolgimento": self.svolta.isoformat(), "esito": "IDONEO",
            "scadenza_medico": (self.svolta - timedelta(days=1)).isoformat(),
        })
        self.assertFalse(form.is_valid())
        self.assertIn("scadenza_medico", form.errors)


class ExportScadenzePrescrizioniTests(TestCase):
    @classmethod
    def setUpClass(cls):
        _ensure_anagrafica_table()
        super().setUpClass()

    def setUp(self):
        self.admin = User.objects.create_superuser("admin_prescr", "admin-prescr@example.invalid", "x")
        dip = AnagraficaDipendente.objects.create(nome="Luca", cognome="Neri", reparto="Ufficio")
        tipo = TipoVisitaMedica.objects.create(nome="Periodica export", durata_mesi=12, obbligatoria=True)
        VisitaMedica.objects.create(legacy_anagrafica_id=dip.id, tipo=tipo,
                                    data_svolgimento=timezone.localdate() - timedelta(days=340),
                                    prescrizioni="PRESCRIZIONE-SINTETICA")

    def _celle(self):
        from openpyxl import load_workbook

        self.client.force_login(self.admin)
        resp = self.client.get(reverse("anagrafica:visite_mediche_export_scadenze"))
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(BytesIO(resp.content))
        return [c.value for row in wb.active.iter_rows() for c in row]

    def test_con_permesso_referti_include_le_prescrizioni(self):
        celle = self._celle()
        self.assertIn("Prescrizioni", celle)
        self.assertIn("PRESCRIZIONE-SINTETICA", celle)

    def test_senza_permesso_referti_le_prescrizioni_non_escono(self):
        with mock.patch("anagrafica.views._can_view_referti", return_value=False):
            celle = self._celle()
        self.assertNotIn("Prescrizioni", celle)
        self.assertNotIn("PRESCRIZIONE-SINTETICA", celle)


class RetentionReportTests(TestCase):
    def _doc(self, legacy_id, tipo="ALTRO"):
        doc = DocumentoDipendente.objects.create(legacy_anagrafica_id=legacy_id, tipo=tipo, nome_originale="x.pdf")
        DocumentoDipendente.objects.filter(pk=doc.pk).update(retention_until=date(2000, 1, 1))
        return doc

    def test_attivo_non_perde_mai_documenti(self):
        from anagrafica.services.retention import documenti_scaduti

        self._doc(10)
        esito = documenti_scaduti()
        self.assertEqual(esito.da_eliminare, [])
        self.assertEqual(esito.saltati_attivi, 1)

    def test_cessato_da_poco_resta_cessato_da_tanto_e_candidato(self):
        from anagrafica.services.retention import documenti_scaduti

        DipendenteAnagraficaAziendale.objects.create(legacy_anagrafica_id=11, data_cessazione=date(1980, 1, 1))
        DipendenteAnagraficaAziendale.objects.create(
            legacy_anagrafica_id=12, data_cessazione=timezone.localdate() - timedelta(days=30))
        vecchio = self._doc(11)
        self._doc(12)
        esito = documenti_scaduti()
        self.assertEqual([d.pk for d, _c in esito.da_eliminare], [vecchio.pk])
        self.assertEqual(esito.saltati_retention, 1)

    def test_report_non_cancella_e_apre_issue(self):
        from monitoring.models import Issue

        from anagrafica.tasks import run_report_retention

        DipendenteAnagraficaAziendale.objects.create(legacy_anagrafica_id=11, data_cessazione=date(1980, 1, 1))
        self._doc(11)
        esito = run_report_retention()
        self.assertEqual(esito["documenti_da_eliminare"], 1)
        self.assertEqual(DocumentoDipendente.objects.count(), 1)
        self.assertTrue(Issue.objects.filter(title__startswith="Retention HR").exists())


class TicketRepartoFallbackTests(TestCase):
    @classmethod
    def setUpClass(cls):
        _ensure_anagrafica_table()
        super().setUpClass()

    def test_reparto_da_anagrafica_legacy(self):
        from tickets.views import _get_user_reparto

        dip = AnagraficaDipendente.objects.create(nome="Ugo", cognome="Bassi", reparto="Officina")
        AnagraficaDipendente.objects.filter(pk=dip.pk).update(utente_id=None)
        with mock.patch("core.legacy_models.AnagraficaDipendente.objects") as manager:
            manager.filter.return_value.only.return_value.first.return_value = dip
            self.assertEqual(_get_user_reparto(999), "Officina")
            manager.filter.assert_called_once_with(utente_id=999)


class ScadenzarioPeriodoTests(TestCase):
    def _periodo(self, **get):
        from django.test import RequestFactory

        from anagrafica.views import _scadenzario_periodo

        return _scadenzario_periodo(RequestFactory().get("/", get), date(2026, 10, 16))

    def test_questo_mese(self):
        self.assertEqual(self._periodo(periodo="mese")[3], (date(2026, 10, 1), date(2026, 10, 31)))

    def test_prossimo_mese_a_cavallo_anno(self):
        from django.test import RequestFactory

        from anagrafica.views import _scadenzario_periodo

        req = RequestFactory().get("/", {"periodo": "prossimo"})
        self.assertEqual(_scadenzario_periodo(req, date(2026, 12, 5))[3], (date(2027, 1, 1), date(2027, 1, 31)))

    def test_intervallo_libero_e_date_invalide(self):
        self.assertEqual(self._periodo(periodo="intervallo", dal="2026-11-01", al="2026-11-15")[3],
                         (date(2026, 11, 1), date(2026, 11, 15)))
        self.assertIsNone(self._periodo(periodo="intervallo", dal="2026-02-30")[3])
        self.assertIsNone(self._periodo(periodo="intervallo", dal="2026-12-01", al="2026-11-01")[3])

    def test_senza_periodo_vale_lo_stato(self):
        self.assertEqual(self._periodo(), ("", None, None, None))


class MioTeamTests(TestCase):
    @classmethod
    def setUpClass(cls):
        _ensure_anagrafica_table()
        super().setUpClass()

    def setUp(self):
        from anagrafica.models import AreaAziendale, Reparto

        self.capo = AnagraficaDipendente.objects.create(nome="Carla", cognome="Capo", reparto="Officina")
        self.membro = AnagraficaDipendente.objects.create(nome="Mino", cognome="Membro", reparto="Officina")
        self.cessato = AnagraficaDipendente.objects.create(nome="Ezio", cognome="Ex", reparto="Officina")
        self.altro = AnagraficaDipendente.objects.create(nome="Olga", cognome="Altro", reparto="Ufficio")
        rep = Reparto.objects.create(nome="Officina", caporeparto_legacy_id=self.capo.id)
        area = AreaAziendale.objects.create(nome="OF1", reparto=rep)
        for dip in (self.capo, self.membro, self.cessato):
            DipendenteAnagraficaAziendale.objects.create(legacy_anagrafica_id=dip.id, area_aziendale=area)
        DipendenteAnagraficaAziendale.objects.filter(legacy_anagrafica_id=self.cessato.id).update(
            data_cessazione=date(2026, 1, 31))

    def _utente(self, username):
        from core.models import UserOnboarding

        user = User.objects.create_user(username, password="x")
        # Onboarding completato: evita il redirect del middleware.
        UserOnboarding.objects.update_or_create(user=user, defaults={"completed": True, "completed_at": timezone.now()})
        return user

    def test_team_del_capo_senza_se_stesso_ne_cessati(self):
        from anagrafica.views_mio_team import membri_team

        self.assertEqual(membri_team(self.capo.id), [self.membro.id])
        self.assertEqual(membri_team(self.altro.id), [])

    def test_pagina_senza_collegamento_mostra_stato_vuoto(self):
        self.client.force_login(self._utente("capo_team_test"))
        with mock.patch("anagrafica.views._current_legacy_anagrafica_id", return_value=None):
            resp = self.client.get(reverse("anagrafica:mio_team"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Account non collegato")

    def test_pagina_del_capo_elenca_il_team_senza_dati_sanitari(self):
        self.client.force_login(self._utente("capo_team_test2"))
        with mock.patch("anagrafica.views._current_legacy_anagrafica_id", return_value=self.capo.id):
            resp = self.client.get(reverse("anagrafica:mio_team"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "MEMBRO")
        self.assertNotContains(resp, "EX EZIO")
        self.assertNotContains(resp, "prescrizion")


class ImportVisiteRichiamoTests(TestCase):
    """L'import Excel scrive la data richiamo del medico nel campo dedicato (sopravvive al ricalcolo)."""

    @classmethod
    def setUpClass(cls):
        _ensure_anagrafica_table()
        super().setUpClass()

    def _xlsx(self, tmpdir, righe):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(["Codice fiscale", "Visita", "Data visita", "Data di richiamo", "Esito", "Note"])
        for r in righe:
            ws.append(r)
        path = f"{tmpdir}/visite.xlsx"
        wb.save(path)
        return path

    def test_creazione_e_overwrite_conservano_il_richiamo(self):
        import tempfile

        from django.core.management import call_command

        from anagrafica.services.scadenze import ricalcola_visite

        dip = AnagraficaDipendente.objects.create(nome="Ida", cognome="Import", reparto="Officina")
        DipendenteAnagraficaCivile.objects.create(legacy_anagrafica_id=dip.id, codice_fiscale="MPRDIA80A41H501X")
        TipoVisitaMedica.objects.create(nome="Periodica import", durata_mesi=12, obbligatoria=True)
        svolta = timezone.localdate() - timedelta(days=10)
        richiamo = svolta + timedelta(days=180)
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
            path = self._xlsx(tmp, [["MPRDIA80A41H501X", "Periodica import", svolta, richiamo, "Idoneo", ""]])
            call_command("importa_visite_mediche_xlsx", path, stdout=StringIO())
            ricalcola_visite()
            v = VisitaMedica.objects.get(legacy_anagrafica_id=dip.id)
            self.assertEqual((v.scadenza_medico, v.data_scadenza), (richiamo, richiamo))

            nuovo = svolta + timedelta(days=90)
            path = self._xlsx(tmp, [["MPRDIA80A41H501X", "Periodica import", svolta, nuovo, "Idoneo", ""]])
            call_command("importa_visite_mediche_xlsx", path, "--overwrite-existing", stdout=StringIO())
            ricalcola_visite()
            v.refresh_from_db()
            self.assertEqual((v.scadenza_medico, v.data_scadenza), (nuovo, nuovo))
