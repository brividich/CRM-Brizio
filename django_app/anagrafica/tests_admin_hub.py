"""Hub «Amministrazione anagrafica» e storico centralizzato degli import (prompt 03, obiettivo C)."""
from __future__ import annotations

from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse

from core.test_excel_export import assert_no_live_formula

from . import views
from .models import ImportazioneRetributiva

User = get_user_model()

CSV_RETRIBUZIONI = (
    "tax_code;nome;pay_item;value;date\n"
    "RSSMRA80A01H501U;ROSSI MARIO;Paga base;1500,00;01/09/2026\n"
    ";BIANCHI ANNA;Paga base;1400,00;01/09/2026\n"
    "VRDLGU85B02F205X;VERDI LUIGI;Paga base;=HYPERLINK(1);01/09/2026\n"
)


@override_settings(LEGACY_AUTH_ENABLED=False, SECURE_SSL_REDIRECT=False)
class AdminHubTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser("hub_admin", "hub_admin@example.invalid", "x")
        cls.utente = User.objects.create_user("hub_utente", "hub_utente@example.invalid", "x")

    def setUp(self):
        self.client.force_login(self.admin)

    def test_tutte_le_schede_si_aprono(self):
        url = reverse("anagrafica:amministrazione_hub")
        for scheda in ("configurazione", "viste", "cataloghi", "import", "manutenzione", "inesistente"):
            with self.subTest(scheda=scheda):
                self.assertEqual(self.client.get(url, {"scheda": scheda}).status_code, 200)
        self.assertContains(self.client.get(url, {"scheda": "cataloghi"}), "?tab=mansioni")
        self.assertContains(self.client.get(url, {"scheda": "manutenzione"}), "pulisci_scadenze_hr")

    def test_negato_a_chi_non_amministra(self):
        self.client.force_login(self.utente)
        r = self.client.get(reverse("anagrafica:amministrazione_hub"))
        self.assertEqual(r.status_code, 302)

    def test_nasconde_le_pagine_negate_dall_acl(self):
        url = reverse("anagrafica:amministrazione_hub")
        with patch("anagrafica.views_admin_hub._consentita", side_effect=lambda request, u: "retribuzioni" not in u):
            html = self.client.get(url, {"scheda": "import"}).content.decode()
        # Il link può restare nella subnav (che ha i suoi cancelli): qui conta la scheda dell'hub.
        self.assertNotIn("<strong>Importa retribuzioni</strong>", html)
        self.assertIn("<strong>Importa cedolini</strong>", html)

    def test_import_registra_gli_scarti_e_lo_storico_li_scarica_senza_formule(self):
        file_obj = SimpleUploadedFile("paghe.csv", CSV_RETRIBUZIONI.encode("utf-8"))
        imp = views._import_csv_retribuzioni(file_obj, self.admin, file_nome="paghe.csv")
        imp.refresh_from_db()
        self.assertEqual((imp.righe_ok, imp.righe_errore), (1, 2))
        self.assertEqual([s["motivo"] for s in imp.scarti], ["Codice fiscale o voce mancante", "Importo o data non validi"])
        html = self.client.get(reverse("anagrafica:amministrazione_hub"), {"scheda": "import"}).content.decode()
        self.assertIn("paghe.csv", html)
        url_scarti = reverse("anagrafica:import_scarti_export", args=["retribuzioni", imp.pk])
        self.assertIn(url_scarti, html)
        xlsx = self.client.get(url_scarti)
        self.assertEqual(xlsx.status_code, 200)
        assert_no_live_formula(self, xlsx.content)
        from openpyxl import load_workbook

        valori = [c.value for row in load_workbook(BytesIO(xlsx.content)).active.iter_rows() for c in row if c.value]
        self.assertIn("VRDLGU85B02F205X", valori)

    def test_storico_hr_e_scarti_riservati_al_permesso_hr(self):
        imp = ImportazioneRetributiva.objects.create(data_competenza="2026-09-01", file_nome="riservato.csv",
                                                     scarti=[{"riga": 2, "chiave": "X", "motivo": "m"}])
        with patch("anagrafica.views._helpers._check_hr_permission", return_value=False):
            html = self.client.get(reverse("anagrafica:amministrazione_hub"), {"scheda": "import"}).content.decode()
            self.assertNotIn("riservato.csv", html)
            r = self.client.get(reverse("anagrafica:import_scarti_export", args=["retribuzioni", imp.pk]))
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.client.get(reverse("anagrafica:import_scarti_export", args=["altro", imp.pk])).status_code, 404)
