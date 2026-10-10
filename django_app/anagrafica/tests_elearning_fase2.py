"""E-learning, fase 2 del prompt 05. Dati sintetici.

Blocco 1: domande tipizzate (singola, multipla, vero/falso) e import da Excel.
"""
from __future__ import annotations

import io

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from .models_formazione import TrainingQuizAttempt, TrainingQuizOption, TrainingQuizQuestion
from .services import elearning_quiz
from .services.elearning_quiz_import import ImportQuizError, leggi
from .tests_elearning_pro import _Pro

User = get_user_model()
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx(righe) -> SimpleUploadedFile:
    from openpyxl import Workbook
    wb = Workbook()
    for r in righe:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile("domande.xlsx", buf.getvalue(), content_type=XLSX)


class DomandeTipizzateTests(_Pro):
    def _domanda(self, tipo, opzioni):
        q = TrainingQuizQuestion.objects.create(corso=self.corso, ordine=9, testo=f"Domanda {tipo}", tipo=tipo)
        for i, (testo, giusta) in enumerate(opzioni, start=1):
            TrainingQuizOption.objects.create(domanda=q, testo=testo, corretta=giusta, ordine=i)
        return q

    def test_validita_per_tipo(self):
        singola_doppia = self._domanda("SINGOLA", [("a", True), ("b", True)])
        multipla = self._domanda("MULTIPLA", [("a", True), ("b", True), ("c", False)])
        senza = self._domanda("MULTIPLA", [("a", False), ("b", False)])
        self.assertIn("una sola", elearning_quiz.problema_domanda(singola_doppia))
        self.assertEqual(elearning_quiz.problema_domanda(multipla), "")
        self.assertIn("Nessuna risposta corretta", elearning_quiz.problema_domanda(senza))
        valide = {d.pk for d in elearning_quiz.domande_valide(self.corso)}
        self.assertIn(multipla.pk, valide)
        self.assertNotIn(singola_doppia.pk, valide)

    def test_multipla_tutte_e_sole_le_giuste(self):
        self.q.delete()
        q = self._domanda("MULTIPLA", [("a", True), ("b", True), ("c", False)])
        a, b, _c = q.opzioni.order_by("ordine")
        self._vedi_tutte()
        url = reverse("anagrafica:formazione_online_quiz", args=[self.corso.pk])
        r = self.client.get(url)
        self.assertContains(r, 'type="checkbox" name="q_%d"' % q.pk)
        self.client.post(url, {"token": r.context["token"], f"q_{q.pk}": [a.pk]})  # incompleta
        self.assertFalse(TrainingQuizAttempt.objects.get(token=r.context["token"]).superato)

    def test_singola_e_vero_falso_con_radio(self):
        vf = self._domanda("VERO_FALSO", [])
        elearning_quiz.imposta_vero_falso(vf, vera=False)
        self.assertEqual([(o.testo, o.corretta) for o in vf.opzioni.order_by("ordine")],
                         [("Vero", False), ("Falso", True)])
        self._vedi_tutte()
        r = self.client.get(reverse("anagrafica:formazione_online_quiz", args=[self.corso.pk]))
        self.assertContains(r, 'type="radio" name="q_%d"' % vf.pk)
        self.assertContains(r, 'type="radio" name="q_%d"' % self.q.pk)
        voce = next(v for v in TrainingQuizAttempt.objects.get(token=r.context["token"]).domande_servite_json
                    if v["id"] == vf.pk)
        self.assertEqual(voce["tipo"], "VERO_FALSO")

    def test_editor_vero_falso_crea_le_opzioni_e_blocca_le_altre(self):
        admin = User.objects.create_superuser("aut.el", "aut@example.invalid", "x")
        self.client.force_login(admin)
        self.client.post(reverse("anagrafica:formazione_question_save", args=[self.corso.pk]),
                         {"testo": "Il casco è un DPI.", "tipo": "VERO_FALSO", "vf_vera": "1", "ordine": 5,
                          "is_active": "on"})
        q = TrainingQuizQuestion.objects.get(testo="Il casco è un DPI.")
        self.assertEqual([(o.testo, o.corretta) for o in q.opzioni.order_by("ordine")],
                         [("Vero", True), ("Falso", False)])
        self.client.post(reverse("anagrafica:formazione_option_save", args=[self.corso.pk, q.pk]),
                         {"testo": "Forse", "ordine": 3})
        self.assertEqual(q.opzioni.count(), 2)
        r = self.client.get(reverse("anagrafica:formazione_corso_elearning", args=[self.corso.pk]))
        self.assertContains(r, "Importa domande da Excel")


class ImportDomandeExcelTests(_Pro):
    def setUp(self):
        super().setUp()
        self.admin = User.objects.create_superuser("imp.el", "imp@example.invalid", "x")
        self.client.force_login(self.admin)
        self.url = reverse("anagrafica:formazione_question_import", args=[self.corso.pk])

    def test_import_dei_tre_tipi(self):
        f = _xlsx([["Tipo", "Domanda", "Risposta 1", "Risposta 2", "Risposta 3", "Corrette"],
                   ["", "Singola implicita?", "sì", "no", None, 1],
                   ["multipla", "Quali?", "a", "b", "c", "1,3"],
                   ["VF", "Vero?", None, None, None, "F"]])
        self.client.post(self.url, {"file": f})
        nuove = list(self.corso.quiz_domande.exclude(pk=self.q.pk).order_by("ordine"))
        self.assertEqual([d.tipo for d in nuove], ["SINGOLA", "MULTIPLA", "VERO_FALSO"])
        self.assertTrue(all(not elearning_quiz.problema_domanda(d) for d in nuove))
        self.assertEqual(nuove[0].ordine, self.q.ordine + 1)

    def test_tutto_o_niente_con_numero_di_riga(self):
        f = _xlsx([["Domanda", "Risposta 1", "Risposta 2", "Corrette"],
                   ["Buona?", "a", "b", "1"],
                   ["Rotta?", "a", "b", "5"]])
        with self.assertRaises(ImportQuizError) as ctx:
            leggi(f)
        self.assertIn("Riga 3", ctx.exception.errori[0])
        f.seek(0)
        r = self.client.post(self.url, {"file": f}, follow=True)
        self.assertContains(r, "nessuna domanda salvata")
        self.assertEqual(self.corso.quiz_domande.count(), 1)

    def test_file_non_excel_rifiutato_e_modello_scaricabile(self):
        finto = SimpleUploadedFile("domande.xlsx", b"<html>no</html>", content_type=XLSX)
        r = self.client.post(self.url, {"file": finto}, follow=True)
        self.assertContains(r, "File rifiutato")
        self.assertEqual(self.corso.quiz_domande.count(), 1)
        modello = self.client.get(self.url)
        self.assertEqual(modello.status_code, 200)
        self.assertEqual(len(leggi(SimpleUploadedFile("m.xlsx", modello.content))), 3)

    def test_discente_non_importa(self):
        self.client.force_login(self.user)
        self.client.post(self.url, {"file": _xlsx([["Domanda", "Risposta 1", "Risposta 2", "Corrette"],
                                                   ["X?", "a", "b", "1"]])})
        self.assertEqual(self.corso.quiz_domande.count(), 1)


class VerificaAttestatoTests(_Pro):
    def setUp(self):
        super().setUp()
        self._vedi_tutte()
        with self.captureOnCommitCallbacks(execute=True):
            self._quiz(self.giusta)
        from .models_formazione import TrainingEmployeeRecord
        self.record = TrainingEmployeeRecord.objects.get(legacy_anagrafica_id=self.lid)

    def _codice(self):
        from .services.attestato_verifica import assegna_codice_verifica
        return assegna_codice_verifica(self.record)

    def test_codice_stabile_e_casuale(self):
        from .services import attestato_verifica as av
        codice = self._codice()
        self.assertEqual(len(codice), av.LUNGHEZZA)
        self.assertTrue(set(codice) <= set(av.ALFABETO))
        self.record.refresh_from_db()
        self.assertEqual(av.assegna_codice_verifica(self.record), codice)

    def test_esiti(self):
        from datetime import date
        from .models_elearning import TrainingElearningCompletamento
        from .services import attestato_verifica as av
        codice = self._codice()
        esito = av.verifica(av.formatta(codice).lower())
        self.assertEqual(esito.stato, av.VALIDO)
        self.assertEqual(esito.nome, "DISCENTE.EL NOME")
        self.assertEqual(av.verifica("ZZZZ-ZZZZ-ZZZZ").stato, av.SCONOSCIUTO)
        self.assertEqual(av.verifica(codice, oggi=date(2999, 1, 1)).stato, av.VALIDO)  # nessuna scadenza
        self.record.refresh_from_db()
        self.record.data_scadenza = date(2020, 1, 1)
        self.record.save()  # dal portale: rifirmato
        self.assertEqual(av.verifica(codice).stato, av.SCADUTO)
        comp = TrainingElearningCompletamento.objects.get(record=self.record)
        TrainingElearningCompletamento.objects.filter(pk=comp.pk).update(verifica_json={"falso": True})
        self.assertEqual(av.verifica(codice).stato, av.ALTERATO)

    def test_pagina_minima_e_solo_autenticati(self):
        from .services.attestato_verifica import formatta
        url = reverse("anagrafica:formazione_verifica_attestato_codice", args=[formatta(self._codice())])
        r = self.client.get(url)
        self.assertContains(r, "Attestato autentico e valido")
        self.assertContains(r, "Sicurezza online")
        self.assertNotContains(r, "discente.el@example.invalid")
        self.assertEqual(r["Referrer-Policy"], "no-referrer")
        self.client.logout()
        self.assertEqual(self.client.get(url).status_code, 302)

    def test_limite_di_verifiche(self):
        from django.core.cache import cache
        from .views_elearning import VERIFICHE_MAX
        cache.set(f"elearning:verifica:{self.user.pk}", VERIFICHE_MAX, 60)
        r = self.client.get(reverse("anagrafica:formazione_verifica_attestato_codice", args=["AAAA-BBBB-CCCC"]))
        self.assertEqual(r.status_code, 429)
        cache.delete(f"elearning:verifica:{self.user.pk}")

    def test_pdf_e_html_portano_il_codice(self):
        from django.test import override_settings
        from .services.attestato_pdf import build_attestato_context, build_attestato_pdf_bytes
        with override_settings(SITE_URL="https://hub.example.invalid"):
            ctx = build_attestato_context(self.record)
            self.assertTrue(ctx["verifica_url"].startswith("https://hub.example.invalid/anagrafica/"))
            self.assertIn("<svg", ctx["verifica_qr_svg"])
            self.assertTrue(build_attestato_pdf_bytes(self.record).startswith(b"%PDF"))
        admin = User.objects.create_superuser("att.el", "att@example.invalid", "x")
        self.client.force_login(admin)
        r = self.client.get(reverse("anagrafica:attestato_formazione", args=[self.record.pk]))
        self.assertContains(r, ctx["codice_verifica"])


class ScalettaPromemoriaTests(_Pro):
    def setUp(self):
        super().setUp()
        from datetime import date
        from core.legacy_models import AnagraficaDipendente
        from .models import AreaAziendale, DipendenteAnagraficaAziendale, Reparto
        from .models_formazione import TrainingAssignment
        from .tests_elearning_sicurezza import _discente
        self.oggi = date(2026, 10, 12)  # lunedì
        _u, self.resp_lid, self.resp_uid = _discente("resp.el")
        AnagraficaDipendente.objects.filter(pk=self.resp_lid).update(email_notifica="resp@example.invalid")
        area = AreaAziendale.objects.create(nome="Linea QA", reparto=Reparto.objects.create(nome="Rep QA"),
                                            responsabile_legacy_id=self.resp_lid)
        DipendenteAnagraficaAziendale.objects.create(legacy_anagrafica_id=self.lid, area_aziendale=area)
        self.ass = TrainingAssignment.objects.get(corso=self.corso, legacy_anagrafica_id=self.lid)

    def _scadenza(self, giorni):
        from datetime import timedelta
        type(self.ass).objects.filter(pk=self.ass.pk).update(due_date=self.oggi + timedelta(days=giorni))

    def _esegui(self, **kw):
        from .services.elearning_promemoria import esegui
        return esegui(oggi=kw.pop("oggi", self.oggi), **kw)

    def _notifiche(self, uid, tipo):
        from core.models import Notifica
        return Notifica.objects.filter(legacy_user_id=uid, tipo=tipo).count()

    def test_soglie(self):
        from .services.elearning_promemoria import soglie
        self.assertEqual(soglie("1, 14;7,x,0,7"), [14, 7, 1])

    def test_promemoria_una_volta_per_soglia(self):
        from django.core import mail
        self._scadenza(7)
        self.assertEqual(self._esegui().promemoria, 1)
        self.assertEqual(self._esegui().promemoria, 0)  # stesso giorno, rilanciato
        self.assertEqual(self._notifiche(self.uid, "elearning_promemoria"), 1)
        self._scadenza(5)  # nessuna nuova soglia raggiunta (7 già inviata)
        self.assertEqual(self._esegui().promemoria, 0)
        self._scadenza(1)
        self.assertEqual(self._esegui().promemoria, 1)
        self.assertEqual(len([m for m in mail.outbox if "riepilogo" in m.subject.lower()]), 1)  # lunedì

    def test_sollecito_a_dipendente_e_responsabile(self):
        from django.core import mail
        self._scadenza(-8)  # scaduto da 8 giorni: soglia 7, non arretrati 1 e 7 insieme
        rie = self._esegui()
        self.assertEqual((rie.solleciti, rie.email_responsabili), (1, 1))
        self.assertEqual(self._notifiche(self.uid, "elearning_sollecito"), 1)
        self.assertEqual(self._notifiche(self.resp_uid, "elearning_sollecito_responsabile"), 1)
        self.assertIn("resp@example.invalid", [d for m in mail.outbox for d in m.to])
        self.assertEqual(self._esegui().solleciti, 0)

    def test_completato_cessato_o_dry_run_niente(self):
        from .models import DipendenteAnagraficaAziendale
        from .models_elearning import TrainingElearningAvviso
        self._scadenza(1)
        self.assertEqual(self._esegui(invia=False).promemoria, 1)
        self.assertFalse(TrainingElearningAvviso.objects.exists())
        DipendenteAnagraficaAziendale.objects.filter(legacy_anagrafica_id=self.lid).update(
            data_cessazione=self.oggi)
        self.assertEqual(self._esegui().promemoria, 0)

    def test_digest_settimanale_una_volta(self):
        from datetime import timedelta
        self._scadenza(3)
        self.assertEqual(self._esegui().digest, 1)
        self.assertEqual(self._esegui().digest, 0)
        self.assertEqual(self._esegui(oggi=self.oggi + timedelta(days=1)).digest, 0)  # martedì

    def test_impostazioni_salvano_la_scaletta(self):
        from .forms import ElearningConfigForm
        from .models_formazione import ElearningConfig
        cfg = ElearningConfig.get_instance()
        form = ElearningConfigForm({"quiz_punteggio_minimo_default": 70, "validita_mesi_default": 0,
                                    "max_tentativi_quiz": 0, "promemoria_giorni_prima": "30, 7",
                                    "solleciti_giorni_dopo": "", "digest_responsabile_giorno": 7}, instance=cfg)
        self.assertTrue(form.is_valid(), form.errors)
        form.save()
        cfg.refresh_from_db()
        self.assertEqual((cfg.promemoria_giorni_prima, cfg.solleciti_giorni_dopo, cfg.digest_responsabile_giorno),
                         ("30,7", "", 7))
        self.assertFalse(ElearningConfigForm({"promemoria_giorni_prima": "sette"}, instance=cfg).is_valid())


class GradimentoEfficaciaTests(_Pro):
    def _completa(self):
        self._vedi_tutte()
        with self.captureOnCommitCallbacks(execute=True):
            return self._quiz(self.giusta)

    def test_offerto_dopo_il_quiz_e_una_volta_sola(self):
        from .models_elearning import TrainingElearningGradimento
        r = self._completa()
        url = reverse("anagrafica:formazione_online_gradimento", args=[self.corso.pk])
        self.assertContains(r, url)
        self.assertContains(self.client.get(reverse("anagrafica:formazione_online_catalog")), url)
        incompleto = self.client.post(url, {"v_0": "5"})
        self.assertEqual(incompleto.status_code, 400)
        self.client.post(url, {"v_0": "5", "v_1": "4", "v_2": "4", "v_3": "3", "commento": "Chiaro."})
        g = TrainingElearningGradimento.objects.get()
        self.assertEqual((str(g.media), g.voti_json, len(g.domande_json)), ("4.00", [5, 4, 4, 3], 4))
        self.assertEqual(self.client.get(url).status_code, 302)  # già dato
        self.assertNotContains(self.client.get(reverse("anagrafica:formazione_online_catalog")), url)

    def test_non_completato_non_valuta(self):
        url = reverse("anagrafica:formazione_online_gradimento", args=[self.corso.pk])
        self.assertEqual(self.client.get(url).status_code, 302)

    def test_completamento_elearning_apre_la_valutazione_di_efficacia(self):
        from .models_formazione import TrainingEfficacia
        self._regola(valutazione_efficacia_mesi=3)
        self._completa()
        self.assertEqual(TrainingEfficacia.objects.filter(legacy_anagrafica_id=self.lid).count(), 1)

    def test_cruscotto_medie_e_commenti_senza_nome(self):
        self._completa()
        url = reverse("anagrafica:formazione_online_gradimento", args=[self.corso.pk])
        self.client.post(url, {"v_0": "5", "v_1": "5", "v_2": "4", "v_3": "4", "commento": "Utile davvero."})
        admin = User.objects.create_superuser("cru.el", "cru@example.invalid", "x")
        self.client.force_login(admin)
        r = self.client.get(reverse("anagrafica:elearning_cruscotto"))
        self.assertTrue("4,5 / 5" in r.content.decode() or "4.5 / 5" in r.content.decode())
        # Con meno di 5 giudizi il commento non si mostra (l'autore sarebbe riconoscibile).
        self.assertNotContains(r, "Utile davvero.")
        self.assertNotContains(r, "DISCENTE.EL")


class ModuliEVersioniTests(_Pro):
    def setUp(self):
        super().setUp()
        self.admin = User.objects.create_superuser("ver.el", "ver@example.invalid", "x")

    def _completa(self):
        self._vedi_tutte()
        with self.captureOnCommitCallbacks(execute=True):
            self._quiz(self.giusta)

    def _modifica_slide(self, testo):
        s = self.corso.slides.get(ordine=1)
        self.client.force_login(self.admin)
        self.client.post(reverse("anagrafica:formazione_slide_save", args=[self.corso.pk]),
                         {"slide_id": s.pk, "titolo": s.titolo, "ordine": 1, "contenuto": testo, "is_active": "on"})
        self.client.force_login(self.user)
        self.corso.refresh_from_db()

    def test_indice_per_moduli_nel_player(self):
        from .models_elearning import TrainingElearningModulo
        m1 = TrainingElearningModulo.objects.create(corso=self.corso, ordine=1, titolo="Rischi")
        m2 = TrainingElearningModulo.objects.create(corso=self.corso, ordine=2, titolo="Diritti")
        self.corso.slides.filter(ordine__in=[1, 2]).update(modulo=m1)
        self.corso.slides.filter(ordine=3).update(modulo=m2)
        self.client.get(reverse("anagrafica:formazione_online_player", args=[self.corso.pk]))
        r = self._slide(1)
        self.assertContains(r, "Indice del corso")
        self.assertContains(r, "Rischi · Lezione 1 / 3")
        self.assertContains(r, "fm-indice-modulo")
        # la lezione 3 non è ancora apribile: niente link
        self.assertNotContains(r, reverse("anagrafica:formazione_online_slide", args=[self.corso.pk, 3]))

    def test_editor_moduli(self):
        self.client.force_login(self.admin)
        self.client.post(reverse("anagrafica:formazione_modulo_save", args=[self.corso.pk]), {"titolo": "Introduzione"})
        m = self.corso.moduli_elearning.get()
        s = self.corso.slides.get(ordine=2)
        self.client.post(reverse("anagrafica:formazione_slide_save", args=[self.corso.pk]),
                         {"slide_id": s.pk, "titolo": s.titolo, "ordine": 2, "contenuto": s.contenuto,
                          "modulo": m.pk, "is_active": "on"})
        s.refresh_from_db()
        self.assertEqual(s.modulo_id, m.pk)
        self.client.post(reverse("anagrafica:formazione_modulo_delete", args=[self.corso.pk, m.pk]))
        s.refresh_from_db()
        self.assertIsNone(s.modulo_id)

    def test_nuova_versione_solo_dopo_un_completamento(self):
        from .services import elearning_versioni as ev
        ev.registra(self.corso, alla_pubblicazione=True)
        self.assertEqual(list(self.corso.versioni.values_list("version_label", flat=True)), ["1.0"])
        self._modifica_slide("Testo rivisto prima di ogni completamento")
        self.assertEqual((self.corso.versione, self.corso.versioni.count()), ("1.0", 1))  # aggiornata lì
        self._completa()
        self._modifica_slide("Testo rivisto dopo un completamento")
        self.assertEqual(self.corso.versione, "1.1")
        vecchia = self.corso.versioni.get(version_label="1.0")
        self.assertIsNotNone(vecchia.data_fine_validita)
        from .models_formazione import TrainingEmployeeRecord
        self.assertEqual(TrainingEmployeeRecord.objects.get(legacy_anagrafica_id=self.lid).course_version_snapshot, "1.0")

    def test_riassegna_chi_ha_completato_la_versione_precedente(self):
        from .models_formazione import TrainingAssignment
        from .services import elearning_versioni as ev
        self._regola(el_nuova_versione="RIASSEGNA")
        ev.registra(self.corso, alla_pubblicazione=True)
        self._completa()
        self._modifica_slide("Contenuto aggiornato per norma nuova")
        # La riassegnazione la fa il job notturno, non la modifica.
        self.assertFalse(TrainingAssignment.objects.filter(corso=self.corso, ciclo=2).exists())
        with self.captureOnCommitCallbacks(execute=True):
            self.assertEqual(ev.fissa_e_riassegna()["riassegnati"], 1)
        nuova = TrainingAssignment.objects.get(corso=self.corso, legacy_anagrafica_id=self.lid, ciclo=2)
        self.assertIsNotNone(nuova.due_date)
        self.assertIn("1.1", nuova.note)

    def test_bozza_senza_versioni_e_storico_in_gestione(self):
        from .services import elearning_versioni as ev
        self.corso.stato = "BOZZA"
        self.corso.save()
        self.assertFalse(ev.registra(self.corso).etichetta)
        self.assertFalse(self.corso.versioni.exists())
        self.assertEqual(ev.prossima_etichetta("1.9"), "1.10")
        self.client.force_login(self.admin)
        r = self.client.get(reverse("anagrafica:formazione_elearning_manage", args=[self.corso.pk]))
        self.assertContains(r, "si crea alla pubblicazione")


class ImportInBackgroundTests(_Pro):
    def setUp(self):
        super().setUp()
        self.admin = User.objects.create_superuser("imp2.el", "imp2@example.invalid", "x")
        self.client.force_login(self.admin)

    def _pdf(self):
        from .tests_elearning import _pdf_due_pagine
        return SimpleUploadedFile("lezione.pdf", _pdf_due_pagine(), content_type="application/pdf")

    def test_upload_accoda_e_il_lavoro_crea_le_slide(self):
        from unittest import mock
        from .models_elearning import TrainingElearningImport
        from .services.elearning_import import esegui_import
        with mock.patch("django_q.tasks.async_task") as accoda, self.captureOnCommitCallbacks(execute=True):
            r = self.client.post(reverse("anagrafica:formazione_slide_import", args=[self.corso.pk]),
                                 {"file": self._pdf()}, follow=True)
        job = TrainingElearningImport.objects.get()
        accoda.assert_called_once()
        self.assertEqual(accoda.call_args.args, ("anagrafica.tasks.run_elearning_import", job.pk))
        self.assertContains(r, "in coda")
        self.assertEqual(self.corso.slides.count(), 3)  # niente conversione nella richiesta web
        esito = esegui_import(job.pk)
        job.refresh_from_db()
        self.assertEqual((esito["ok"], job.stato, job.n_slide), (True, "COMPLETATO", 2))
        self.assertEqual(self.corso.slides.count(), 5)
        self.assertFalse(job.file)  # file di origine cancellato
        self.assertEqual(esegui_import(job.pk)["motivo"], "non_in_coda")  # riconsegna del task: niente doppioni
        stato = self.client.get(reverse("anagrafica:formazione_slide_import_stato", args=[self.corso.pk]))
        self.assertContains(stato, "2 slide importate")
        self.assertNotContains(stato, "every 4s")

    def test_file_rotto_errore_leggibile_e_permessi(self):
        from .models_elearning import TrainingElearningImport
        from .services.elearning_import import accoda_import, esegui_import
        with self.captureOnCommitCallbacks(execute=False):
            job = accoda_import(self.corso, SimpleUploadedFile("rotto.pdf", b"non un pdf"), user=self.admin)
        esegui_import(job.pk)
        job.refresh_from_db()
        self.assertEqual(job.stato, TrainingElearningImport.ERRORE)
        self.assertTrue(job.errore)
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse("anagrafica:formazione_slide_import_stato",
                                                 args=[self.corso.pk])).status_code, 403)



class CorrezioniReviewFase2Tests(_Pro):
    def setUp(self):
        super().setUp()
        self.admin = User.objects.create_superuser("rev2.el", "rev2@example.invalid", "x")

    def _record_completato(self):
        self._vedi_tutte()
        with self.captureOnCommitCallbacks(execute=True):
            self._quiz(self.giusta)
        from .models_formazione import TrainingEmployeeRecord
        return TrainingEmployeeRecord.objects.get(legacy_anagrafica_id=self.lid)

    def test_dati_alterati_nel_database_non_verificabili(self):
        from datetime import date
        from .services import attestato_verifica as av
        record = self._record_completato()
        codice = av.assegna_codice_verifica(record)
        record.refresh_from_db()
        record.data_scadenza = date(2030, 1, 1)
        record.save()  # modifica dal portale: rifirmata, resta valida
        self.assertEqual(av.verifica(codice).stato, av.VALIDO)
        type(record).objects.filter(pk=record.pk).update(data_scadenza=date(2099, 1, 1))  # modifica diretta
        self.assertEqual(av.verifica(codice).stato, av.ALTERATO)

    def test_import_slide_verificato_dal_contenuto(self):
        self.client.force_login(self.admin)
        finto = SimpleUploadedFile("lezione.pptx", b"<html>no</html>")
        r = self.client.post(reverse("anagrafica:formazione_slide_import", args=[self.corso.pk]),
                             {"file": finto}, follow=True)
        self.assertContains(r, "File rifiutato")
        from .models_elearning import TrainingElearningImport
        self.assertFalse(TrainingElearningImport.objects.exists())

    def test_import_bloccato_scade(self):
        from datetime import timedelta
        from django.utils import timezone
        from .models_elearning import TrainingElearningImport
        from .services.elearning_import import scadi_import_bloccati
        job = TrainingElearningImport.objects.create(corso=self.corso, nome_file="x.pdf", stato="IN_CORSO")
        TrainingElearningImport.objects.filter(pk=job.pk).update(creato_il=timezone.now() - timedelta(hours=1))
        self.assertEqual(scadi_import_bloccati(), 1)
        job.refresh_from_db()
        self.assertEqual(job.stato, "ERRORE")
        self.client.force_login(self.user)
        self.assertIn(self.client.get(reverse("anagrafica:formazione_slide_import_stato",
                                              args=[self.corso.pk])).status_code, (286, 403))

    def test_excel_con_booleani_e_decimali(self):
        from .services.elearning_quiz_import import leggi
        righe = leggi(_xlsx([["Tipo", "Domanda", "Risposta 1", "Risposta 2", "Risposta 3", "Corrette"],
                             ["vero / falso", "Vero?", None, None, None, False],
                             ["multipla", "Quali?", "a", "b", "c", 1.3]]))
        self.assertEqual([(r.tipo, r.corrette) for r in righe], [("VERO_FALSO", {2}), ("MULTIPLA", {1, 3})])

    def test_passaggio_a_vero_falso_chiede_la_risposta(self):
        self.client.force_login(self.admin)
        url = reverse("anagrafica:formazione_question_save", args=[self.corso.pk])
        self.client.post(url, {"question_id": self.q.pk, "testo": self.q.testo, "tipo": "VERO_FALSO",
                               "ordine": 1, "is_active": "on"})
        self.q.refresh_from_db()
        self.assertEqual(self.q.tipo, "SINGOLA")  # rifiutato: manca vera/falsa
        self.client.post(url, {"question_id": self.q.pk, "testo": self.q.testo, "tipo": "VERO_FALSO",
                               "vf_vera": "0", "ordine": 1, "is_active": "on"})
        self.q.refresh_from_db()
        self.assertEqual([(o.testo, o.corretta) for o in self.q.opzioni.order_by("ordine")],
                         [("Vero", False), ("Falso", True)])

    def test_niente_solleciti_per_scaduti_storici(self):
        from datetime import date, timedelta
        from .models_formazione import TrainingAssignment
        from .services.elearning_promemoria import esegui
        oggi = date(2026, 10, 13)
        TrainingAssignment.objects.filter(corso=self.corso, legacy_anagrafica_id=self.lid).update(
            due_date=oggi - timedelta(days=120))
        self.assertEqual(esegui(oggi=oggi).solleciti, 0)
        TrainingAssignment.objects.filter(corso=self.corso, legacy_anagrafica_id=self.lid).update(
            due_date=oggi - timedelta(days=16))
        self.assertEqual(esegui(oggi=oggi).solleciti, 1)

    def test_versione_di_base_per_i_corsi_gia_pubblicati(self):
        from .services import elearning_versioni as ev
        self.assertFalse(self.corso.versioni.exists())
        self.assertEqual(ev.fissa_e_riassegna()["fissate"], 1)
        self.assertEqual(self.corso.versioni.get().version_label, "1.0")
