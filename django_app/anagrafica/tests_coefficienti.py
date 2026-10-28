"""Classificazione Diretto/Indiretto 1T/2T, coefficienti (somma 1) e costo orario (ottobre 2026)."""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.urls import reverse
from django.utils import timezone

from .models import (
    AbilitazioneMacchina, CompetenzaSkm, CostoOrarioMensile, ImportazioneRetributiva, ProfiloMansioniPersona,
    VoceRetributiva,
)
from .services import costo_orario as costo
from .services import mansionario_struttura as svc
from .services import mansioni_coefficienti as coeff
from .services import struttura_excel as xl
from .tests_capi_mansione import _ScenarioMansioni
from .tests_struttura import _Base, _xlsx

User = get_user_model()
D = Decimal


class _Mansioni(_Base):
    def setUp(self):
        super().setUp()
        self.m2 = CompetenzaSkm.objects.create(competenza_key="P-LAV", display="Lavaggio", tipo=CompetenzaSkm.TIPO_PROCESSO)
        self.m3 = CompetenzaSkm.objects.create(competenza_key="P-IMB", display="Imballo", tipo=CompetenzaSkm.TIPO_PROCESSO)

    def _ab(self, mansione, persona=None):
        return AbilitazioneMacchina.objects.get(legacy_anagrafica_id=(persona or self.p1).pk, competenza=mansione)

    def _coeff(self, persona=None):
        return {ab.competenza.competenza_key: ab.coefficiente
                for ab in AbilitazioneMacchina.objects.filter(legacy_anagrafica_id=(persona or self.p1).pk)
                .select_related("competenza")}


class RipartizioneTests(_Mansioni):
    def test_una_mansione_vale_uno_e_la_seconda_divide(self):
        svc.assegna(self.processo, legacy_anagrafica_id=self.p1.pk, livello="U")
        self.assertEqual(self._coeff(), {"P-SBA": D("1")})
        self.assertFalse(coeff.stato(self.p1.pk)["da_rivedere"])
        svc.assegna(self.m2, legacy_anagrafica_id=self.p1.pk, livello="L")
        self.assertEqual(self._coeff(), {"P-SBA": D("0.5"), "P-LAV": D("0.5")})
        st = coeff.stato(self.p1.pk)
        self.assertTrue(st["in_regola"])
        self.assertTrue(st["da_rivedere"])
        self.assertIn("Lavaggio", st["motivo"])

    def test_la_nuova_prende_un_terzo_le_altre_in_proporzione(self):
        a = svc.assegna(self.processo, legacy_anagrafica_id=self.p1.pk, livello="U")
        b = svc.assegna(self.m2, legacy_anagrafica_id=self.p1.pk, livello="U")
        coeff.salva(self.p1.pk, {a.pk: ("", "0,75"), b.pk: ("", "0.25")})
        svc.assegna(self.m3, legacy_anagrafica_id=self.p1.pk, livello="U")
        valori = self._coeff()
        self.assertEqual(sum(valori.values()), D("1"))
        self.assertEqual(valori["P-SBA"], D("0.5"))
        self.assertEqual(valori["P-IMB"], D("0.3333"))
        self.assertEqual(valori["P-LAV"], D("0.1667"))

    def test_tolta_lascia_il_peso_alle_altre(self):
        a = svc.assegna(self.processo, legacy_anagrafica_id=self.p1.pk, livello="U")
        b = svc.assegna(self.m2, legacy_anagrafica_id=self.p1.pk, livello="U")
        c = svc.assegna(self.m3, legacy_anagrafica_id=self.p1.pk, livello="U")
        coeff.salva(self.p1.pk, {a.pk: ("", "0.6"), b.pk: ("", "0.2"), c.pk: ("", "0.2")})
        svc.togli(self._ab(self.processo))
        valori = self._coeff()
        self.assertIsNone(valori["P-SBA"])
        self.assertEqual((valori["P-LAV"], valori["P-IMB"]), (D("0.5"), D("0.5")))
        self.assertTrue(coeff.stato(self.p1.pk)["da_rivedere"])

    def test_sospesa_non_conta(self):
        a = svc.assegna(self.processo, legacy_anagrafica_id=self.p1.pk, livello="U")
        svc.assegna(self.m2, legacy_anagrafica_id=self.p1.pk, livello="U")
        a.stato = AbilitazioneMacchina.STATO_SOSPESA
        a.save()
        self.assertEqual(self._coeff(), {"P-SBA": None, "P-LAV": D("1")})

    def test_arrotonda_a_uno(self):
        self.assertEqual(sum(coeff.arrotonda_a_uno([D(1) / 3] * 3)), D("1"))
        self.assertEqual(coeff.ripartisci([None, None, None, None]), [D("0.25")] * 4)


class SalvaTests(_Mansioni):
    def setUp(self):
        super().setUp()
        self.a = svc.assegna(self.processo, legacy_anagrafica_id=self.p1.pk, livello="U")
        self.b = svc.assegna(self.m2, legacy_anagrafica_id=self.p1.pk, livello="U")

    def test_somma_diversa_da_uno_non_salva(self):
        with self.assertRaises(ValueError):
            coeff.salva(self.p1.pk, {self.a.pk: ("D1T", "0.5"), self.b.pk: ("D1T", "0.4")})
        self.assertEqual(self._ab(self.processo).classificazione, "")
        with self.assertRaises(ValueError):  # manca una mansione
            coeff.salva(self.p1.pk, {self.a.pk: ("D1T", "1")})
        with self.assertRaises(ValueError):
            coeff.salva(self.p1.pk, {self.a.pk: ("X9", "0.5"), self.b.pk: ("D1T", "0.5")})

    def test_tolleranza_e_prevalente(self):
        c = svc.assegna(self.m3, legacy_anagrafica_id=self.p1.pk, livello="U")
        coeff.salva(self.p1.pk, {self.a.pk: ("D2T", "0.3333"), self.b.pk: ("diretto 2t", "0,3333"),
                                 c.pk: ("D2T", "33.33%")}, prevalente_scelta="D2T")
        self.assertEqual(sum(self._coeff().values()), D("1"))
        st = coeff.stato(self.p1.pk)
        self.assertFalse(st["da_rivedere"])
        self.assertEqual(st["classe_principale"], "D2T")
        # Una mansione nuova prende la classificazione prevalente.
        nuova = CompetenzaSkm.objects.create(competenza_key="P-NEW", display="Nuova", tipo=CompetenzaSkm.TIPO_PROCESSO)
        self.assertEqual(svc.assegna(nuova, legacy_anagrafica_id=self.p1.pk, livello="I").classificazione, "D2T")
        self.assertEqual(AbilitazioneMacchina.objects.get(pk=self._ab(nuova).pk).classificazione, "D2T")

    def test_turni_misti(self):
        coeff.salva(self.p1.pk, {self.a.pk: ("D1T", "0.5"), self.b.pk: ("I2T", "0.5")})
        self.assertTrue(coeff.stato(self.p1.pk)["turni_misti"])
        coeff.salva(self.p1.pk, {self.a.pk: ("D1T", "0.5"), self.b.pk: ("I1T", "0.5")})
        self.assertFalse(coeff.stato(self.p1.pk)["turni_misti"])

    def test_pagina(self):
        url = reverse("anagrafica:mansioni_coefficienti", args=[self.p1.pk])
        r = self.client.get(url)
        self.assertContains(r, "Lavaggio")
        self.assertContains(r, "Coefficienti da rivedere")
        r = self.client.post(url, {f"classe_{self.a.pk}": "D1T", f"coeff_{self.a.pk}": "0.5",
                                   f"classe_{self.b.pk}": "D1T", f"coeff_{self.b.pk}": "0.6", "prevalente": "D1T"},
                             follow=True)
        self.assertContains(r, "devono fare 1")
        self.client.post(url, {f"classe_{self.a.pk}": "D1T", f"coeff_{self.a.pk}": "0.7",
                               f"classe_{self.b.pk}": "I1T", f"coeff_{self.b.pk}": "0.3", "prevalente": "D1T"})
        self.assertEqual((self._ab(self.processo).coefficiente, self._ab(self.m2).classificazione), (D("0.7"), "I1T"))
        self.assertEqual(ProfiloMansioniPersona.objects.get(legacy_anagrafica_id=self.p1.pk).classificazione_prevalente,
                         "D1T")
        self.client.post(url, {"azione": "uguale"})
        self.assertEqual(self._ab(self.processo).coefficiente, D("0.5"))
        r = self.client.get(reverse("anagrafica:dipendente_detail", args=[self.p1.pk]))
        self.assertContains(r, "Classificazione e coefficienti")


class InBloccoEExcelTests(_Mansioni):
    def _file(self, mansioni):
        from .services import assegnazioni_excel as ax

        return _xlsx({ax.FOGLIO_MANSIONI: [ax.COL_MANSIONI, *mansioni]})

    def test_in_blocco_con_classificazione(self):
        from .services import mansioni_in_blocco as mb

        self._col_ruolo(self.processo, persone=[self.p1, self.p2])
        mb.esegui([self.p1.pk, self.p2.pk], mansione=self.processo, modo=mb.ASSEGNA, livello="U",
                  classificazione="I1T", applica=True)
        self.assertEqual(self._ab(self.processo, self.p2).classificazione, "I1T")
        esiti = mb.esegui([self.p1.pk], mansione=self.processo, modo=mb.ASSEGNA, livello="U", classificazione="D1T")
        self.assertEqual(esiti[0].azione, "cambia")
        self.assertIn("Diretto 1T", esiti[0].dettaglio)

    def test_excel_coefficienti(self):
        from .services import assegnazioni_excel as ax

        svc.assegna(self.processo, legacy_anagrafica_id=self.p1.pk, livello="U")
        svc.assegna(self.m2, legacy_anagrafica_id=self.p1.pk, livello="U")
        # Scaricato e ricaricato: niente cambia, nemmeno «da rivedere».
        contenuto = ax.esporta_xlsx()
        esito = ax.importa(contenuto, applica=True)
        self.assertEqual((esito.cambi, esito.errori), (0, 0))
        self.assertTrue(coeff.stato(self.p1.pk)["da_rivedere"])
        # Somma sbagliata: la persona si scarta.
        esito = ax.importa(self._file([[self.p1.pk, "", "", "P-SBA", "", "U", "D1T", 0.9, ""],
                                       [self.p1.pk, "", "", "P-LAV", "", "U", "D1T", 0.3, ""]]), applica=True)
        self.assertEqual(esito.errori, 2)
        self.assertEqual(self._ab(self.processo).coefficiente, D("0.5"))
        # Giusta: salva, classificazione compresa, e «da rivedere» si spegne.
        esito = ax.importa(self._file([[self.p1.pk, "", "", "P-SBA", "", "U", "D1T", 0.7, ""],
                                       [self.p1.pk, "", "", "P-LAV", "", "U", "Indiretto 1T", "0,3", ""]]),
                           applica=True)
        self.assertEqual(esito.errori, 0)
        self.assertEqual([r.azione for r in esito.righe], [xl.MODIFICATO, xl.MODIFICATO])
        self.assertEqual((self._ab(self.processo).coefficiente, self._ab(self.m2).classificazione), (D("0.7"), "I1T"))
        self.assertFalse(coeff.stato(self.p1.pk)["da_rivedere"])
        # Nella prova non si scrive niente.
        ax.importa(self._file([[self.p1.pk, "", "", "P-SBA", "", "U", "", 0.5, ""],
                               [self.p1.pk, "", "", "P-LAV", "", "U", "", 0.5, ""]]), applica=False)
        self.assertEqual(self._ab(self.processo).coefficiente, D("0.7"))

    def test_excel_colonne_esportate(self):
        from openpyxl import load_workbook

        from .services import assegnazioni_excel as ax

        a = svc.assegna(self.processo, legacy_anagrafica_id=self.p1.pk, livello="U")
        coeff.salva(self.p1.pk, {a.pk: ("D2T", "1")})
        ws = load_workbook(BytesIO(ax.esporta_xlsx()))[ax.FOGLIO_MANSIONI]
        intest = [c.value for c in ws[1]]
        riga = dict(zip(intest, [c.value for c in ws[2]]))
        self.assertEqual((riga["Classificazione"], riga["Coefficiente"]), ("D2T", 1))


class CostoOrarioTests(_Base):
    CF = "SNTTST80A01H501Z"

    def _voce(self, importo, *, mese=None, manuale=False, imp=None):
        mese = mese or timezone.localdate().replace(day=1)
        imp = imp or ImportazioneRetributiva.objects.create(data_competenza=mese)
        return VoceRetributiva.objects.create(importazione=imp, tax_code=self.CF, legacy_anagrafica_id=self.p1.pk,
                                              data_competenza=mese, pay_item="RML", pay_item_key="rml",
                                              categoria="totale", importo=D(importo), manuale=manuale)

    def test_calcolo_e_correzione_manuale(self):
        self._voce("3463.20")
        c = CostoOrarioMensile.objects.get(tax_code=self.CF)
        self.assertEqual(c.costo_orario, D("20.0000"))  # 3463,20 ÷ 26 ÷ 6,66
        manuale = self._voce("1731.60", manuale=True)
        self.assertEqual(CostoOrarioMensile.objects.get(tax_code=self.CF).costo_orario, D("10.0000"))
        self.assertTrue(CostoOrarioMensile.objects.get(tax_code=self.CF).rml_manuale)
        manuale.delete()
        self.assertEqual(CostoOrarioMensile.objects.get(tax_code=self.CF).costo_orario, D("20.0000"))

    def test_vale_l_import_piu_recente_e_mese_per_mese(self):
        oggi = timezone.localdate().replace(day=1)
        prima = (oggi - timezone.timedelta(days=1)).replace(day=1)
        self._voce("3463.20", mese=prima)
        self._voce("3463.20")
        self._voce("5194.80")  # reimport dello stesso mese
        costo.ricalcola()
        mesi = dict(CostoOrarioMensile.objects.values_list("mese", "costo_orario"))
        self.assertEqual(mesi, {prima: D("20.0000"), oggi: D("30.0000")})
        self.assertEqual(costo.ultimi([self.p1.pk])[self.p1.pk].mese, oggi)

    def test_divisori(self):
        self._voce("2600")
        costo.imposta_divisori("26", "10")
        costo.ricalcola()
        self.assertEqual(CostoOrarioMensile.objects.get().costo_orario, D("10.0000"))
        with self.assertRaises(ValueError):
            costo.imposta_divisori("0", "6,66")

    def test_import_csv(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from .views.retribuzioni import _import_csv_retribuzioni

        data = timezone.localdate().replace(day=1).strftime("%d/%m/%Y")
        csv = f"tax_code;nome;pay_item;value;date\n{self.CF};STRUTTURA UNO;RML;3463,20;{data}\n"
        _import_csv_retribuzioni(SimpleUploadedFile("p.csv", csv.encode()), self.admin, "p.csv")
        self.assertEqual(CostoOrarioMensile.objects.get(tax_code=self.CF).costo_orario, D("20.0000"))

    def test_pagina_solo_hr(self):
        self._voce("3463.20")
        r = self.client.get(reverse("anagrafica:costo_orario"))
        self.assertContains(r, "20,00")
        r = self.client.get(reverse("anagrafica:export", args=["costo_orario"]), {"format": "xlsx", "scope": "full"})
        self.assertIn("spreadsheetml", r["Content-Type"])
        from .exports import EXPORT_SPECS

        righe = EXPORT_SPECS["costo_orario"].dataset(self.client.get(reverse("anagrafica:costo_orario")).wsgi_request,
                                                     "filtered")
        self.assertEqual([x["costo"] for x in righe], [D("20.0000")])
        normale = User.objects.create_user("co_normale", "co@example.invalid", "x")
        self.client.force_login(normale)
        r = self.client.get(reverse("anagrafica:costo_orario"))
        self.assertEqual(r.status_code, 302)


class ReportCoefficientiTests(_ScenarioMansioni):
    def test_sezione(self):
        from datetime import timedelta

        from .reportistica import sezioni as catalogo
        from .reportistica.dati import Contesto, Perimetro

        m1, m2 = self._mansione("Taglio sintetico"), self._mansione("Pulizia sintetica")
        a = svc.assegna(m1, legacy_anagrafica_id=self.neri, livello="U")
        b = svc.assegna(m2, legacy_anagrafica_id=self.neri, livello="U")
        coeff.salva(self.neri, {a.pk: ("D1T", "0.5"), b.pk: ("I2T", "0.5")})
        oggi = timezone.localdate()
        ctx = Contesto(date_from=oggi - timedelta(days=365), date_to=oggi, perimetro=Perimetro())
        res = catalogo.get("classificazione_coefficienti").calcola(ctx, {})
        righe = [r for r in res.righe if r["mansione"] in {"Taglio sintetico", "Pulizia sintetica"}]
        self.assertEqual(len(righe), 2)
        self.assertTrue(all("turni misti" in r["avvisi"] for r in righe))
        self.assertEqual({r["coefficiente"] for r in righe}, {"0,5000"})
        res = catalogo.get("ruoli_mansioni_persona").calcola(ctx, {"una_riga_per": "mansione"})
        self.assertIn("Diretto 1T", {r["classificazione"] for r in res.righe})


class AltreSchermateTests(_Mansioni):
    """Classificazione, coefficienti e costo orario anche nell'elenco dipendenti e nella scheda."""

    CF = "SNTTST80A01H501Z"

    def setUp(self):
        super().setUp()
        a = svc.assegna(self.processo, legacy_anagrafica_id=self.p1.pk, livello="U")
        b = svc.assegna(self.m2, legacy_anagrafica_id=self.p1.pk, livello="U")
        coeff.salva(self.p1.pk, {a.pk: ("D2T", "0.7"), b.pk: ("D2T", "0.3")})
        mese = timezone.localdate().replace(day=1)
        imp = ImportazioneRetributiva.objects.create(data_competenza=mese)
        VoceRetributiva.objects.create(importazione=imp, tax_code=self.CF, legacy_anagrafica_id=self.p1.pk,
                                       data_competenza=mese, pay_item="RML", pay_item_key="rml",
                                       categoria="totale", importo=D("3463.20"))

    def test_colonne_dell_elenco(self):
        from .services import colonne_elenco as ce

        righe = [{"id": self.p1.pk}, {"id": self.p2.pk}]
        colonne = [ce.REGISTRO[k] for k in ("n_mansioni", "classificazione", "coefficienti", "costo_orario")]
        ce.arricchisci(righe, colonne)
        v = righe[0]["valori"]
        self.assertEqual((v["n_mansioni"][0], v["classificazione"][0], v["coefficienti"][0], v["costo_orario"][0]),
                         ("2", "Diretto 2T", "in regola", "20,00"))
        self.assertEqual(righe[1]["valori"]["classificazione"][0], "")
        self.assertEqual(ce.REGISTRO["costo_orario"].permesso, "hr")

    def test_costo_orario_solo_con_permesso_hr(self):
        from types import SimpleNamespace

        from .services import colonne_elenco as ce

        req = SimpleNamespace(_colonne_permessi=set())
        chiavi = {d["colonna"].chiave for d in ce.colonne_da_specifica(["classificazione", "costo_orario"], req)}
        self.assertEqual(chiavi, {"classificazione"})

    def test_elenco_vista_e_scheda(self):
        from .services.viste_elenco import PRESET

        self.assertIn("costo_orario", PRESET["mansioni_costi"][2])
        r = self.client.get(reverse("anagrafica:dipendenti_list"), {"vista": "mansioni_costi"})
        self.assertEqual(r.status_code, 200)
        r = self.client.get(reverse("anagrafica:dipendente_detail", args=[self.p1.pk]))
        self.assertContains(r, "Diretto 2T")
        self.assertContains(r, "€ 20,00/h")


class StoricoCostoOrarioTests(_Base):
    CF = "SNTTST80A01H501Z"

    def setUp(self):
        super().setUp()
        oggi = timezone.localdate().replace(day=1)
        self.mesi = []
        m = oggi
        for importo in ("3636.36", "3463.20", "3463.20"):  # dal più recente: 21 · 20 · 20 €/h
            self.mesi.append(m)
            imp = ImportazioneRetributiva.objects.create(data_competenza=m)
            VoceRetributiva.objects.create(importazione=imp, tax_code=self.CF, legacy_anagrafica_id=self.p1.pk,
                                           data_competenza=m, pay_item="RML", pay_item_key="rml",
                                           categoria="totale", importo=D(importo))
            m = (m - timezone.timedelta(days=1)).replace(day=1)
        self.url = reverse("anagrafica:costo_orario_storico")

    def test_tabella_mesi_variazione(self):
        from types import SimpleNamespace

        from django.http import QueryDict

        from .views_coefficienti import _storico

        ctx = _storico(SimpleNamespace(GET=QueryDict()))
        self.assertEqual(ctx["mesi"], sorted(self.mesi))
        r = ctx["righe"][0]
        self.assertEqual([x["c"].costo_orario for x in r["celle"]], [D("20.0000"), D("20.0000"), D("21.0000")])
        self.assertTrue(r["celle"][2]["su"])
        self.assertEqual((r["var"], r["var_pct"]), (D("1.0000"), D("5.0")))
        self.assertEqual(ctx["kpi"]["aumenti"], 1)
        q = QueryDict(mutable=True)
        q["dal"], q["al"] = self.mesi[1].isoformat(), self.mesi[0].isoformat()
        self.assertEqual(len(_storico(SimpleNamespace(GET=q))["mesi"]), 2)

    def test_pagina_export_e_permesso(self):
        r = self.client.get(self.url, {"persona": self.p1.pk})
        self.assertContains(r, "21,00")
        self.assertContains(r, "+5,0%")
        r = self.client.get(reverse("anagrafica:export", args=["costo_orario_storico"]),
                            {"format": "xlsx", "scope": "full"})
        self.assertIn("spreadsheetml", r["Content-Type"])
        normale = User.objects.create_user("cs_normale", "cs@example.invalid", "x")
        self.client.force_login(normale)
        self.assertEqual(self.client.get(self.url).status_code, 302)


class ReportCostoOrarioTests(_ScenarioMansioni):
    """Reportistica: costo orario per persona, ripartito con i coefficienti, andamento; solo HR."""

    def setUp(self):
        from datetime import timedelta

        super().setUp()
        m1, m2 = self._mansione("Taglio sintetico"), self._mansione("Pulizia sintetica")
        a = svc.assegna(m1, legacy_anagrafica_id=self.neri, livello="U")
        b = svc.assegna(m2, legacy_anagrafica_id=self.neri, livello="U")
        coeff.salva(self.neri, {a.pk: ("D1T", "0.75"), b.pk: ("I1T", "0.25")})
        self.mese = timezone.localdate().replace(day=1)
        self.prima = (self.mese - timedelta(days=1)).replace(day=1)
        for mese, rml in ((self.prima, D("3000.00")), (self.mese, D("3463.20"))):
            CostoOrarioMensile.objects.create(
                tax_code="SNTTST80A01H501Z", legacy_anagrafica_id=self.neri, mese=mese, rml=rml,
                divisore_giorni=D("26"), divisore_ore=D("6.66"),
                costo_orario=(rml / D("26") / D("6.66")).quantize(D("0.0001")))

    def _ctx(self):
        from .reportistica.dati import Contesto, Perimetro

        return Contesto(date_from=self.prima, date_to=timezone.localdate(), perimetro=Perimetro(persone=[self.neri]))

    def test_per_persona(self):
        from .reportistica import sezioni as catalogo
        from .reportistica.motore import fmt

        res = catalogo.get("costo_orario_persone").calcola(self._ctx(), {})
        (r,) = res.righe
        self.assertEqual((fmt(r["rml"]), fmt(r["costo_orario"]), fmt(r["precedente"]), fmt(r["variazione"])),
                         ("3.463,20", "20,00", "17,33", "2,68"))
        self.assertEqual(r["classificazione"], "Diretto 1T")
        (r,) = catalogo.get("costo_orario_persone").calcola(self._ctx(), {"mese": self.prima.isoformat()}).righe
        self.assertIsNone(r["variazione"])

    def test_ripartito_con_i_coefficienti(self):
        from .reportistica import sezioni as catalogo

        res = catalogo.get("costo_orario_ripartito").calcola(self._ctx(), {})
        per_voce = {r["voce"]: r for r in res.righe}
        self.assertEqual(set(per_voce), {"Diretto 1T", "Indiretto 1T"})
        self.assertEqual((per_voce["Diretto 1T"]["fte"], per_voce["Diretto 1T"]["rml"]), (D("0.75"), D("2597.40")))
        self.assertEqual((per_voce["Indiretto 1T"]["quota"], per_voce["Indiretto 1T"]["costo_medio"]),
                         (D("25.0"), D("20.00")))
        res = catalogo.get("costo_orario_ripartito").calcola(self._ctx(), {"per": "mansione"})
        self.assertEqual({r["voce"] for r in res.righe}, {"Taglio sintetico", "Pulizia sintetica"})

    def test_andamento(self):
        from report_conformita.registry import TONE_WARN

        from .reportistica import sezioni as catalogo

        res = catalogo.get("costo_orario_andamento").calcola(self._ctx(), {})
        self.assertEqual([r["mese"] for r in res.righe], [f"{self.prima:%m/%Y}", f"{self.mese:%m/%Y}"])
        self.assertEqual(res.righe[1]["variazione_pct"], D("15.4"))
        self.assertEqual(res.toni[1], TONE_WARN)
        self.assertIn("Variazione nel periodo", [k.label for k in res.kpis])

    def test_solo_hr_e_modello_pronto(self):
        from django.test import RequestFactory

        from .models_reportistica import ReportModello
        from .reportistica import sezioni as catalogo

        rf = RequestFactory()
        normale, capo = rf.get("/"), rf.get("/")
        normale.user = User.objects.create_user("rep_normale", "rn@example.invalid", "x")
        capo.user = User.objects.create_superuser("rep_hr", "rh@example.invalid", "x")
        for key in ("costo_orario_persone", "costo_orario_ripartito", "costo_orario_andamento"):
            self.assertFalse(catalogo.get(key).consentita(normale))
            self.assertTrue(catalogo.get(key).consentita(capo))
        from .reportistica.conversazione import catalogo_per_ai

        chiavi = {s["chiave"] for s in catalogo_per_ai(capo, {})["sezioni"]}
        self.assertFalse(chiavi & {"costo_orario_persone", "costo_orario_ripartito", "costo_orario_andamento"})
        modello = ReportModello.objects.get(codice_sistema="costo_personale")
        self.assertEqual(modello.riservatezza, "DATI_PERSONALI")
        self.assertEqual(modello.blocchi.count(), 5)