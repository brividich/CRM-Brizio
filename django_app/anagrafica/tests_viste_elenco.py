"""Colonne configurabili dell'elenco dipendenti (prompt 03, obiettivo A).

Dati sintetici nella tabella legacy (helper condivisi dei test export). Le colonne sensibili non
arrivano mai a chi non ha il permesso: né a video (non sono nell'HTML), né in export, né forzandole
dalla querystring o mettendole in una vista condivisa.
"""
from __future__ import annotations

import json
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.messages.storage.fallback import FallbackStorage
from django.contrib.sessions.middleware import SessionMiddleware
from django.db import connection
from django.test import RequestFactory, TestCase, override_settings
from django.test.utils import CaptureQueriesContext

from core.test_excel_export import assert_no_live_formula

from . import views, views_elenco
from .exports import build_export_response
from .models import DipendenteAnagraficaAziendale, DipendenteAnagraficaCivile
from .models_elenco import PreferenzaVistaUtente, VistaElenco
from .services import colonne_elenco, viste_elenco
from .tests import _ensure_anagrafica_table
from .tests_exports_persone import _insert_dipendente

User = get_user_model()


def _req(user, path="/anagrafica/dipendenti/", data=None, method="get"):
    request = getattr(RequestFactory(), method)(path, data=data or {})
    request.user = user
    SessionMiddleware(lambda r: None).process_request(request)
    request.session.save()
    request._messages = FallbackStorage(request)
    return request


@override_settings(LEGACY_AUTH_ENABLED=False, SECURE_SSL_REDIRECT=False)
class VisteElencoBase(TestCase):
    @classmethod
    def setUpTestData(cls):
        _ensure_anagrafica_table()
        cls.admin = User.objects.create_superuser("admin_viste", "admin_viste@example.invalid", "x")
        cls.utente = User.objects.create_user("utente_viste", "utente_viste@example.invalid", "x")
        cls.id_a = _insert_dipendente(7001, "Mario", "Rossi", reparto="PRODUZIONE", mansione="=CMD()", email_notifica="m.r@example.invalid")
        cls.id_b = _insert_dipendente(7002, "Anna", "Bianchi", reparto="UFFICIO", mansione="Impiegata")
        DipendenteAnagraficaCivile.objects.create(legacy_anagrafica_id=cls.id_a, codice_fiscale="RSSMRA80A01H501U")
        DipendenteAnagraficaAziendale.objects.create(legacy_anagrafica_id=cls.id_a, badge="B-001", tipologia_contratto="IND")


class RegistroEPermessiTests(VisteElencoBase):
    def test_registro_dichiara_permesso_e_sensibile(self):
        cf = colonne_elenco.REGISTRO["codice_fiscale"]
        self.assertEqual((cf.permesso, cf.sensibile), ("hr", True))
        self.assertTrue(all(c.gruppo for c in colonne_elenco.REGISTRO.values()))
        self.assertGreaterEqual(len(colonne_elenco.REGISTRO), 25)

    def test_colonna_sensibile_forzata_da_querystring_sparisce_senza_permesso(self):
        richieste = viste_elenco.colonne_richieste(_req(self.utente, data={"colonne": "reparto,codice_fiscale,chiave_inventata"}), None)
        self.assertEqual([c["colonna"].chiave for c in richieste], ["reparto"])
        concesse = viste_elenco.colonne_richieste(_req(self.admin, data={"colonne": "reparto,codice_fiscale"}), None)
        self.assertEqual([c["colonna"].chiave for c in concesse], ["reparto", "codice_fiscale"])

    def test_vista_condivisa_con_colonna_sensibile_non_la_espone_nell_html(self):
        vista = VistaElenco.objects.create(codice="cf", nome="Con CF", colonne=viste_elenco.specifica(["reparto", "codice_fiscale"]))
        negato = views.dipendenti_list(_req(self.utente, data={"vista": vista.codice})).content.decode()
        self.assertNotIn("RSSMRA80A01H501U", negato)
        self.assertNotIn("Codice fiscale", negato)
        concesso = views.dipendenti_list(_req(self.admin, data={"vista": vista.codice})).content.decode()
        self.assertIn("RSSMRA80A01H501U", concesso)


class VisteEPreferenzeTests(VisteElencoBase):
    def test_preset_creati_dalla_migrazione(self):
        codici = set(VistaElenco.objects.filter(preset=True).values_list("codice", flat=True))
        self.assertTrue({"standard", "iso9001", "iso45001", "hr", "sicurezza", "altro"} <= codici)
        self.assertTrue(VistaElenco.objects.get(codice="standard").predefinita)

    def test_vista_scelta_ricordata_per_utente(self):
        hr = VistaElenco.objects.get(codice="hr")
        self.assertEqual(viste_elenco.vista_corrente(_req(self.utente, data={"vista": "hr"})), hr)
        self.assertEqual(PreferenzaVistaUtente.objects.get(user=self.utente).vista, hr)
        self.assertEqual(viste_elenco.vista_corrente(_req(self.utente)), hr)

    def test_predefinita_di_ruolo_prima_della_globale(self):
        ruolo = VistaElenco.objects.create(codice="ruolo-x", nome="Ruolo X", ambito=VistaElenco.AMBITO_RUOLO, ruolo_legacy_id=9, predefinita=True,
                                           colonne=viste_elenco.specifica(["mansione"]))
        with patch("anagrafica.services.viste_elenco._ruolo_legacy_id", return_value=9):
            self.assertEqual(viste_elenco.vista_corrente(_req(self.utente)), ruolo)
        self.assertEqual(viste_elenco.vista_corrente(_req(self.utente)).codice, "standard")

    def test_elenco_mostra_selettore_e_colonne_della_vista(self):
        html = views.dipendenti_list(_req(self.admin, data={"vista": "hr"})).content.decode()
        self.assertIn('data-table-id="anagrafica.dipendenti.list:hr"', html)
        self.assertIn("Tipo contratto", html)
        self.assertIn('name="vista"', html)


class FiltriOrdinamentoTests(VisteElencoBase):
    def _righe(self, html, *nomi):
        return [html.index(n) for n in nomi]

    def test_filtro_della_vista_solo_all_apertura_e_reset_lo_toglie(self):
        VistaElenco.objects.create(codice="prod", nome="Produzione", filtri={"reparto": "PRODUZIONE"}, colonne=viste_elenco.specifica(["reparto"]))
        aperta = views.dipendenti_list(_req(self.admin, data={"vista": "prod"})).content.decode()
        self.assertNotIn("UFFICIO", aperta.split("<tbody", 1)[1])
        reset = views.dipendenti_list(_req(self.admin, data={"vf": "0"})).content.decode()
        self.assertIn("UFFICIO", reset.split("<tbody", 1)[1])
        tutti = views.dipendenti_list(_req(self.admin, data={"reparto": ""})).content.decode()
        self.assertIn("UFFICIO", tutti.split("<tbody", 1)[1])

    def test_ordinamento_della_vista(self):
        VistaElenco.objects.create(codice="ord", nome="Ordinata", ordinamento="reparto", colonne=viste_elenco.specifica(["reparto"]))
        corpo = views.dipendenti_list(_req(self.admin, data={"vista": "ord"})).content.decode().split("<tbody", 1)[1]
        self.assertLess(corpo.index("PRODUZIONE"), corpo.index("UFFICIO"))

    def test_colonne_duplicate_e_vista_bloccata(self):
        self.assertEqual(len(viste_elenco.specifica(["reparto", "reparto"])), 1)
        bloccata = VistaElenco.objects.create(codice="lock", nome="Bloccata", blocca_personalizzazione=True,
                                              colonne=viste_elenco.specifica(["mansione"]))
        cols = viste_elenco.colonne_richieste(_req(self.admin, data={"colonne": "reparto"}), bloccata)
        self.assertEqual([c["colonna"].chiave for c in cols], ["mansione"])

    def test_form_vista_con_json_malformato_non_va_in_errore(self):
        for grezzo in ("non json", json.dumps(["x", 3, None]), json.dumps({"a": 1})):
            with self.subTest(grezzo=grezzo):
                r = views_elenco.vista_elenco_modifica(_req(self.admin, path="/x/", method="post",
                                                            data={"nome": "", "colonne_json": grezzo, "ambito": "GLOBALE"}))
                self.assertEqual(r.status_code, 200)

    def test_vista_eliminata_torna_alla_predefinita(self):
        hr = VistaElenco.objects.get(codice="hr")
        viste_elenco.vista_corrente(_req(self.utente, data={"vista": "hr"}))
        hr.delete()
        self.assertEqual(viste_elenco.vista_corrente(_req(self.utente)).codice, "standard")


class QueryCostantiTests(VisteElencoBase):
    def test_numero_di_query_non_cresce_con_righe_e_colonne(self):
        tutte = list(colonne_elenco.REGISTRO.values())
        pochi = [{"id": 7001, "utente_id": None}, {"id": 7002, "utente_id": None}]
        molti = [{"id": 8000 + n, "utente_id": None} for n in range(200)]
        for righe in (pochi, molti):
            for r in righe:
                r["ruoli"] = []
        with CaptureQueriesContext(connection) as a:
            colonne_elenco.arricchisci(pochi, tutte)
        with CaptureQueriesContext(connection) as b:
            colonne_elenco.arricchisci(molti, tutte)
        self.assertEqual(len(a), len(b))
        self.assertGreaterEqual(len(tutte), 25)


class ExportVistaTests(VisteElencoBase):
    def test_export_rispetta_vista_permessi_e_formule(self):
        response = build_export_response(_req(self.utente, path="/anagrafica/export/dipendenti/", data={"colonne": "mansione,codice_fiscale"}),
                                          "dipendenti", "xlsx", "full")
        assert_no_live_formula(self, response.content)
        from openpyxl import load_workbook
        valori = [str(c.value) for row in load_workbook(BytesIO(response.content)).active.iter_rows() for c in row if c.value]
        self.assertIn("Mansione", valori)
        self.assertNotIn("Codice fiscale", valori)
        self.assertNotIn("RSSMRA80A01H501U", valori)
        self.assertIn("=CMD()", valori)  # testo, non formula viva

    def test_audit_dell_export_dichiara_le_colonne_sensibili(self):
        from core.models import AuditLog

        build_export_response(_req(self.admin, path="/anagrafica/export/dipendenti/", data={"colonne": "mansione,codice_fiscale"}),
                              "dipendenti", "xlsx", "full")
        log = AuditLog.objects.filter(azione="export", modulo="anagrafica").order_by("-id").first()
        self.assertIn("v:codice_fiscale", log.dettaglio["colonne"])
        self.assertEqual(log.dettaglio["colonne_sensibili"], ["codice_fiscale"])


class AdminVisteTests(VisteElencoBase):
    def test_form_accetta_solo_chiavi_del_registro_e_predefinita_unica(self):
        dati = {"nome": "Mia vista", "descrizione": "", "ambito": "GLOBALE", "predefinita": "on", "ordine": 50, "attiva": "on",
                "colonne_json": json.dumps([{"chiave": "reparto", "visibile": True}, {"chiave": "__class__", "visibile": True},
                                            {"chiave": "reparto", "visibile": True}])}
        form = views_elenco.VistaElencoForm(dati, instance=VistaElenco(), request=_req(self.admin))
        self.assertTrue(form.is_valid(), form.errors)
        vista = form.save()
        self.assertEqual(vista.chiavi, ["reparto"])
        self.assertEqual(VistaElenco.objects.filter(predefinita=True, ambito="GLOBALE").get(), vista)

    def test_elimina_richiede_conferma_e_ripristina_preset(self):
        hr = VistaElenco.objects.get(codice="hr")
        hr.colonne = viste_elenco.specifica(["reparto"])
        hr.save()
        views_elenco.vista_elenco_azione(_req(self.admin, path="/x/", data={"azione": "ripristina"}, method="post"), hr.pk)
        hr.refresh_from_db()
        self.assertIn("tipologia_contratto", hr.chiavi)
        views_elenco.vista_elenco_azione(_req(self.admin, path="/x/", data={"azione": "elimina", "conferma": "no"}, method="post"), hr.pk)
        self.assertTrue(VistaElenco.objects.filter(pk=hr.pk).exists())
        views_elenco.vista_elenco_azione(_req(self.admin, path="/x/", data={"azione": "elimina", "conferma": "hr"}, method="post"), hr.pk)
        self.assertFalse(VistaElenco.objects.filter(pk=hr.pk).exists())

    def test_pagine_admin_negate_a_chi_non_amministra(self):
        self.assertEqual(views_elenco.viste_elenco(_req(self.utente)).status_code, 302)
        negata = views_elenco.vista_elenco_anteprima(_req(self.utente))
        self.assertEqual((negata.status_code, negata["Content-Type"]), (403, "application/json"))
        self.assertContains(views_elenco.vista_elenco_anteprima(_req(self.admin, data={"colonne_json": json.dumps(viste_elenco.specifica(["reparto"]))})),
                            "PRODUZIONE")
