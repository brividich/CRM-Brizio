"""Test della reportistica v2: opzioni per sezione, nuove sezioni, impaginazione, Excel tipizzato."""
from __future__ import annotations

import io
from datetime import date, timedelta

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from anagrafica.models import ReportBlocco, ReportModello
from anagrafica.reportistica import motore
from anagrafica.reportistica import sezioni as catalogo
from anagrafica.tests_reportistica import OGGI, _ConPersone

User = get_user_model()


class OpzioniTest(TestCase):
    def test_normalizzazione_per_tipo(self):
        o_int = catalogo.Opzione("g", "G", catalogo.INTERO, predefinito=30, minimo=1, massimo=90)
        self.assertEqual(o_int.normalizza("500"), 90)
        self.assertEqual(o_int.normalizza("x"), 30)
        o_si = catalogo.Opzione("s", "S", catalogo.SI_NO, predefinito=True)
        self.assertTrue(o_si.normalizza(["1"]))
        self.assertFalse(o_si.normalizza([]))
        o_multi = catalogo.Opzione("m", "M", catalogo.MULTI, (("a", "A"), ("b", "B")), ("a",))
        self.assertEqual(o_multi.normalizza(["b", "z"]), ["b"])
        self.assertEqual(o_multi.normalizza(None), ["a"])
        o_scelta = catalogo.Opzione("c", "C", catalogo.SCELTA, (("x", "X"), ("y", "Y")), "x")
        self.assertEqual(o_scelta.normalizza("nonvalido"), "x")

    def test_opzioni_generali_presenti_solo_dove_hanno_senso(self):
        statica = {o.nome for o in catalogo.get("personale_elenco").tutte_le_opzioni()}
        self.assertTrue({"ordina_per", "raggruppa_per", "solo_criticita", "max_righe"} <= statica)
        matrice = {o.nome for o in catalogo.get("matrice_qualifiche").tutte_le_opzioni()}
        self.assertNotIn("ordina_per", matrice)
        self.assertIn("solo_criticita", matrice)


class GranularitaTest(_ConPersone):
    def setUp(self):
        super().setUp()
        from anagrafica.models import DipendenteQualifica, TipoQualifica

        self.patentino = TipoQualifica.objects.create(nome="Patentino sintetico", categoria="PROFESSIONALE")
        self.ple = TipoQualifica.objects.create(nome="PLE sintetica", categoria="SICUREZZA")
        DipendenteQualifica.objects.create(legacy_anagrafica_id=901, tipo=self.patentino, numero="S-1",
                                           data_conseguimento=OGGI - timedelta(days=300),
                                           data_scadenza=OGGI + timedelta(days=20))
        DipendenteQualifica.objects.create(legacy_anagrafica_id=902, tipo=self.ple, numero="S-2",
                                           data_conseguimento=OGGI - timedelta(days=900),
                                           data_scadenza=OGGI - timedelta(days=3))
        DipendenteQualifica.objects.create(legacy_anagrafica_id=903, tipo=self.ple, numero="S-3",
                                           data_conseguimento=OGGI - timedelta(days=100),
                                           data_scadenza=OGGI + timedelta(days=700))

    def test_filtri_di_sezione_qualifiche(self):
        sez = catalogo.get("personale_qualifiche")
        tutte = sez.calcola(self.ctx())
        self.assertEqual(len(tutte.righe), 3)
        solo_ple = sez.calcola(self.ctx(), {"tipi": [str(self.ple.pk)]})
        self.assertEqual({r["numero"] for r in solo_ple.righe}, {"S-2", "S-3"})
        scadute = sez.calcola(self.ctx(), {"stati": ["scaduta"]})
        self.assertEqual([r["numero"] for r in scadute.righe], ["S-2"])
        # Preavviso corto: la scadenza a 20 giorni non è più «in scadenza».
        corto = sez.calcola(self.ctx(), {"preavviso": 10})
        self.assertEqual(next(r for r in corto.righe if r["numero"] == "S-1")["stato"], "Valida")
        self.assertIsInstance(tutte.righe[0]["scadenza"], date, "valori grezzi: date vere per l'Excel")

    def test_matrice_qualifiche_toni_per_cella(self):
        res = catalogo.get("matrice_qualifiche").calcola(self.ctx())
        self.assertEqual([lbl for _k, lbl in res.colonne][2:], ["Patentino sintetico", "PLE sintetica"])
        bianchi = next(r for r in res.righe if r["nominativo"] == "BIANCHI LUCA")
        self.assertEqual(bianchi[f"q{self.ple.pk}"], "X")
        self.assertEqual(bianchi["_toni"][f"q{self.ple.pk}"], "danger")
        self.assertEqual(bianchi[f"q{self.patentino.pk}"], "—")
        con_date = catalogo.get("matrice_qualifiche").calcola(self.ctx(), {"cella": "scadenza"})
        verdi = next(r for r in con_date.righe if r["nominativo"] == "VERDI GIULIA")
        self.assertEqual(verdi[f"q{self.ple.pk}"], OGGI + timedelta(days=700))

    def test_scadenzario_orizzonte(self):
        sez = catalogo.get("scadenzario_unico")
        res = sez.calcola(self.ctx(), {"giorni": 30, "tipi": ["qualifiche"]})
        self.assertEqual([r["voce"] for r in res.righe], ["PLE sintetica", "Patentino sintetico"])
        senza_scadute = sez.calcola(self.ctx(), {"giorni": 30, "tipi": ["qualifiche"], "includi_scadute": False})
        self.assertEqual([r["voce"] for r in senza_scadute.righe], ["Patentino sintetico"])

    def test_perimetro_esteso(self):
        self.assertEqual({d.id for d in self.ctx(contratti=["DETERMINATO"]).dipendenti()}, {902})
        self.assertEqual({d.id for d in self.ctx(escludi=[901]).dipendenti()}, {902, 903})
        self.assertEqual({d.id for d in self.ctx(livelli=["3"]).dipendenti()}, {901})
        self.assertEqual({d.id for d in self.ctx(qualifiche=[self.ple.pk]).dipendenti()}, {903})
        recenti = self.ctx(assunti_dal=(OGGI - timedelta(days=100)).isoformat()).dipendenti()
        self.assertEqual({d.id for d in recenti}, {902})

    def test_soglia_anonimato(self):
        res = catalogo.get("organico_indicatori").calcola(self.ctx(), {"dimensioni": ["livello"], "soglia_anonimato": 3})
        self.assertTrue(res.righe)
        self.assertTrue(all(r["n"] == "<3" for r in res.righe))


class MotoreGranulareTest(_ConPersone):
    def setUp(self):
        super().setUp()
        self.richiesta = type("R", (), {"user": User.objects.create_superuser("rp_mot", "rp_mot@example.com", "x")})()

    def _modello(self, opzioni: dict, sezione: str = "personale_elenco", **campi) -> ReportModello:
        m = ReportModello.objects.create(nome="Prova", titolo_documento="Prova", **campi)
        ReportBlocco.objects.create(modello=m, ordine=10, tipo="SEZIONE", sezione=sezione, opzioni=opzioni)
        return m

    def _doc(self, m):
        return motore.componi(m, motore.Parametri.dal_modello(m), self.richiesta)

    def test_ordinamento_raggruppamento_e_limite(self):
        m = self._modello({"colonne": ["nominativo", "reparto"],
                           "valori": {"ordina_per": "nominativo", "ordine": "desc", "raggruppa_per": "reparto"}})
        e = self._doc(m).sezioni_calcolate[0]
        self.assertEqual([g.etichetta for g in e.gruppi], ["Produzione", "Qualità"])
        self.assertEqual([r.valori[0] for r in e.gruppi[0].righe], ["ROSSI ANNA", "BIANCHI LUCA"])
        self.assertEqual([k for k, _l in e.colonne], ["nominativo"], "la colonna di raggruppamento esce dalla tabella")
        e2 = self._doc(self._modello({"valori": {"max_righe": 2}})).sezioni_calcolate[0]
        self.assertEqual(e2.n_righe, 2)
        self.assertIn("Mostrate le prime 2 righe su 3.", e2.note)

    def test_solo_criticita_e_sezione_vuota_omessa(self):
        m = self._modello({"valori": {"solo_criticita": True, "nascondi_se_vuota": True}})
        self.assertEqual(self._doc(m).sezioni_calcolate, [])

    def test_escludi_blocco_in_generazione(self):
        m = self._modello({})
        parametri = motore.Parametri.dal_modello(m)
        parametri.escludi_blocchi = {m.blocchi.first().pk}
        self.assertEqual(motore.componi(m, parametri, self.richiesta).elementi, [])

    def test_pdf_verticale_con_filigrana_e_indice(self):
        m = self._modello({}, orientamento="VERTICALE", filigrana="BOZZA", mostra_indice=True,
                          codice_documento="MOD.999", revisione="2", piede_pagina="{{codice}} rev. {{revisione}}")
        doc = self._doc(m)
        self.assertEqual(doc.intestazione, "MOD.999 · Rev. 2")
        self.assertEqual(doc.piede_pagina, "MOD.999 rev. 2")
        self.assertTrue(motore.render_pdf(doc).startswith(b"%PDF"))

    def test_excel_tipizzato_con_fogli(self):
        from openpyxl import load_workbook

        m = self._modello({"colonne": ["nominativo", "assunzione", "anzianita"]})
        wb = load_workbook(io.BytesIO(motore.render_xlsx(self._doc(m))))
        self.assertEqual(wb.sheetnames, ["Documento", "Indicatori", "Elenco del personale"])
        ws = wb["Elenco del personale"]
        intestazione = next(r for r in range(1, 12) if ws.cell(row=r, column=1).value == "Nominativo")
        assunzione = ws.cell(row=intestazione + 1, column=2)
        self.assertTrue(assunzione.is_date)
        self.assertEqual(assunzione.number_format, "DD/MM/YYYY")
        self.assertIsInstance(ws.cell(row=intestazione + 1, column=3).value, int)

    def test_matrice_in_excel_con_colori_per_cella(self):
        from openpyxl import load_workbook

        from anagrafica.models import DipendenteQualifica, TipoQualifica

        tipo = TipoQualifica.objects.create(nome="Qualifica matrice", categoria="ALTRO")
        DipendenteQualifica.objects.create(legacy_anagrafica_id=902, tipo=tipo, data_scadenza=OGGI - timedelta(days=1))
        m = self._modello({}, sezione="matrice_qualifiche", excel_foglio_documento=False, excel_foglio_indicatori=False)
        doc = self._doc(m)
        self.assertTrue(doc.sezioni_calcolate[0].matrice)
        ws = load_workbook(io.BytesIO(motore.render_xlsx(doc))).active
        celle = [c for row in ws.iter_rows() for c in row if c.value == "X"]
        self.assertEqual(len(celle), 1)
        self.assertEqual(celle[0].fill.fgColor.rgb[-6:], "FDE8E8")

    def test_nome_file_da_schema(self):
        m = self._modello({}, nome_file="{codice}_{destinatario}", codice_documento="MOD.230")
        parametri = motore.Parametri.dal_modello(m)
        parametri.destinatario = "Cliente Demo S.p.A."
        doc = motore.componi(m, parametri, self.richiesta)
        nome = motore.nome_file(m, doc, "xlsx")
        self.assertTrue(nome.startswith("mod230_cliente-demo-spa_"), nome)
        self.assertTrue(nome.endswith(".xlsx"))


class EditorOpzioniTest(_ConPersone):
    def setUp(self):
        super().setUp()
        self.admin = User.objects.create_superuser("rp_ed", "rp_ed@example.com", "x")
        self.client.force_login(self.admin)

    def test_opzioni_di_sezione_salvate_dal_form(self):
        dati = {
            "nome": "Opzioni", "titolo_documento": "X", "riservatezza": "INTERNO", "periodo_tipo": "ULTIMI_12_MESI",
            "orientamento": "VERTICALE", "formato_predefinito": "xlsx", "mostra_frontespizio": "on",
            "contratti": ["INDETERMINATO"],
            "blocchi-TOTAL_FORMS": "1", "blocchi-INITIAL_FORMS": "0",
            "blocchi-MIN_NUM_FORMS": "0", "blocchi-MAX_NUM_FORMS": "1000",
            "blocchi-0-tipo": "SEZIONE", "blocchi-0-ordine": "10", "blocchi-0-sezione": "sicurezza_sorveglianza",
            "blocchi-0-o__sicurezza_sorveglianza__preavviso": "45",
            "blocchi-0-o__sicurezza_sorveglianza__stati": ["scaduta", "in_scadenza"],
            "blocchi-0-o__sicurezza_sorveglianza__solo_criticita": "1",
            # Opzioni di un'altra sezione (gruppo nascosto nel form): vanno ignorate.
            "blocchi-0-o__personale_elenco__max_righe": "7",
        }
        resp = self.client.post(reverse("anagrafica:reportistica_modello_create"), dati)
        self.assertEqual(resp.status_code, 302)
        m = ReportModello.objects.get(nome="Opzioni")
        self.assertEqual((m.orientamento, m.formato_predefinito), ("VERTICALE", "xlsx"))
        self.assertEqual(m.filtri["contratti"], ["INDETERMINATO"])
        valori = m.blocchi.get().opzioni["valori"]
        self.assertEqual(valori["preavviso"], 45)
        self.assertEqual(valori["stati"], ["scaduta", "in_scadenza"])
        self.assertTrue(valori["solo_criticita"])
        self.assertFalse(valori["includi_senza_visita"], "checkbox non inviata = disattivata")
        self.assertEqual(valori["max_righe"], 0)
        pagina = self.client.get(reverse("anagrafica:reportistica_modello_edit", args=[m.pk]))
        self.assertContains(pagina, 'name="blocchi-0-o__sicurezza_sorveglianza__preavviso" value="45"')

    def test_predefiniti_nuovi_e_blocchi_esclusi(self):
        self.assertTrue(ReportModello.objects.filter(codice_sistema="matrice_commessa").exists())
        scad = ReportModello.objects.get(codice_sistema="scadenzario_mensile")
        self.assertEqual(scad.blocchi.get().opzioni["valori"]["raggruppa_per"], "mese")
        url = reverse("anagrafica:reportistica_genera", args=[scad.pk])
        self.assertEqual(self.client.get(url).status_code, 200)
        # Tutti i blocchi tolti (form inviato senza spunte): documento senza sezioni.
        resp = self.client.post(url, {"titolo": "S", "periodo_tipo": "ULTIMI_12_MESI", "azione": "anteprima",
                                      "blocchi_inviati": "1"})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context["documento"].elementi, [])
