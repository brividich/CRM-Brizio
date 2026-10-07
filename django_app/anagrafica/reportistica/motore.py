"""Composizione e rendering dei documenti (anteprima HTML, PDF, Excel).

Un documento si compone una sola volta (:func:`componi`) e lo stesso oggetto
alimenta anteprima, PDF ed Excel: i numeri non possono divergere fra i formati.
Le righe restano con i valori *grezzi* (date, numeri) fino all'uscita: il PDF e
l'anteprima li formattano all'italiana, l'Excel li scrive tipizzati.

Testi liberi: markdown minimo e prevedibile — ``## titolo``, righe con ``-`` per
gli elenchi, ``**grassetto**``, riga vuota per andare a capo — piu' i segnaposto
``{{oggi}}``, ``{{periodo}}``, ``{{destinatario}}``, ``{{n_persone}}``…
"""
from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from html import escape

from django.utils import timezone
from django.utils.text import slugify

from . import sezioni as catalogo
from .dati import Contesto, Perimetro, periodo_da_tipo

logger = logging.getLogger(__name__)

SEGNAPOSTO = (
    ("{{oggi}}", "data di generazione"),
    ("{{periodo}}", "periodo del report"),
    ("{{periodo_da}}", "inizio periodo"),
    ("{{periodo_a}}", "fine periodo"),
    ("{{destinatario}}", "destinatario"),
    ("{{n_persone}}", "persone nel perimetro"),
    ("{{perimetro}}", "descrizione del perimetro"),
    ("{{norme}}", "norme di riferimento"),
    ("{{codice}}", "codice documento"),
    ("{{revisione}}", "revisione"),
    ("{{redatto_da}}", "redattore"),
    ("{{approvato_da}}", "approvatore"),
)
SEGNAPOSTO_FILE = ("{titolo}", "{destinatario}", "{data}", "{codice}", "{revisione}", "{modello}")

_BOLD = re.compile(r"\*\*(.+?)\*\*")
TONI_EVIDENZA = ("warn", "danger")


def _inline(text: str) -> str:
    """Testo -> markup inline sicuro (escape + solo <b>), valido per HTML e ReportLab."""
    return _BOLD.sub(r"<b>\1</b>", escape(text, quote=False))


def paragrafi(testo: str) -> list[tuple[str, str]]:
    """Markdown minimo -> [(tipo, markup)] con tipo in {"h", "p", "li"}."""
    out: list[tuple[str, str]] = []
    buffer: list[str] = []

    def _flush() -> None:
        if buffer:
            out.append(("p", "<br/>".join(_inline(x) for x in buffer)))
            buffer.clear()

    for raw in (testo or "").replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if not line:
            _flush()
        elif line.startswith("#"):
            _flush()
            out.append(("h", _inline(line.lstrip("#").strip())))
        elif line[:2] in ("- ", "* ", "• "):
            _flush()
            out.append(("li", _inline(line[2:].strip())))
        else:
            buffer.append(line)
    _flush()
    return out


def fmt(valore) -> str:
    """Valore grezzo -> testo all'italiana (anteprima e PDF)."""
    if valore is None:
        return ""
    if isinstance(valore, bool):
        return "Sì" if valore else "No"
    if isinstance(valore, datetime):
        return timezone.localtime(valore).strftime("%d/%m/%Y %H:%M") if timezone.is_aware(valore) else valore.strftime("%d/%m/%Y %H:%M")
    if isinstance(valore, date):
        return valore.strftime("%d/%m/%Y")
    if isinstance(valore, (Decimal, float)):
        testo = f"{Decimal(str(valore)).quantize(Decimal('0.1'))}"
        return testo[:-2] if testo.endswith(".0") else testo.replace(".", ",")
    return str(valore)


def _chiave_ordinamento(valore):
    """Ordina insieme numeri, date e testi senza TypeError; i vuoti in fondo."""
    if valore is None or valore == "":
        return (2, "")
    if isinstance(valore, (int, float, Decimal)) and not isinstance(valore, bool):
        return (0, float(valore))
    if isinstance(valore, date):
        return (0, valore.toordinal())
    return (1, str(valore).casefold())


@dataclass
class Riga:
    valori: list
    tono: str = ""
    toni_celle: list[str] = field(default_factory=list)

    @property
    def celle(self) -> list[tuple[str, str]]:
        toni = self.toni_celle or [""] * len(self.valori)
        return [(fmt(v), toni[i] if i < len(toni) else "") for i, v in enumerate(self.valori)]


@dataclass
class Gruppo:
    etichetta: str
    righe: list[Riga]


@dataclass
class Elemento:
    tipo: str  # testo | sezione | pagina | firme | negata | errore
    blocco_id: int | None = None
    titolo: str = ""
    paragrafi: list[tuple[str, str]] = field(default_factory=list)
    sezione: object = None
    risultato: object = None
    colonne: list[tuple[str, str]] = field(default_factory=list)
    gruppi: list[Gruppo] = field(default_factory=list)
    mostra_indicatori: bool = True
    mostra_tabella: bool = True
    mostra_note: bool = True
    riferimenti: str = ""
    note_extra: list[str] = field(default_factory=list)
    matrice: bool = False

    @property
    def righe(self) -> list[Riga]:
        return [r for g in self.gruppi for r in g.righe]

    @property
    def n_righe(self) -> int:
        return sum(len(g.righe) for g in self.gruppi)

    @property
    def note(self) -> list[str]:
        base = list(self.risultato.note) if self.risultato is not None and self.mostra_note else []
        return base + self.note_extra


@dataclass
class Documento:
    titolo: str
    sottotitolo: str
    destinatario: str
    riservatezza: str
    norme: list[str]
    periodo_label: str
    perimetro_label: str
    date_from: date
    date_to: date
    generato_il: object
    generato_da: str
    redatto_da: str = ""
    verificato_da: str = ""
    approvato_da: str = ""
    codice: str = ""
    revisione: str = ""
    orientamento: str = "ORIZZONTALE"
    filigrana: str = ""
    piede_pagina: str = ""
    mostra_frontespizio: bool = True
    mostra_indice: bool = False
    excel_foglio_indicatori: bool = True
    excel_foglio_documento: bool = True
    elementi: list[Elemento] = field(default_factory=list)

    @property
    def contiene_dati_personali(self) -> bool:
        return any(
            e.tipo == "sezione" and e.mostra_tabella and e.n_righe and getattr(e.sezione, "nominativa", True)
            for e in self.elementi
        )

    @property
    def sezioni_calcolate(self) -> list[Elemento]:
        return [e for e in self.elementi if e.tipo == "sezione"]

    @property
    def indice(self) -> list[str]:
        return [e.titolo for e in self.elementi if e.titolo and e.tipo in ("testo", "sezione", "firme")]

    @property
    def intestazione(self) -> str:
        parti = [p for p in (self.codice, f"Rev. {self.revisione}" if self.revisione else "") if p]
        return " · ".join(parti)


@dataclass
class Parametri:
    """Scelte fatte in generazione: prevalgono su quelle salvate nel modello."""

    periodo_tipo: str
    data_da: date | None
    data_a: date | None
    perimetro: Perimetro
    titolo: str = ""
    sottotitolo: str = ""
    destinatario: str = ""
    testi: dict[int, str] = field(default_factory=dict)
    escludi_blocchi: set[int] = field(default_factory=set)

    @classmethod
    def dal_modello(cls, modello) -> "Parametri":
        return cls(
            periodo_tipo=modello.periodo_tipo, data_da=modello.data_da, data_a=modello.data_a,
            perimetro=Perimetro.from_dict(modello.filtri), titolo=modello.titolo_documento,
            sottotitolo=modello.sottotitolo, destinatario=modello.destinatario,
            testi={b.pk: b.testo for b in modello.blocchi.all() if b.tipo in ("TESTO", "SEZIONE")},
        )

    def as_dict(self) -> dict:
        return {
            "periodo_tipo": self.periodo_tipo,
            "data_da": self.data_da.isoformat() if self.data_da else "",
            "data_a": self.data_a.isoformat() if self.data_a else "",
            "perimetro": self.perimetro.as_dict(),
            "titolo": self.titolo, "destinatario": self.destinatario,
            "escludi_blocchi": sorted(self.escludi_blocchi),
        }


def _nome_utente(user) -> str:
    if user is None:
        return ""
    return ((getattr(user, "get_full_name", lambda: "")() or getattr(user, "username", "")) or "").strip()


def componi(modello, parametri: Parametri, request, *, blocchi=None) -> Documento:
    """Documento dal modello. ``blocchi`` (non salvati) sostituisce i blocchi del modello:
    serve ai report a conversazione, che non creano record finche' non li si salva."""
    oggi = timezone.localdate()
    date_from, date_to = periodo_da_tipo(parametri.periodo_tipo, oggi=oggi,
                                         data_da=parametri.data_da, data_a=parametri.data_a)
    ctx = Contesto(date_from=date_from, date_to=date_to, perimetro=parametri.perimetro, request=request, today=oggi)

    nomi = {d.id: d.nominativo for d in ctx._tutti()} if parametri.perimetro.persone else {}
    perimetro_label = parametri.perimetro.label(nomi)
    norme = [n for n in (modello.norme or []) if isinstance(n, str)]
    sostituzioni = {
        "{{oggi}}": oggi.strftime("%d/%m/%Y"),
        "{{periodo}}": ctx.periodo_label,
        "{{periodo_da}}": date_from.strftime("%d/%m/%Y"),
        "{{periodo_a}}": date_to.strftime("%d/%m/%Y"),
        "{{destinatario}}": parametri.destinatario or "",
        "{{perimetro}}": perimetro_label,
        "{{norme}}": ", ".join(norme),
        "{{codice}}": modello.codice_documento or "",
        "{{revisione}}": modello.revisione or "",
        "{{redatto_da}}": modello.redatto_da or "",
        "{{approvato_da}}": modello.approvato_da or "",
    }

    def _sost(text: str) -> str:
        text = text or ""
        if "{{n_persone}}" in text:
            text = text.replace("{{n_persone}}", str(len(ctx.dipendenti())))
        for chiave, valore in sostituzioni.items():
            text = text.replace(chiave, valore)
        return text

    doc = Documento(
        titolo=_sost(parametri.titolo or modello.titolo_documento or modello.nome),
        sottotitolo=_sost(parametri.sottotitolo),
        destinatario=parametri.destinatario,
        riservatezza=modello.get_riservatezza_display(),
        norme=norme,
        periodo_label=ctx.periodo_label,
        perimetro_label=perimetro_label,
        date_from=date_from, date_to=date_to,
        generato_il=timezone.localtime(),
        generato_da=_nome_utente(getattr(request, "user", None)),
        redatto_da=modello.redatto_da, verificato_da=modello.verificato_da, approvato_da=modello.approvato_da,
        codice=modello.codice_documento, revisione=modello.revisione, orientamento=modello.orientamento,
        filigrana=modello.filigrana, piede_pagina=_sost(modello.piede_pagina),
        mostra_frontespizio=modello.mostra_frontespizio, mostra_indice=modello.mostra_indice,
        excel_foglio_indicatori=modello.excel_foglio_indicatori,
        excel_foglio_documento=modello.excel_foglio_documento,
    )

    for blocco in (blocchi if blocchi is not None else modello.blocchi.all()):
        if blocco.pk is not None and blocco.pk in parametri.escludi_blocchi:
            continue
        testo = parametri.testi.get(blocco.pk, blocco.testo)
        if blocco.tipo == "TESTO":
            doc.elementi.append(Elemento("testo", blocco.pk, _sost(blocco.titolo), paragrafi(_sost(testo))))
        elif blocco.tipo == "PAGINA":
            doc.elementi.append(Elemento("pagina", blocco.pk))
        elif blocco.tipo == "FIRME":
            doc.elementi.append(Elemento("firme", blocco.pk, _sost(blocco.titolo) or "Approvazione"))
        elif blocco.tipo == "SEZIONE":
            elemento = _elemento_sezione(blocco, testo, ctx, request, _sost, modello)
            if elemento is not None:
                doc.elementi.append(elemento)
    return doc


def _elemento_sezione(blocco, commento: str, ctx: Contesto, request, sost, modello) -> Elemento | None:
    sezione = catalogo.get(blocco.sezione)
    opzioni_salvate = blocco.opzioni if isinstance(blocco.opzioni, dict) else {}
    titolo = sost(blocco.titolo) or (sezione.titolo if sezione else blocco.sezione)
    if sezione is None:
        return Elemento("errore", blocco.pk, titolo, paragrafi("Sezione non più disponibile nel catalogo."))
    if not sezione.consentita(request):
        return Elemento("negata", blocco.pk, titolo,
                        paragrafi("Sezione omessa: non hai i permessi per questi dati."), sezione=sezione)
    valori = sezione.valori_opzioni(opzioni_salvate.get("valori"))
    ctx.raccogli_avvisi()
    try:
        risultato = sezione.builder(ctx, valori)
    except Exception:
        logger.exception("reportistica: sezione %s non calcolabile", sezione.key)
        return Elemento("errore", blocco.pk, titolo,
                        paragrafi("Sezione non calcolabile in questo momento: vedi log applicativo."), sezione=sezione)

    # Colonne: scelte dall'utente (sezioni statiche) o decise dai dati (matrici, report conformità).
    if sezione.colonne and not sezione.colonne_dinamiche:
        etichette = dict(sezione.colonne)
        chiavi = sezione.colonne_effettive(opzioni_salvate.get("colonne"))
    else:
        etichette = dict(risultato.colonne)
        chiavi = [k for k, _l in risultato.colonne]

    # Opzioni generali di tabella, nello stesso ordine per tutte le sezioni.
    righe = list(zip(risultato.righe, risultato.toni))
    # Gli avvisi (dati non leggibili) restano visibili anche con le note di calcolo nascoste.
    note_extra: list[str] = ctx.raccogli_avvisi()
    if valori.get("solo_criticita"):
        righe = [(r, t) for r, t in righe
                 if t in TONI_EVIDENZA or any(v in TONI_EVIDENZA for v in (r.get("_toni") or {}).values())]
        note_extra.append("Mostrate solo le righe critiche (scadute, in scadenza, mancanti).")
    ordina = valori.get("ordina_per")
    if ordina:
        righe.sort(key=lambda rt: _chiave_ordinamento(rt[0].get(ordina)), reverse=valori.get("ordine") == "desc")
    gruppo_per = valori.get("raggruppa_per") or ""
    if gruppo_per:
        # Ordinamento stabile: prima i gruppi, dentro il gruppo resta l'ordine scelto.
        righe.sort(key=lambda rt: _chiave_ordinamento(rt[0].get(gruppo_per)))
        chiavi = [k for k in chiavi if k != gruppo_per]
    massimo = valori.get("max_righe") or 0
    if massimo and len(righe) > massimo:
        note_extra.append(f"Mostrate le prime {massimo} righe su {len(righe)}.")
        righe = righe[:massimo]
    if valori.get("nascondi_se_vuota") and not righe:
        return None

    gruppi: list[Gruppo] = []
    for r, tono in righe:
        etichetta = fmt(r.get(gruppo_per)) or "—" if gruppo_per else ""
        if not gruppi or gruppi[-1].etichetta != etichetta:
            gruppi.append(Gruppo(etichetta, []))
        toni_celle = r.get("_toni") or {}
        gruppi[-1].righe.append(Riga([r.get(k, "") for k in chiavi], tono, [toni_celle.get(k, "") for k in chiavi]))
    if not gruppi:
        gruppi = [Gruppo("", [])]

    return Elemento(
        "sezione", blocco.pk, titolo, paragrafi(sost(commento)), sezione=sezione, risultato=risultato,
        colonne=[(k, etichette.get(k, k)) for k in chiavi], gruppi=gruppi,
        mostra_indicatori=opzioni_salvate.get("mostra_indicatori", True) is not False,
        mostra_tabella=opzioni_salvate.get("mostra_tabella", True) is not False,
        mostra_note=opzioni_salvate.get("mostra_note", True) is not False and modello.mostra_note_calcolo,
        riferimenti=" · ".join(sezione.riferimenti) if modello.mostra_riferimenti else "",
        note_extra=note_extra,
        matrice=sezione.colonne_dinamiche and sezione.usa_perimetro,
    )


def nome_file(modello, documento: Documento, formato: str) -> str:
    stamp = timezone.localtime().strftime("%Y%m%d_%H%M")
    schema = (modello.nome_file or "").strip() or "{titolo}_{destinatario}_{data}"
    valori = {
        "{titolo}": documento.titolo, "{destinatario}": documento.destinatario, "{data}": stamp,
        "{codice}": documento.codice, "{revisione}": documento.revisione, "{modello}": modello.nome,
    }
    for chiave, valore in valori.items():
        schema = schema.replace(chiave, slugify(valore or "")[:60])
    base = re.sub(r"[_\-]{2,}", "_", schema).strip("_-") or "report"
    if "{data}" not in (modello.nome_file or "{data}"):
        base = f"{base}_{stamp}"
    return f"{base[:150]}.{formato}"


# ═══════════════════════════════════════════════════════════════════════════
# PDF
# ═══════════════════════════════════════════════════════════════════════════

_TONE_BG = {"danger": "#fde8e8", "warn": "#fdf3d8", "ok": "#e7f6ec"}


def render_pdf(doc: Documento) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, PageBreak, Paragraph, Spacer

    from core.pdf import PdfTheme, build_styles, data_table, header_footer_callback, make_document, section_heading
    # Riuso dei renderer di Report conformità: stessa resa dei riquadri indicatori e delle larghezze.
    from report_conformita.exports import _col_widths, _kpi_table

    theme = PdfTheme.from_branding()
    styles = build_styles(theme)
    bullet = ParagraphStyle("rp_li", parent=styles["body"], leftIndent=10, bulletIndent=2)
    sub = ParagraphStyle("rp_h", parent=styles["body"], fontName="Helvetica-Bold", spaceBefore=4)
    gruppo_st = ParagraphStyle("rp_g", parent=styles["body"], fontName="Helvetica-Bold", spaceBefore=6, spaceAfter=2)
    nota = ParagraphStyle("rp_nota", parent=styles["body"], fontSize=7, textColor=colors.HexColor("#64748b"))
    cella_c = ParagraphStyle("rp_cc", parent=styles["cell"], alignment=1)
    buf = io.BytesIO()
    pdf = make_document(buf, title=doc.titolo, landscape=doc.orientamento != "VERTICALE")
    width = pdf.pagesize[0] - pdf.leftMargin - pdf.rightMargin
    el: list = []

    def _testo(pars: list[tuple[str, str]]) -> None:
        for kind, markup in pars:
            if kind == "h":
                el.append(Paragraph(markup, sub))
            elif kind == "li":
                el.append(Paragraph(markup, bullet, bulletText="•"))
            else:
                el.append(Paragraph(markup, styles["body"]))
        if pars:
            el.append(Spacer(1, 3 * mm))

    def _tabella(e: Elemento, righe: list[Riga]):
        intestazioni = [label for _k, label in e.colonne]
        testi = [[c for c, _t in r.celle] for r in righe]
        rows = [[Paragraph(escape(h), styles["table_header"]) for h in intestazioni]]
        extra = [("FONTSIZE", (0, 0), (-1, -1), 7), ("VALIGN", (0, 0), (-1, -1), "TOP")]
        for i, r in enumerate(righe, start=1):
            stile = [cella_c if (e.matrice and j >= 2) else styles["cell"] for j in range(len(r.valori))]
            rows.append([Paragraph(escape(t).replace("\n", "<br/>"), stile[j]) for j, t in enumerate(testi[i - 1])])
            if r.tono in TONI_EVIDENZA and not e.matrice:
                extra.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor(_TONE_BG[r.tono])))
            for j, tono in enumerate(r.toni_celle):
                if tono in _TONE_BG:
                    extra.append(("BACKGROUND", (j, i), (j, i), colors.HexColor(_TONE_BG[tono])))
        return data_table(rows, theme, col_widths=_col_widths(intestazioni, testi, width), repeat_rows=1, extra_style=extra)

    if doc.mostra_frontespizio:
        el.append(Paragraph(escape(doc.titolo), styles["title"]))
        if doc.sottotitolo:
            el.append(Paragraph(escape(doc.sottotitolo), styles["subtitle"]))
        el.append(Spacer(1, 4 * mm))
        info = [("Destinatario", doc.destinatario or "—"), ("Periodo di riferimento", doc.periodo_label),
                ("Perimetro", doc.perimetro_label), ("Norme di riferimento", ", ".join(doc.norme) or "—")]
        if doc.intestazione:
            info.append(("Documento", doc.intestazione))
        info += [("Classificazione", doc.riservatezza),
                 ("Generato", f"{doc.generato_il:%d/%m/%Y %H:%M}" + (f" da {doc.generato_da}" if doc.generato_da else ""))]
        el.append(data_table(
            [[Paragraph(f"<b>{escape(k)}</b>", styles["cell"]), Paragraph(escape(v), styles["cell"])] for k, v in info],
            theme, col_widths=[width * 0.24, width * 0.76], header=False,
        ))
        el.append(Spacer(1, 6 * mm))
    else:
        el.append(Paragraph(escape(doc.titolo), styles["subtitle"]))
    if doc.mostra_indice and doc.indice:
        el.extend(section_heading("Indice", theme, styles))
        for i, voce in enumerate(doc.indice, start=1):
            el.append(Paragraph(f"{i}. {escape(voce)}", styles["body"]))
        el.append(Spacer(1, 6 * mm))

    for e in doc.elementi:
        if e.tipo == "pagina":
            el.append(PageBreak())
            continue
        if e.tipo == "testo":
            if e.titolo:
                el.extend(section_heading(escape(e.titolo), theme, styles))
            _testo(e.paragrafi)
            continue
        if e.tipo == "firme":
            ruoli = _ruoli_firma(doc)
            griglia = [
                [Paragraph(f"<b>{escape(r)}</b>", styles["table_header"]) for r, _n in ruoli],
                [Paragraph(escape(n or ""), styles["cell"]) for _r, n in ruoli],
                [Paragraph("<br/><br/>Firma ____________________<br/><br/>Data ____________", styles["cell"])
                 for _ in ruoli],
            ]
            el.append(KeepTogether([
                *section_heading(escape(e.titolo), theme, styles),
                data_table(griglia, theme, col_widths=[width / len(ruoli)] * len(ruoli)),
                Spacer(1, 4 * mm),
            ]))
            continue
        if e.tipo in ("negata", "errore"):
            el.extend(section_heading(escape(e.titolo), theme, styles))
            _testo(e.paragrafi)
            continue
        head = list(section_heading(escape(e.titolo), theme, styles))
        if e.riferimenti:
            head.append(Paragraph(f"<b>Riferimenti:</b> {escape(e.riferimenti)}", styles["body"]))
            head.append(Spacer(1, 2 * mm))
        if e.mostra_indicatori and e.risultato.kpis:
            head.append(_kpi_table(e.risultato.kpis, theme, styles, width))
            head.append(Spacer(1, 3 * mm))
        el.append(KeepTogether(head))
        if e.mostra_tabella and e.colonne:
            if e.n_righe:
                for g in e.gruppi:
                    if g.etichetta:
                        el.append(Paragraph(f"{escape(g.etichetta)} <font size=7 color='#64748b'>({len(g.righe)})</font>", gruppo_st))
                    el.append(_tabella(e, g.righe))
            else:
                el.append(Paragraph("Nessun dato nel perimetro selezionato.", styles["body"]))
            el.append(Spacer(1, 2 * mm))
        for n in e.note:
            el.append(Paragraph(f"• {escape(n)}", nota))
        _testo(e.paragrafi)
        el.append(Spacer(1, 4 * mm))

    sottotitolo = " | ".join(p for p in (doc.intestazione, doc.riservatezza, f"Periodo {doc.periodo_label}",
                                         doc.destinatario) if p)
    draw = header_footer_callback(theme, title=doc.titolo.upper()[:90], subtitle=sottotitolo[:140])

    def _pagina(canvas, documento):
        draw(canvas, documento)
        w, h = documento.pagesize
        if doc.piede_pagina:
            canvas.saveState()
            canvas.setFont("Helvetica", 7)
            canvas.setFillColor(colors.HexColor("#64748b"))
            canvas.drawCentredString(w / 2, 7 * mm, doc.piede_pagina[:180])
            canvas.restoreState()
        if doc.filigrana:
            canvas.saveState()
            canvas.setFont("Helvetica-Bold", 60)
            canvas.setFillColor(colors.Color(0.6, 0.6, 0.6, alpha=0.18))
            canvas.translate(w / 2, h / 2)
            canvas.rotate(35)
            canvas.drawCentredString(0, 0, doc.filigrana[:40].upper())
            canvas.restoreState()

    pdf.build(el, onFirstPage=_pagina, onLaterPages=_pagina)
    return buf.getvalue()


def _ruoli_firma(doc: Documento) -> list[tuple[str, str]]:
    return [("Redatto da", doc.redatto_da), ("Verificato da", doc.verificato_da), ("Approvato da", doc.approvato_da)]


# ═══════════════════════════════════════════════════════════════════════════
# Excel
# ═══════════════════════════════════════════════════════════════════════════

_TAG = re.compile(r"<[^>]+>")
_SHEET_BAD = re.compile(r"[\[\]\*\?/\\:]")
_XLS_FILL = {"danger": "FDE8E8", "warn": "FDF3D8", "ok": "E7F6EC"}


def _plain(markup: str) -> str:
    return _TAG.sub("", markup.replace("<br/>", "\n")).replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")


def _xls_valore(valore):
    """Valore grezzo -> valore Excel tipizzato (Decimal come float, datetime senza fuso)."""
    if isinstance(valore, Decimal):
        return float(valore)
    if isinstance(valore, datetime):
        return timezone.localtime(valore).replace(tzinfo=None) if timezone.is_aware(valore) else valore
    return valore


def render_xlsx(doc: Documento) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    from core.excel_export import append_row, write_cell

    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    titolo_font = Font(bold=True, size=14)
    head_fill = PatternFill("solid", fgColor="0C2545")
    head_font = Font(bold=True, color="FFFFFF")
    gruppo_fill = PatternFill("solid", fgColor="E8EEF6")
    sottile = Side(style="thin", color="D5DEE8")
    bordo = Border(left=sottile, right=sottile, top=sottile, bottom=sottile)
    # Nomi dei fogli fissati prima: il foglio «Documento» li cita come rimando.
    usati: set[str] = {"Documento", "Indicatori"}
    fogli = {id(e): _nome_foglio(e.titolo, usati) for e in doc.sezioni_calcolate if e.mostra_tabella and e.colonne}

    def _impagina(ws, intestazione_righe: int | None = None) -> None:
        ws.page_setup.orientation = "landscape" if doc.orientamento != "VERTICALE" else "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddHeader.center.text = doc.titolo[:200]
        ws.oddFooter.left.text = doc.riservatezza
        ws.oddFooter.right.text = "Pagina &P di &N"
        if intestazione_righe:
            ws.print_title_rows = f"{intestazione_righe}:{intestazione_righe}"

    def _intestazione(ws, titolo: str) -> None:
        append_row(ws, [titolo])
        ws.cell(row=1, column=1).font = titolo_font
        meta = " | ".join(p for p in (doc.intestazione, f"Periodo {doc.periodo_label}", doc.destinatario,
                                      doc.riservatezza) if p)
        append_row(ws, [meta])
        append_row(ws, [f"Perimetro: {doc.perimetro_label}"])
        append_row(ws, [])

    if doc.excel_foglio_documento:
        ws = wb.create_sheet("Documento")
        _intestazione(ws, doc.titolo)
        if doc.sottotitolo:
            append_row(ws, [doc.sottotitolo])
        for k, v in (("Destinatario", doc.destinatario or "—"), ("Norme", ", ".join(doc.norme) or "—"),
                     ("Generato", f"{doc.generato_il:%d/%m/%Y %H:%M} {doc.generato_da}".strip())):
            append_row(ws, [k, v])
            ws.cell(row=ws.max_row, column=1).font = bold
        append_row(ws, [])
        for e in doc.elementi:
            if e.tipo == "pagina":
                continue
            if e.titolo:
                append_row(ws, [e.titolo])
                ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
            if e.tipo == "firme":
                for ruolo, nome in _ruoli_firma(doc):
                    append_row(ws, [ruolo, nome or "", "Firma:", "Data:"])
            if id(e) in fogli:
                append_row(ws, ["Dettaglio nel foglio", fogli[id(e)], f"{e.n_righe} righe"])
            for _kind, markup in e.paragrafi:
                append_row(ws, [_plain(markup)])
                ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
            append_row(ws, [])
        ws.column_dimensions["A"].width = 70
        ws.column_dimensions["B"].width = 40
        ws.column_dimensions["C"].width = 18
        _impagina(ws)

    if doc.excel_foglio_indicatori and any(e.risultato and e.risultato.kpis for e in doc.sezioni_calcolate):
        ws = wb.create_sheet("Indicatori")
        _intestazione(ws, "Indicatori")
        append_row(ws, ["Sezione", "Indicatore", "Valore", "Nota", "Riferimenti"])
        riga_h = ws.max_row
        for cell in ws[riga_h]:
            cell.fill, cell.font = head_fill, head_font
        for e in doc.sezioni_calcolate:
            if not e.mostra_indicatori:
                continue
            for kpi in e.risultato.kpis:
                cells = append_row(ws, [e.titolo, kpi.label, _xls_valore(kpi.value), kpi.hint, e.riferimenti])
                if kpi.tone in _XLS_FILL:
                    for c in cells[1:3]:
                        c.fill = PatternFill("solid", fgColor=_XLS_FILL[kpi.tone])
        for col, w in zip("ABCDE", (40, 40, 16, 34, 40)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = ws.cell(row=riga_h + 1, column=1)
        _impagina(ws, riga_h)

    for e in doc.sezioni_calcolate:
        if id(e) not in fogli:
            continue
        ws = wb.create_sheet(fogli[id(e)])
        _intestazione(ws, e.titolo)
        if e.riferimenti:
            append_row(ws, [f"Riferimenti: {e.riferimenti}"])
            append_row(ws, [])
        append_row(ws, [label for _k, label in e.colonne])
        riga_h = ws.max_row
        for cell in ws[riga_h]:
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        n_col = len(e.colonne)
        larghezze = [len(label) for _k, label in e.colonne]
        raggruppata = any(g.etichetta for g in e.gruppi)
        for g in e.gruppi:
            if g.etichetta:
                riga = ws.max_row + 1
                write_cell(ws, riga, 1, f"{g.etichetta} ({len(g.righe)})")
                ws._current_row = riga
                for j in range(1, n_col + 1):
                    ws.cell(row=riga, column=j).fill = gruppo_fill
                ws.cell(row=riga, column=1).font = bold
            for r in g.righe:
                riga = ws.max_row + 1
                for j, valore in enumerate(r.valori, start=1):
                    c = write_cell(ws, riga, j, _xls_valore(valore))
                    c.border = bordo
                    if isinstance(valore, date):
                        c.number_format = "DD/MM/YYYY"
                    elif isinstance(valore, (Decimal, float)):
                        c.number_format = "0.0"
                    if e.matrice and j > 2:
                        c.alignment = Alignment(horizontal="center")
                    tono = (r.toni_celle[j - 1] if j - 1 < len(r.toni_celle) else "") or (
                        r.tono if r.tono in TONI_EVIDENZA and not e.matrice else "")
                    if tono in _XLS_FILL:
                        c.fill = PatternFill("solid", fgColor=_XLS_FILL[tono])
                    larghezze[j - 1] = max(larghezze[j - 1], len(fmt(valore)))
                ws._current_row = riga
        for j, w in enumerate(larghezze, start=1):
            ws.column_dimensions[ws.cell(row=riga_h, column=j).column_letter].width = min(max(w + 2, 8), 60)
        ws.freeze_panes = ws.cell(row=riga_h + 1, column=3 if e.matrice else 1)
        if e.n_righe and not raggruppata:
            ws.auto_filter.ref = f"A{riga_h}:{ws.cell(row=ws.max_row, column=n_col).coordinate}"
        if e.note:
            append_row(ws, [])
            for n in e.note:
                append_row(ws, [f"• {n}"])
        _impagina(ws, riga_h)

    if not wb.sheetnames:
        ws = wb.create_sheet("Documento")
        append_row(ws, [doc.titolo])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _nome_foglio(titolo: str, usati: set[str]) -> str:
    """Nome foglio Excel valido (max 31 caratteri, niente []*?/\\:) e unico."""
    base = _SHEET_BAD.sub(" ", titolo or "Sezione").strip()[:28] or "Sezione"
    nome, i = base, 2
    while nome in usati:
        nome = f"{base[:26]} {i}"
        i += 1
    usati.add(nome)
    return nome
