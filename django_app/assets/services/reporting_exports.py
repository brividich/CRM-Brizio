"""PDF con tema condiviso del portale ed Excel formattato dello stesso snapshot."""
from html import escape
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.platypus import Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors

from core.pdf import PdfTheme, build_styles, data_table, make_document, header_footer_callback
from .reporting import display_snapshot

ASSET_COLUMNS = [("tag", "Asset"), ("nome", "Nome"), ("tipo", "Tipo"), ("categoria", "Categoria"),
                 ("reparto", "Reparto"), ("stato", "Stato"), ("seriale", "Seriale"),
                 ("produttore", "Produttore"), ("modello", "Modello"), ("ubicazione", "Ubicazione"), ("ip", "IP")]
MFC_COLUMNS = [("asset", "Asset"), ("nome", "MFC"), ("stato", "Stato SNMP"), ("mese", "Mese lettura"),
               ("data", "Rilevata il"), ("a4_bn", "A4 BN"), ("a3_bn", "A3 BN"), ("a4_col", "A4 COL"), ("a3_col", "A3 COL")]
METRICS = [("assets", "Asset totali"), ("in_use", "In uso"), ("in_repair", "In riparazione"), ("snmp_errors", "Errori SNMP")]


def export_pdf(snapshot):
    snapshot = display_snapshot(snapshot)
    theme = PdfTheme.from_branding()
    styles = build_styles(theme)
    stream = BytesIO()
    doc = make_document(stream, title=snapshot["name"], landscape=True)
    text = lambda value: Paragraph(escape(str(value if value is not None else "—")), styles["body"])
    header_style = ParagraphStyle("report_header_cell", parent=styles["body"], textColor=colors.white, fontSize=9)
    heading = lambda label: Paragraph(escape(label), header_style)
    title_style = ParagraphStyle("report_section", parent=styles["body"], fontSize=14, leading=19, textColor=colors.HexColor(theme.primary), spaceBefore=16, spaceAfter=8, keepWithNext=True)
    section = lambda label: Paragraph(escape(label), title_style)
    story = [section(snapshot["name"]), text(snapshot.get("scope_label", "")), text(f"Dati rilevati: {snapshot['captured_at']}"), Spacer(1, 12)]
    story.append(data_table([[label for _, label in METRICS], [snapshot["metrics"][key] for key, _ in METRICS]], theme, col_widths=[doc.width / 4] * 4))
    trend = snapshot.get("trend", [])
    if trend:
        story += [section("Andamento delle estrazioni — stesso perimetro")]
        rows = [["Data"] + [label for _, label in METRICS]]
        rows += [[row["date"]] + [row[key] for key, _ in METRICS] for row in trend]
        story.append(data_table(rows, theme, col_widths=[doc.width / 5] * 5))
    sections = [("Inventario asset", snapshot["assets"], ASSET_COLUMNS[:6]),
                ("Dati tecnici e ubicazioni", snapshot["assets"], [ASSET_COLUMNS[0], *ASSET_COLUMNS[6:]]),
                ("Letture MFC", snapshot["mfc"], [c for c in MFC_COLUMNS if c[0] != "nome"])]
    for title, records, columns in sections:
        if not records:
            continue
        story += [section(title)]
        rows = [[heading(label) for _, label in columns]]
        rows += [[text(row.get(key)) for key, _ in columns] for row in records]
        story.append(data_table(rows, theme, col_widths=[doc.width / len(columns)] * len(columns)))
    if snapshot["devices"]:
        story += [section("Monitor SNMP e sonde")]
        rows = [["Asset / dispositivo", "Stato / ultimo controllo", "Sonde"]]
        for device in snapshot["devices"]:
            for probe in device["valori"] or [None]:
                value = f"{probe['sonda']}: {probe['valore']} ({probe['stato']})" if probe else "Nessuna misura nell'ultima interrogazione"
                rows.append([text(f"{device['asset']} · {device['nome']} · {device['host']}"), text(f"{device['stato']} · {device['controllo']}"), text(value)])
        story.append(data_table(rows, theme, col_widths=[doc.width * .3, doc.width * .25, doc.width * .45]))
    draw = header_footer_callback(theme, title="REPORTISTICA ASSET", subtitle="Archivio storico · NOVICROM HUB")
    doc.build(story, onFirstPage=draw, onLaterPages=draw)
    return stream.getvalue()


def export_excel(snapshot):
    snapshot = display_snapshot(snapshot)
    theme = PdfTheme.from_branding()
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, columns, rows):
        ws = wb.create_sheet(name)
        ws.append([theme.portal_name + " · " + snapshot["name"]])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, len(columns)))
        ws.append(["Rilevato il", snapshot["captured_at"]])
        ws.append([label for _, label in columns])
        for row in rows:
            ws.append([row.get(key) for key, _ in columns])
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    # Valori da inventario e SNMP sempre testo, mai formule Excel.
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row in (1, 3):
                    cell.fill = PatternFill("solid", fgColor=theme.primary.lstrip("#"))
                    cell.font = Font(color="FFFFFF", bold=True)
                elif cell.row > 3 and cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F1F5F9")
        from openpyxl.utils import get_column_letter
        for col in range(1, len(columns) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 24 if col < 3 else 19
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{get_column_letter(len(columns))}{max(3, ws.max_row)}"
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.print_title_rows = "1:3"
        return ws

    sheet("Riepilogo", [("label", "Indicatore"), ("value", "Valore")], [{"label": label, "value": snapshot["metrics"][key]} for key, label in METRICS] + [{"label": "Perimetro", "value": snapshot.get("scope_label", "")}])
    sheet("Asset", ASSET_COLUMNS, snapshot["assets"])
    sheet("MFC", MFC_COLUMNS, snapshot["mfc"])
    devices = []
    for device in snapshot["devices"]:
        for value in device["valori"] or [{"sonda": "", "valore": "", "stato": device["stato"]}]:
            devices.append({**device, **value, "device_stato": device["stato"]})
    sheet("SNMP", [("asset", "Asset"), ("nome", "Dispositivo"), ("host", "IP"), ("device_stato", "Stato dispositivo"), ("controllo", "Ultimo controllo"), ("sonda", "Sonda"), ("valore", "Valore"), ("stato", "Stato sonda")], devices)
    trend = sheet("Andamento", [("date", "Data rilevazione"), *METRICS], snapshot.get("trend", []))
    if trend.max_row > 4:
        chart = LineChart()
        chart.title = "Andamento asset — stesso perimetro"
        chart.y_axis.title = "Numero"
        chart.add_data(Reference(trend, min_col=2, max_col=5, min_row=3, max_row=trend.max_row), titles_from_data=True)
        chart.set_categories(Reference(trend, min_col=1, min_row=4, max_row=trend.max_row))
        trend.add_chart(chart, "G3")
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
