"""Riallineamento dei piani di manutenzione allo scadenziario del gestionale X-MINE.

Il problema che risolve: i piani storici sono applicati a un'intera **categoria**
(es. «Macchine CNC»), quindi ogni macchina della categoria riceve manutenzioni
che non le spettano. X-MINE invece elenca, macchina per macchina, quali
interventi fa, con l'ultima esecuzione e la prossima data pianificata.

Il riallineamento, per i piani che corrispondono alle causali X-MINE
(codice ``caus-<codice test>``):

1. abbina ogni macchina X-MINE a un asset (numero interno, oppure mappatura
   manuale per i codici ambigui);
2. crea/aggiorna un'applicazione **per singolo asset** con la periodicita' di
   X-MINE e disattiva le applicazioni di categoria degli stessi piani;
3. registra l'ultima esecuzione X-MINE come storico (se piu' recente di quella
   del portale) e porta la scadenza aperta alla data pianificata X-MINE;
4. annulla le scadenze aperte che X-MINE non prevede (con motivo).

``build_plan`` calcola tutto senza scrivere (il dry-run gira anche col profilo
di sola lettura sulla produzione); ``apply_plan`` esegue in un'unica transazione.
"""

from __future__ import annotations

import csv
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from dateutil.relativedelta import relativedelta
from django.db import transaction
from django.utils import timezone

from ..models import (
    Asset,
    MaintenanceInterventionTemplate,
    MaintenanceOccurrence,
    MaintenancePlanAssignment,
)

PLAN_CODE_PREFIX = "caus-"
# Codici A* (verifiche di legge, corsi, qualifiche, visite) restano fuori: hanno
# il loro registro. A16 (controllo perdita freon) e' oggi un piano di categoria.
DEFAULT_A_CODES = frozenset({"A16"})
# Centri di lavoro di servizio: il codice macchina e' locale al centro.
SERVICE_CENTER_PREFIX = "9"
CANCEL_REASON_NOT_DUE = "Non prevista per questo asset nello scadenziario X-MINE (riallineamento)"
CANCEL_REASON_RETIRED = "Asset dismesso (riallineamento X-MINE)"
DONE_NOTE = "Ultima esecuzione registrata in X-MINE (riallineamento)"
_N_INT_RE = re.compile(r"n\W{0,2}\s*int\.?\s*(\d+)", re.IGNORECASE)
_RETIRED_STATUSES = {"RETIRED", "DISPOSED", "SCRAPPED"}


class XmineFileError(Exception):
    """File X-MINE illeggibile o senza le colonne attese."""


# ---------------------------------------------------------------------------
# Lettura del file
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class XmineRow:
    line: int
    center: str
    center_label: str
    machine: str
    machine_label: str
    test_code: str
    test_label: str
    planned: date | None
    last_done: date | None

    @property
    def machine_key(self) -> tuple[str, str]:
        return (self.center, self.machine)

    @property
    def period_days(self) -> int | None:
        if self.planned and self.last_done and self.planned > self.last_done:
            return (self.planned - self.last_done).days
        return None


_HEADERS = {
    "planned": "Data Pianif.",
    "last_done": "Dt.Ult.Coll.",
    "test_code": "Cod. Test",
    "test_label": "Descrizione Test di Collaudo",
    "center": "C.di Lav.",
    "center_label": "Descrizione Centro di Lavoro",
    "machine": "Macc.Ut.",
    "machine_label": "Descrizione Macchina Utensile",
}


def _cell_text(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value if value is not None else "").strip()


def read_xmine(path: str | Path) -> list[XmineRow]:
    """Legge l'export «SCADENZE X-MINE» (.xls o .xlsx)."""
    path = Path(path)
    suffix = path.suffix.lower()
    try:
        if suffix == ".xls":
            import xlrd

            book = xlrd.open_workbook(str(path))
            sheet = book.sheet_by_index(0)

            def cell(r, c):
                item = sheet.cell(r, c)
                if item.ctype == xlrd.XL_CELL_DATE:
                    return xlrd.xldate.xldate_as_datetime(item.value, book.datemode).date()
                return item.value

            raw = [[cell(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
        else:
            from openpyxl import load_workbook

            book = load_workbook(str(path), read_only=True, data_only=True)
            raw = [list(row) for row in book.worksheets[0].iter_rows(values_only=True)]
    except Exception as exc:  # file corrotto o formato sconosciuto
        raise XmineFileError(f"Impossibile leggere {path.name}: {exc}") from exc
    if not raw:
        raise XmineFileError("Il file e' vuoto.")

    header = [_cell_text(h) for h in raw[0]]
    try:
        index = {key: header.index(label) for key, label in _HEADERS.items()}
    except ValueError as exc:
        raise XmineFileError(f"Colonna mancante nel file X-MINE: {exc}") from exc

    def as_date(value):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        return None

    rows: list[XmineRow] = []
    for number, values in enumerate(raw[1:], start=2):
        values = list(values) + [None] * (len(header) - len(values))
        test_code = _cell_text(values[index["test_code"]]).upper()
        machine = _cell_text(values[index["machine"]])
        if not test_code or not machine:
            continue
        rows.append(XmineRow(
            line=number,
            center=_cell_text(values[index["center"]]).upper(),
            center_label=_cell_text(values[index["center_label"]]),
            machine=machine,
            machine_label=_cell_text(values[index["machine_label"]]),
            test_code=test_code,
            test_label=_cell_text(values[index["test_label"]]),
            planned=as_date(values[index["planned"]]),
            last_done=as_date(values[index["last_done"]]),
        ))
    return rows


def read_mapping(path: str | Path | None) -> dict[tuple[str, str], str]:
    """Mappatura manuale ``centro;macchina;asset_tag`` (CSV, separatore ; o ,)."""
    if not path:
        return {}
    text = Path(path).read_text(encoding="utf-8-sig")
    dialect = csv.Sniffer().sniff(text.splitlines()[0] if text else ";", delimiters=";,")
    mapping: dict[tuple[str, str], str] = {}
    for row in csv.reader(text.splitlines(), dialect):
        if len(row) < 3 or not row[0].strip() or row[0].strip().lower() == "centro":
            continue
        mapping[(row[0].strip().upper(), row[1].strip())] = row[2].strip()
    return mapping


# ---------------------------------------------------------------------------
# Abbinamento macchine -> asset
# ---------------------------------------------------------------------------

def normalize_internal(value: str | None) -> str:
    text = re.sub(r"(?i)^int\.?\s*", "", (value or "").strip())
    text = text.strip()
    if text.isdigit():
        return text.lstrip("0") or "0"
    return text.upper()


def asset_internal_keys(asset: Asset) -> set[str]:
    keys = set()
    extra = asset.extra_columns if isinstance(asset.extra_columns, dict) else {}
    for raw in (extra.get("n_interno"), asset.internal_number):
        key = normalize_internal(str(raw) if raw is not None else "")
        if key:
            keys.add(key)
    return keys


def _words(text: str) -> set[str]:
    return {w for w in re.findall(r"[A-Z0-9]+", (text or "").upper()) if len(w) >= 3 and not w.isdigit()}


@dataclass
class MachineMatch:
    center: str
    machine: str
    label: str
    asset: Asset | None = None
    method: str = ""
    candidates: list[Asset] = field(default_factory=list)
    problem: str = ""


def _machine_lookup_key(center: str, machine: str, label: str) -> tuple[str | None, str]:
    """Chiave di numero interno per la macchina, o ``None`` se non affidabile."""
    is_service = center.startswith(SERVICE_CENTER_PREFIX)
    found = _N_INT_RE.search(label or "")
    if found:
        return normalize_internal(found.group(1)), "n°int. nella descrizione"
    key = normalize_internal(machine)
    if not is_service:
        return key, "numero interno"
    # Nei centri di servizio i codici brevi (001, 002, 010...) si ripetono fra
    # centri diversi: valgono solo i codici alfanumerici (C1, E10) o >= 100.
    if not key.isdigit() or int(key) >= 100:
        return key, "numero interno"
    return None, ""


def match_machines(rows: list[XmineRow], assets: list[Asset], mapping: dict[tuple[str, str], str]) -> dict:
    by_key: dict[str, list[Asset]] = defaultdict(list)
    by_tag = {asset.asset_tag.upper(): asset for asset in assets}
    for asset in assets:
        for key in asset_internal_keys(asset):
            by_key[key].append(asset)

    matches: dict[tuple[str, str], MachineMatch] = {}
    for row in rows:
        if row.machine_key in matches:
            continue
        match = MachineMatch(center=row.center, machine=row.machine, label=row.machine_label)
        matches[row.machine_key] = match
        manual = mapping.get(row.machine_key)
        if manual:
            asset = by_tag.get(manual.upper())
            if asset is None:
                match.problem = f"Mappatura: asset «{manual}» inesistente"
            else:
                match.asset, match.method = asset, "mappatura manuale"
            continue
        key, method = _machine_lookup_key(row.center, row.machine, row.machine_label)
        if key is None:
            match.problem = "Codice macchina locale al centro: serve la mappatura manuale"
            continue
        candidates = by_key.get(key, [])
        active = [a for a in candidates if a.status not in _RETIRED_STATUSES]
        pool = active or candidates
        match.candidates = candidates
        if len(pool) > 1:
            # Spareggio: il nome dell'asset deve condividere una parola con la
            # descrizione X-MINE (AKEBONO, TNZ...). Se non basta resta ambiguo.
            words = _words(row.machine_label)
            named = [a for a in pool if _words(a.name) & words]
            if len(named) == 1:
                pool, method = named, f"{method} + nome"
        if not pool:
            match.problem = "Nessun asset con questo numero interno"
        elif len(pool) > 1:
            match.problem = "Ambiguo: piu' asset con lo stesso numero interno"
        else:
            match.asset, match.method = pool[0], method
    return matches


# ---------------------------------------------------------------------------
# Periodicita'
# ---------------------------------------------------------------------------

def infer_recurrence(last_done: date | None, planned: date | None) -> tuple[str, int] | None:
    """Periodicita' X-MINE come (frequenza, intervallo): mesi/anni se coincidono
    con le date (±3 giorni), altrimenti giorni esatti."""
    if not last_done or not planned or planned <= last_done:
        return None
    days = (planned - last_done).days
    months = round(days / 30.4375)
    if months >= 1 and abs((last_done + relativedelta(months=months) - planned).days) <= 3:
        if months % 12 == 0:
            return MaintenancePlanAssignment.FREQ_YEARS, months // 12
        return MaintenancePlanAssignment.FREQ_MONTHS, months
    return MaintenancePlanAssignment.FREQ_DAYS, days


def recurrence_text(frequency: str, interval: int) -> str:
    unit = {"DAYS": "gg", "WEEKS": "sett.", "MONTHS": "mesi", "YEARS": "anni"}.get(frequency, frequency)
    if interval == 1:
        unit = {"DAYS": "giorno", "WEEKS": "settimana", "MONTHS": "mese", "YEARS": "anno"}.get(frequency, unit)
    return f"{interval} {unit}"


# ---------------------------------------------------------------------------
# Piano di riallineamento (nessuna scrittura)
# ---------------------------------------------------------------------------

@dataclass
class PairAction:
    plan: MaintenanceInterventionTemplate
    asset: Asset
    row: XmineRow
    frequency: str
    interval: int
    warning_days: int
    existing_assignment: MaintenancePlanAssignment | None
    assignment_change: str            # "crea" | "aggiorna" | "invariata"
    old_recurrence: str
    portal_last_done: date | None
    record_done: bool                 # registra Dt.Ult.Coll. come storico
    target_due: date | None
    open_occurrence: MaintenanceOccurrence | None
    due_change: str                   # "sposta" | "invariata" | "nessuna aperta" | ...
    note: str = ""


@dataclass
class CancelAction:
    occurrence: MaintenanceOccurrence
    reason: str


@dataclass
class AlignmentPlan:
    rows_total: int = 0
    rows_in_scope: int = 0
    rows_out_of_scope: list[XmineRow] = field(default_factory=list)
    rows_unmatched: list[tuple[XmineRow, MachineMatch]] = field(default_factory=list)
    rows_without_plan: list[XmineRow] = field(default_factory=list)
    duplicates: list[XmineRow] = field(default_factory=list)
    matches: dict = field(default_factory=dict)
    pairs: list[PairAction] = field(default_factory=list)
    category_assignments_off: list[MaintenancePlanAssignment] = field(default_factory=list)
    cancellations: list[CancelAction] = field(default_factory=list)
    assets_left_without_plans: list[Asset] = field(default_factory=list)

    def summary(self) -> dict:
        changes = defaultdict(int)
        for pair in self.pairs:
            changes[f"applicazioni {pair.assignment_change}"] += 1
            changes[f"scadenze {pair.due_change}"] += 1
            if pair.record_done:
                changes["esecuzioni X-MINE da registrare"] += 1
        return {
            "righe X-MINE": self.rows_total,
            "righe nel perimetro": self.rows_in_scope,
            "righe fuori perimetro (codici A*)": len(self.rows_out_of_scope),
            "righe con macchina non abbinata": len(self.rows_unmatched),
            "righe senza piano nel portale": len(self.rows_without_plan),
            "righe duplicate (ignorate)": len(self.duplicates),
            "coppie asset/piano da X-MINE": len(self.pairs),
            **dict(sorted(changes.items())),
            "applicazioni di categoria da disattivare": len(self.category_assignments_off),
            "scadenze aperte da annullare": len(self.cancellations),
            "  di cui dentro un OdL": sum(1 for c in self.cancellations if c.occurrence.work_order_id),
            "asset che restano senza piani X-MINE": len(self.assets_left_without_plans),
        }


def _plan_code(test_code: str) -> str:
    return f"{PLAN_CODE_PREFIX}{test_code.lower()}"


def build_plan(
    rows: list[XmineRow],
    *,
    mapping: dict[tuple[str, str], str] | None = None,
    a_codes: frozenset[str] = DEFAULT_A_CODES,
    today: date | None = None,
) -> AlignmentPlan:
    today = today or timezone.localdate()
    result = AlignmentPlan(rows_total=len(rows))

    in_scope: list[XmineRow] = []
    for row in rows:
        if row.test_code.startswith("A") and row.test_code not in a_codes:
            result.rows_out_of_scope.append(row)
        else:
            in_scope.append(row)
    result.rows_in_scope = len(in_scope)

    assets = list(Asset.objects.select_related("asset_category").order_by("id"))
    result.matches = match_machines(in_scope, assets, mapping or {})
    plans = {
        plan.code: plan
        for plan in MaintenanceInterventionTemplate.objects.filter(code__startswith=PLAN_CODE_PREFIX)
    }
    xmine_plan_ids = {plan.id for plan in plans.values()}

    # (plan_id, asset_id) -> riga X-MINE (in caso di doppioni vince la data piu' recente)
    wanted: dict[tuple[int, int], XmineRow] = {}
    for row in in_scope:
        match = result.matches[row.machine_key]
        if match.asset is None:
            result.rows_unmatched.append((row, match))
            continue
        plan = plans.get(_plan_code(row.test_code))
        if plan is None:
            result.rows_without_plan.append(row)
            continue
        key = (plan.id, match.asset.id)
        previous = wanted.get(key)
        if previous is not None:
            keep, drop = (row, previous) if (row.last_done or date.min) > (previous.last_done or date.min) else (previous, row)
            wanted[key] = keep
            result.duplicates.append(drop)
            continue
        wanted[key] = row

    assets_by_id = {asset.id: asset for asset in assets}
    plans_by_id = {plan.id: plan for plan in plans.values()}
    asset_assignments = {
        (a.plan_id, a.asset_id): a
        for a in MaintenancePlanAssignment.objects.filter(
            target_type=MaintenancePlanAssignment.TARGET_ASSET, plan_id__in=xmine_plan_ids
        )
    }
    category_assignments = list(
        MaintenancePlanAssignment.objects.select_related("asset_category", "plan").filter(
            target_type=MaintenancePlanAssignment.TARGET_CATEGORY, plan_id__in=xmine_plan_ids, is_active=True
        )
    )
    cat_by_plan_category = {(a.plan_id, a.asset_category_id): a for a in category_assignments}

    last_done: dict[tuple[int, int], date] = {}
    for plan_id, asset_id, completed_on in MaintenanceOccurrence.objects.filter(
        status=MaintenanceOccurrence.STATUS_DONE, plan_id__in=xmine_plan_ids, completed_on__isnull=False
    ).values_list("plan_id", "asset_id", "completed_on"):
        if completed_on > last_done.get((plan_id, asset_id), date.min):
            last_done[(plan_id, asset_id)] = completed_on
    done_dates = set(
        MaintenanceOccurrence.objects.filter(plan_id__in=xmine_plan_ids).values_list("plan_id", "asset_id", "due_date")
    )
    open_occ: dict[tuple[int, int], list[MaintenanceOccurrence]] = defaultdict(list)
    for occ in MaintenanceOccurrence.objects.select_related("plan", "asset", "work_order").filter(
        status=MaintenanceOccurrence.STATUS_OPEN, plan_id__in=xmine_plan_ids
    ).order_by("due_date", "id"):
        open_occ[(occ.plan_id, occ.asset_id)].append(occ)

    for (plan_id, asset_id), row in sorted(wanted.items(), key=lambda item: (item[1].center, item[1].machine, item[1].test_code)):
        plan, asset = plans_by_id[plan_id], assets_by_id[asset_id]
        existing = asset_assignments.get((plan_id, asset_id))
        inherited = cat_by_plan_category.get((plan_id, asset.asset_category_id))
        recurrence = infer_recurrence(row.last_done, row.planned)
        note = ""
        if recurrence is None:
            source = existing if existing and not existing.is_excluded else inherited
            if source is not None:
                recurrence = (source.frequency, source.interval)
                note = "Periodicita' non ricavabile da X-MINE: tenuta quella del portale"
            else:
                recurrence = (MaintenancePlanAssignment.FREQ_YEARS, 1)
                note = "Periodicita' non ricavabile: impostato 1 anno, da verificare"
        frequency, interval = recurrence
        reference = existing if existing and not existing.is_excluded else inherited
        old_recurrence = recurrence_text(reference.frequency, reference.interval) if reference else ""
        warning_days = int((reference.warning_days if reference else 30) or 30)

        if existing is None:
            change = "crea"
        elif (existing.is_excluded or not existing.is_active
              or (existing.frequency, existing.interval) != (frequency, interval)
              or existing.first_due_date != row.planned):
            change = "aggiorna"
        else:
            change = "invariata"

        portal_last = last_done.get((plan_id, asset_id))
        record_done = bool(
            row.last_done and row.last_done <= today
            and (portal_last is None or portal_last < row.last_done)
            and (plan_id, asset_id, row.last_done) not in done_dates
        )
        if portal_last and row.last_done and portal_last > row.last_done:
            target = portal_last + _delta(frequency, interval)
            note = (note + "; " if note else "") + "Esecuzione nel portale piu' recente di X-MINE: scadenza calcolata dal portale"
        else:
            target = row.planned

        current = open_occ.get((plan_id, asset_id), [])
        open_first = current[0] if current else None
        if open_first is None:
            due_change = "nessuna aperta (la crea lo scheduler)"
        elif target is None or open_first.due_date == target:
            due_change = "invariata"
        else:
            due_change = "sposta"
        # Eventuali altre aperte sulla stessa coppia sono doppioni.
        for extra in current[1:]:
            result.cancellations.append(CancelAction(extra, "Scadenza doppia sullo stesso piano (riallineamento X-MINE)"))

        result.pairs.append(PairAction(
            plan=plan, asset=asset, row=row, frequency=frequency, interval=interval,
            warning_days=warning_days, existing_assignment=existing, assignment_change=change,
            old_recurrence=old_recurrence, portal_last_done=portal_last, record_done=record_done,
            target_due=target, open_occurrence=open_first, due_change=due_change, note=note,
        ))

    # Applicazioni di categoria dei piani X-MINE: si disattivano tutte, il perimetro
    # reale lo danno ora le applicazioni per asset.
    result.category_assignments_off = list(category_assignments)
    covered_categories = defaultdict(set)
    for assignment in category_assignments:
        covered_categories[assignment.plan_id].add(assignment.asset_category_id)

    already_cancelled = {c.occurrence.id for c in result.cancellations}
    for (plan_id, asset_id), occurrences in open_occ.items():
        if (plan_id, asset_id) in wanted:
            continue
        asset = assets_by_id.get(asset_id)
        if asset is None:
            continue
        explicit = asset_assignments.get((plan_id, asset_id))
        from_category = asset.asset_category_id in covered_categories.get(plan_id, set())
        retired = asset.status in _RETIRED_STATUSES
        # Si toccano solo le scadenze nate dalla categoria (o su asset dismessi):
        # un'applicazione esplicita per asset non prevista da X-MINE si segnala ma resta.
        if not (from_category or retired) or (explicit and explicit.is_active and not explicit.is_excluded and not retired):
            continue
        reason = CANCEL_REASON_RETIRED if retired else CANCEL_REASON_NOT_DUE
        for occ in occurrences:
            if occ.id not in already_cancelled:
                result.cancellations.append(CancelAction(occ, reason))

    wanted_assets = {asset_id for _, asset_id in wanted}
    category_ids = {a.asset_category_id for a in category_assignments}
    result.assets_left_without_plans = [
        asset for asset in assets
        if asset.asset_category_id in category_ids and asset.id not in wanted_assets
        and asset.status not in _RETIRED_STATUSES
    ]
    return result


def _delta(frequency: str, interval: int):
    if frequency == MaintenancePlanAssignment.FREQ_YEARS:
        return relativedelta(years=interval)
    if frequency == MaintenancePlanAssignment.FREQ_MONTHS:
        return relativedelta(months=interval)
    if frequency == MaintenancePlanAssignment.FREQ_WEEKS:
        return timedelta(weeks=interval)
    return timedelta(days=interval)


# ---------------------------------------------------------------------------
# Applicazione
# ---------------------------------------------------------------------------

@transaction.atomic
def apply_plan(plan: AlignmentPlan, *, user=None, source_name: str = "X-MINE") -> dict:
    """Esegue il riallineamento calcolato da :func:`build_plan` (tutto o niente)."""
    from .maintenance_domain import cancel_occurrence

    stamp = timezone.localdate().strftime("%d/%m/%Y")
    counts = defaultdict(int)

    for pair in plan.pairs:
        assignment = pair.existing_assignment
        note = f"Da {source_name} (centro {pair.row.center}, macchina {pair.row.machine}) il {stamp}"
        if assignment is None:
            assignment = MaintenancePlanAssignment(
                plan=pair.plan,
                target_type=MaintenancePlanAssignment.TARGET_ASSET,
                asset=pair.asset,
            )
        if pair.assignment_change != "invariata" or assignment.pk is None:
            assignment.frequency = pair.frequency
            assignment.interval = pair.interval
            assignment.weekday = assignment.week_of_month = None
            assignment.day_of_month = assignment.month_of_year = None
            assignment.warning_days = pair.warning_days
            assignment.first_due_date = pair.row.planned
            assignment.is_excluded = False
            assignment.is_active = True
            assignment.auto_generate = True
            assignment.notes = note[:2000]
            assignment.full_clean()
            assignment.save()
            counts["applicazioni create" if pair.assignment_change == "crea" else "applicazioni aggiornate"] += 1

        if pair.record_done:
            MaintenanceOccurrence.objects.create(
                plan=pair.plan,
                assignment=assignment,
                asset=pair.asset,
                due_date=pair.row.last_done,
                completed_on=pair.row.last_done,
                status=MaintenanceOccurrence.STATUS_DONE,
                warning_days=pair.warning_days,
                schedule_anchor=assignment.effective_schedule_anchor,
                completion_notes=DONE_NOTE,
                source=MaintenanceOccurrence.SOURCE_IMPORT,
            )
            counts["esecuzioni X-MINE registrate"] += 1

        occ = pair.open_occurrence
        if occ is not None:
            fields = []
            if occ.assignment_id != assignment.id:
                occ.assignment = assignment
                fields.append("assignment")
            if pair.due_change == "sposta" and pair.target_due:
                clash = MaintenanceOccurrence.objects.filter(
                    plan_id=occ.plan_id, asset_id=occ.asset_id, due_date=pair.target_due
                ).exclude(pk=occ.pk).exists()
                if clash:
                    counts["scadenze non spostate (data gia' occupata)"] += 1
                else:
                    occ.due_date = pair.target_due
                    occ.original_due_date = None
                    occ.warning_days = pair.warning_days
                    fields += ["due_date", "original_due_date", "warning_days"]
                    counts["scadenze spostate alla data X-MINE"] += 1
            if fields:
                occ.save(update_fields=fields + ["updated_at"])

    for assignment in plan.category_assignments_off:
        assignment.is_active = False
        assignment.notes = (f"{assignment.notes}\n" if assignment.notes else "") + (
            f"Disattivata il {stamp}: sostituita da applicazioni per singolo asset da {source_name}."
        )
        assignment.save(update_fields=["is_active", "notes", "updated_at"])
        counts["applicazioni di categoria disattivate"] += 1

    for action in plan.cancellations:
        action.occurrence.refresh_from_db()
        if action.occurrence.status != MaintenanceOccurrence.STATUS_OPEN:
            continue
        cancel_occurrence(action.occurrence, reason=action.reason, user=user)
        counts["scadenze annullate"] += 1

    return dict(counts)


# ---------------------------------------------------------------------------
# Report Excel
# ---------------------------------------------------------------------------

def write_report(plan: AlignmentPlan, path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    def fmt(value):
        return value.strftime("%d/%m/%Y") if isinstance(value, date) else value

    book = Workbook()
    sheet = book.active
    sheet.title = "Riepilogo"
    sheet.append(["Voce", "Numero"])
    for key, value in plan.summary().items():
        sheet.append([key, value])

    def add(title, header, rows):
        ws = book.create_sheet(title[:31])
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([fmt(v) for v in row])
        ws.freeze_panes = "A2"

    add("Abbinamenti macchine",
        ["Centro", "Macchina", "Descrizione X-MINE", "Asset", "Nome asset", "Categoria", "Stato asset", "Metodo", "Problema", "Candidati"],
        [[m.center, m.machine, m.label,
          m.asset.asset_tag if m.asset else "", m.asset.name if m.asset else "",
          getattr(getattr(m.asset, "asset_category", None), "label", "") if m.asset else "",
          m.asset.status if m.asset else "", m.method, m.problem,
          ", ".join(f"{a.asset_tag} ({a.name}, {a.status})" for a in m.candidates)]
         for m in sorted(plan.matches.values(), key=lambda m: (m.problem == "", m.center, m.machine))])
    add("Coppie asset-piano",
        ["Asset", "Nome asset", "Piano", "Applicazione", "Periodicita' portale", "Periodicita' X-MINE",
         "Ult. esecuz. X-MINE", "Ult. esecuz. portale", "Registra esecuz.", "Scadenza aperta attuale",
         "Nuova scadenza", "Esito scadenza", "OdL", "Note"],
        [[p.asset.asset_tag, p.asset.name, p.plan.label, p.assignment_change, p.old_recurrence,
          recurrence_text(p.frequency, p.interval), p.row.last_done, p.portal_last_done,
          "si" if p.record_done else "", p.open_occurrence.due_date if p.open_occurrence else None,
          p.target_due, p.due_change,
          p.open_occurrence.work_order_id if p.open_occurrence and p.open_occurrence.work_order_id else "",
          p.note] for p in plan.pairs])
    add("Scadenze da annullare",
        ["Asset", "Nome asset", "Piano", "Scadenza", "OdL", "Motivo"],
        [[c.occurrence.asset.asset_tag, c.occurrence.asset.name, c.occurrence.plan.label,
          c.occurrence.due_date, c.occurrence.work_order_id or "", c.reason] for c in plan.cancellations])
    add("Categorie disattivate",
        ["Piano", "Categoria", "Periodicita'"],
        [[a.plan.label, a.asset_category.label if a.asset_category_id else "", recurrence_text(a.frequency, a.interval)]
         for a in plan.category_assignments_off])
    add("Asset senza piani X-MINE",
        ["Asset", "Nome", "Categoria", "Numero interno"],
        [[a.asset_tag, a.name, getattr(a.asset_category, "label", ""), a.internal_number]
         for a in plan.assets_left_without_plans])
    add("Righe non abbinate",
        ["Riga", "Centro", "Macchina", "Descrizione", "Test", "Descrizione test", "Data pianif.", "Ult. esecuz.", "Problema"],
        [[r.line, r.center, r.machine, r.machine_label, r.test_code, r.test_label, r.planned, r.last_done, m.problem]
         for r, m in plan.rows_unmatched])
    add("Senza piano nel portale",
        ["Riga", "Centro", "Macchina", "Test", "Descrizione test", "Codice piano atteso"],
        [[r.line, r.center, r.machine, r.test_code, r.test_label, _plan_code(r.test_code)] for r in plan.rows_without_plan])
    add("Fuori perimetro",
        ["Riga", "Centro", "Descrizione centro", "Macchina", "Descrizione", "Test", "Descrizione test", "Data pianif."],
        [[r.line, r.center, r.center_label, r.machine, r.machine_label, r.test_code, r.test_label, r.planned]
         for r in plan.rows_out_of_scope])

    path = Path(path)
    book.save(path)
    return path
