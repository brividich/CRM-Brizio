"""Dossier PART 145: processo, attrezzature, personale, scadenze e situazioni aperte.

Una sola costruzione (``costruisci``) per pagina, Excel (un foglio per sezione) e PDF:
le tre uscite mostrano gli stessi dati. Il personale (dati MOD.128) compare solo a chi
ha il permesso di vedere i processi qualificati: senza, le sezioni lo omettono.

Allineato all'SGI: perimetro = processo Part 145 + processi speciali autorizzati
(MO-ID-009.24, Annesso 2); attrezzi classificati come il MOE §2.4 (a vista / con
manutenzione / con taratura); documenti di riferimento dal registro SGI.
"""
from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from datetime import timedelta
from html import escape

from django.db.models import Min, Q
from django.urls import reverse
from django.utils import timezone

from automazioni.managed_flows import event_flow

from .part145 import REGIME_PART145, processi_perimetro

logger = logging.getLogger(__name__)

GIORNI_SCADENZE = 90

ESITI_IDONEITA = {"ok": "Idoneo", "warn": "In scadenza", "incompleto": "Incompleto", "ko": "Non idoneo",
                  "nessuno": "Nessun requisito"}

# Classi di controllo degli attrezzi (MOE C750-03 §2.4.2), dedotte dai piani di manutenzione.
CLASSE_TARATURA = "Taratura"
CLASSE_MANUTENZIONE = "Manutenzione periodica"
CLASSE_A_VISTA = "A vista"

# Documenti SGI di riferimento: (codice nel registro SGI, perché serve). Si cerca il codice
# esatto, poi il primo che inizia così (i MOE hanno nel codice anche edizione e data).
DOCUMENTI_SGI = [
    ("MO-ID-009.24", "Accordo con l'organizzazione di manutenzione (ATC)"),
    ("MO-ID-009.24 Annesso 1", "P/N autorizzati"),
    ("MO-ID-009.24 Annesso 2", "Processi speciali autorizzati (DQP)"),
    ("MT CN 70", "Memoria tecnica PART 145"),
    ("MT CN 17", "Attrezzature, utensili, strumenti (controllo semestrale)"),
    ("MT CN 68", "Gestione apparecchiature per misurazione"),
    ("MOE Civile C750-03", "Manuale dell'organizzazione di manutenzione — civile"),
    ("MOE Militare C750-29", "Manuale dell'organizzazione di manutenzione — militare"),
]


# Colore del badge in pagina per i valori di stato (rosso / ambra / verde).
BADGE = {v: "ko" for v in ("Scaduta", "Alta", "Mai frequentata", "Non idoneo", "Non trovato nel registro SGI")}
BADGE.update({v: "warn" for v in ("Entro 30 giorni", "Entro 90 giorni", "Media", "Senza scadenza", "In scadenza",
                                   "Incompleto")})
BADGE.update({v: "ok" for v in ("Valida", "Illimitata", "Idoneo")})

PRIORITA_ALTA = "Alta"
PRIORITA_MEDIA = "Media"
# Situazioni che bloccano l'operatività o l'audit: in cima, in rosso.
TIPI_PRIORITA_ALTA = {"Processo scaduto", "Processo non attivo", "Contratto scaduto", "Requisito di processo scaduto",
                      "Taratura scaduta", "Timbro sospeso", "Abilitazione revocata", "Abilitazione sospesa",
                      "Formazione scaduta", "Azione da audit scaduta", "NC Part 145 aperta"}


@dataclass
class Sezione:
    codice: str
    titolo: str
    intestazioni: list[str]
    righe: list[list] = field(default_factory=list)
    nota: str = ""
    vuoto: str = "Niente da segnalare."
    #: Id asset per riga (link alla scheda), allineato a ``righe``; vuoto = nessun link.
    asset_ids: list = field(default_factory=list)
    #: Azione per riga ``(testo, url)`` o None: dove si risolve, solo in pagina (non negli export).
    azioni: list = field(default_factory=list)
    #: In pagina la sezione parte chiusa (consultazione, non lavoro quotidiano).
    chiusa: bool = False

    def aggiungi(self, riga, azione=None, asset_id=None):
        self.righe.append(riga)
        self.azioni.append(azione)
        if asset_id is not None or self.asset_ids:
            self.asset_ids += [None] * (len(self.righe) - 1 - len(self.asset_ids))
            self.asset_ids.append(asset_id)

    @property
    def ha_azioni(self) -> bool:
        return any(self.azioni)

    @property
    def righe_link(self):
        ids = self.asset_ids + [None] * (len(self.righe) - len(self.asset_ids))
        azioni = self.azioni + [None] * (len(self.righe) - len(self.azioni))
        return list(zip(self.righe, ids[:len(self.righe)], azioni[:len(self.righe)], strict=True))

    def ordina(self, chiave):
        """Riordina righe, link e azioni insieme."""
        terne = sorted(self.righe_link, key=lambda t: chiave(t[0]))
        self.righe = [t[0] for t in terne]
        self.asset_ids = [t[1] for t in terne] if any(t[1] for t in terne) else []
        self.azioni = [t[2] for t in terne]


def puo_vedere_personale(request) -> bool:
    """Il personale abilitato (dati MOD.128) solo a chi vede i processi qualificati."""
    try:
        from anagrafica.views_mpq import _check_mpq_permission

        return _check_mpq_permission(request)
    except Exception:
        return False


def _d(valore) -> str:
    return valore.strftime("%d/%m/%Y") if valore else ""


def _stato_scadenza(giorno, oggi) -> str:
    if giorno is None:
        return "Senza scadenza"
    if giorno < oggi:
        return "Scaduta"
    if giorno <= oggi + timedelta(days=30):
        return "Entro 30 giorni"
    return "Valida"


def _sicuro(avvisi: list, cosa: str, funzione, predefinito):
    """Una parte accessoria che si rompe non deve far cadere il dossier: avviso in pagina e log."""
    try:
        return funzione()
    except Exception:
        logger.exception("Dossier PART 145: %s non disponibile", cosa)
        avvisi.append(f"{cosa}: dati non disponibili in questo momento (errore registrato nei log).")
        return predefinito


def _attrezzature():
    from ..models import Asset

    return (Asset.objects.filter(Q(part_145=True) | Q(processi_qualificati__regime=REGIME_PART145)
                                 | Q(processi_qualificati__perimetro_part145=True))
            .distinct().select_related("asset_category").order_by("asset_tag"))


def _controllo_vuoto() -> dict:
    return {"classe": CLASSE_A_VISTA, "enti": [], "ultima": None, "prossima": None, "scaduta": False}


def _controlli_attrezzi(ids, oggi=None) -> dict:
    """Per asset: classe di controllo, ente di taratura, ultima e prossima taratura.

    La classe viene dai piani applicati all'asset (anche ereditati da gruppo/categoria):
    piano di taratura → «Taratura»; altro piano non amministrativo → «Manutenzione
    periodica»; nessun piano → «A vista» (controllo visivo prima dell'uso).
    """
    from ..models import Asset, MaintenanceInterventionTemplate as Piano, MaintenanceOccurrence
    from .maintenance_domain import build_plan_resolutions

    oggi = oggi or timezone.localdate()
    esito = {i: _controllo_vuoto() for i in ids}
    if not ids:
        return esito
    for (_, asset_id), r in build_plan_resolutions(asset_queryset=Asset.objects.filter(pk__in=ids)).items():
        if asset_id not in esito or not r.is_applied or r.plan.maintenance_type == Piano.TYPE_ADMINISTRATIVE:
            continue
        voce = esito[asset_id]
        if r.plan.maintenance_type == Piano.TYPE_CALIBRATION:
            voce["classe"] = CLASSE_TARATURA
            if r.assignment.effective_execution_mode == Piano.MODE_INTERNAL:
                ente = "Interna"
            else:
                fornitore = r.assignment.effective_supplier
                ente = fornitore.ragione_sociale if fornitore else "Esterna (fornitore non indicato)"
            if ente not in voce["enti"]:
                voce["enti"].append(ente)
        elif voce["classe"] == CLASSE_A_VISTA:
            voce["classe"] = CLASSE_MANUTENZIONE
    for o in (MaintenanceOccurrence.objects.filter(asset_id__in=ids, plan__maintenance_type=Piano.TYPE_CALIBRATION)
              .exclude(status=MaintenanceOccurrence.STATUS_CANCELED).values("asset_id", "status", "due_date", "completed_on")):
        voce = esito[o["asset_id"]]
        voce["classe"] = CLASSE_TARATURA
        if o["status"] == MaintenanceOccurrence.STATUS_DONE and o["completed_on"]:
            voce["ultima"] = max(filter(None, [voce["ultima"], o["completed_on"]]))
        elif o["status"] == MaintenanceOccurrence.STATUS_OPEN:
            voce["prossima"] = min(filter(None, [voce["prossima"], o["due_date"]]))
    for voce in esito.values():
        voce["scaduta"] = bool(voce["prossima"] and voce["prossima"] < oggi)
    return esito


def _documenti_sgi() -> list[dict]:
    """Revisione corrente dei documenti SGI di riferimento (registro procedure), in una query."""
    from django.db.models.functions import Upper

    from procedure_refresh.models import ProcedureRevision

    filtro = Q()
    for codice, _ in DOCUMENTI_SGI:
        filtro |= Q(document__code__istartswith=codice)
    correnti = list(ProcedureRevision.objects.filter(filtro, is_current=True).select_related("document")
                    .order_by(Upper("document__code"), "-revision_date"))
    righe = []
    for codice, perche in DOCUMENTI_SGI:
        chiave = codice.upper()
        esatte = [r for r in correnti if r.document.code.upper() == chiave]
        rev = (esatte or [r for r in correnti if r.document.code.upper().startswith(chiave)] or [None])[0]
        righe.append({"codice": codice, "perche": perche, "rev": rev})
    return righe


def visibilita(request) -> dict:
    """Cosa può vedere chi apre il dossier: ogni modulo collegato tiene il suo permesso."""
    from core.acl_v2 import request_has_permission_code

    def ha(codice):
        try:
            return bool(request_has_permission_code(request, codice))
        except Exception:
            return False

    from anomalie.acl_bootstrap import PERM_ANOMALIE_NC
    from sistema_gestione.acl_bootstrap import PERM_AUDIT_VIEW

    return {"con_personale": puo_vedere_personale(request), "con_audit": ha(PERM_AUDIT_VIEW),
            "con_nc": ha(PERM_ANOMALIE_NC)}


def _formazione_richiesta(processi, abilitazioni) -> tuple[list, list]:
    """Scadenze dei corsi richiesti dai processi del perimetro per gli abilitati (es. Safety Training / Human Factor).

    Ritorna (scadenze presenti, coppie (persona, corso) senza scadenza registrata).
    """
    from anagrafica.models_formazione import TrainingCourse, TrainingDeadline

    corsi_per_processo = {p.pk: {c.pk for c in p.corsi_richiesti.all()} for p in processi}
    coppie = {(ab.legacy_anagrafica_id, c) for ab in abilitazioni
              if ab.stato == ab.STATO_ATTIVA and not ab.nominativo_esterno and ab.legacy_anagrafica_id
              for c in corsi_per_processo.get(ab.processo_id, ())}
    if not coppie:
        return [], []
    scadenze = [d for d in TrainingDeadline.objects.filter(legacy_anagrafica_id__in={p for p, _ in coppie},
                                                           corso_id__in={c for _, c in coppie}).select_related("corso")
                if (d.legacy_anagrafica_id, d.corso_id) in coppie]
    trovate = {(d.legacy_anagrafica_id, d.corso_id) for d in scadenze}
    corsi = {c.pk: c for c in TrainingCourse.objects.filter(pk__in={c for _, c in coppie - trovate})}
    mancanti = sorted(((p, corsi[c]) for p, c in coppie - trovate if c in corsi), key=lambda t: (t[0], str(t[1])))
    return scadenze, mancanti


def _audit_part145(anno) -> tuple[list, list]:
    """Audit interni Part 145 dell'anno e rilievi ancora aperti (di qualsiasi anno)."""
    from gestione_specifiche.models import RegistroOFI
    from sistema_gestione.models import Audit, AuditEsito

    audit = list(Audit.objects.exclude(part145="").filter(data_inizio__year=anno)
                 .exclude(stato=Audit.STATO_ANNULLATO).order_by("data_inizio"))
    rilievi = list(AuditEsito.objects.filter(audit__part145__gt="", esito__in=AuditEsito.ESITI_RILIEVO)
                   .exclude(audit__stato=Audit.STATO_ANNULLATO).exclude(ofi__fase=RegistroOFI.FASE_CHIUSO)
                   .select_related("audit").order_by("scadenza_azione", "audit__data_inizio"))
    return audit, rilievi


def _nc_part145() -> list:
    from anomalie.quality_models import AnomaliaNC

    return list(AnomaliaNC.objects.filter(part145=True).exclude(stato=AnomaliaNC.Stato.CHIUSA).order_by("created_at"))


def solo_attrezzature(*, oggi=None) -> dict:
    """Dati minimi per gli elenchi operativi (strumenti tarati, controllo semestrale)."""
    from ..models import Asset

    oggi = oggi or timezone.localdate()
    attrezzature = [a for a in _attrezzature() if a.status != Asset.STATUS_RETIRED]
    return {"oggi": oggi, "attrezzature": attrezzature, "stati_asset": dict(Asset.STATUS_CHOICES),
            "controlli": _controlli_attrezzi([a.pk for a in attrezzature], oggi)}


def costruisci(*, oggi=None, con_personale: bool = True, con_audit: bool = False, con_nc: bool = False,
                giorni: int = GIORNI_SCADENZE) -> dict:
    """Tutto il dossier. Personale, audit e NC compaiono solo con il permesso del rispettivo modulo."""
    from anagrafica.models_mpq import AbilitazioneProcesso, CertificazioneIndividuale, RequisitoQualifica

    from ..models import Asset, MaintenanceInterventionTemplate, MaintenanceOccurrence, WorkOrder

    oggi = oggi or timezone.localdate()
    limite = oggi + timedelta(days=giorni)
    avvisi: list[str] = []
    stati_asset = dict(Asset.STATUS_CHOICES)

    processi = list(processi_perimetro().select_related("cliente", "ente_certificatore")
                    .prefetch_related("riferimenti", "corsi_richiesti").order_by("-regime", "nome"))
    requisiti = _sicuro(avvisi, "Contratti e requisiti dei processi", lambda: list(
        RequisitoQualifica.objects.filter(processo__in=processi, attivo=True)
        .select_related("processo").order_by("data_scadenza", "processo__nome", "ordine")), [])
    # Le attrezzature dismesse escono dal dossier (restano nell'inventario con il loro storico).
    tutte = list(_attrezzature())
    attrezzature = [a for a in tutte if a.status != Asset.STATUS_RETIRED]
    dismesse = len(tutte) - len(attrezzature)
    ids = [a.pk for a in attrezzature]
    controlli = _sicuro(avvisi, "Tarature delle attrezzature", lambda: _controlli_attrezzi(ids, oggi),
                        {i: _controllo_vuoto() for i in ids})

    # Dati collegati alle attrezzature, in poche query.
    ultimo_chiuso: dict = {}
    for r in (WorkOrder.objects.filter(asset_id__in=ids, status=WorkOrder.STATUS_DONE)
              .order_by("asset_id", "-closed_at").values("asset_id", "closed_at", "kind")):
        ultimo_chiuso.setdefault(r["asset_id"], r)
    odl_aperti = list(WorkOrder.objects.filter(asset_id__in=ids, status=WorkOrder.STATUS_OPEN)
                      .select_related("asset").order_by("opened_at"))
    aperti_per_asset: dict = {}
    for wo in odl_aperti:
        aperti_per_asset[wo.asset_id] = aperti_per_asset.get(wo.asset_id, 0) + 1
    prossime = dict(MaintenanceOccurrence.objects.filter(asset_id__in=ids, status=MaintenanceOccurrence.STATUS_OPEN)
                    .order_by().values("asset_id").annotate(p=Min("due_date")).values_list("asset_id", "p"))
    tipi_odl = dict(WorkOrder.KIND_CHOICES)
    processi_per_asset: dict = {}
    collegati_per_processo: dict = {}
    for asset_id, processo_id, nome in (Asset.processi_qualificati.through.objects.filter(asset_id__in=ids)
                                        .values_list("asset_id", "processoqualificato_id", "processoqualificato__nome")):
        processi_per_asset.setdefault(asset_id, []).append(nome)
        collegati_per_processo.setdefault(processo_id, set()).add(asset_id)
    con_flag = {a.pk for a in attrezzature if a.part_145}

    # ── Personale ────────────────────────────────────────────────────────────
    abilitazioni: list = []
    nomi: dict = {}
    timbri: dict = {}

    def _carica_personale():
        from anagrafica.views_mpq import _nomi_map

        elenco = list(
            AbilitazioneProcesso.objects.filter(Q(is_part145=True) | Q(processo__regime=REGIME_PART145)
                                                | Q(processo__perimetro_part145=True))
            .select_related("processo").prefetch_related("certificazioni").order_by("processo__nome", "legacy_anagrafica_id")
        )
        return elenco, _nomi_map([a.legacy_anagrafica_id for a in elenco])

    if con_personale:
        abilitazioni, nomi = _sicuro(avvisi, "Personale abilitato", _carica_personale, ([], {}))

        def _carica_timbri():
            from timbri.models import RegistroTimbro

            per_ab: dict = {}
            for t in RegistroTimbro.objects.filter(abilitazione_processo__in=abilitazioni, is_archived=False):
                per_ab.setdefault(t.abilitazione_processo_id, []).append(t)
            return per_ab

        timbri = _sicuro(avvisi, "Timbri", _carica_timbri, {}) if abilitazioni else {}

    def nome_persona(ab):
        if ab.nominativo_esterno:
            return f"{ab.nominativo_esterno} (esterno)"
        return nomi.get(ab.legacy_anagrafica_id) or f"#{ab.legacy_anagrafica_id}"

    stati_ab = dict(AbilitazioneProcesso.STATO_CHOICES)

    def apri_mod128(processo):
        return ("Apri in MOD.128", reverse("anagrafica:mpq_processo_detail", args=[processo.pk])) if con_personale else None

    # ── Processi ─────────────────────────────────────────────────────────────
    s_processi = Sezione("processi", "Processi PART 145",
                         ["Processo", "Ambito", "Cliente / ente", "Livello", "Stato", "Scadenza", "Esito scadenza",
                          "Persone abilitate", "Attrezzature"],
                         nota="Il processo in regime PART 145 e i processi speciali con la casella «Rientra nel "
                              "perimetro Part 145» (MOD.128).",
                         vuoto="Nessun processo in regime PART 145 in MOD.128.")
    for p in processi:
        scad = p.scadenza_effettiva
        attrezzi_p = set(collegati_per_processo.get(p.pk, set()))
        if p.regime == REGIME_PART145:
            attrezzi_p |= con_flag
        s_processi.aggiungi([
            p.nome, "Regime PART 145" if p.regime == REGIME_PART145 else "Processo speciale nel perimetro",
            str(p.cliente) if p.cliente_id else "", p.livello, p.get_stato_display(), _d(scad),
            _stato_scadenza(scad, oggi) if p.tipo_validita != p.VALIDITA_ILLIMITATA else "Illimitata",
            sum(1 for ab in abilitazioni if ab.processo_id == p.pk and ab.stato == ab.STATO_ATTIVA) if con_personale else "—",
            len(attrezzi_p),
        ], apri_mod128(p))

    # ── Contratti e requisiti ────────────────────────────────────────────────
    s_req = Sezione("contratti", "Contratti e requisiti dei processi",
                    ["Processo", "Tipo", "Descrizione", "Obbligatorio", "Stato", "Scadenza", "Esito scadenza"],
                    nota="Requisiti dei processi del perimetro in MOD.128: es. il contratto quadro con la sua "
                         "scadenza (tipo «Contratto / accordo»), audit, certificati.",
                    vuoto="Nessun requisito: in MOD.128 aggiungi al processo PART 145 il contratto quadro "
                          "(tipo «Contratto / accordo») con la scadenza.")
    for r in requisiti:
        s_req.aggiungi([r.processo.nome, r.get_tipo_display(), r.descrizione, "Sì" if r.obbligatorio else "No",
                        r.get_stato_display(), _d(r.data_scadenza),
                        _stato_scadenza(r.data_scadenza, oggi) if r.data_scadenza else ""], apri_mod128(r.processo))

    # ── Attrezzature ─────────────────────────────────────────────────────────
    s_attr = Sezione("attrezzature", "Attrezzature",
                     ["Asset", "Nome", "Reparto", "Stato", "Seriale", "Processi", "Controllo", "Ente taratura",
                      "Ultima taratura", "Prossima taratura", "Prossima manutenzione", "Ultimo intervento", "OdL aperti"],
                     nota="Asset con la casella «Rientra in PART 145» e asset collegati a un processo del perimetro, "
                          "esclusi i dismessi. «Controllo» come il MOE §2.4: «Taratura» se l'asset ha un piano di "
                          "taratura, «Manutenzione periodica» se ha altri piani, altrimenti «A vista» (esame visivo "
                          "prima dell'uso).",
                     vuoto="Nessuna attrezzatura PART 145.")
    for a in attrezzature:
        chiuso = ultimo_chiuso.get(a.pk)
        processi_a = processi_per_asset.get(a.pk, [])
        if a.part_145 and not processi_a:
            processi_a = ["PART 145"]
        c = controlli.get(a.pk) or _controllo_vuoto()
        s_attr.aggiungi([
            a.asset_tag, a.name, a.reparto, stati_asset.get(a.status, a.status), a.serial_number or "",
            ", ".join(processi_a), c["classe"], ", ".join(c["enti"]), _d(c["ultima"]), _d(c["prossima"]),
            _d(prossime.get(a.pk)),
            (f"{timezone.localtime(chiuso['closed_at']):%d/%m/%Y} · {tipi_odl.get(chiuso['kind'], chiuso['kind'])}"
             if chiuso and chiuso["closed_at"] else ""),
            aperti_per_asset.get(a.pk, 0),
        ], asset_id=a.pk)

    # ── Personale ────────────────────────────────────────────────────────────
    s_pers = Sezione("personale", "Personale abilitato",
                     ["Persona", "Processo", "Ruoli", "Stato abilitazione", "Certificazioni", "Prima scadenza certificati",
                      "Idoneità requisiti", "Timbro"],
                     nota="Abilitazioni MOD.128 con ruolo «Part 145» o su un processo del perimetro.",
                     vuoto="Nessuna persona abilitata: in MOD.128 aggiungi le abilitazioni con ruolo «Part 145».")
    if con_personale and abilitazioni:
        from anagrafica.services.mpq_conformita import verifica_requisiti

        for ab in abilitazioni:
            ruoli = [n for flag, n in ((ab.is_qualificato, "Qualificato"), (ab.is_addetto, "Addetto"),
                                       (ab.is_controllore, "Controllore"), (ab.is_part145, "Part 145")) if flag]
            certs = list(ab.certificazioni.all())
            scad_cert = min((c.data_scadenza for c in certs if c.data_scadenza), default=None)
            idoneita = "—"
            if not ab.nominativo_esterno and ab.legacy_anagrafica_id:
                try:
                    idoneita = ESITI_IDONEITA.get(verifica_requisiti(ab.processo, ab.legacy_anagrafica_id, oggi=oggi)["esito"], "—")
                except Exception:
                    logger.exception("Dossier PART 145: idoneità non calcolabile per abilitazione %s", ab.pk)
                    idoneita = "non calcolabile"
            timbro = ", ".join(f"{t.codice_timbro or 'timbro'} ({'sospeso' if t.is_sospeso else 'attivo' if t.is_attivo else 'non attivo'})"
                               for t in timbri.get(ab.pk, []))
            s_pers.aggiungi([nome_persona(ab), ab.processo.nome, ", ".join(ruoli), stati_ab.get(ab.stato, ab.stato),
                             len(certs), _d(scad_cert), idoneita, timbro], apri_mod128(ab.processo))

    # ── Formazione richiesta dai processi (MOD.128 → Formazione) ─────────────
    formazione, formazione_mancante = [], []
    if con_personale and abilitazioni:
        formazione, formazione_mancante = _sicuro(avvisi, "Formazione del personale",
                                                  lambda: _formazione_richiesta(processi, abilitazioni), ([], []))
    scadenzario = ("Apri lo scadenzario", reverse("anagrafica:formazione_scadenzario"))
    stati_formazione = {"VALIDO": "Valida", "IN_SCADENZA_30": "Entro 30 giorni", "IN_SCADENZA_90": "Entro 90 giorni",
                        "SCADUTO": "Scaduta", "MAI_FREQUENTATO": "Mai frequentata", "UNA_TANTUM": "Valida"}
    s_form = Sezione("formazione", "Formazione richiesta",
                     ["Persona", "Corso", "Stato", "Ultimo completamento", "Scadenza"],
                     nota="Corsi richiesti dai processi del perimetro in MOD.128 (es. Safety Training / Human Factor "
                          "ogni 2 anni, formazione MOE) per le persone abilitate, dallo scadenzario della formazione.",
                     vuoto="Nessun corso richiesto: in MOD.128 aggiungi al processo PART 145 i corsi richiesti.")
    for d in sorted(formazione, key=lambda d: (d.data_scadenza or oggi.replace(year=oggi.year + 50), d.legacy_anagrafica_id)):
        s_form.aggiungi([nomi.get(d.legacy_anagrafica_id) or f"#{d.legacy_anagrafica_id}", str(d.corso),
                         stati_formazione.get(d.stato_scadenza, d.stato_scadenza), _d(d.data_ultimo_completamento),
                         _d(d.data_scadenza)], scadenzario)
    for persona, corso in formazione_mancante:
        s_form.aggiungi([nomi.get(persona) or f"#{persona}", str(corso), "Mai frequentata", "", ""], scadenzario)

    # ── Audit interni Part 145 (Sistema di gestione) ─────────────────────────
    audit_anno, rilievi = ([], [])
    if con_audit:
        audit_anno, rilievi = _sicuro(avvisi, "Audit interni", lambda: _audit_part145(oggi.year), ([], []))
    s_audit = Sezione("audit", f"Audit interni Part 145 ({oggi.year})",
                      ["Audit", "Perimetro", "Data", "Stato", "Rilievi aperti"],
                      nota="Audit con «Perimetro Part 145» nel Sistema di gestione: l'anno ne richiede almeno uno di "
                           "sistema e uno di prodotto (MT CN 70 §5.3). I rilievi aperti valgono anche per gli anni prima.",
                      vuoto=f"Nessun audit Part 145 nel {oggi.year}: pianificali nel programma audit.")
    rilievi_per_audit: dict = {}
    for r in rilievi:
        rilievi_per_audit[r.audit_id] = rilievi_per_audit.get(r.audit_id, 0) + 1
    for au in audit_anno:
        s_audit.aggiungi([au.numero, au.get_part145_display().replace("Sì — ", ""), _d(au.data_inizio),
                          au.get_stato_display(), rilievi_per_audit.get(au.pk, 0)],
                         ("Apri l'audit", reverse("sistema_gestione:audit_dettaglio", args=[au.pk])))

    # ── Non conformità Part 145 (Anomalie) ───────────────────────────────────
    nc_aperte = _sicuro(avvisi, "Non conformità", _nc_part145, []) if con_nc else []
    s_nc = Sezione("nc", "Non conformità Part 145 aperte", ["NC", "OP", "P/N", "Stato", "Aperta il", "Verifica prevista"],
                   nota="NC segnate «Componente Part 145» nel modulo Anomalie (MT CN 70 §6.13): vanno comunicate "
                        "all'organizzazione di manutenzione, che decide le azioni.",
                   vuoto="Nessuna NC Part 145 aperta.")
    for nc in nc_aperte:
        s_nc.aggiungi([nc.protocollo, nc.op_titolo, nc.part_number, nc.get_stato_display(),
                       _d(timezone.localtime(nc.created_at).date()) if nc.created_at else "", _d(nc.verifica_prevista)],
                      ("Apri la NC", reverse("anomalie_nc_dettaglio", args=[nc.pk])))

    # ── Scadenze ─────────────────────────────────────────────────────────────
    s_scad = Sezione("scadenze", f"Scadenze entro {giorni} giorni (e già scadute)",
                     ["Data", "Tipo", "Riferimento", "Dettaglio", "Stato"], vuoto=f"Nessuna scadenza nei prossimi {giorni} giorni.")
    scadenze: list = []
    for p in processi:
        scad = p.scadenza_effettiva
        if scad and scad <= limite:
            scadenze.append(([scad, "Processo", p.nome, str(p.cliente) if p.cliente_id else "",
                              _stato_scadenza(scad, oggi)], apri_mod128(p), None))
    for r in requisiti:
        if r.data_scadenza and r.data_scadenza <= limite:
            scadenze.append(([r.data_scadenza, "Contratto" if r.tipo == r.TIPO_CONTRATTO else "Requisito di processo",
                              r.descrizione, r.processo.nome, _stato_scadenza(r.data_scadenza, oggi)],
                             apri_mod128(r.processo), None))
    for o in (MaintenanceOccurrence.objects.filter(asset_id__in=ids, status=MaintenanceOccurrence.STATUS_OPEN, due_date__lte=limite)
              .select_related("asset", "plan").order_by("due_date")):
        tipo = "Taratura" if o.plan.maintenance_type == MaintenanceInterventionTemplate.TYPE_CALIBRATION else "Manutenzione"
        scadenze.append(([o.due_date, tipo, f"{o.asset.asset_tag} — {o.asset.name}", getattr(o.plan, "label", "") or "",
                          _stato_scadenza(o.due_date, oggi)], ("Apri la scheda", reverse("assets:asset_view", args=[o.asset_id])),
                         None))
    if con_personale and abilitazioni:
        for c in (CertificazioneIndividuale.objects.filter(abilitazione__in=abilitazioni, data_scadenza__lte=limite)
                  .select_related("abilitazione", "abilitazione__processo").order_by("data_scadenza")):
            scadenze.append(([c.data_scadenza, "Certificazione", nome_persona(c.abilitazione),
                              f"{c.schema} {c.livello}".strip() + f" · {c.abilitazione.processo.nome}",
                              _stato_scadenza(c.data_scadenza, oggi)], apri_mod128(c.abilitazione.processo), None))
    for d in formazione:
        if d.data_scadenza and d.data_scadenza <= limite:
            scadenze.append(([d.data_scadenza, "Formazione", nomi.get(d.legacy_anagrafica_id) or f"#{d.legacy_anagrafica_id}",
                              str(d.corso), _stato_scadenza(d.data_scadenza, oggi)], scadenzario, None))
    for r in rilievi:
        if r.scadenza_azione and r.scadenza_azione <= limite:
            scadenze.append(([r.scadenza_azione, "Azione da audit", r.audit.numero, r.get_esito_display(),
                              _stato_scadenza(r.scadenza_azione, oggi)],
                             ("Apri l'audit", reverse("sistema_gestione:audit_dettaglio", args=[r.audit_id])), None))
    for nc in nc_aperte:
        if nc.verifica_prevista and nc.verifica_prevista <= limite:
            scadenze.append(([nc.verifica_prevista, "Verifica NC", nc.protocollo, nc.part_number or nc.op_titolo,
                              _stato_scadenza(nc.verifica_prevista, oggi)],
                             ("Apri la NC", reverse("anomalie_nc_dettaglio", args=[nc.pk])), None))
    scadenze.sort(key=lambda t: t[0][0])
    for riga, azione, _ in scadenze:
        riga[0] = _d(riga[0])
        s_scad.aggiungi(riga, azione)

    # ── Situazioni aperte (Da sistemare) ─────────────────────────────────────
    s_aperte = Sezione("aperte", "Da sistemare", ["Priorità", "Tipo", "Riferimento", "Dettaglio"],
                       nota="Situazioni aperte, prima quelle a priorità alta (bloccano l'operatività o l'audit).",
                       vuoto="Nessuna situazione aperta.")

    def aperta(tipo, riferimento, dettaglio, azione):
        priorita = PRIORITA_ALTA if tipo in TIPI_PRIORITA_ALTA else PRIORITA_MEDIA
        s_aperte.aggiungi([priorita, tipo, riferimento, dettaglio], azione)

    for p in processi:
        if p.stato != p.STATO_ATTIVO:
            aperta("Processo non attivo", p.nome, p.get_stato_display(), apri_mod128(p))
        scad = p.scadenza_effettiva
        if scad and scad < oggi:
            aperta("Processo scaduto", p.nome, f"scaduto il {_d(scad)}", apri_mod128(p))
        elif scad is None and p.tipo_validita != p.VALIDITA_ILLIMITATA:
            aperta("Processo senza scadenza", p.nome,
                   "validità a data o a periodo ma nessuna data: imposta la scadenza o «Illimitata»", apri_mod128(p))
        if con_personale and not any(ab.processo_id == p.pk and ab.stato == ab.STATO_ATTIVA for ab in abilitazioni):
            aperta("Processo senza persone abilitate", p.nome, "nessuna abilitazione attiva in MOD.128", apri_mod128(p))
    for r in requisiti:
        if r.obbligatorio and (r.stato == r.STATO_SCADUTO or (r.data_scadenza and r.data_scadenza < oggi)):
            aperta("Contratto scaduto" if r.tipo == r.TIPO_CONTRATTO else "Requisito di processo scaduto", r.descrizione,
                   r.processo.nome + (f" · scaduto il {_d(r.data_scadenza)}" if r.data_scadenza else ""),
                   apri_mod128(r.processo))
    for a in attrezzature:
        scheda = ("Apri la scheda", reverse("assets:asset_view", args=[a.pk]))
        c = controlli.get(a.pk) or _controllo_vuoto()
        if c["scaduta"]:
            aperta("Taratura scaduta", f"{a.asset_tag} — {a.name}",
                   f"scaduta il {_d(c['prossima'])}: segregare fino alla taratura (MOE §2.4.3.4)", scheda)
        if a.status == a.STATUS_IN_REPAIR:
            aperta("Attrezzatura in riparazione", f"{a.asset_tag} — {a.name}", a.reparto or "", scheda)
        mancanti = [n for v, n in ((a.serial_number, "seriale"), (a.reparto, "reparto")) if not v]
        if mancanti:
            aperta("Attrezzatura con dati mancanti", f"{a.asset_tag} — {a.name}", "manca " + " e ".join(mancanti),
                   ("Completa la scheda", reverse("assets:asset_edit", args=[a.pk])))
    for wo in odl_aperti:
        aperta(f"OdL aperto ({tipi_odl.get(wo.kind, wo.kind)})", f"{wo.asset.asset_tag} — {wo.asset.name}",
               f"{wo.title} · aperto il {timezone.localtime(wo.opened_at):%d/%m/%Y}",
               ("Apri l'OdL", reverse("assets:wo_view", args=[wo.pk])))
    for ab in abilitazioni:
        if ab.stato != ab.STATO_ATTIVA:
            aperta(f"Abilitazione {stati_ab.get(ab.stato, ab.stato).lower()}", nome_persona(ab), ab.processo.nome,
                   apri_mod128(ab.processo))
        for t in timbri.get(ab.pk, []):
            if t.is_sospeso:
                aperta("Timbro sospeso", nome_persona(ab), t.sospeso_motivo or ab.processo.nome, apri_mod128(ab.processo))
    for d in formazione:
        if d.stato_scadenza in ("SCADUTO", "MAI_FREQUENTATO"):
            aperta("Formazione scaduta" if d.stato_scadenza == "SCADUTO" else "Formazione mancante",
                   nomi.get(d.legacy_anagrafica_id) or f"#{d.legacy_anagrafica_id}",
                   str(d.corso) + (f" · scaduta il {_d(d.data_scadenza)}" if d.data_scadenza else ""), scadenzario)
    for persona, corso in formazione_mancante:
        aperta("Formazione mancante", nomi.get(persona) or f"#{persona}", str(corso), scadenzario)
    if con_audit:
        tipi_fatti = {au.part145 for au in audit_anno}
        from sistema_gestione.models import Audit

        for codice, nome in ((Audit.PART145_SISTEMA, "di sistema"), (Audit.PART145_PRODOTTO, "di prodotto")):
            if codice not in tipi_fatti:
                aperta(f"Audit Part 145 {nome} da pianificare", str(oggi.year),
                       "MT CN 70 §5.3: almeno un audit di sistema e uno di prodotto all'anno",
                       ("Apri gli audit", reverse("sistema_gestione:audit_index")))
        for r in rilievi:
            if r.scadenza_azione and r.scadenza_azione < oggi:
                aperta("Azione da audit scaduta", r.audit.numero,
                       f"{r.get_esito_display()} · azione scaduta il {_d(r.scadenza_azione)}",
                       ("Apri l'audit", reverse("sistema_gestione:audit_dettaglio", args=[r.audit_id])))
    for nc in nc_aperte:
        aperta("NC Part 145 aperta", nc.protocollo, f"{nc.part_number or nc.op_titolo} · {nc.get_stato_display()}",
               ("Apri la NC", reverse("anomalie_nc_dettaglio", args=[nc.pk])))
    s_aperte.ordina(lambda riga: 0 if riga[0] == PRIORITA_ALTA else 1)

    # ── Documenti di riferimento ─────────────────────────────────────────────
    s_doc = Sezione("documenti", "Documenti di riferimento",
                    ["Documento", "A cosa serve", "Titolo nel registro SGI", "Revisione", "Data revisione"],
                    nota="Revisione corrente nel registro SGI (Procedure) e riferimenti registrati sui processi MOD.128.",
                    chiusa=True)
    for d in _sicuro(avvisi, "Documenti SGI", _documenti_sgi, []):
        rev = d["rev"]
        if rev is None:
            s_doc.aggiungi([d["codice"], d["perche"], "Non trovato nel registro SGI", "", ""])
            continue
        s_doc.aggiungi([rev.document.code, d["perche"], rev.document.title, rev.revision_code or "—",
                        _d(rev.revision_date)],
                       ("Apri il documento", reverse("procedure_refresh:document_open", args=[rev.pk])))
    tipi_rif = None
    for p in processi:
        for rif in p.riferimenti.all():
            tipi_rif = tipi_rif or dict(type(rif).TIPO_CHOICES)
            s_doc.aggiungi([rif.codice, f"{tipi_rif.get(rif.tipo, rif.tipo)} · {p.nome}", rif.note, "", ""],
                           apri_mod128(p))

    sezioni = ([s_aperte, s_scad, s_processi, s_req, s_attr] + ([s_pers, s_form] if con_personale else [])
               + ([s_audit] if con_audit else []) + ([s_nc] if con_nc else []) + [s_doc])

    # ── Riepilogo ────────────────────────────────────────────────────────────
    in_uso = sum(1 for a in attrezzature if a.status == a.STATUS_IN_USE)
    tarate = sum(1 for c in controlli.values() if c["classe"] == CLASSE_TARATURA)
    tarature_scadute = sum(1 for c in controlli.values() if c["scaduta"])
    alte = sum(1 for r in s_aperte.righe if r[0] == PRIORITA_ALTA)
    scadute = sum(1 for r in s_scad.righe if r[-1] == "Scaduta")
    persone = len({(a.legacy_anagrafica_id, a.nominativo_esterno) for a in abilitazioni if a.stato == a.STATO_ATTIVA})
    kpi = [
        {"etichetta": "Da sistemare", "valore": len(s_aperte.righe), "sezione": "aperte",
         "nota": f"{alte} a priorità alta" if alte else "nessuna urgente",
         "stato": "ko" if alte else ("warn" if s_aperte.righe else "ok")},
        {"etichetta": f"Scadenze {giorni} giorni", "valore": len(s_scad.righe), "sezione": "scadenze",
         "nota": f"{scadute} già scadute" if scadute else "nessuna scaduta",
         "stato": "ko" if scadute else ("warn" if s_scad.righe else "ok")},
        {"etichetta": "Processi PART 145", "valore": len(processi), "sezione": "processi",
         "nota": "regime + perimetro", "stato": "" if processi else "warn"},
        {"etichetta": "Attrezzature", "valore": len(attrezzature), "sezione": "attrezzature",
         "nota": f"{in_uso} in uso · {tarate} con taratura", "stato": "ko" if tarature_scadute else ""},
    ]
    if con_personale:
        formazione_ko = sum(1 for r in s_form.righe if r[2] in ("Scaduta", "Mai frequentata"))
        kpi.append({"etichetta": "Persone abilitate", "valore": persone, "sezione": "personale",
                    "nota": f"{formazione_ko} con formazione da fare" if formazione_ko else "formazione in regola",
                    "stato": "warn" if not persone or formazione_ko else ""})
    if con_audit:
        mancano = 2 - len({au.part145 for au in audit_anno})
        kpi.append({"etichetta": f"Audit Part 145 {oggi.year}", "valore": len(audit_anno), "sezione": "audit",
                    "nota": f"{len(rilievi)} rilievi aperti" if rilievi else ("manca sistema o prodotto" if mancano else "sistema e prodotto"),
                    "stato": "warn" if mancano or rilievi else "ok"})
    if con_nc:
        kpi.append({"etichetta": "NC Part 145 aperte", "valore": len(nc_aperte), "sezione": "nc",
                    "nota": "da comunicare all'AMO" if nc_aperte else "nessuna", "stato": "ko" if nc_aperte else "ok"})
    riepilogo = [("Processi PART 145", len(processi)), ("Attrezzature", len(attrezzature)), ("In uso", in_uso),
                 ("Con taratura", tarate)]
    if con_personale:
        riepilogo.append(("Persone abilitate", persone))
    riepilogo += [(f"Scadenze {giorni} giorni", len(s_scad.righe)), ("Situazioni aperte", len(s_aperte.righe))]
    if con_audit:
        riepilogo.append((f"Audit Part 145 {oggi.year}", len(audit_anno)))
    if con_nc:
        riepilogo.append(("NC Part 145 aperte", len(nc_aperte)))
    return {"oggi": oggi, "riepilogo": riepilogo, "kpi": kpi, "sezioni": sezioni, "con_personale": con_personale,
            "avvisi": avvisi, "dismesse": dismesse, "badge": BADGE,
            "attrezzature": attrezzature, "controlli": controlli, "stati_asset": stati_asset}


def _titolo_file(oggi) -> str:
    return f"dossier-part145-{oggi:%Y%m%d}"


def _intestazione(utente, con_personale) -> str:
    chi = ""
    if utente is not None and getattr(utente, "is_authenticated", False):
        chi = f" da {utente.get_full_name() or utente.get_username()}"
    testo = f"Estratto il {timezone.localtime():%d/%m/%Y %H:%M}{chi}"
    if not con_personale:
        testo += " · senza dati del personale (serve l'accesso ai processi qualificati MOD.128)"
    return testo


def esporta_xlsx(dati: dict, utente=None) -> tuple[bytes, str]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from core.excel_export import _brand_header, write_cell

    wb = Workbook()
    ws = wb.active
    ws.title = "Riepilogo"
    riga = _brand_header(ws, title="Dossier PART 145", subtitle=_intestazione(utente, dati["con_personale"]),
                         filters_label="", logo=True)
    for etichetta, valore in dati["riepilogo"]:
        write_cell(ws, riga, 1, etichetta).font = Font(bold=True)
        write_cell(ws, riga, 2, valore)
        riga += 1
    ws.column_dimensions["A"].width = 30
    fill = PatternFill("solid", fgColor="0C2545")
    for s in dati["sezioni"]:
        foglio = wb.create_sheet(s.titolo[:31].replace("/", "-"))
        write_cell(foglio, 1, 1, s.titolo).font = Font(bold=True, size=13, color="0C2545")
        if s.nota:
            write_cell(foglio, 2, 1, s.nota).font = Font(size=9, color="5B6B7C")
        for c, h in enumerate(s.intestazioni, start=1):
            cella = write_cell(foglio, 4, c, h)
            cella.fill, cella.font, cella.alignment = fill, Font(bold=True, color="FFFFFF"), Alignment(vertical="center")
        for r, valori in enumerate(s.righe, start=5):
            for c, v in enumerate(valori, start=1):
                write_cell(foglio, r, c, v)
        if not s.righe:
            write_cell(foglio, 5, 1, s.vuoto)
        for c, h in enumerate(s.intestazioni, start=1):
            larghezza = max([len(str(h))] + [len(str(r[c - 1])) for r in s.righe if c - 1 < len(r)])
            foglio.column_dimensions[get_column_letter(c)].width = min(60, max(10, larghezza + 2))
        foglio.freeze_panes = "A5"
        if s.righe:
            foglio.auto_filter.ref = f"A4:{get_column_letter(len(s.intestazioni))}{4 + len(s.righe)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), _titolo_file(dati["oggi"]) + ".xlsx"


def esporta_pdf(dati: dict, utente=None) -> tuple[bytes, str]:
    from reportlab.platypus import Paragraph, Spacer

    from core.pdf import PdfTheme, build_styles, data_table, header_footer_callback, make_document, section_heading

    theme = PdfTheme.from_branding()
    styles = build_styles(theme)
    buf = io.BytesIO()
    doc = make_document(buf, title="Dossier PART 145", landscape=True)
    larghezza = doc.pagesize[0] - doc.leftMargin - doc.rightMargin
    elementi: list = []
    riepilogo = [[Paragraph(escape(str(e)), styles["table_header"]) for e, _ in dati["riepilogo"]],
                 [Paragraph(escape(str(v)), styles["cell"]) for _, v in dati["riepilogo"]]]
    elementi.append(data_table(riepilogo, theme, col_widths=[larghezza / len(dati["riepilogo"])] * len(dati["riepilogo"])))
    for s in dati["sezioni"]:
        elementi.append(Spacer(1, 10))
        elementi += section_heading(escape(s.titolo), theme, styles)
        if s.nota:
            elementi.append(Paragraph(escape(s.nota), styles["body"]))
        if not s.righe:
            elementi.append(Paragraph(escape(s.vuoto), styles["body"]))
            continue
        righe = [[Paragraph(escape(str(h)), styles["table_header"]) for h in s.intestazioni]]
        righe += [[Paragraph(escape("" if v is None else str(v)), styles["cell"]) for v in r] for r in s.righe]
        elementi.append(data_table(righe, theme, col_widths=[larghezza / len(s.intestazioni)] * len(s.intestazioni),
                                   repeat_rows=1, extra_style=[("FONTSIZE", (0, 0), (-1, -1), 7),
                                                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    draw = header_footer_callback(theme, title="DOSSIER PART 145", subtitle=_intestazione(utente, dati["con_personale"]))
    doc.build(elementi, onFirstPage=draw, onLaterPages=draw)
    return buf.getvalue(), _titolo_file(dati["oggi"]) + ".pdf"


# ── Elenchi operativi ─────────────────────────────────────────────────────────


def esporta_tarature_xlsx(dati: dict, utente=None) -> tuple[bytes, str]:
    """Strumenti con taratura: ente, ultima e prossima (MT CN 70 §6.12: elenco da comunicare all'AMO)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    from core.excel_export import _brand_header, write_cell

    wb = Workbook()
    ws = wb.active
    ws.title = "Strumenti tarati"
    riga = _brand_header(ws, title="PART 145 — Strumenti con taratura",
                         subtitle=_intestazione(utente, True), filters_label="", logo=True)
    intestazioni = ["Asset", "Nome", "Reparto", "Seriale", "Costruttore / modello", "Stato", "Ente taratura",
                    "Ultima taratura", "Prossima taratura", "Esito"]
    fill = PatternFill("solid", fgColor="0C2545")
    for c, h in enumerate(intestazioni, start=1):
        cella = write_cell(ws, riga, c, h)
        cella.fill, cella.font = fill, Font(bold=True, color="FFFFFF")
    inizio = riga
    oggi = dati["oggi"]
    for a in dati["attrezzature"]:
        c = dati["controlli"][a.pk]
        if c["classe"] != CLASSE_TARATURA:
            continue
        riga += 1
        valori = [a.asset_tag, a.name, a.reparto, a.serial_number or "",
                  " ".join(filter(None, [a.manufacturer, a.model])), dati["stati_asset"].get(a.status, a.status),
                  ", ".join(c["enti"]), _d(c["ultima"]), _d(c["prossima"]),
                  _stato_scadenza(c["prossima"], oggi) if c["prossima"] else "Nessuna taratura pianificata"]
        for col, v in enumerate(valori, start=1):
            write_cell(ws, riga, col, v)
    if riga == inizio:
        write_cell(ws, riga + 1, 1, "Nessuno strumento con piano di taratura tra le attrezzature PART 145.")
    for col, larghezza in enumerate([14, 34, 16, 18, 24, 14, 30, 15, 16, 22], start=1):
        ws.column_dimensions[get_column_letter(col)].width = larghezza
    ws.freeze_panes = ws.cell(row=inizio + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"part145-strumenti-tarati-{oggi:%Y%m%d}.xlsx"


def esporta_controllo_semestrale_pdf(dati: dict, utente=None) -> tuple[bytes, str]:
    """Elenco per il controllo semestrale delle attrezzature (MT CN 17 §6.1.3, Mod.072).

    Una pagina per reparto: il capo reparto esegue l'esame visivo (idoneità, completezza,
    identificazione/targhetta), timbra e firma; poi CQF verifica.
    """
    from reportlab.platypus import PageBreak, Paragraph, Spacer

    from core.pdf import PdfTheme, build_styles, data_table, header_footer_callback, make_document, section_heading

    theme = PdfTheme.from_branding()
    styles = build_styles(theme)
    buf = io.BytesIO()
    doc = make_document(buf, title="Controllo semestrale attrezzature PART 145", landscape=True)
    larghezza = doc.pagesize[0] - doc.leftMargin - doc.rightMargin
    per_reparto: dict = {}
    for a in dati["attrezzature"]:
        if a.status != a.STATUS_RETIRED:
            per_reparto.setdefault(a.reparto or "Senza reparto", []).append(a)
    elementi: list = []
    colonne = ["Codice", "Nome", "Seriale", "Controllo", "Targhetta presente", "Integra e completa", "Esito", "Note"]
    pesi = [0.11, 0.25, 0.12, 0.12, 0.09, 0.09, 0.09, 0.13]
    for n, (reparto, assets) in enumerate(sorted(per_reparto.items())):
        if n:
            elementi.append(PageBreak())
        elementi += section_heading(escape(f"Reparto: {reparto}"), theme, styles)
        elementi.append(Paragraph(
            "Controllo semestrale (gennaio / giugno) — data: ____________ · Esame visivo: idoneità, completezza "
            "e identificazione. Senza targhetta l'attrezzatura è non conforme e non va immagazzinata.", styles["body"]))
        elementi.append(Spacer(1, 6))
        righe = [[Paragraph(escape(h), styles["table_header"]) for h in colonne]]
        for a in sorted(assets, key=lambda x: x.asset_tag):
            righe.append([Paragraph(escape(str(v)), styles["cell"]) for v in
                          [a.asset_tag, a.name, a.serial_number or "", dati["controlli"][a.pk]["classe"],
                           "[ ] Sì  [ ] No", "[ ] Sì  [ ] No", "", ""]])
        elementi.append(data_table(righe, theme, col_widths=[larghezza * p for p in pesi], repeat_rows=1,
                                   extra_style=[("FONTSIZE", (0, 0), (-1, -1), 8), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        elementi.append(Spacer(1, 18))
        elementi.append(Paragraph("Capo reparto (timbro e firma): ______________________________"
                                  "&nbsp;&nbsp;&nbsp;&nbsp;Verifica CQF: ______________________________",
                                  styles["body"]))
    if not elementi:
        elementi.append(Paragraph("Nessuna attrezzatura PART 145 da controllare.", styles["body"]))
    draw = header_footer_callback(theme, title="CONTROLLO SEMESTRALE ATTREZZATURE PART 145",
                                  subtitle="Rif. MT CN 17 §6.1.3 · Mod.072 · " + _intestazione(utente, True))
    doc.build(elementi, onFirstPage=draw, onLaterPages=draw)
    return buf.getvalue(), f"part145-controllo-semestrale-{dati['oggi']:%Y%m%d}.pdf"


# ── Avviso scadenze ───────────────────────────────────────────────────────────


@event_flow("part145_scadenze", skipped_result=0)
def notifica_scadenze(*, oggi=None, override=None, fail_silently: bool = True) -> int:
    """Mail settimanale a Qualità: scadenze PART 145 entro 90 giorni e già scadute.

    Contratto, processi, requisiti, tarature/manutenzioni e certificazioni del personale:
    l'accordo con l'organizzazione di manutenzione chiede di avvisare almeno 3 mesi prima
    (MO-ID-009.24 §2). Nessuna scadenza = nessuna mail. Ritorna le mail inviate (0 o 1).
    """
    from anagrafica.services.mpq_timbri import MSM_CONFIG_KEY
    from anagrafica.services.reminders import get_reminder_recipients
    from core.email_utils import send_hub_mail

    recipients = get_reminder_recipients(MSM_CONFIG_KEY, override=override)
    if not recipients:
        return 0
    dati = costruisci(oggi=oggi, con_personale=True, con_audit=True, con_nc=True)
    scadenze = next(s for s in dati["sezioni"] if s.codice == "scadenze")
    if not scadenze.righe:
        return 0
    cella = "padding:8px 12px;border-top:1px solid #e2e8f0;"
    intestazione = "text-align:left;padding:9px 12px;color:#64748b;font-size:12px;"
    righe = "".join("<tr>" + "".join(f"<td style='{cella}'>{escape(str(v))}</td>" for v in r) + "</tr>"
                    for r in scadenze.righe)
    tabella = (
        "<table role='presentation' width='100%' cellpadding='0' cellspacing='0' "
        "style='width:100%;border-collapse:collapse;background:#f6f9fc;border:1px solid #d8e1ec;'>"
        "<tr style='background:#eef3f8;'>"
        + "".join(f"<th style='{intestazione}'>{escape(h)}</th>" for h in scadenze.intestazioni)
        + "</tr>" + righe + "</table>"
    )
    scadute = sum(1 for r in scadenze.righe if r[-1] == "Scaduta")
    testo = (f"{len(scadenze.righe)} scadenze PART 145 entro {GIORNI_SCADENZE} giorni"
             + (f", di cui {scadute} già scadute" if scadute else "")
             + ". Le scadenze della formazione vanno comunicate al Maintenance Manager almeno 3 mesi prima.")
    return send_hub_mail(
        subject=f"[PART 145] Scadenze entro {GIORNI_SCADENZE} giorni ({len(scadenze.righe)})",
        body_text=testo,
        recipients=recipients,
        title="PART 145 — scadenze",
        email_type="Qualità",
        badge="PART 145",
        section_label="Dossier PART 145",
        body_html_fragment=f"<p style='margin:0 0 14px;'>{escape(testo)}</p>{tabella}",
        fail_silently=fail_silently,
    )
