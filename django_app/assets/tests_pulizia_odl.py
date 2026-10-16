from datetime import timedelta
from io import StringIO

from django.core.management import call_command
from django.test import TestCase
from django.utils import timezone

from assets.models import (
    Asset,
    AssetCategory,
    MaintenanceInterventionTemplate,
    MaintenanceOccurrence,
    WorkOrder,
    WorkOrderLog,
)


class PuliziaOdlVecchioMetodoTests(TestCase):
    def setUp(self):
        category = AssetCategory.objects.create(
            code="cnc-pulizia-odl", label="CNC pulizia OdL", base_asset_type=Asset.TYPE_CNC, is_active=True
        )
        self.plan = MaintenanceInterventionTemplate.objects.create(
            code="pulizia-odl-olio", label="Cambio olio", asset_category=category
        )
        self.asset = Asset.objects.create(
            asset_tag="CNC-PUL-01", name="Centro sintetico", asset_type=Asset.TYPE_CNC,
            asset_category=category, status=Asset.STATUS_IN_USE,
        )

    def _wo(self, title, **kwargs):
        kwargs.setdefault("origin", WorkOrder.ORIGIN_PERIODIC)
        return WorkOrder.objects.create(asset=self.asset, title=title, **kwargs)

    def _occ(self, wo, status):
        # Una scadenza per data: (piano, asset, data) e' univoco.
        days = 5 + MaintenanceOccurrence.objects.count()
        return MaintenanceOccurrence.objects.create(
            plan=self.plan, asset=self.asset, due_date=timezone.localdate() + timedelta(days=days),
            work_order=wo, status=status,
        )

    def _run(self, *args):
        out = StringIO()
        call_command("pulisci_odl_vecchio_metodo", *args, stdout=out)
        return out.getvalue()

    def test_dry_run_non_modifica_e_classifica(self):
        senza = self._wo("Vecchio senza scadenze")
        annullate = self._wo("Scadenze annullate")
        self._occ(annullate, MaintenanceOccurrence.STATUS_CANCELED)
        vivo = self._wo("Ancora valido")
        self._occ(vivo, MaintenanceOccurrence.STATUS_OPEN)
        manuale = self._wo("Guasto", origin=WorkOrder.ORIGIN_MANUAL)
        iniziato = self._wo("Gia' iniziato", started_at=timezone.now())

        out = self._run()

        self.assertIn("DRY-RUN", out)
        self.assertIn("da annullare: 2", out)
        self.assertIn("da guardare a mano (lavoro registrato): 1", out)
        self.assertIn(f"#{senza.id} ", out)
        self.assertIn(f"#{annullate.id} ", out)
        self.assertIn(f"#{iniziato.id} ", out)
        self.assertNotIn(f"#{vivo.id} ", out)
        self.assertNotIn(f"#{manuale.id} ", out)
        self.assertEqual(WorkOrder.objects.filter(status=WorkOrder.STATUS_OPEN).count(), 5)

    def test_applica_annulla_solo_i_puliti_e_rispetta_escludi(self):
        senza = self._wo("Vecchio senza scadenze")
        escluso = self._wo("Da tenere")
        eseguito = self._wo("Con scadenza eseguita")
        self._occ(eseguito, MaintenanceOccurrence.STATUS_DONE)
        vivo = self._wo("Ancora valido")
        self._occ(vivo, MaintenanceOccurrence.STATUS_OPEN)

        out = self._run("--escludi", str(escluso.id), "--applica")

        self.assertIn("Annullati: 1", out)
        senza.refresh_from_db()
        self.assertEqual(senza.status, WorkOrder.STATUS_CANCELED)
        self.assertIsNotNone(senza.closed_at)
        self.assertTrue(WorkOrderLog.objects.filter(work_order=senza, note__startswith="Intervento annullato").exists())
        for wo in (escluso, eseguito, vivo):
            wo.refresh_from_db()
            self.assertEqual(wo.status, WorkOrder.STATUS_OPEN)

    def test_report_excel(self):
        self._wo("Vecchio senza scadenze")
        import tempfile
        from pathlib import Path

        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "odl.xlsx"
            self._run("--report", str(path))
            rows = list(load_workbook(path).active.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "OdL")
        self.assertEqual(rows[1][7], "da annullare")
