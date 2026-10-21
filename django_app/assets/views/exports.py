"""Esportazioni XLSX/PDF di asset, interventi, macchine e schede.

Estratto meccanicamente da assets/views.py (refactor C4): nessun cambio di comportamento.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

if TYPE_CHECKING:
    # Solo per le annotazioni di tipo (nessun import a runtime).
    from django.db.models import QuerySet

import io

from django.contrib.auth.decorators import login_required
from django.db import DatabaseError
from django.db.models import Q
from django.http import HttpRequest, HttpResponse
from django.shortcuts import get_object_or_404, redirect
from django.utils import timezone


# Etichetta con cui le voci di audit si agganciano all'asset: e' la stessa che
# `core.audit.storico_oggetto` usa per ripescarle sulla scheda. Costante e non
# stringa sparsa perche' un refuso qui non romperebbe nulla — semplicemente la
# voce non comparirebbe mai nello storico, e non se ne accorgerebbe nessuno.
AUDIT_OGGETTO_ASSET = "assets.asset"
from core.pdf import (
    PdfTheme,
)
from core.table_pdf import render_table_pdf as _report_table_pdf
from ..forms import (
    AssetFilterForm,
    PRODUCTION_ASSET_TYPES,
    WorkMachineFilterForm,
)
from ..models import (
    Asset,
    AssetReportTemplate,
    WorkOrder,
)

from ._helpers import (  # noqa: F401
    _clean_string,
    _model_table_exists,
)
from ._helpers_b import (  # noqa: F401
    _build_asset_report_snapshot,
)
from ._helpers_c import (  # noqa: F401
    _apply_workorder_list_filters,
    _apply_workorder_operational_view,
    _category_subtree_ids,
    _resolve_workorder_operational_view,
    _xl_write_sheet,
)

def _active_asset_report_template(report_code: str) -> AssetReportTemplate | None:
    report_value = _clean_string(report_code)
    if not report_value or not _model_table_exists(AssetReportTemplate):
        return None
    try:
        return (
            AssetReportTemplate.objects.filter(report_code=report_value, is_active=True)
            .select_related("uploaded_by")
            .order_by("-updated_at", "-id")
            .first()
        )
    except DatabaseError:
        return None


def render_asset_report_pdf(asset: Asset, snapshot: dict[str, object], *, template_name: str = "") -> bytes:
    """Report PDF della scheda asset col template standard del portale (``core.pdf``).

    Allinea la veste a tutti gli altri report (header/footer branded via
    ``header_footer_callback``, tabelle zebra col colore primario) e mette in
    evidenza il flag PART 145.
    """
    import io as _io
    from html import escape as _escape

    from reportlab.lib import colors as _colors
    from reportlab.lib.styles import ParagraphStyle as _ParagraphStyle
    from reportlab.lib.units import mm as _mm
    from reportlab.platypus import Paragraph as _Paragraph, Spacer as _Spacer, Table as _Table, TableStyle as _TableStyle

    from core.pdf import build_styles, data_table, header_footer_callback, make_document, section_heading

    theme = PdfTheme.from_branding()
    styles = build_styles(theme)
    buf = _io.BytesIO()
    doc = make_document(buf, title=f"Report asset {asset.asset_tag or ''}".strip())
    page_w = doc.pagesize[0] - doc.leftMargin - doc.rightMargin
    story: list = []

    # Evidenza PART 145 (rosso) in testa quando applicabile.
    if getattr(asset, "part_145", False):
        badge_style = _ParagraphStyle(
            "p145Badge", fontName="Helvetica-Bold", fontSize=11, textColor=_colors.white, alignment=1,
        )
        note_style = _ParagraphStyle(
            "p145Note", fontName="Helvetica-Bold", fontSize=9.5, textColor=_colors.HexColor("#7f1d1d"),
        )
        banner = _Table(
            [[
                _Paragraph("PART 145", badge_style),
                _Paragraph("Asset soggetto al regolamento aeronautico PART 145", note_style),
            ]],
            colWidths=[30 * _mm, page_w - 30 * _mm],
        )
        banner.setStyle(_TableStyle([
            ("BACKGROUND", (0, 0), (0, 0), _colors.HexColor("#dc2626")),
            ("BACKGROUND", (1, 0), (1, 0), _colors.HexColor("#fef2f2")),
            ("BOX", (0, 0), (-1, -1), 0.8, _colors.HexColor("#dc2626")),
            ("INNERGRID", (0, 0), (-1, -1), 0.8, _colors.HexColor("#dc2626")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        story.append(banner)
        story.append(_Spacer(1, 5 * _mm))

    def kv_section(title: str, rows) -> None:
        rows = list(rows or [])
        if not rows:
            return
        story.extend(section_heading(title, theme, styles))
        table_rows = [
            [
                _Paragraph(_escape(str(key)), styles["label"]),
                _Paragraph(_escape(str(value)).replace("\n", "<br/>"), styles["value"]),
            ]
            for key, value in rows
        ]
        story.append(data_table(table_rows, theme, header=False, col_widths=[54 * _mm, page_w - 54 * _mm]))
        story.append(_Spacer(1, 3 * _mm))

    def list_section(title: str, headers, rows) -> None:
        rows = list(rows or [])
        if not rows:
            return
        story.extend(section_heading(title, theme, styles))
        col_w = page_w / max(1, len(headers))
        table_rows = [[_Paragraph(_escape(str(h)), styles["table_header"]) for h in headers]]
        for row in rows:
            table_rows.append([_Paragraph(_escape(str(cell)), styles["cell"]) for cell in row])
        story.append(data_table(table_rows, theme, col_widths=[col_w] * len(headers), repeat_rows=1))
        story.append(_Spacer(1, 3 * _mm))

    kv_section("Riepilogo asset", snapshot.get("summary_rows"))
    kv_section("Dati tecnici", snapshot.get("technical_rows"))
    kv_section("Campi categoria", snapshot.get("category_rows"))
    kv_section("Campi personalizzati", snapshot.get("custom_rows"))
    list_section("Documenti", ["Categoria", "File", "Data", "Peso"], snapshot.get("document_rows"))
    list_section("Work order recenti", ["Data", "Tipo", "Titolo", "Stato"], snapshot.get("workorder_rows"))
    list_section("Ticket collegati", ["Ticket", "Tipo", "Titolo", "Stato"], snapshot.get("ticket_rows"))
    list_section("Manutenzione periodica", ["Piano", "Prossima data", "Fornitore"], snapshot.get("periodic_rows"))

    if not story:
        story.append(_Paragraph("Nessun dato disponibile per questo asset.", styles["body"]))

    subtitle = f"{asset.asset_tag or '-'} — {asset.name or '-'}"
    if template_name:
        subtitle = f"{subtitle} | Template: {template_name}"
    draw = header_footer_callback(theme, title="REPORT ASSET", subtitle=subtitle)
    doc.build(story, onFirstPage=draw, onLaterPages=draw)
    return buf.getvalue()


@login_required
def asset_part_145_export_pdf(request: HttpRequest) -> HttpResponse:
    """Dossier PART 145 in PDF: le stesse sezioni della pagina."""
    from ..services import part145_dossier

    dati = part145_dossier.costruisci(**part145_dossier.visibilita(request))
    data, nome = part145_dossier.esporta_pdf(dati, request.user)
    resp = HttpResponse(data, content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{nome}"'
    return resp


@login_required
def asset_part_145_export_excel(request: HttpRequest) -> HttpResponse:
    """Dossier PART 145 in Excel: un foglio di riepilogo e uno per sezione."""
    from ..services import part145_dossier

    dati = part145_dossier.costruisci(**part145_dossier.visibilita(request))
    data, nome = part145_dossier.esporta_xlsx(dati, request.user)
    resp = HttpResponse(data, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{nome}"'
    return resp


@login_required
def asset_part_145_export_tarature(request: HttpRequest) -> HttpResponse:
    """Strumenti PART 145 con taratura (ente, ultima, prossima): l'elenco da comunicare all'AMO."""
    from ..services import part145_dossier

    dati = part145_dossier.solo_attrezzature()
    data, nome = part145_dossier.esporta_tarature_xlsx(dati, request.user)
    resp = HttpResponse(data, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{nome}"'
    return resp


@login_required
def asset_part_145_export_controllo(request: HttpRequest) -> HttpResponse:
    """Elenco per reparto del controllo semestrale delle attrezzature PART 145 (MT CN 17, Mod.072)."""
    from ..services import part145_dossier

    dati = part145_dossier.solo_attrezzature()
    data, nome = part145_dossier.esporta_controllo_semestrale_pdf(dati, request.user)
    resp = HttpResponse(data, content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{nome}"'
    return resp


def _apply_asset_export_filters(get_params) -> "QuerySet[Asset]":
    form = AssetFilterForm(get_params or None)
    qs = Asset.objects.select_related("asset_category").prefetch_related("endpoints")
    if form.is_valid():
        q = _clean_string(form.cleaned_data.get("q"))
        asset_type = _clean_string(form.cleaned_data.get("asset_type"))
        asset_category = form.cleaned_data.get("asset_category")
        reparto = _clean_string(form.cleaned_data.get("reparto"))
        vlan = form.cleaned_data.get("vlan")
        ip = _clean_string(form.cleaned_data.get("ip"))
        if q:
            qs = qs.filter(
                Q(asset_tag__icontains=q) | Q(name__icontains=q) | Q(serial_number__icontains=q)
                | Q(manufacturer__icontains=q) | Q(model__icontains=q)
                | Q(endpoints__endpoint_name__icontains=q) | Q(endpoints__ip__icontains=q)
            )
        if asset_type:
            qs = qs.filter(asset_type=asset_type)
        if asset_category:
            qs = qs.filter(asset_category_id__in=_category_subtree_ids(asset_category))
        if reparto:
            qs = qs.filter(reparto__icontains=reparto)
        if vlan is not None:
            qs = qs.filter(endpoints__vlan=vlan)
        if ip:
            qs = qs.filter(endpoints__ip__icontains=ip)
    return qs.distinct().order_by("name", "asset_tag")


# ---------------------------------------------------------------------------
# EXPORT: asset list
# ---------------------------------------------------------------------------

_ASSET_EXPORT_HEADERS = [
    "Tag", "Nome", "Tipo", "Categoria", "Reparto", "Stato",
    "Produttore", "Modello", "Matricola", "Assegnato a", "Posizione", "Note",
]


def _asset_export_row(a: "Asset") -> list:
    return [
        a.asset_tag or "",
        a.name or "",
        a.get_asset_type_display(),
        a.asset_category.label if a.asset_category_id else "",
        a.reparto or "",
        a.get_status_display(),
        a.manufacturer or "",
        a.model or "",
        a.serial_number or "",
        a.assignment_to or "",
        a.assignment_location or "",
        a.notes or "",
    ]


@login_required
def asset_list_export(request: HttpRequest) -> HttpResponse:
    scope = request.GET.get("scope", "filtered")
    fmt = request.GET.get("format", "xlsx")
    if scope == "full":
        assets = list(Asset.objects.select_related("asset_category").distinct().order_by("name", "asset_tag"))
    else:
        assets = list(_apply_asset_export_filters(request.GET))

    today = timezone.localdate().strftime("%Y%m%d")
    rows = [_asset_export_row(a) for a in assets]

    if fmt == "pdf":
        pdf_bytes = _report_table_pdf(title="Inventario asset", headers=_ASSET_EXPORT_HEADERS, rows=rows)
        resp = HttpResponse(pdf_bytes, content_type="application/pdf")
        resp["Content-Disposition"] = f'attachment; filename="asset_{today}.pdf"'
        return resp

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Asset"
    _xl_write_sheet(ws, _ASSET_EXPORT_HEADERS, rows)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="asset_{today}.xlsx"'
    return resp


# ---------------------------------------------------------------------------
# EXPORT: workorder list
# ---------------------------------------------------------------------------

_WO_EXPORT_HEADERS = [
    "ID", "Asset Tag", "Asset Nome", "Reparto", "Categoria", "Titolo", "Tipo", "Origine", "Stato",
    "Assegnato a", "Eseguito da", "Fornitore", "Contratto", "Coperto da contratto",
    "Aperto il", "Chiuso il", "Durata min", "Fermo min", "Costo manodopera",
    "Costo materiali", "Costo totale", "Descrizione", "Risoluzione",
]


def _wo_export_row(wo: "WorkOrder") -> list:
    asset = wo.asset if wo.asset_id else None
    assigned = wo.assigned_to.get_full_name() or wo.assigned_to.username if wo.assigned_to_id else ""
    executed = wo.executed_by.get_full_name() or wo.executed_by.username if wo.executed_by_id else ""
    return [
        str(wo.id),
        asset.asset_tag if asset else "",
        asset.name if asset else "",
        asset.reparto if asset else "",
        asset.asset_category.label if asset and asset.asset_category_id else "",
        wo.title or "",
        wo.get_kind_display(),
        wo.get_origin_display(),
        wo.get_status_display(),
        assigned,
        executed,
        str(wo.supplier) if wo.supplier_id else "",
        wo.assistance_contract.title if wo.assistance_contract_id else "",
        "Si" if wo.covered_by_contract else "No",
        wo.opened_at.strftime("%d-%m-%Y %H:%M") if wo.opened_at else "",
        wo.closed_at.strftime("%d-%m-%Y %H:%M") if wo.closed_at else "",
        str(wo.intervention_duration_minutes or 0),
        str(wo.downtime_minutes or 0),
        str(wo.labor_cost_eur or ""),
        str(wo.materials_cost_eur or ""),
        str(wo.resolved_total_cost_eur or ""),
        wo.description or "",
        wo.resolution or "",
    ]


@login_required
def workorder_list_export(request: HttpRequest) -> HttpResponse:
    scope = request.GET.get("scope", "filtered")
    fmt = request.GET.get("format", "xlsx")
    qs = _apply_workorder_list_filters(request.GET, include_filters=(scope == "filtered"))
    if scope == "filtered" and _clean_string(request.GET.get("view")):
        qs = _apply_workorder_operational_view(
            qs,
            operational_view=_resolve_workorder_operational_view(request.GET),
            user=request.user,
        )
    workorders = list(qs.order_by("-opened_at"))
    today = timezone.localdate().strftime("%Y%m%d")
    rows = [_wo_export_row(wo) for wo in workorders]

    if fmt == "pdf":
        pdf_bytes = _report_table_pdf(title="Interventi / Work Orders", headers=_WO_EXPORT_HEADERS, rows=rows)
        resp = HttpResponse(pdf_bytes, content_type="application/pdf")
        resp["Content-Disposition"] = f'attachment; filename="workorders_{today}.pdf"'
        return resp

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Interventi"
    _xl_write_sheet(ws, _WO_EXPORT_HEADERS, rows)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="workorders_{today}.xlsx"'
    return resp


# ---------------------------------------------------------------------------
# EXPORT: work machine list PDF (XLSX già presente)
# ---------------------------------------------------------------------------

@login_required
def work_machine_export_pdf(request: HttpRequest) -> HttpResponse:
    scope = request.GET.get("scope", "filtered")
    if scope == "full":
        qs = Asset.objects.filter(asset_type__in=PRODUCTION_ASSET_TYPES).select_related("work_machine")
    else:
        form = WorkMachineFilterForm(request.GET or None)
        qs = Asset.objects.filter(asset_type__in=PRODUCTION_ASSET_TYPES).select_related("work_machine")
        if form.is_valid():
            q = _clean_string(form.cleaned_data.get("q"))
            reparto = _clean_string(form.cleaned_data.get("reparto"))
            status_f = _clean_string(form.cleaned_data.get("status"))
            cnc_only = bool(form.cleaned_data.get("cnc_only"))
            five_axes_only = bool(form.cleaned_data.get("five_axes_only"))
            tcr_only = bool(form.cleaned_data.get("tcr_only"))
            if q:
                qs = qs.filter(
                    Q(asset_tag__icontains=q) | Q(name__icontains=q) | Q(reparto__icontains=q)
                    | Q(manufacturer__icontains=q) | Q(model__icontains=q) | Q(serial_number__icontains=q)
                )
            if reparto:
                qs = qs.filter(reparto__icontains=reparto)
            if status_f:
                qs = qs.filter(status=status_f)
            if cnc_only:
                qs = qs.filter(work_machine__cnc_controlled=True)
            if five_axes_only:
                qs = qs.filter(work_machine__five_axes=True)
            if tcr_only:
                qs = qs.filter(work_machine__tcr_enabled=True)

    machines = list(qs.order_by("reparto", "name", "asset_tag"))
    headers = ["Tag", "Nome", "Reparto", "Stato", "Produttore", "Anno", "CNC", "5 assi", "TCR", "X mm", "Y mm", "Z mm"]
    rows = []
    for a in machines:
        wm = getattr(a, "work_machine", None)
        rows.append([
            a.asset_tag or "", a.name or "", a.reparto or "", a.get_status_display(),
            a.manufacturer or "",
            str(wm.year) if wm and wm.year else "",
            "Sì" if wm and wm.cnc_controlled else "No",
            "Sì" if wm and wm.five_axes else "No",
            "Sì" if wm and wm.tcr_enabled else "No",
            str(wm.x_mm) if wm and wm.x_mm else "",
            str(wm.y_mm) if wm and wm.y_mm else "",
            str(wm.z_mm) if wm and wm.z_mm else "",
        ])
    today = timezone.localdate().strftime("%Y%m%d")
    pdf_bytes = _report_table_pdf(title="Macchine di lavoro", headers=headers, rows=rows)
    resp = HttpResponse(pdf_bytes, content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="macchine_{today}.pdf"'
    return resp


@login_required
def work_machine_export_excel(request: HttpRequest) -> HttpResponse:
    """Esporta la lista macchine filtrata in formato Excel (.xlsx)."""
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    from core.excel_export import write_cell, write_row

    form = WorkMachineFilterForm(request.GET or None)
    machines_qs = Asset.objects.filter(asset_type__in=PRODUCTION_ASSET_TYPES).select_related("work_machine")

    if form.is_valid():
        q = _clean_string(form.cleaned_data.get("q"))
        reparto = _clean_string(form.cleaned_data.get("reparto"))
        status = _clean_string(form.cleaned_data.get("status"))
        cnc_only = bool(form.cleaned_data.get("cnc_only"))
        five_axes_only = bool(form.cleaned_data.get("five_axes_only"))
        tcr_only = bool(form.cleaned_data.get("tcr_only"))

        if q:
            machines_qs = machines_qs.filter(
                Q(asset_tag__icontains=q)
                | Q(name__icontains=q)
                | Q(reparto__icontains=q)
                | Q(manufacturer__icontains=q)
                | Q(model__icontains=q)
                | Q(serial_number__icontains=q)
            )
        if reparto:
            machines_qs = machines_qs.filter(reparto__icontains=reparto)
        if status:
            machines_qs = machines_qs.filter(status=status)
        if cnc_only:
            machines_qs = machines_qs.filter(work_machine__cnc_controlled=True)
        if five_axes_only:
            machines_qs = machines_qs.filter(work_machine__five_axes=True)
        if tcr_only:
            machines_qs = machines_qs.filter(work_machine__tcr_enabled=True)

    machines = list(machines_qs.order_by("reparto", "name", "asset_tag"))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Macchine officina"

    headers = [
        "Tag", "Nome macchina", "Reparto", "Stato asset",
        "Produttore", "Modello", "N. seriale",
        "Anno", "TMC (mesi)", "Prossima manutenzione",
        "CNC", "5 assi", "TCR",
        "Cursa X (mm)", "Corsa Y (mm)", "Corsa Z (mm)",
        "Diametro (mm)", "Mandrino (mm)", "Pressione (bar)", "Accuracy",
    ]

    header_fill = PatternFill(fill_type="solid", fgColor="2563EB")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)

    for col_idx, header in enumerate(headers, 1):
        cell = write_cell(ws, 1, col_idx, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 20

    for row_idx, asset in enumerate(machines, 2):
        wm = getattr(asset, "work_machine", None)
        values = [
            asset.asset_tag,
            asset.name,
            asset.reparto or "",
            asset.get_status_display(),
            asset.manufacturer or "",
            asset.model or "",
            asset.serial_number or "",
            wm.year if wm and wm.year else "",
            wm.tmc if wm and wm.tmc else "",
            wm.next_maintenance_date.strftime("%d-%m-%Y") if wm and wm.next_maintenance_date else "",
            "Sì" if wm and wm.cnc_controlled else "No",
            "Sì" if wm and wm.five_axes else "No",
            "Sì" if wm and wm.tcr_enabled else "No",
            wm.x_mm if wm and wm.x_mm else "",
            wm.y_mm if wm and wm.y_mm else "",
            wm.z_mm if wm and wm.z_mm else "",
            wm.diameter_mm if wm and wm.diameter_mm else "",
            wm.spindle_mm if wm and wm.spindle_mm else "",
            str(wm.pressure_bar) if wm and wm.pressure_bar else "",
            wm.accuracy_from if wm and wm.accuracy_from else "",
        ]
        write_row(ws, row_idx, values)

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 3, 10), 45)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    today_str = timezone.localdate().strftime("%Y%m%d")
    filename = f"macchine_officina_{today_str}.xlsx"
    response = HttpResponse(
        buf.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required
def asset_report_pdf(request: HttpRequest, id: int | None = None) -> HttpResponse:
    if id is None:
        return redirect("assets:asset_list")
    asset = get_object_or_404(
        Asset.objects.select_related("asset_category", "it_details", "work_machine")
        .prefetch_related(
            "workorders",
            "documents",
            "tickets",
            "periodic_verifications",
            "periodic_verifications__supplier",
            "asset_category__category_fields",
        ),
        pk=id,
    )
    active_template = _active_asset_report_template(AssetReportTemplate.REPORT_ASSET_DETAIL)
    snapshot = _build_asset_report_snapshot(asset)
    pdf_bytes = render_asset_report_pdf(
        asset,
        snapshot,
        template_name=active_template.name if active_template else "",
    )
    response = HttpResponse(pdf_bytes, content_type="application/pdf")
    response["Content-Disposition"] = f'inline; filename="{asset.asset_tag or "asset"}-report.pdf"'
    return response


@login_required
def asset_detail_export_xlsx(request: HttpRequest, id: int) -> HttpResponse:
    import openpyxl
    import openpyxl.styles as xlst

    from core.excel_export import write_cell

    asset = get_object_or_404(
        Asset.objects.select_related("asset_category", "it_details", "work_machine"),
        pk=id,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scheda asset"

    hfill = xlst.PatternFill(fill_type="solid", fgColor="2563EB")
    hfont = xlst.Font(bold=True, color="FFFFFF", name="Calibri", size=10)
    sfill = xlst.PatternFill(fill_type="solid", fgColor="EFF6FF")
    sfont = xlst.Font(bold=True, name="Calibri", size=9)
    kfont = xlst.Font(bold=True, name="Calibri", size=9)
    vfont = xlst.Font(name="Calibri", size=9)

    r = 1
    for ci, h in enumerate(["Attributo", "Valore"], 1):
        cell = write_cell(ws, r, ci, h)
        cell.fill = hfill
        cell.font = hfont
        cell.alignment = xlst.Alignment(horizontal="center", vertical="center")
    r += 1

    def _sec(label):
        nonlocal r
        c = write_cell(ws, r, 1, label)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c.fill = sfill
        c.font = sfont
        r += 1

    def _kv(key, val):
        nonlocal r
        write_cell(ws, r, 1, key).font = kfont
        write_cell(ws, r, 2, str(val) if val is not None else "").font = vfont
        r += 1

    _sec("Anagrafica")
    _kv("Tag", asset.asset_tag)
    _kv("Nome", asset.name)
    _kv("Tipo", asset.get_asset_type_display())
    _kv("Categoria", asset.asset_category.label if asset.asset_category_id else "")
    _kv("Reparto", asset.reparto)
    _kv("Stato", asset.get_status_display())
    _kv("Produttore", asset.manufacturer)
    _kv("Modello", asset.model)
    _kv("Matricola", asset.serial_number)
    _kv("Note", asset.notes)

    _sec("Assegnazione")
    _kv("Assegnato a", asset.assignment_to)
    _kv("Reparto assegnazione", asset.assignment_reparto)
    _kv("Ubicazione", asset.assignment_location)

    it = getattr(asset, "it_details", None)
    if it:
        _sec("Dettagli IT")
        _kv("OS", it.os)
        _kv("CPU", it.cpu)
        _kv("RAM", it.ram)
        _kv("Disco", it.disco)
        _kv("Dominio", "Sì" if it.domain_joined else "No")
        _kv("EDR", "Sì" if it.edr_enabled else "No")
        _kv("2FA Office", "Sì" if it.office_2fa_enabled else "No")

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 55

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    tag = asset.asset_tag or str(asset.pk)
    today = timezone.localdate().strftime("%Y%m%d")
    resp = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="asset_{tag}_{today}.xlsx"'
    return resp
