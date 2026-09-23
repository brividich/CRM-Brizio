"""Manutenzione, round 2 di UI/UX (docs/manutenzione/CHECKLIST_UX_2026-09.md)."""

from __future__ import annotations

from datetime import timedelta
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from assets.models import (
    Asset,
    AssetCategory,
    MaintenanceInterventionTemplate,
    MaintenanceOccurrence,
    WorkOrder,
)

User = get_user_model()


@override_settings(LEGACY_AUTH_ENABLED=False)
class Ux2TestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.today = timezone.localdate()
        cls.admin = User.objects.create_superuser(username="ux2-admin", password="x", email="u@b.c")
        cls.altro = User.objects.create_user(username="ux2-altro", password="x")
        cls.category = AssetCategory.objects.create(code="ux2-cnc", label="CNC ux2")
        cls.asset = Asset.objects.create(asset_tag="UX2-01", name="Tornio", asset_category=cls.category)
        cls.asset2 = Asset.objects.create(asset_tag="UX2-02", name="Fresa", asset_category=cls.category)
        cls.plan = MaintenanceInterventionTemplate.objects.create(code="ux2-lub", label="Lubrificazione ux2")
        due = cls.today + timedelta(days=2)
        cls.libera = MaintenanceOccurrence.objects.create(plan=cls.plan, asset=cls.asset, due_date=due, warning_days=30)
        cls.di_altri = MaintenanceOccurrence.objects.create(plan=cls.plan, asset=cls.asset2, due_date=due, warning_days=30)
        wo = WorkOrder.objects.create(asset=cls.asset2, title="OdL di un altro", assigned_to=cls.altro)
        MaintenanceOccurrence.objects.filter(pk=cls.di_altri.pk).update(work_order=wo)

    def setUp(self):
        self.client.force_login(self.admin)


class DaFareManutentoreTests(Ux2TestCase):
    def _ids(self, response):
        return {row["occurrence"].id for block in response.context["blocks"] for row in block["rows"]}

    def test_chi_pianifica_vede_tutto_di_default(self):
        response = self.client.get(reverse("assets:maintenance_da_fare"))
        self.assertFalse(response.context["only_mine"])
        self.assertEqual(self._ids(response), {self.libera.id, self.di_altri.id})

    def test_il_mio_lavoro_esclude_quello_assegnato_ad_altri(self):
        response = self.client.get(reverse("assets:maintenance_da_fare"), {"mio": "1"})
        self.assertTrue(response.context["only_mine"])
        self.assertEqual(self._ids(response), {self.libera.id})
        self.assertContains(response, "Il mio lavoro")

    def test_chi_esegue_senza_pianificare_apre_sul_suo_lavoro(self):
        with mock.patch("assets.views_maintenance.can_plan_maintenance", return_value=False), \
                mock.patch("assets.views_maintenance.can_execute_maintenance", return_value=True):
            response = self.client.get(reverse("assets:maintenance_da_fare"))
            self.assertTrue(response.context["only_mine"])
            tutto = self.client.get(reverse("assets:maintenance_da_fare"), {"mio": "0"})
            self.assertFalse(tutto.context["only_mine"])


class CalendarioOdlTests(Ux2TestCase):
    def test_json_segnala_cosa_si_puo_pianificare(self):
        response = self.client.get(
            reverse("assets:calendario_asset_json"),
            {"start": (self.today - timedelta(days=5)).isoformat(), "end": (self.today + timedelta(days=10)).isoformat()},
        )
        events = {e["id"]: e for e in response.json()["events"]}
        libera = events[f"occ-{self.libera.id}"]
        self.assertTrue(libera["plannable"])
        self.assertEqual(libera["occurrence_id"], self.libera.id)
        self.assertFalse(events[f"occ-{self.di_altri.id}"]["plannable"])  # ha gia' un OdL

    def test_vassoio_solo_per_chi_pianifica_e_crea_l_odl(self):
        page = self.client.get(reverse("assets:calendario_asset"))
        self.assertContains(page, 'id="cm-tray"', html=False)
        response = self.client.post(
            reverse("assets:occurrence_create_workorder"),
            {"occurrence_ids": [self.libera.id], "next": reverse("assets:calendario_asset")},
        )
        self.libera.refresh_from_db()
        self.assertIsNotNone(self.libera.work_order_id)
        self.assertRedirects(response, reverse("assets:wo_view", args=[self.libera.work_order_id]), fetch_redirect_response=False)
        with mock.patch("assets.views_maintenance.can_plan_maintenance", return_value=False):
            self.assertNotContains(self.client.get(reverse("assets:calendario_asset")), 'id="cm-tray"', html=False)


class ScadenzarioExportTests(Ux2TestCase):
    def test_excel_con_gli_stessi_filtri(self):
        import io

        from openpyxl import load_workbook

        response = self.client.get(reverse("assets:maintenance_scadenze"), {"window": "30", "q": "UX2-01", "format": "xlsx"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("scadenzario_manutenzione_", response["Content-Disposition"])
        sheet = load_workbook(io.BytesIO(response.content)).active
        valori = [str(cell.value or "") for row in sheet.iter_rows() for cell in row]
        self.assertIn("UX2-01", valori)
        self.assertNotIn("UX2-02", valori)  # filtrato fuori dalla ricerca

    def test_pdf_e_pulsanti_in_testata(self):
        response = self.client.get(reverse("assets:maintenance_scadenze"), {"window": "30", "format": "pdf"})
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))
        page = self.client.get(reverse("assets:maintenance_scadenze"), {"window": "30"})
        self.assertContains(page, "format=xlsx")
        self.assertContains(page, "format=pdf")
