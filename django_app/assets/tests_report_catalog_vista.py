"""Catalogo report asset: report nuovi, riepilogo/grafico e vista (ricerca, ordinamento, colonne).

Dati sintetici: IP 192.0.2.x, nomi fittizi.
"""
from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO

from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .models import Asset, AssetEndpoint, AssetSavedReport, SoftwareLicense, WorkOrder
from .services import report_catalog
from .tests_report_catalog import _user


@override_settings(LEGACY_AUTH_ENABLED=False, SECURE_SSL_REDIRECT=False)
class CatalogoNuoviReportTests(TestCase):
    def setUp(self):
        self.admin = _user("rep2.admin", superuser=True)
        self.client.force_login(self.admin)
        self.switch = Asset.objects.create(asset_tag="SYN-SW-1", name="Switch sintetico", asset_type="FIREWALL",
                                           switch_prefisso_porte="A0", reparto="CED")
        self.pc1 = Asset.objects.create(asset_tag="SYN-PC-1", name="PC uno", asset_type="PC", reparto="CED",
                                        serial_number="SN1", manufacturer="Acme", model="M1",
                                        assignment_to="Utente Uno", purchase_date=date(2024, 1, 1))
        self.pc2 = Asset.objects.create(asset_tag="SYN-PC-2", name="PC due", asset_type="PC", reparto="Officina")
        self.pc3 = Asset.objects.create(asset_tag="SYN-PC-3", name="PC senza rete", asset_type="PC", reparto="Officina")
        AssetEndpoint.objects.create(asset=self.pc1, endpoint_name="LAN", ip="192.0.2.10", vlan=10, switch_port="A0.1.3")
        AssetEndpoint.objects.create(asset=self.pc2, endpoint_name="LAN", ip="192.0.2.10", vlan=20, switch_port="B9.1.1")
        SoftwareLicense.objects.create(product_name="Suite sintetica", asset=self.pc1, seats_used=1, seats_total=5,
                                       expiry_date=timezone.localdate() - timedelta(days=1))
        SoftwareLicense.objects.create(product_name="Antivirus sintetico", category="ANTIVIRUS", seats_used=3,
                                       assigned_to_display="Utente Due",
                                       expiry_date=timezone.localdate() + timedelta(days=200))
        adesso = timezone.now()
        for asset, kind, costo in [(self.pc1, "PREVENTIVE", "100.50"), (self.pc1, "CORRECTIVE", "50.00"),
                                   (self.pc2, "CORRECTIVE", "20.00")]:
            WorkOrder.objects.create(asset=asset, title="Intervento sintetico", kind=kind, status="DONE",
                                     closed_at=adesso, cost_eur=costo, downtime_minutes=90)

    def _run(self, code, query=""):
        response = self.client.get(reverse("assets:report_catalog_run", args=[code]) + query)
        self.assertEqual(response.status_code, 200, code)
        return response

    def test_tutti_i_report_hanno_area_e_sintesi(self):
        for d in report_catalog.definizioni().values():
            self.assertIn(d.area, report_catalog.AREE, d.code)
            self.assertIsNotNone(d.sintesi, d.code)
            self._run(d.code)

    def test_parco_it_switch_da_prefisso_ip_duplicato_e_senza_rete(self):
        risultato = self._run("parco-it-rete").context["risultato"]
        per_asset = {r[0]: r for r in risultato.righe}
        self.assertEqual(per_asset["SYN-PC-1"][9], "SYN-SW-1")
        self.assertIn("IP duplicato", per_asset["SYN-PC-1"][12])
        self.assertIn("switch non riconosciuto", per_asset["SYN-PC-2"][12])
        self.assertEqual(per_asset["SYN-PC-3"][12], "Nessuna interfaccia di rete")
        self.assertEqual(dict(risultato.riepilogo)["IP duplicati"], "2")
        filtrato = self._run("parco-it-rete", f"?switch={self.switch.pk}").context["risultato"]
        self.assertEqual([r[0] for r in filtrato.righe], ["SYN-PC-1"])
        senza = self._run("parco-it-rete", "?problema=senza_rete").context["risultato"]
        # Anche lo switch senza interfacce censite è un dispositivo IT senza rete.
        self.assertEqual([r[0] for r in senza.righe], ["SYN-PC-3", "SYN-SW-1"])

    def test_licenze_stato_e_posti(self):
        risultato = self._run("licenze-software").context["risultato"]
        self.assertEqual({r[1]: r[11] for r in risultato.righe},
                         {"Suite sintetica": "Scaduta", "Antivirus sintetico": "Attiva"})
        self.assertEqual(risultato.totali[7], 4)
        scadute = self._run("licenze-software", "?stato=scadute").context["risultato"]
        self.assertEqual([r[1] for r in scadute.righe], ["Suite sintetica"])

    def test_costi_raggruppati_e_totali(self):
        risultato = self._run("costi-manutenzione").context["risultato"]
        self.assertEqual(risultato.totali[1], 3)
        self.assertEqual(risultato.totali[9], 170.5)
        self.assertEqual(risultato.righe[0][0], "SYN-PC-1 — PC uno")  # il più costoso in alto
        per_tipo = self._run("costi-manutenzione", "?raggruppa=tipo").context["risultato"]
        self.assertEqual({r[0]: r[1] for r in per_tipo.righe}, {"Preventiva": 1, "Correttiva": 2})
        self.assertEqual(per_tipo.grafico["barre"][0]["etichetta"], "Preventiva")

    def test_qualita_dati(self):
        risultato = self._run("qualita-dati").context["risultato"]
        per_asset = {r[0]: r for r in risultato.righe}
        self.assertIn("Numero di serie", per_asset["SYN-PC-2"][7])
        self.assertIn("Indirizzo IP", per_asset["SYN-PC-3"][7])
        self.assertNotIn("Indirizzo IP", per_asset["SYN-PC-2"][7])
        solo_ip = self._run("qualita-dati", "?manca=Indirizzo+IP").context["risultato"]
        self.assertEqual([r[0] for r in solo_ip.righe], ["SYN-PC-3", "SYN-SW-1"])

    def test_ricerca_ricalcola_totali_e_riepilogo(self):
        risultato = self._run("licenze-software", "?cerca=antivirus").context["risultato"]
        self.assertEqual(len(risultato.righe), 1)
        self.assertEqual(risultato.totali[0], "Totale 1 righe")
        self.assertEqual(risultato.totali[7], 3)
        self.assertEqual(dict(risultato.riepilogo)["Licenze"], "1")

    def test_ordinamento_e_colonne_anche_in_excel(self):
        query = "?ordina=7&verso=desc&col=1&col=7"
        risultato = self._run("licenze-software", query).context["risultato"]
        self.assertEqual(risultato.intestazioni, ["Prodotto", "Posti usati"])
        self.assertEqual([r[0] for r in risultato.righe], ["Antivirus sintetico", "Suite sintetica"])
        xlsx = self.client.get(reverse("assets:report_catalog_run", args=["licenze-software"]) + query + "&formato=xlsx")
        valori = [[c.value for c in row] for row in load_workbook(BytesIO(xlsx.content)).active.iter_rows()]
        intestazione = next(r for r in valori if r[:2] == ["Prodotto", "Posti usati"])
        riga = valori[valori.index(intestazione) + 1]
        self.assertEqual(riga[:2], ["Antivirus sintetico", 3])
        self.assertTrue(all(v is None for v in riga[2:]))
        # Le date italiane si ordinano come date, non come testo.
        self.assertGreater(report_catalog._chiave("02/01/2026"), report_catalog._chiave("01/02/2025"))

    def test_vista_salvata_col_report(self):
        self.client.post(reverse("assets:report_catalog_save", args=["licenze-software"]),
                         {"nome": "Solo prodotti", "cerca": "suite", "col": ["1", "11"], "ordina": "1"})
        salvato = AssetSavedReport.objects.get(nome="Solo prodotti")
        self.assertEqual(salvato.filtri["_vista"], {"cerca": "suite", "col": [1, 11], "ordina": 1})
        risultato = self._run("licenze-software", f"?salvato={salvato.pk}").context["risultato"]
        self.assertEqual(risultato.intestazioni, ["Prodotto", "Stato"])
        self.assertEqual(len(risultato.righe), 1)

    def test_catalogo_per_aree_e_schede_report(self):
        body = self.client.get(reverse("assets:report_catalog")).content.decode()
        for area in report_catalog.AREE:
            self.assertIn(area, body)
        self.assertIn('id="rp-cerca"', body)
        cruscotto = self.client.get(reverse("assets:reports"), follow=True).content.decode()
        self.assertIn(reverse("assets:report_catalog"), cruscotto)
        self.assertIn('aria-current="page">Cruscotto manutenzione', cruscotto)
