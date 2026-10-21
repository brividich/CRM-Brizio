"""PART 145: legame attrezzature ↔ processi MOD.128, avviso cambio stato, dossier.

Dati sintetici (nomi e codici fittizi, email example.test).
"""
from __future__ import annotations

from datetime import timedelta
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core import mail
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from anagrafica.models_mpq import AbilitazioneProcesso, ClienteQualificante, ProcessoQualificato
from core.models import UserOnboarding

from .models import Asset, WorkOrder
from .services import part145, part145_dossier

User = get_user_model()


def _admin(username="p145.admin"):
    user = User.objects.create_user(username=username, password="x", email=f"{username}@example.test",
                                    is_superuser=True, is_staff=True)
    UserOnboarding.objects.update_or_create(
        user=user, defaults={"completed": True, "skipped": False, "completed_at": timezone.now()})
    return user


@override_settings(LEGACY_AUTH_ENABLED=False, SECURE_SSL_REDIRECT=False)
class Part145Tests(TestCase):
    def setUp(self):
        self.admin = _admin()
        self.client.force_login(self.admin)
        cliente = ClienteQualificante.objects.create(nome="Cliente sintetico")
        self.p145 = ProcessoQualificato.objects.create(nome="Manutenzione componenti sintetica", regime="PART145",
                                                       cliente=cliente,
                                                       data_scadenza=timezone.localdate() + timedelta(days=20))
        self.nadcap = ProcessoQualificato.objects.create(nome="Trattamento sintetico", regime="NADCAP", cliente=cliente)
        self.flag = Asset.objects.create(asset_tag="SYN-145-A", name="Banco prova", part_145=True, reparto="CQF",
                                         serial_number="SN-1")
        self.altro = Asset.objects.create(asset_tag="SYN-145-B", name="Forno sintetico", reparto="Forni")

    def test_attrezzature_processo_unisce_flag_e_collegate(self):
        self.altro.processi_qualificati.add(self.nadcap)
        self.assertEqual(list(part145.attrezzature_processo(self.p145)), [self.flag])
        self.assertEqual(list(part145.attrezzature_processo(self.nadcap)), [self.altro])
        self.assertEqual(part145.processi_dell_asset(self.flag), [self.p145])

    def test_collega_e_scollega_dal_dettaglio_mod128(self):
        url = reverse("anagrafica:mpq_attrezzatura_add", args=[self.nadcap.pk])
        self.client.post(url, {"asset": "syn-145-b"})
        self.assertTrue(self.altro.processi_qualificati.filter(pk=self.nadcap.pk).exists())
        dettaglio = self.client.get(reverse("anagrafica:mpq_processo_detail", args=[self.nadcap.pk]))
        self.assertContains(dettaglio, "SYN-145-B")
        self.assertContains(dettaglio, "Collegata qui")
        self.client.post(reverse("anagrafica:mpq_attrezzatura_remove", args=[self.nadcap.pk, self.altro.pk]))
        self.assertFalse(self.altro.processi_qualificati.exists())
        # Tag inesistente: nessun collegamento, messaggio d'errore.
        risposta = self.client.post(url, {"asset": "NON-ESISTE"}, follow=True)
        self.assertContains(risposta, "Nessun asset")

    def test_collegare_richiede_gestione_mod128(self):
        utente = User.objects.create_user(username="p145.base", password="x")
        UserOnboarding.objects.update_or_create(
            user=utente, defaults={"completed": True, "skipped": False, "completed_at": timezone.now()})
        self.client.force_login(utente)
        with patch("anagrafica.views_mpq._check_mpq_permission", return_value=False):
            self.client.post(reverse("anagrafica:mpq_attrezzatura_add", args=[self.nadcap.pk]), {"asset": "SYN-145-B"})
        self.assertFalse(self.altro.processi_qualificati.exists())

    def test_scheda_asset_mostra_i_processi(self):
        risposta = self.client.get(reverse("assets:asset_view", args=[self.flag.pk]))
        self.assertContains(risposta, reverse("anagrafica:mpq_processo_detail", args=[self.p145.pk]))

    @override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
    def test_avviso_cambio_stato_solo_attrezzature_qualificate(self):
        with patch("anagrafica.services.reminders.get_reminder_recipients", return_value=["qualita@example.test"]):
            with self.captureOnCommitCallbacks(execute=True):
                self.flag.status = Asset.STATUS_IN_REPAIR
                self.flag.save()
                self.altro.status = Asset.STATUS_IN_REPAIR
                self.altro.save()
        self.assertEqual(len(mail.outbox), 1)
        corpo = mail.outbox[0].alternatives[0][0] if mail.outbox[0].alternatives else mail.outbox[0].body
        self.assertIn("SYN-145-A", corpo)
        self.assertNotIn("SYN-145-B", corpo)
        # Salvataggio senza cambio di stato: nessuna mail.
        mail.outbox.clear()
        with patch("anagrafica.services.reminders.get_reminder_recipients", return_value=["qualita@example.test"]):
            with self.captureOnCommitCallbacks(execute=True):
                self.flag.name = "Banco prova 2"
                self.flag.save()
        self.assertEqual(len(mail.outbox), 0)

    def test_dossier_sezioni_e_situazioni_aperte(self):
        AbilitazioneProcesso.objects.create(legacy_anagrafica_id=990001, processo=self.p145, is_part145=True)
        WorkOrder.objects.create(asset=self.flag, title="Controllo sintetico", kind="CORRECTIVE")
        Asset.objects.filter(pk=self.flag.pk).update(status=Asset.STATUS_IN_REPAIR)
        dati = part145_dossier.costruisci()
        sezioni = {s.codice: s for s in dati["sezioni"]}
        self.assertEqual([r[0] for r in sezioni["attrezzature"].righe], ["SYN-145-A"])
        self.assertEqual(len(sezioni["personale"].righe), 1)
        self.assertEqual(sezioni["scadenze"].righe[0][1], "Processo")
        tipi = {r[1] for r in sezioni["aperte"].righe}
        self.assertIn("Attrezzatura in riparazione", tipi)
        self.assertIn("OdL aperto (Correttiva)", tipi)
        senza = part145_dossier.costruisci(con_personale=False)
        self.assertNotIn("personale", {s.codice for s in senza["sezioni"]})

    def test_situazioni_aperte_hanno_dove_si_risolve(self):
        wo = WorkOrder.objects.create(asset=self.flag, title="Controllo sintetico", kind="CORRECTIVE")
        Asset.objects.filter(pk=self.flag.pk).update(status=Asset.STATUS_IN_REPAIR)
        aperte = {s.codice: s for s in part145_dossier.costruisci()["sezioni"]}["aperte"]
        azioni = {riga[1]: azione for riga, _, azione in aperte.righe_link}
        self.assertEqual(azioni["Attrezzatura in riparazione"][1], reverse("assets:asset_view", args=[self.flag.pk]))
        self.assertEqual(azioni["OdL aperto (Correttiva)"][1], reverse("assets:wo_view", args=[wo.pk]))
        self.assertEqual(azioni["Processo senza persone abilitate"][1],
                         reverse("anagrafica:mpq_processo_detail", args=[self.p145.pk]))
        # Senza permesso MOD.128 niente link verso MOD.128.
        processi = {s.codice: s for s in part145_dossier.costruisci(con_personale=False)["sezioni"]}["processi"]
        self.assertFalse(processi.ha_azioni)
        pagina = self.client.get(reverse("assets:part_145_list"))
        self.assertContains(pagina, "Cosa fare")
        self.assertContains(pagina, reverse("assets:wo_view", args=[wo.pk]))
        dettaglio = self.client.get(reverse("anagrafica:mpq_processo_detail", args=[self.p145.pk]))
        self.assertContains(dettaglio, "Dossier PART 145")

    def test_pagina_ed_export_del_dossier(self):
        pagina = self.client.get(reverse("assets:part_145_list"))
        self.assertContains(pagina, "Dossier PART 145")
        self.assertContains(pagina, "Banco prova")
        xlsx = self.client.get(reverse("assets:part_145_export_excel"))
        wb = load_workbook(BytesIO(xlsx.content))
        self.assertIn("Riepilogo", wb.sheetnames)
        self.assertIn("Attrezzature", wb.sheetnames)
        pdf = self.client.get(reverse("assets:part_145_export_pdf"))
        self.assertTrue(pdf.content.startswith(b"%PDF"))

    def test_dossier_nascosto_il_personale_senza_permesso_mod128(self):
        with patch("assets.services.part145_dossier.puo_vedere_personale", return_value=False):
            pagina = self.client.get(reverse("assets:part_145_list"))
        self.assertContains(pagina, "servono i permessi sui processi qualificati")
        self.assertNotIn("personale", {s.codice for s in pagina.context["dossier"]["sezioni"]})


@override_settings(LEGACY_AUTH_ENABLED=False, SECURE_SSL_REDIRECT=False)
class Part145AllineamentoSgiTests(TestCase):
    """Perimetro DQP, attrezzi come il MOE §2.4, contratto, documenti SGI, avviso scadenze."""

    def setUp(self):
        from anagrafica.models import Fornitore

        from .models import MaintenanceInterventionTemplate, MaintenancePlanAssignment

        self.admin = _admin("p145.sgi")
        self.client.force_login(self.admin)
        self.oggi = timezone.localdate()
        cliente = ClienteQualificante.objects.create(nome="Cliente sintetico SGI")
        self.p145 = ProcessoQualificato.objects.create(nome="Riparazione sintetica", regime="PART145", cliente=cliente,
                                                       tipo_validita=ProcessoQualificato.VALIDITA_ILLIMITATA)
        self.dqp = ProcessoQualificato.objects.create(nome="Ritocco sintetico DQP", regime="SPECIALE", cliente=cliente,
                                                      perimetro_part145=True)
        self.fuori = ProcessoQualificato.objects.create(nome="Processo fuori perimetro", regime="SPECIALE",
                                                        cliente=cliente)
        self.strumento = Asset.objects.create(asset_tag="SYN-CAL-1", name="Chiave dinamometrica sintetica",
                                              part_145=True, reparto="Aggiustaggio", serial_number="SN-C1")
        self.pennello = Asset.objects.create(asset_tag="SYN-DQP-1", name="Kit ritocco sintetico", reparto="Aggiustaggio",
                                             serial_number="SN-D1")
        self.pennello.processi_qualificati.add(self.dqp)
        lab = Fornitore.objects.create(ragione_sociale="Laboratorio LAT sintetico")
        self.piano = MaintenanceInterventionTemplate.objects.create(
            code="syn-taratura", label="Taratura sintetica",
            maintenance_type=MaintenanceInterventionTemplate.TYPE_CALIBRATION,
            execution_mode=MaintenanceInterventionTemplate.MODE_EXTERNAL, default_supplier=lab)
        MaintenancePlanAssignment.objects.create(plan=self.piano, target_type=MaintenancePlanAssignment.TARGET_ASSET,
                                                 asset=self.strumento, frequency=MaintenancePlanAssignment.FREQ_MONTHS,
                                                 interval=12)

    def _sezioni(self, **kw):
        return {s.codice: s for s in part145_dossier.costruisci(**kw)["sezioni"]}

    def test_perimetro_include_i_processi_speciali_con_la_casella(self):
        sezioni = self._sezioni()
        nomi = [r[0] for r in sezioni["processi"].righe]
        self.assertIn("Ritocco sintetico DQP", nomi)
        self.assertNotIn("Processo fuori perimetro", nomi)
        self.assertEqual({r[0] for r in sezioni["attrezzature"].righe}, {"SYN-CAL-1", "SYN-DQP-1"})
        # Il flag dell'asset non lo lega ai processi speciali del perimetro.
        self.assertEqual(part145.processi_dell_asset(self.strumento), [self.p145])

    def test_classificazione_e_taratura_scaduta(self):
        from .models import MaintenanceOccurrence

        MaintenanceOccurrence.objects.create(plan=self.piano, asset=self.strumento,
                                             due_date=self.oggi - timedelta(days=3))
        sezioni = self._sezioni()
        attr = sezioni["attrezzature"]
        col = attr.intestazioni.index
        righe = {r[0]: r for r in attr.righe}
        self.assertEqual(righe["SYN-CAL-1"][col("Controllo")], part145_dossier.CLASSE_TARATURA)
        self.assertEqual(righe["SYN-CAL-1"][col("Ente taratura")], "Laboratorio LAT sintetico")
        self.assertEqual(righe["SYN-DQP-1"][col("Controllo")], part145_dossier.CLASSE_A_VISTA)
        self.assertIn("Taratura scaduta", {r[1] for r in sezioni["aperte"].righe})
        self.assertIn("Taratura", {r[1] for r in sezioni["scadenze"].righe})

    def test_contratto_quadro_in_scadenza_e_scaduto(self):
        from anagrafica.models_mpq import RequisitoQualifica

        RequisitoQualifica.objects.create(processo=self.p145, tipo=RequisitoQualifica.TIPO_CONTRATTO,
                                          descrizione="Contratto quadro sintetico",
                                          data_scadenza=self.oggi + timedelta(days=60))
        sezioni = self._sezioni()
        self.assertEqual(sezioni["contratti"].righe[0][1], "Contratto / accordo")
        self.assertIn("Contratto", {r[1] for r in sezioni["scadenze"].righe})
        RequisitoQualifica.objects.filter(processo=self.p145).update(data_scadenza=self.oggi - timedelta(days=1))
        self.assertIn("Contratto scaduto", {r[1] for r in self._sezioni()["aperte"].righe})

    def test_processo_senza_scadenza_va_segnalato(self):
        # Il DQP è «a data» ma senza data: va completato in MOD.128. Il Part 145 illimitato no.
        aperte = {(r[1], r[2]) for r in self._sezioni()["aperte"].righe}
        self.assertIn(("Processo senza scadenza", "Ritocco sintetico DQP"), aperte)
        self.assertNotIn(("Processo senza scadenza", "Riparazione sintetica"), aperte)

    def test_documenti_di_riferimento_dal_registro_sgi(self):
        from procedure_refresh.models import ProcedureDocument, ProcedureRevision

        doc = ProcedureDocument.objects.create(code="MT CN 70", title="PART 145 sintetica", document_type="MT")
        rev = ProcedureRevision.objects.create(document=doc, revision_code="1", revision_date=self.oggi,
                                               effective_date=self.oggi, file_name="mt70.pdf", is_current=True)
        documenti = self._sezioni()["documenti"]
        righe = {r[0]: (r, azione) for r, _, azione in documenti.righe_link}
        self.assertEqual(righe["MT CN 70"][1][1], reverse("procedure_refresh:document_open", args=[rev.pk]))
        self.assertEqual(righe["MO-ID-009.24"][0][2], "Non trovato nel registro SGI")

    def test_export_tarature_e_controllo_semestrale(self):
        xlsx = self.client.get(reverse("assets:part_145_export_tarature"))
        ws = load_workbook(BytesIO(xlsx.content)).active
        valori = [c.value for riga in ws.iter_rows() for c in riga]
        self.assertIn("SYN-CAL-1", valori)
        self.assertNotIn("SYN-DQP-1", valori)
        pdf = self.client.get(reverse("assets:part_145_export_controllo"))
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        pagina = self.client.get(reverse("assets:part_145_list"))
        self.assertContains(pagina, "Strumenti tarati")
        self.assertContains(pagina, "Documenti di riferimento")

    def test_avviso_scadenze_solo_se_ce_qualcosa(self):
        from .models import MaintenanceOccurrence

        with self.captureOnCommitCallbacks(execute=True):
            self.assertEqual(part145_dossier.notifica_scadenze(override=["qualita@example.test"]), 0)
        MaintenanceOccurrence.objects.create(plan=self.piano, asset=self.strumento,
                                             due_date=self.oggi + timedelta(days=10))
        mail.outbox.clear()
        self.assertEqual(part145_dossier.notifica_scadenze(override=["qualita@example.test"]), 1)
        self.assertIn("[PART 145] Scadenze", mail.outbox[-1].subject)


@override_settings(LEGACY_AUTH_ENABLED=False, SECURE_SSL_REDIRECT=False)
class Part145IntegrazioniTests(TestCase):
    """Formazione, audit e NC collegati al dossier; robustezza; report del catalogo. Dati sintetici."""

    def setUp(self):
        self.admin = _admin("p145.int")
        self.client.force_login(self.admin)
        self.oggi = timezone.localdate()
        cliente = ClienteQualificante.objects.create(nome="Cliente sintetico INT")
        self.p145 = ProcessoQualificato.objects.create(nome="Riparazione INT", regime="PART145", cliente=cliente,
                                                       tipo_validita=ProcessoQualificato.VALIDITA_ILLIMITATA)
        self.ab = AbilitazioneProcesso.objects.create(legacy_anagrafica_id=990077, processo=self.p145, is_part145=True)
        self.attrezzo = Asset.objects.create(asset_tag="SYN-INT-1", name="Calibro sintetico", part_145=True,
                                             reparto="CQF", serial_number="SN-I1")

    def _sezioni(self, **kw):
        return {s.codice: s for s in part145_dossier.costruisci(**kw)["sezioni"]}

    def test_formazione_richiesta_scaduta_e_mancante(self):
        from anagrafica.models_formazione import TrainingCourse, TrainingDeadline, TrainingPlan

        piano = TrainingPlan.objects.create(codice="PL-HF", nome="Piano sintetico")
        hf = TrainingCourse.objects.create(piano=piano, codice="HF-SYN", titolo="Human Factor sintetico", durata_ore_teorica=1)
        moe = TrainingCourse.objects.create(piano=piano, codice="MOE-SYN", titolo="MOE sintetico", durata_ore_teorica=1)
        self.p145.corsi_richiesti.add(hf, moe)
        TrainingDeadline.objects.create(corso=hf, legacy_anagrafica_id=990077, stato_scadenza="SCADUTO",
                                        data_scadenza=self.oggi - timedelta(days=5), is_required=True)
        sezioni = self._sezioni()
        self.assertEqual(len(sezioni["formazione"].righe), 2)
        tipi = {r[1] for r in sezioni["aperte"].righe}
        self.assertIn("Formazione scaduta", tipi)
        self.assertIn("Formazione mancante", tipi)
        self.assertEqual(sezioni["aperte"].righe[0][0], part145_dossier.PRIORITA_ALTA)
        self.assertIn("Formazione", {r[1] for r in sezioni["scadenze"].righe})
        # Senza permesso MOD.128 niente formazione per persona.
        self.assertNotIn("formazione", self._sezioni(con_personale=False))

    def test_audit_part145_mancanti_e_rilievo_scaduto(self):
        from sistema_gestione.models import Audit, AuditEsito, Auditor

        auditor = Auditor.objects.create(user=self.admin, interno=True)
        audit = Audit.objects.create(numero="RAIS-SYN-01", lead_auditor=auditor, processi="Part 145",
                                     data_inizio=self.oggi, part145=Audit.PART145_SISTEMA, created_by=self.admin)
        AuditEsito.objects.create(audit=audit, esito=AuditEsito.ESITO_NC, scostamento="Rilievo sintetico",
                                  scadenza_azione=self.oggi - timedelta(days=2))
        self.assertNotIn("audit", self._sezioni())  # serve il permesso audit
        sezioni = self._sezioni(con_audit=True)
        self.assertEqual([r[0] for r in sezioni["audit"].righe], ["RAIS-SYN-01"])
        tipi = {r[1] for r in sezioni["aperte"].righe}
        self.assertIn("Audit Part 145 di prodotto da pianificare", tipi)
        self.assertNotIn("Audit Part 145 di sistema da pianificare", tipi)
        self.assertIn("Azione da audit scaduta", tipi)

    def test_nc_part145_aperta_e_segnalazione_dal_dettaglio(self):
        from anomalie.quality_models import AnomaliaNC

        nc = AnomaliaNC.objects.create(protocollo="NC-SYN-145", op_titolo="OP-SYN-1", part_number="PN-SYN")
        self.assertEqual(self._sezioni(con_nc=True)["nc"].righe, [])
        self.client.post(reverse("anomalie_nc_dettaglio", args=[nc.pk]), {"action": "part145", "part145": "1"})
        nc.refresh_from_db()
        self.assertTrue(nc.part145)
        sezioni = self._sezioni(con_nc=True)
        self.assertEqual([r[0] for r in sezioni["nc"].righe], ["NC-SYN-145"])
        self.assertIn("NC Part 145 aperta", {r[1] for r in sezioni["aperte"].righe})
        self.assertNotIn("nc", self._sezioni())

    def test_parte_accessoria_rotta_non_fa_cadere_il_dossier(self):
        with patch("assets.services.part145_dossier._documenti_sgi", side_effect=RuntimeError("sintetico")):
            dati = part145_dossier.costruisci()
            pagina = self.client.get(reverse("assets:part_145_list"))
        self.assertTrue(any("Documenti SGI" in a for a in dati["avvisi"]))
        self.assertEqual(pagina.status_code, 200)
        self.assertContains(pagina, "Alcune parti non sono disponibili")

    def test_attrezzature_dismesse_escluse(self):
        Asset.objects.filter(pk=self.attrezzo.pk).update(status=Asset.STATUS_RETIRED)
        dati = part145_dossier.costruisci()
        self.assertEqual({s.codice: s for s in dati["sezioni"]}["attrezzature"].righe, [])
        self.assertEqual(dati["dismesse"], 1)

    def test_pagina_ricerca_unica_e_kpi(self):
        pagina = self.client.get(reverse("assets:part_145_list"))
        self.assertContains(pagina, 'id="p145-cerca"')
        self.assertContains(pagina, 'data-fm-table-skip="1"')
        self.assertContains(pagina, "Da sistemare")
        self.assertEqual(pagina.context["dossier"]["kpi"][0]["sezione"], "aperte")

    def test_report_catalogo_part145(self):
        from .services import report_catalog

        definizioni = report_catalog.definizioni()
        for code in ("part145-tarature", "part145-da-sistemare", "part145-scadenze"):
            self.assertEqual(definizioni[code].area, "Qualità e PART 145")
        tarature = definizioni["part145-tarature"].esegui({})
        self.assertEqual([r[0] for r in tarature.righe], ["SYN-INT-1"])
        aperte = definizioni["part145-da-sistemare"].esegui({"priorita": ""})
        # Senza dati del personale il catalogo non segnala i processi senza persone: è un dato MOD.128.
        self.assertNotIn("Processo senza persone abilitate", {r[1] for r in aperte.righe})
        self.assertEqual(aperte.intestazioni[0], "Priorità")
        scadenze = definizioni["part145-scadenze"].esegui({"giorni": "365"})
        self.assertIn("365", scadenze.titolo)
        risposta = self.client.get(reverse("assets:report_catalog_run", args=["part145-tarature"]))
        self.assertEqual(risposta.status_code, 200)
        self.assertContains(risposta, "SYN-INT-1")
