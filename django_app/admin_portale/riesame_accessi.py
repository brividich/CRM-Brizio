"""Riesame degli accessi (ISO/IEC 27001 A.5.18): un file Excel da far firmare.

Fotografia di chi ha quali ruoli del portale e perche', piu' cosa concede ogni
ruolo. Lo prepara il referente 27001 (o la Qualita') e lo approva la Direzione;
il file firmato si archivia nel Sistema di gestione. Sola lettura: l'export non
cambia nulla e finisce nel registro modifiche.
"""
from __future__ import annotations

from io import BytesIO

from django.utils import timezone

from core.excel_export import write_cell


def _intestazioni(ws, titoli: list[str]) -> None:
    from openpyxl.styles import Font

    for colonna, titolo in enumerate(titoli, start=1):
        write_cell(ws, 1, colonna, titolo).font = Font(bold=True)
    ws.freeze_panes = "A2"


def _foglio_account_m365(wb) -> None:
    """Gli account di Entra con l'esito del confronto. Se Graph non risponde il foglio lo dice."""
    from core import account_m365, m365_graph

    ws = wb.create_sheet("Account Microsoft 365")
    if not m365_graph.configurato():
        write_cell(ws, 1, 1, "Microsoft Graph non configurato: confronto con Microsoft 365 non disponibile.")
        return
    try:
        dati = account_m365.esegui()
    except m365_graph.GraphError as exc:
        write_cell(ws, 1, 1, f"Confronto con Microsoft 365 non disponibile: {exc}")
        return
    _intestazioni(ws, ["Esito", "Persona in anagrafica", "Cessato il", "Account", "UPN", "Stato",
                       "Ultimo accesso", "Licenze", "Note", "Decisione (conferma / disattiva)"])
    for r, riga in enumerate(dati["righe"], start=2):
        valori = [
            account_m365.ESITI.get(riga["esito"], riga["esito"]), riga["persona"], riga["cessato_il"],
            riga["nome_account"], riga["upn"], "Abilitato" if riga["abilitato"] else "Disattivato",
            riga["ultimo_accesso"][:10] if riga["ultimo_accesso"] else ("mai" if dati["con_accessi"] else "n/d"),
            riga["licenze"], ", ".join(riga["note"]), "",
        ]
        for c, valore in enumerate(valori, start=1):
            write_cell(ws, r, c, valore)
    for col, larghezza in zip("ABCDEFGHIJ", (30, 30, 12, 30, 38, 12, 14, 9, 26, 30)):
        ws.column_dimensions[col].width = larghezza


def _foglio_variazioni(wb, oggi) -> None:
    from datetime import timedelta

    from core.legacy_models import AnagraficaDipendente
    from core.models import AuditLog
    from core.permessi.variazioni import AZIONE, OGGETTO_TIPO

    ws = wb.create_sheet("Variazioni da organigramma")
    _intestazioni(ws, ["Data", "Persona", "Causa", "Ruoli del portale ricevuti", "Ruoli del portale persi", "Modificato da"])
    voci = list(AuditLog.objects.filter(azione=AZIONE, oggetto_tipo=OGGETTO_TIPO,
                                        created_at__date__gte=oggi - timedelta(days=365)).order_by("-created_at"))
    ids = {int(v.oggetto_id) for v in voci if str(v.oggetto_id).isdigit()}
    nomi = {p.id: f"{p.cognome or ''} {p.nome or ''}".strip() for p in AnagraficaDipendente.objects.filter(id__in=ids)}
    for r, v in enumerate(voci, start=2):
        d = v.dettaglio or {}
        pid = int(v.oggetto_id) if str(v.oggetto_id).isdigit() else None
        valori = [timezone.localtime(v.created_at).strftime("%d/%m/%Y %H:%M"), nomi.get(pid) or f"#{v.oggetto_id}",
                  d.get("causa", ""), ", ".join(d.get("ricevuti") or []), ", ".join(d.get("tolti") or []),
                  v.utente_display]
        for c, valore in enumerate(valori, start=1):
            write_cell(ws, r, c, valore)
    if not voci:
        write_cell(ws, 2, 1, "Nessuna variazione registrata negli ultimi 12 mesi.")
    for col, larghezza in zip("ABCDEF", (16, 30, 48, 36, 36, 26)):
        ws.column_dimensions[col].width = larghezza


def genera_riesame_xlsx(*, preparato_da: str, sistema_attivo: bool) -> bytes:
    from openpyxl import Workbook

    from anagrafica.models import RuoloOperativo
    from core.legacy_models import UtenteLegacy
    from core.models import AccessGroup, AccessGroupMembership, RuoloPermessoModulo
    from core.permessi.organigramma import provenienze_per_utenti
    from core.permessi.registro import moduli_registrati
    from core.permessi.team import _assegnazioni_attive

    oggi = timezone.localdate()
    wb = Workbook()

    # 1. Foglio da firmare.
    ws = wb.active
    ws.title = "Riesame"
    righe = [
        ("Riesame degli accessi al portale NOVICROM HUB", ""),
        ("Riferimento", "ISO/IEC 27001:2022 A.5.15, A.5.18"),
        ("Data dell'estrazione", oggi.strftime("%d/%m/%Y")),
        ("Estratto da", preparato_da),
        ("Sistema dei ruoli", "acceso" if sistema_attivo else "spento (valgono ancora le regole storiche)"),
        ("", ""),
        ("Esito del riesame (accessi confermati / da togliere)", ""),
        ("Azioni decise", ""),
        ("Preparato da (referente 27001)", ""),
        ("Approvato dalla Direzione", ""),
        ("Data e firma", ""),
    ]
    for r, (etichetta, valore) in enumerate(righe, start=1):
        write_cell(ws, r, 1, etichetta)
        write_cell(ws, r, 2, valore)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 60

    # 2. Persone e ruoli, con la provenienza.
    ws = wb.create_sheet("Persone")
    _intestazioni(ws, ["Persona", "Email", "Ruolo del portale", "Ordine", "Provenienza", "Valido fino al", "Stato"])
    utenti = {u.id: u for u in UtenteLegacy.objects.filter(attivo=True)}
    gruppi = {g.id: g for g in AccessGroup.objects.all()}
    righe_persone = []
    for m in AccessGroupMembership.objects.filter(legacy_user_id__in=list(utenti)):
        gruppo = gruppi.get(m.group_id)
        if gruppo is None:
            continue
        scaduto = bool(m.scade_il and m.scade_il < oggi)
        stato = "scaduto" if scaduto else ("ruolo disattivato" if not gruppo.is_active else "attivo")
        righe_persone.append((utenti[m.legacy_user_id], gruppo, "a mano", m.scade_il, stato))
    for utente_id, per_gruppo in provenienze_per_utenti(list(utenti)).items():
        for gruppo_id, nomi in per_gruppo.items():
            gruppo = gruppi.get(gruppo_id)
            if gruppo is not None:
                righe_persone.append(
                    (utenti[utente_id], gruppo, "organigramma: " + ", ".join(sorted(nomi)), None, "attivo")
                )
    righe_persone.sort(key=lambda r: (str(r[0].nome or "").upper(), -int(r[1].priority or 0)))
    for i, (utente, gruppo, provenienza, scade_il, stato) in enumerate(righe_persone, start=2):
        valori = [utente.nome, utente.email, gruppo.label, int(gruppo.priority or 0), provenienza,
                  scade_il.strftime("%d/%m/%Y") if scade_il else "", stato]
        for c, valore in enumerate(valori, start=1):
            write_cell(ws, i, c, valore)
    for col, larghezza in zip("ABCDEFG", (34, 34, 30, 8, 46, 14, 16)):
        ws.column_dimensions[col].width = larghezza

    # 3. Cosa concede ogni ruolo.
    ws = wb.create_sheet("Matrice")
    _intestazioni(ws, ["Ruolo", "Ordine", "Tipo", "Modulo", "Azione", "Valore"])
    nomi_moduli = {spec.chiave: spec for spec in moduli_registrati().values()}
    r = 2
    for gruppo in sorted(gruppi.values(), key=lambda g: (-int(g.priority or 0), g.label)):
        if not gruppo.is_active:
            continue
        tipo = "amministratore" if gruppo.amministratore else ("vede tutto" if gruppo.vede_tutto else "")
        if gruppo.amministratore:
            for c, valore in enumerate([gruppo.label, int(gruppo.priority or 0), tipo, "tutti", "tutte", "tutti"], 1):
                write_cell(ws, r, c, valore)
            r += 1
            continue
        for cella in RuoloPermessoModulo.objects.filter(ruolo=gruppo).order_by("modulo", "azione"):
            spec = nomi_moduli.get(cella.modulo)
            valori = [gruppo.label, int(gruppo.priority or 0), tipo, spec.nome if spec else cella.modulo,
                      spec.etichetta_azione(cella.azione) if spec else cella.azione,
                      cella.ambito if cella.consentito else "No"]
            for c, valore in enumerate(valori, start=1):
                write_cell(ws, r, c, valore)
            r += 1
    for col, larghezza in zip("ABCDEF", (30, 8, 16, 32, 30, 10)):
        ws.column_dimensions[col].width = larghezza

    # 4. Collegamenti con l'organigramma.
    ws = wb.create_sheet("Organigramma")
    _intestazioni(ws, ["Ruolo del portale", "Ruolo nell'organigramma", "Ambito", "Persone che lo ricoprono oggi"])
    persone_per_ruolo: dict[int, int] = {}
    for ruolo_id in _assegnazioni_attive().values_list("ruolo_id", flat=True):
        persone_per_ruolo[ruolo_id] = persone_per_ruolo.get(ruolo_id, 0) + 1
    r = 2
    through = RuoloOperativo.profili_permessi.through
    for ruolo_id, gruppo_id in through.objects.order_by("accessgroup_id").values_list("ruolooperativo_id", "accessgroup_id"):
        ruolo = RuoloOperativo.objects.select_related("ambito").filter(pk=ruolo_id).first()
        gruppo = gruppi.get(gruppo_id)
        if ruolo is None or gruppo is None:
            continue
        valori = [gruppo.label + (" (amministratore: ignorato)" if gruppo.amministratore else ""),
                  ruolo.nome, ruolo.ambito_label, persone_per_ruolo.get(ruolo.id, 0)]
        for c, valore in enumerate(valori, start=1):
            write_cell(ws, r, c, valore)
        r += 1
    for col, larghezza in zip("ABCD", (34, 40, 26, 14)):
        ws.column_dimensions[col].width = larghezza

    # 5. Variazioni degli accessi dovute all'organigramma negli ultimi 12 mesi (core.permessi.variazioni).
    _foglio_variazioni(wb, oggi)

    # 6. Account Microsoft 365 confrontati con l'anagrafica (core.account_m365).
    _foglio_account_m365(wb)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
