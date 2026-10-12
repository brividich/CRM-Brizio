"""Test dei kickoff: creazione, allegati, VRF, readiness, fasi, portfolio."""
from __future__ import annotations

from datetime import timedelta
import io
from pathlib import Path
import shutil
from unittest.mock import patch
from uuid import uuid4

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.http import HttpResponse
from django.test import RequestFactory, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone


from .models import (
    MeetingStatus,
    Project,
    ProjectComment,
    Task,
    TaskAttachment,
    TaskComment,
    TaskEvent,
    TaskEventType,
    TaskPriority,
    TaskStatus,
    VRFDocStatus,
)
from .forms import ProjectKickoffForm
from .views import project_list
from .tests_utils import (
    TasksBaseTestCase,
    _create_user_with_legacy,
    _ensure_role,
    _grant_role_actions,
    make_project,
)

User = get_user_model()

@override_settings(LEGACY_AUTH_ENABLED=False, SECURE_SSL_REDIRECT=False)
class TaskProjectsAndAttachmentsTests(TasksBaseTestCase):
    def setUp(self):
        super().setUp()
        _ensure_role(2, "utente")
        _grant_role_actions(2, ["tasks_view", "tasks_create", "tasks_edit"])
        self._refresh_acl_cache()

        self.user = _create_user_with_legacy(username="projectuser", legacy_user_id=5001, role_id=2, role_name="utente")

        tmp_root = Path(__file__).resolve().parents[1] / ".tmp_tests"
        tmp_root.mkdir(parents=True, exist_ok=True)
        self._media_root = str(tmp_root / f"tasks_test_media_{uuid4().hex}")
        Path(self._media_root).mkdir(parents=True, exist_ok=True)
        self._media_override = override_settings(MEDIA_ROOT=self._media_root)
        self._media_override.enable()
        today = timezone.localdate()
        Path(self._media_root, "tasks_vrf", today.strftime("%Y"), today.strftime("%m")).mkdir(parents=True, exist_ok=True)
        Path(self._media_root, "tasks_attachments", today.strftime("%Y"), today.strftime("%m")).mkdir(parents=True, exist_ok=True)
        self.addCleanup(self._media_override.disable)
        self.addCleanup(shutil.rmtree, self._media_root, True)

    def _base_task_payload(self, title: str) -> dict:
        return {
            "title": title,
            "description": "Descrizione test",
            "status": TaskStatus.TODO,
            "priority": TaskPriority.MEDIUM,
            "task_scope": "single",
        }

    def _make_vrf_workbook_upload(
        self,
        *,
        part_number: str,
        description: str = "Descrizione VRF",
        version: str = "1.0",
        client_name: str = "Cliente Test",
        filename: str = "kickoff-vrf.xlsx",
    ) -> SimpleUploadedFile:
        import openpyxl

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "VRF"
        worksheet["O2"] = "Q-001"
        worksheet["P2"] = version
        worksheet["B3"] = part_number
        worksheet["I3"] = description
        worksheet["P3"] = "ESP"
        worksheet["B4"] = client_name

        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        return SimpleUploadedFile(
            filename,
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def _read_vrf_cell(self, project: Project, cell_ref: str) -> str:
        import openpyxl

        project.vrf_file.open("rb")
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(project.vrf_file.read()), data_only=True)
        finally:
            project.vrf_file.close()
        try:
            worksheet = workbook["VRF"] if "VRF" in workbook.sheetnames else workbook.active
            return str(worksheet[cell_ref].value or "")
        finally:
            workbook.close()

    def test_cannot_create_standalone_task(self):
        """Modulo solo-kickoff: un nuovo task senza kickoff è rifiutato, anche via POST forzato."""
        self.client.force_login(self.user)
        payload = self._base_task_payload("Task senza kickoff")  # task_scope=single, nessun kickoff
        response = self.client.post(reverse("tasks:create"), payload)
        self.assertEqual(response.status_code, 200)  # re-render con errore, nessun redirect
        self.assertFalse(Task.objects.filter(title="Task senza kickoff").exists())

    def test_edit_legacy_orphan_keeps_no_kickoff(self):
        """L'edit di un'attività legacy senza kickoff la lascia orfana (nessun kickoff forzato)."""
        self.client.force_login(self.user)
        orphan = Task.objects.create(title="Legacy orfano", created_by=self.user)
        self.assertIsNone(orphan.project_id)
        payload = self._base_task_payload("Legacy orfano")  # task_scope=single
        response = self.client.post(reverse("tasks:edit", args=[orphan.id]), payload, follow=True)
        self.assertEqual(response.status_code, 200)
        orphan.refresh_from_db()
        self.assertIsNone(orphan.project_id)

    def test_create_task_with_new_project(self):
        self.client.force_login(self.user)
        payload = self._base_task_payload("Task con nuovo progetto")
        payload.update(
            {
                "task_scope": "project",
                "project_new_description": "Descrizione progetto A",
            }
        )
        response = self.client.post(reverse("tasks:create"), payload)
        self.assertEqual(response.status_code, 302)

        task = Task.objects.get(title="Task con nuovo progetto")
        self.assertIsNotNone(task.project_id)
        self.assertTrue(task.project.name.startswith("KICK-OFF "))
        self.assertIsNotNone(task.project.kickoff_number)
        self.assertEqual(task.project.created_by_id, self.user.id)

    def test_create_task_with_new_project_metadata(self):
        self.client.force_login(self.user)
        project_manager = User.objects.create_user(username="pm_user", password="pass12345")
        capo_commessa = User.objects.create_user(username="capo_user", password="pass12345")
        programmatore = User.objects.create_user(username="prog_user", password="pass12345")
        similar_project = make_project(name="Commessa simile", created_by=self.user)

        payload = self._base_task_payload("Task con metadati progetto")
        payload.update(
            {
                "task_scope": "project",
                "project_new_description": "Descrizione completa",
                "project_new_client": "Cliente Alfa",
                "project_new_manager": str(project_manager.id),
                "project_new_capo_commessa": str(capo_commessa.id),
                "project_new_programmer": str(programmatore.id),
                "project_new_control_method": "Checklist e test collaudo",
                "project_new_part_number": "PN-001",
                "project_similar_choice": str(similar_project.id),
            }
        )
        response = self.client.post(reverse("tasks:create"), payload)
        self.assertEqual(response.status_code, 302)

        task = Task.objects.get(title="Task con metadati progetto")
        self.assertIsNotNone(task.project_id)
        project = task.project
        self.assertEqual(project.client_name, "Cliente Alfa")
        self.assertEqual(list(project.project_managers.all()), [project_manager])
        self.assertEqual(list(project.capi_commessa.all()), [capo_commessa])
        self.assertEqual(list(project.programmers.all()), [programmatore])
        self.assertEqual(project.control_method, "Checklist e test collaudo")
        self.assertEqual(project.part_number, "PN-001")
        self.assertEqual(project.similar_project_id, similar_project.id)

    def test_create_task_reuses_existing_project_with_same_part_number_identity(self):
        self.client.force_login(self.user)
        project = make_project(
            name="Commessa storica",
            client_name="Cliente Legacy",
            part_number="PN-777",
            revisione="A",
            versione="1.0",
            created_by=self.user,
        )

        payload = self._base_task_payload("Task agganciata a progetto esistente per PN")
        payload.update(
            {
                "task_scope": "project",
                "project_new_client": "Cliente Nuovo",
                "project_new_part_number": "PN-777",
                "project_new_revisione": "A",
                "project_new_versione": "1.0",
            }
        )
        response = self.client.post(reverse("tasks:create"), payload)
        self.assertEqual(response.status_code, 302)

        task = Task.objects.get(title="Task agganciata a progetto esistente per PN")
        self.assertEqual(task.project_id, project.id)
        self.assertEqual(
            Project.objects.filter(part_number="PN-777", revisione="A", versione="1.0").count(),
            1,
        )
        project.refresh_from_db()
        self.assertEqual(project.client_name, "Cliente Nuovo")

    def test_new_project_requires_part_number_when_revision_or_version_is_set(self):
        # Regression: la validazione deve restare attiva sul ramo legacy
        # `project_link_mode=new` del TaskForm (il template ora usa il flow
        # dedicato /tasks/projects/new/ ma il backend mantiene compat).
        self.client.force_login(self.user)
        payload = self._base_task_payload("Task con revisione senza pn")
        payload.update(
            {
                "task_scope": "project",
                "project_link_mode": "new",
                "project_new_revisione": "B",
            }
        )
        response = self.client.post(reverse("tasks:create"), payload)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Compila il P/N prima di usare revisione o versione.")
        self.assertFalse(Task.objects.filter(title="Task con revisione senza pn").exists())

    def test_create_task_with_existing_project(self):
        self.client.force_login(self.user)
        project = make_project(name="Progetto Esistente", created_by=self.user)
        payload = self._base_task_payload("Task su progetto esistente")
        payload.update(
            {
                "task_scope": "project",
                "project_choice": str(project.id),
            }
        )
        response = self.client.post(reverse("tasks:create"), payload)
        self.assertEqual(response.status_code, 302)

        task = Task.objects.get(title="Task su progetto esistente")
        self.assertEqual(task.project_id, project.id)
        self.assertEqual(Project.objects.count(), 1)

    def test_project_model_auto_generates_kickoff_name_and_number(self):
        first = make_project(name="", created_by=self.user)
        second = make_project(name="", created_by=self.user)

        self.assertTrue(first.name.startswith("KICK-OFF "))
        self.assertTrue(second.name.startswith("KICK-OFF "))
        self.assertIsNotNone(first.kickoff_number)
        self.assertIsNotNone(second.kickoff_number)
        self.assertNotEqual(first.kickoff_number, second.kickoff_number)

    def test_existing_kickoff_name_is_not_renamed(self):
        legacy = make_project(
            name="Legacy Kickoff",
            kickoff_number=77,
            client_name="Cliente Legacy",
            created_by=self.user,
        )
        legacy.description = "Aggiornato"
        legacy.save()
        legacy.refresh_from_db()

        self.assertEqual(legacy.name, "Legacy Kickoff")
        self.assertEqual(legacy.kickoff_number, 77)

    def test_create_form_no_longer_shows_removed_vrf_summary_section(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("tasks:create"))

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "Scheda VRF")
        self.assertNotContains(response, "Contenuto operativo")
        self.assertContains(response, "Dettaglio attivita")

    def test_create_form_from_kickoff_context_hides_kickoff_selection(self):
        self.client.force_login(self.user)
        project = make_project(
            name="",
            client_name="Cliente Context",
            part_number="PN-CONTEXT-01",
            created_by=self.user,
        )

        response = self.client.get(f"{reverse('tasks:create')}?project={project.id}")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Attivita gia collegata al kickoff corrente")
        self.assertContains(response, project.name)
        self.assertNotContains(response, "Come nasce questa attivita")
        self.assertNotContains(response, "Aggancio e anagrafica kickoff")

    def test_create_form_from_kickoff_context_forces_task_into_that_kickoff(self):
        self.client.force_login(self.user)
        project = make_project(
            name="",
            client_name="Cliente Locked",
            part_number="PN-LOCK-01",
            created_by=self.user,
        )

        payload = self._base_task_payload("Task dal kickoff contestuale")
        payload.update(
            {
                "task_scope": "single",
                "project_link_mode": "new",
                "project_new_client": "Cliente alternativo",
            }
        )

        response = self.client.post(f"{reverse('tasks:create')}?project={project.id}", payload)

        self.assertEqual(response.status_code, 302)
        task = Task.objects.get(title="Task dal kickoff contestuale")
        self.assertEqual(task.project_id, project.id)
        self.assertEqual(Project.objects.count(), 1)

    def test_project_list_shows_copy_buttons(self):
        self.client.force_login(self.user)
        make_project(name="", created_by=self.user)

        response = self.client.get(reverse("tasks:project_list"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Copia VRF")
        self.assertContains(response, "Copia senza P/N")

    def test_project_list_query_budget_with_many_projects(self):
        """§10.4: annotate_open_action_counts() usa sottoquery correlate, non
        Count() incatenati su relazioni diverse — il budget query non deve
        crescere con il numero di commesse nel portfolio. Chiama la view
        direttamente (come ProjectActionsViewTests) per isolare il costo dalla
        query di ACL/middleware/branding condivisa da ogni richiesta HTTP."""
        for i in range(12):
            make_project(
                name=f"Commessa budget {i}", client_name="Cliente Budget", created_by=self.user
            )
        admin = User.objects.create_superuser(username="portfolio-budget-admin", password="pass12345")
        request = RequestFactory().get(reverse("tasks:project_list"))
        request.user = admin

        def render_without_template_queries(request, template_name, context):
            return HttpResponse(f"{template_name}:{len(context['projects'])}")

        # Budget: 1 query per le righe progetto (readiness Exists() e i conteggi
        # azioni di annotate_open_action_counts() sono sottoquery correlate nella
        # stessa query, non round-trip aggiuntivi) + 4 per il prefetch dei ruoli
        # del team (M2M) + 1 per client_choices + 6 per
        # TaskImpostazioni.get_singleton() (get_or_create con audit log). Nessuna
        # di queste cresce con il numero di commesse nel portfolio.
        with patch("tasks.views.projects.render", side_effect=render_without_template_queries):
            with self.assertNumQueries(12):
                response = project_list(request)

        self.assertEqual(response.status_code, 200)

    def test_project_list_builds_client_filter_from_distinct_names(self):
        self.client.force_login(self.user)
        make_project(name="Kickoff Alfa", client_name="Cliente Alfa", created_by=self.user)
        make_project(name="Kickoff Beta", client_name="Cliente Beta", created_by=self.user)
        make_project(name="Kickoff vuoto", client_name="", created_by=self.user)

        response = self.client.get(reverse("tasks:project_list"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Cliente Alfa")
        self.assertContains(response, "Cliente Beta")

    def test_copy_project_with_vrf_duplicates_kickoff_metadata_and_file(self):
        self.client.force_login(self.user)
        project_manager = User.objects.create_user(username="copy_pm", password="pass12345")
        source_project = make_project(
            name="",
            description="Kickoff sorgente",
            client_name="Cliente Copy",
            project_manager=project_manager,
            control_method="Controllo finale",
            part_number="PN-COPY-001",
            revisione="A",
            versione="2.0",
            vrf_status=VRFDocStatus.UPLOADED,
            vrf_original_name="source-vrf.xlsx",
            vrf_quote_number="Q-500",
            vrf_description="Scheda sorgente",
            vrf_esp="ESP-1",
            created_by=self.user,
        )
        source_project.vrf_file.save(
            "source-vrf.xlsx",
            self._make_vrf_workbook_upload(
                part_number="PN-COPY-001",
                description="Scheda sorgente",
                version="2.0",
                client_name="Cliente Copy",
                filename="source-vrf.xlsx",
            ),
            save=True,
        )
        source_project.refresh_from_db()
        source_task = Task.objects.create(title="Task sorgente", created_by=self.user, project=source_project)
        ProjectComment.objects.create(project=source_project, author=self.user, body="Commento kickoff")
        TaskComment.objects.create(task=source_task, author=self.user, body="Commento task")
        TaskAttachment.objects.create(
            project=source_project,
            uploaded_by=self.user,
            file="tasks_attachments/mock.txt",
            original_name="mock.txt",
        )

        response = self.client.post(reverse("tasks:copy_project_with_vrf", args=[source_project.id]))
        self.assertEqual(response.status_code, 302)

        copied_project = Project.objects.exclude(id=source_project.id).get()
        # Il nome proposto porta cliente e P/N davanti al progressivo.
        self.assertIn("KICK-OFF ", copied_project.name)
        self.assertNotEqual(copied_project.kickoff_number, source_project.kickoff_number)
        self.assertEqual(copied_project.description, source_project.description)
        self.assertEqual(copied_project.client_name, source_project.client_name)
        self.assertEqual(
            list(copied_project.project_managers.all()),
            list(source_project.project_managers.all()),
        )
        self.assertEqual(copied_project.control_method, source_project.control_method)
        self.assertEqual(copied_project.part_number, "PN-COPY-001")
        self.assertEqual(copied_project.revisione, "A")
        self.assertEqual(copied_project.versione, "2.0")
        self.assertEqual(copied_project.vrf_status, VRFDocStatus.UPLOADED)
        self.assertEqual(copied_project.vrf_original_name, "source-vrf.xlsx")
        self.assertEqual(self._read_vrf_cell(source_project, "B3"), "PN-COPY-001")
        self.assertEqual(self._read_vrf_cell(copied_project, "B3"), "PN-COPY-001")
        self.assertEqual(Task.objects.filter(project=copied_project).count(), 0)
        self.assertEqual(ProjectComment.objects.filter(project=copied_project).count(), 0)
        self.assertEqual(TaskAttachment.objects.filter(project=copied_project).count(), 0)

    def test_copy_project_with_vrf_without_pn_clears_project_and_excel_part_number(self):
        self.client.force_login(self.user)
        source_project = make_project(
            name="",
            client_name="Cliente No PN",
            part_number="PN-COPY-002",
            revisione="B",
            versione="3.1",
            vrf_status=VRFDocStatus.UPLOADED,
            vrf_original_name="source-no-pn.xlsx",
            created_by=self.user,
        )
        source_project.vrf_file.save(
            "source-no-pn.xlsx",
            self._make_vrf_workbook_upload(
                part_number="PN-COPY-002",
                version="3.1",
                client_name="Cliente No PN",
                filename="source-no-pn.xlsx",
            ),
            save=True,
        )
        source_project.refresh_from_db()

        response = self.client.post(reverse("tasks:copy_project_with_vrf_without_pn", args=[source_project.id]))
        self.assertEqual(response.status_code, 302)

        copied_project = Project.objects.exclude(id=source_project.id).get()
        self.assertEqual(copied_project.part_number, "")
        self.assertEqual(copied_project.revisione, "B")
        self.assertEqual(copied_project.versione, "3.1")
        self.assertEqual(self._read_vrf_cell(source_project, "B3"), "PN-COPY-002")
        self.assertEqual(self._read_vrf_cell(copied_project, "B3"), "")

    def test_assignment_conflict_alert_on_create_with_keep_priority(self):
        self.client.force_login(self.user)
        assignee = User.objects.create_user(username="planner_keep", password="pass12345")
        today = timezone.localdate()
        Task.objects.create(
            title="Impegno esistente",
            created_by=self.user,
            assigned_to=assignee,
            status=TaskStatus.IN_PROGRESS,
            priority=TaskPriority.MEDIUM,
            next_step_due=today + timedelta(days=2),
            due_date=today + timedelta(days=5),
        )

        conflict_project = make_project(
            name="", client_name="Cliente Conf", part_number="PN-CONF-1", created_by=self.user
        )
        payload = self._base_task_payload("Task nuova con conflitto")
        payload.update(
            {
                "assigned_to": str(assignee.id),
                "priority": TaskPriority.LOW,
                "next_step_due": (today + timedelta(days=3)).isoformat(),
                "due_date": (today + timedelta(days=6)).isoformat(),
                "assignment_conflict_action": "keep_priority",
            }
        )
        response = self.client.post(f"{reverse('tasks:create')}?project={conflict_project.id}", payload, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Impegni sovrapposti")
        task = Task.objects.get(title="Task nuova con conflitto")
        self.assertEqual(task.priority, TaskPriority.LOW)

    def test_assignment_conflict_can_raise_priority_to_high(self):
        self.client.force_login(self.user)
        assignee = User.objects.create_user(username="planner_raise", password="pass12345")
        today = timezone.localdate()
        Task.objects.create(
            title="Impegno esistente raise",
            created_by=self.user,
            assigned_to=assignee,
            status=TaskStatus.TODO,
            priority=TaskPriority.MEDIUM,
            next_step_due=today + timedelta(days=1),
            due_date=today + timedelta(days=4),
        )

        raise_project = make_project(
            name="", client_name="Cliente Raise", part_number="PN-RAISE-1", created_by=self.user
        )
        payload = self._base_task_payload("Task nuova con priorita auto")
        payload.update(
            {
                "assigned_to": str(assignee.id),
                "priority": TaskPriority.LOW,
                "next_step_due": (today + timedelta(days=2)).isoformat(),
                "due_date": (today + timedelta(days=5)).isoformat(),
                "assignment_conflict_action": "raise_to_high",
            }
        )
        response = self.client.post(f"{reverse('tasks:create')}?project={raise_project.id}", payload, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Priorita aggiornata automaticamente a High")
        task = Task.objects.get(title="Task nuova con priorita auto")
        self.assertEqual(task.priority, TaskPriority.HIGH)

    def test_upload_attachment_to_task_creates_event(self):
        self.client.force_login(self.user)
        task = Task.objects.create(title="Task allegato", created_by=self.user)
        file_obj = SimpleUploadedFile("task-note.txt", b"contenuto allegato task", content_type="text/plain")
        with patch("tasks.forms.validate_extension_and_mime", return_value="text/plain"):
            response = self.client.post(
                reverse("tasks:add_attachment", args=[task.id]),
                {"attach_to": "task", "file": file_obj},
            )
        self.assertEqual(response.status_code, 302)

        attachment = TaskAttachment.objects.get(task=task)
        self.assertEqual(attachment.original_name, "task-note.txt")
        event = TaskEvent.objects.filter(task=task, type=TaskEventType.ATTACHMENT_ADDED).first()
        self.assertIsNotNone(event)
        self.assertEqual(event.payload.get("target"), "task")
        self.assertEqual(event.payload.get("attachment_id"), attachment.id)

    def test_upload_attachment_to_project_creates_event(self):
        self.client.force_login(self.user)
        project = make_project(name="Project Attach", created_by=self.user)
        task = Task.objects.create(title="Task con progetto allegato", created_by=self.user, project=project)
        file_obj = SimpleUploadedFile("project-note.txt", b"contenuto allegato progetto", content_type="text/plain")
        with patch("tasks.forms.validate_extension_and_mime", return_value="text/plain"):
            response = self.client.post(
                reverse("tasks:add_attachment", args=[task.id]),
                {"attach_to": "project", "file": file_obj},
            )
        self.assertEqual(response.status_code, 302)

        attachment = TaskAttachment.objects.get(project=project, task__isnull=True)
        self.assertEqual(attachment.original_name, "project-note.txt")
        event = TaskEvent.objects.filter(task=task, type=TaskEventType.ATTACHMENT_ADDED).first()
        self.assertIsNotNone(event)
        self.assertEqual(event.payload.get("target"), "project")
        self.assertEqual(event.payload.get("project_id"), project.id)


@override_settings(LEGACY_AUTH_ENABLED=False, SECURE_SSL_REDIRECT=False)
class ProjectCreateFlowTests(TasksBaseTestCase):
    def setUp(self):
        super().setUp()
        _ensure_role(2, "utente")
        _grant_role_actions(2, ["tasks_view", "tasks_create", "tasks_edit"])
        self._refresh_acl_cache()
        self.user = _create_user_with_legacy(
            username="kickoffuser", legacy_user_id=7001, role_id=2, role_name="utente",
        )

    def test_get_project_create_page(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("tasks:project_create"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Nuovo kickoff")
        self.assertContains(response, 'name="part_number"')
        self.assertContains(response, 'name="client_name"')
        self.assertContains(response, 'name="safety_impact"')
        self.assertContains(response, "Impatto sulla sicurezza")

    def test_post_creates_project_with_first_meeting_and_redirects_there(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("tasks:project_create"),
            {
                "client_name": "Cliente Test",
                "part_number": "PN-NEW-001",
                "revisione": "A",
                "versione": "1.0",
                "description": "Desc",
                "control_method": "",
                "vrf_quote_number": "Q-1",
                "vrf_description": "",
                "vrf_esp": "02",
                "project_managers": [self.user.id],
            },
        )
        project = Project.objects.get(part_number="PN-NEW-001")
        meeting = project.meetings.get()
        self.assertEqual(meeting.numero, 1)
        self.assertEqual(meeting.stato, MeetingStatus.PIANIFICATO)
        self.assertEqual(meeting.data, timezone.localdate())
        self.assertEqual(meeting.created_by, self.user)
        # Il team indicato nel form entra fra i partecipanti dell'incontro.
        self.assertIn(self.user, meeting.partecipanti_utenti.all())
        self.assertRedirects(
            response,
            reverse("tasks:project_meeting_edit", args=[project.id, meeting.id]),
        )
        self.assertEqual(project.client_name, "Cliente Test")
        self.assertEqual(project.revisione, "A")
        self.assertEqual(project.created_by, self.user)
        self.assertFalse(project.safety_impact)

    def test_project_safety_impact_defaults_false(self):
        project = make_project(name="Kickoff safety default", created_by=self.user)

        self.assertFalse(project.safety_impact)

    def test_post_creates_project_with_safety_impact_true_without_changing_status_or_priority(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("tasks:project_create"),
            {
                "client_name": "Cliente Safety",
                "part_number": "PN-SAFE-001",
                "revisione": "A",
                "versione": "1.0",
                "description": "Desc",
                "control_method": "",
                "safety_impact": "on",
                "vrf_quote_number": "Q-SAFE",
                "vrf_description": "",
                "vrf_esp": "",
                "project_managers": [self.user.id],
            },
        )

        project = Project.objects.get(part_number="PN-SAFE-001")
        meeting = project.meetings.get()
        self.assertRedirects(
            response, reverse("tasks:project_meeting_edit", args=[project.id, meeting.id])
        )
        self.assertTrue(project.safety_impact)
        self.assertEqual(project.vrf_status, VRFDocStatus.PENDING)

    def test_project_kickoff_form_edit_post_can_clear_safety_impact(self):
        project = make_project(
            client_name="Cliente Edit",
            part_number="PN-EDIT-SAFE",
            revisione="A",
            versione="1.0",
            safety_impact=True,
            created_by=self.user,
        )

        form = ProjectKickoffForm(
            {
                "client_name": "Cliente Edit",
                "part_number": "PN-EDIT-SAFE-2",
                "revisione": "A",
                "versione": "1.0",
                "description": project.description,
                "control_method": project.control_method,
                "vrf_quote_number": project.vrf_quote_number,
                "vrf_description": project.vrf_description,
                "vrf_esp": project.vrf_esp,
                "project_managers": [self.user.id],
            },
            instance=project,
            project_queryset=Project.objects.exclude(pk=project.pk),
        )

        self.assertTrue(form.is_valid(), form.errors)
        updated = form.save()
        self.assertFalse(updated.safety_impact)

    def test_revisione_without_part_number_rejected(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("tasks:project_create"),
            {"client_name": "X", "part_number": "", "revisione": "A", "versione": "", "project_managers": [self.user.id]},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            "Revisione e versione sono valide solo con un P/N indicato.",
        )
        self.assertFalse(Project.objects.filter(client_name="X").exists())

    def test_duplicate_identity_reuses_existing_kickoff(self):
        existing = make_project(
            part_number="DUP-001", revisione="B", versione="2.0",
            client_name="Cliente Precedente", created_by=self.user,
        )
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("tasks:project_create"),
            {
                "client_name": "Cliente Nuovo",
                "part_number": "DUP-001",
                "revisione": "B",
                "versione": "2.0",
                "project_managers": [self.user.id],
            },
        )
        self.assertRedirects(
            response,
            reverse("tasks:project_vrf_compile", args=[existing.id]),
        )
        self.assertEqual(Project.objects.filter(part_number="DUP-001").count(), 1)
        # Riusare un kickoff non e' crearlo: nessun incontro nuovo.
        self.assertEqual(existing.meetings.count(), 0)
        existing.refresh_from_db()
        self.assertEqual(existing.client_name, "Cliente Precedente")

    def test_safety_impact_badge_visible_on_project_detail_and_list_only_when_true(self):
        safety_project = make_project(
            name="Kickoff safety",
            part_number="PN-SAFE-LIST",
            safety_impact=True,
            created_by=self.user,
        )
        make_project(
            name="Kickoff normal",
            part_number="PN-NORMAL-LIST",
            safety_impact=False,
            created_by=self.user,
        )
        self.client.force_login(self.user)

        list_response = self.client.get(reverse("tasks:project_list"))
        self.assertEqual(list_response.status_code, 200)
        self.assertContains(list_response, "Impatto sicurezza", count=1)

        detail_response = self.client.get(reverse("tasks:project_gantt", args=[safety_project.id]))
        self.assertEqual(detail_response.status_code, 200)
        self.assertContains(detail_response, "Impatto sicurezza")

        normal_response = self.client.get(reverse("tasks:project_gantt", args=[Project.objects.get(part_number="PN-NORMAL-LIST").id]))
        self.assertEqual(normal_response.status_code, 200)
        self.assertNotContains(normal_response, "Impatto sicurezza")


@override_settings(LEGACY_AUTH_ENABLED=False, SECURE_SSL_REDIRECT=False)
class VRFCompileOnlineTests(TasksBaseTestCase):
    def setUp(self):
        super().setUp()
        _ensure_role(2, "utente")
        _grant_role_actions(2, ["tasks_view", "tasks_create", "tasks_edit"])
        self._refresh_acl_cache()
        self.user = _create_user_with_legacy(
            username="vrfuser", legacy_user_id=6001, role_id=2, role_name="utente",
        )
        tmp_root = Path(__file__).resolve().parents[1] / ".tmp_tests"
        tmp_root.mkdir(parents=True, exist_ok=True)
        self._media_root = str(tmp_root / f"vrf_compile_{uuid4().hex}")
        Path(self._media_root).mkdir(parents=True, exist_ok=True)
        self._media_override = override_settings(MEDIA_ROOT=self._media_root)
        self._media_override.enable()
        today = timezone.localdate()
        Path(self._media_root, "tasks_vrf", today.strftime("%Y"), today.strftime("%m")).mkdir(parents=True, exist_ok=True)
        self.addCleanup(self._media_override.disable)
        self.addCleanup(shutil.rmtree, self._media_root, True)
        self.project = make_project(
            name="TBD",
            client_name="Test Client",
            part_number="PN-TEST-1",
            vrf_description="Particolare di test",
            vrf_esp="02",
            vrf_quote_number="Q-TEST",
            created_by=self.user,
        )

    def test_compile_page_renders_full_matrix(self):
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("tasks:project_vrf_compile", args=[self.project.id])
        )
        self.assertEqual(response.status_code, 200)
        for r_code in ("R1", "R2", "R3", "R4", "R5", "R6", "R7", "R8", "R9"):
            self.assertContains(response, f'data-risk="{r_code}"')
        self.assertContains(response, 'name="score_R1_a_p"')
        self.assertContains(response, 'name="k_R2_c"')

    def test_save_draft_persists_scores_and_recomputes_totals(self):
        from .models import VRFRiskAssessment

        self.client.force_login(self.user)
        payload = {
            "action": "save_draft",
            "score_R1_a_p": "3", "score_R1_b_p": "2",
            "score_R2_a_p": "1", "score_R2_b_p": "2", "score_R2_c_p": "3",
            "k_R1_p": "3", "k_R1_i": "3", "k_R1_c": "3",
            "k_R2_p": "3", "k_R2_i": "3", "k_R2_c": "5",
        }
        response = self.client.post(
            reverse("tasks:project_vrf_compile", args=[self.project.id]), payload
        )
        self.assertEqual(response.status_code, 302)
        assessment = VRFRiskAssessment.objects.get(project=self.project)
        self.assertEqual(assessment.data["risks"]["R1"]["subs"]["a"]["p"], 3)
        self.assertEqual(assessment.data["risks"]["R2"]["subs"]["c"]["p"], 3)
        self.assertGreater(assessment.total_p, 0.0)
        self.assertEqual(assessment.total_i, 0.0)
        self.project.refresh_from_db()
        self.assertEqual(self.project.vrf_status, VRFDocStatus.PENDING)

    def test_confirm_generates_xlsx_and_marks_uploaded(self):
        self.client.force_login(self.user)
        payload = {
            "action": "confirm",
            "score_R1_a_p": "2", "score_R1_b_p": "3",
            "k_R1_p": "3",
        }
        response = self.client.post(
            reverse("tasks:project_vrf_compile", args=[self.project.id]), payload
        )
        self.assertEqual(response.status_code, 302)
        self.project.refresh_from_db()
        self.assertEqual(self.project.vrf_status, VRFDocStatus.UPLOADED)
        self.assertTrue(self.project.vrf_file.name)
        self.assertTrue(self.project.vrf_original_name.endswith(".xlsx"))
        import openpyxl
        self.project.vrf_file.open("rb")
        try:
            wb = openpyxl.load_workbook(io.BytesIO(self.project.vrf_file.read()))
        finally:
            self.project.vrf_file.close()
        ws = wb["VRF"]
        self.assertEqual(ws["B3"].value, "PN-TEST-1")
        self.assertEqual(ws["B4"].value, "Test Client")
        self.assertEqual(ws["U8"].value, 2)
        self.assertEqual(ws["U9"].value, 3)
        wb.close()

    def test_skip_with_reminder_saves_draft_and_redirects_to_gantt(self):
        from .models import VRFRiskAssessment

        self.client.force_login(self.user)
        payload = {
            "action": "skip_with_reminder",
            "score_R1_a_p": "2",
            "k_R1_p": "3",
        }
        response = self.client.post(
            reverse("tasks:project_vrf_compile", args=[self.project.id]), payload
        )
        self.assertRedirects(
            response,
            reverse("tasks:project_gantt", args=[self.project.id]),
        )
        # Stato rimane PENDING (non UPLOADED): il reminder progressivo resta attivo
        self.project.refresh_from_db()
        self.assertEqual(self.project.vrf_status, VRFDocStatus.PENDING)
        # La bozza e' comunque stata salvata
        assessment = VRFRiskAssessment.objects.get(project=self.project)
        self.assertEqual(assessment.data["risks"]["R1"]["subs"]["a"]["p"], 2)

    def test_dig_threshold_triggered_when_totals_high(self):
        from .models import VRFRiskAssessment
        from . import vrf_catalog

        self.client.force_login(self.user)
        payload = {"action": "save_draft"}
        for risk in vrf_catalog.RISKS:
            for sub in risk["sub_parameters"]:
                payload[f"score_{risk['code']}_{sub['code']}_p"] = "3"
            payload[f"k_{risk['code']}_p"] = "5"
            payload[f"k_{risk['code']}_i"] = str(risk["k_default"]["i"])
            payload[f"k_{risk['code']}_c"] = str(risk["k_default"]["c"])
        response = self.client.post(
            reverse("tasks:project_vrf_compile", args=[self.project.id]), payload
        )
        self.assertEqual(response.status_code, 302)
        assessment = VRFRiskAssessment.objects.get(project=self.project)
        self.assertGreaterEqual(assessment.total_p, vrf_catalog.DIG_THRESHOLD)
        self.assertTrue(assessment.dig_triggered)


class ReadinessComputeTests(TestCase):
    """F2 Task 1 — calcolo prontezza all'avvio (gate 4 criteri)."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(username="rd_owner", password="x")

    def _project(self, **kw):

        return make_project(name=kw.pop("name", "P"), created_by=self.user, **kw)

    def test_all_criteria_met_is_ready(self):
        from tasks.models import VRFDocStatus
        from tasks.readiness import compute_project_readiness

        p = self._project(
            vrf_status=VRFDocStatus.UPLOADED,
            project_manager=self.user, capo_commessa=self.user, programmer=self.user,
        )
        p.rd_has_meeting = True
        p.rd_has_planned = True
        r = compute_project_readiness(p)
        assert r.met == 4 and r.total == 4
        assert r.level == "ready" and r.label == "Pronto"
        assert all(c.action_url is None for c in r.criteria)

    def test_nothing_met_is_notready_with_action_urls(self):
        from tasks.readiness import compute_project_readiness

        p = self._project()  # vrf PENDING, no team, no meeting, no plan
        p.rd_has_meeting = False
        p.rd_has_planned = False
        r = compute_project_readiness(p)
        assert r.met == 0 and r.level == "notready" and r.label == "Non pronto"
        by = {c.key: c for c in r.criteria}
        assert by["vrf"].action_url and f"/tasks/projects/{p.id}/vrf/" in by["vrf"].action_url
        assert by["meeting"].action_url and f"/tasks/projects/{p.id}/incontri/new/" in by["meeting"].action_url
        assert by["plan"].action_url and f"project={p.id}" in by["plan"].action_url
        assert by["team"].action_url is None

    def test_two_met_is_partial(self):
        from tasks.models import VRFDocStatus
        from tasks.readiness import compute_project_readiness

        p = self._project(
            vrf_status=VRFDocStatus.NOT_REQUIRED,
            project_manager=self.user, capo_commessa=self.user, programmer=self.user,
        )
        p.rd_has_meeting = False
        p.rd_has_planned = False
        r = compute_project_readiness(p)
        assert r.met == 2 and r.level == "partial" and r.label == "Quasi pronto"

    def test_team_partial_is_not_ok(self):
        from tasks.models import VRFDocStatus
        from tasks.readiness import compute_project_readiness

        p = self._project(vrf_status=VRFDocStatus.UPLOADED, project_manager=self.user)
        p.rd_has_meeting = False
        p.rd_has_planned = False
        r = compute_project_readiness(p)
        by = {c.key: c for c in r.criteria}
        assert by["team"].ok is False


class ReadinessQuerysetTests(TestCase):
    """F2 Task 2 — annotazioni queryset (Exists), no-N+1, aggregato."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(username="rd_qs", password="x")

    def test_annotations_mark_meeting_and_plan(self):
        from tasks.models import Project, KickoffMeeting, Task
        from tasks.readiness import annotate_readiness_qs

        from tasks.models import MeetingStatus

        p_full = make_project(name="full", created_by=self.user)
        KickoffMeeting.objects.create(
            project=p_full, titolo="k", data=timezone.localdate(),
            created_by=self.user, stato=MeetingStatus.SVOLTO,
        )
        Task.objects.create(
            title="t", created_by=self.user, project=p_full, due_date=timezone.localdate()
        )
        p_empty = make_project(name="empty", created_by=self.user)

        rows = {p.id: p for p in annotate_readiness_qs(Project.objects.all())}
        assert rows[p_full.id].rd_has_meeting is True
        assert rows[p_full.id].rd_has_planned is True
        assert rows[p_empty.id].rd_has_meeting is False
        assert rows[p_empty.id].rd_has_planned is False

    def test_planned_meeting_does_not_count_as_done(self):
        from tasks.models import KickoffMeeting, MeetingStatus, Project
        from tasks.readiness import annotate_readiness_qs, compute_project_readiness

        p = make_project(name="solo pianificato", created_by=self.user)
        KickoffMeeting.objects.create(
            project=p, titolo="futuro", data=timezone.localdate() + timedelta(days=7),
            created_by=self.user, stato=MeetingStatus.PIANIFICATO,
        )
        row = annotate_readiness_qs(Project.objects.filter(pk=p.pk)).get()
        assert row.rd_has_meeting is False
        assert row.rd_has_any_meeting is True
        criteria = {c.key: c for c in compute_project_readiness(row).criteria}
        assert criteria["meeting"].ok is False
        # La CTA porta alla lista incontri (dove si registra l'esito), non a un doppione.
        assert f"/tasks/projects/{p.id}/incontri/" in criteria["meeting"].action_url
        assert not criteria["meeting"].action_url.endswith("/new/")

    def test_no_n_plus_one(self):
        from tasks.models import Project
        from tasks.readiness import annotate_readiness_qs, compute_project_readiness

        for i in range(3):
            make_project(name=f"P{i}", created_by=self.user)
        with self.assertNumQueries(1):
            projects = list(annotate_readiness_qs(Project.objects.all()))
            _ = [compute_project_readiness(p) for p in projects]

    def test_summary_counts_levels(self):
        from tasks.models import Project, VRFDocStatus
        from tasks.readiness import annotate_readiness_qs, readiness_summary

        make_project(name="n", created_by=self.user)
        make_project(
            name="p", created_by=self.user, vrf_status=VRFDocStatus.UPLOADED,
            project_manager=self.user, capo_commessa=self.user, programmer=self.user,
        )
        summary = readiness_summary(list(annotate_readiness_qs(Project.objects.all())))
        assert summary["notready"] >= 1 and summary["partial"] >= 1
        assert set(summary.keys()) == {"ready", "partial", "notready"}


class ReadinessPortfolioRenderTests(TestCase):
    """F2 Task 3 — badge readiness nelle card portfolio."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser(
            username="rd_admin", email="a@x.local", password="x"
        )

    def setUp(self):

        self.client.force_login(self.admin)
        make_project(name="Vuota", created_by=self.admin)

    def test_portfolio_shows_readiness_badge(self):
        r = self.client.get(reverse("tasks:project_list"))
        assert r.status_code == 200
        assert b"tk-readiness" in r.content


class ReadinessProjectHeaderRenderTests(TestCase):
    """F2 Task 4 — checklist readiness nell'header commessa."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser(
            username="rd_hdr", email="h@x.local", password="x"
        )

    def setUp(self):

        self.client.force_login(self.admin)
        self.project = make_project(name="Commessa", created_by=self.admin)

    def test_project_mode_shows_readiness_checklist(self):
        r = self.client.get(reverse("tasks:list") + f"?project={self.project.id}")
        assert r.status_code == 200
        assert b"tk-rchecklist" in r.content
        assert "Prontezza all'avvio".encode() in r.content


class ReadinessDashboardAggregateTests(TestCase):
    """F2 Task 5 — riepilogo aggregato + badge riga backlog."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser(
            username="rd_agg", email="g@x.local", password="x"
        )

    def setUp(self):
        from tasks.models import Task

        self.client.force_login(self.admin)
        p = make_project(name="C1", created_by=self.admin)
        Task.objects.create(
            title="t1", created_by=self.admin, project=p, due_date=timezone.localdate()
        )

    def test_dashboard_shows_readiness_summary_and_row_badge(self):
        r = self.client.get(reverse("tasks:list") + "?mine=0")
        assert r.status_code == 200
        assert b"tk-readiness-summary" in r.content
        assert b"tk-readiness" in r.content


class KickoffDaGestireTests(TestCase):
    """Centro 'Da gestire' — logica delle 4 sezioni."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser(
            username="dg_admin", email="dg@x.local", password="x"
        )

    def setUp(self):
        self.client.force_login(self.admin)

    def _req(self):
        # Request reale processato dal middleware (ha legacy_user, ecc.):
        # gli helper _scoped_* chiamano has_task_permission che tocca attributi
        # settati dal middleware, quindi NON usare RequestFactory.
        return self.client.get(reverse("tasks:list")).wsgi_request

    def _sections(self, data):
        return {s["key"]: s for s in data["sections"]}

    def test_vrf_pending_section(self):
        # NB: Project.save() sovrascrive il name con "KICK-OFF <n>" per i nuovi
        # progetti senza kickoff_number → verificare via id/url, non via name.
        from tasks.models import VRFDocStatus
        from tasks.da_gestire import build_kickoff_da_gestire

        p_pending = make_project(name="x", created_by=self.admin)  # PENDING
        make_project(
            name="y", created_by=self.admin, vrf_status=VRFDocStatus.UPLOADED
        )
        secs = self._sections(build_kickoff_da_gestire(self._req(), "portfolio"))
        urls = [i["url"] for i in secs["vrf"]["items"]]
        assert len(secs["vrf"]["items"]) == 1
        assert any(f"/tasks/projects/{p_pending.id}/vrf/" in u for u in urls)

    def test_not_ready_section(self):
        from tasks.da_gestire import build_kickoff_da_gestire

        p = make_project(name="x", created_by=self.admin)  # 0/4 -> notready
        secs = self._sections(build_kickoff_da_gestire(self._req(), "portfolio"))
        urls = [i["url"] for i in secs["not_ready"]["items"]]
        assert any(f"?project={p.id}" in u for u in urls)

    def test_critical_tasks_section_reasons(self):
        from tasks.models import Task, TaskStatus
        from tasks.da_gestire import build_kickoff_da_gestire

        p = make_project(name="C", created_by=self.admin)
        Task.objects.create(
            title="Scaduta", created_by=self.admin, project=p,
            due_date=timezone.localdate() - timedelta(days=2),
            status=TaskStatus.TODO, assigned_to=self.admin,
        )
        Task.objects.create(
            title="NonAssegnata", created_by=self.admin, project=p,
            due_date=timezone.localdate() + timedelta(days=5), status=TaskStatus.TODO,
        )
        secs = self._sections(build_kickoff_da_gestire(self._req(), "portfolio"))
        titles = [i["label"] for i in secs["critical"]["items"]]
        assert "Scaduta" in titles and "NonAssegnata" in titles

    def test_meetings_open_issues_section(self):
        from tasks.models import KickoffMeeting, MeetingIssue, MeetingIssueStatus
        from tasks.da_gestire import build_kickoff_da_gestire

        p = make_project(name="M", created_by=self.admin)
        m = KickoffMeeting.objects.create(
            project=p, titolo="Inc1", data=timezone.localdate(), created_by=self.admin
        )
        MeetingIssue.objects.create(
            project=p, source_meeting=m, title="Problema", status=MeetingIssueStatus.OPEN
        )
        secs = self._sections(build_kickoff_da_gestire(self._req(), "portfolio"))
        assert len(secs["meetings"]["items"]) == 1
        assert f"/tasks/projects/{p.id}/incontri/{m.id}/" in secs["meetings"]["items"][0]["url"]

    def test_total_and_scope_normalized(self):
        from tasks.da_gestire import build_kickoff_da_gestire

        data = build_kickoff_da_gestire(self._req(), "garbage")
        assert data["scope"] == "portfolio"
        assert data["total"] == sum(len(s["items"]) for s in data["sections"])


class KickoffDaGestireRenderTests(TestCase):
    """Pagina /tasks/da-gestire/ — render + toggle."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser(
            username="dg_render", email="r@x.local", password="x"
        )

    def setUp(self):

        self.client.force_login(self.admin)
        make_project(name="Vuota", created_by=self.admin)

    def test_page_renders_with_sections_and_toggle(self):
        r = self.client.get(reverse("tasks:da_gestire"))
        assert r.status_code == 200
        assert b"dg-section" in r.content
        assert b"Commesse non pronte" in r.content
        assert b"?scope=mine" in r.content

    def test_scope_mine_renders(self):
        r = self.client.get(reverse("tasks:da_gestire") + "?scope=mine")
        assert r.status_code == 200


class KickoffDaGestireSubnavTests(TestCase):
    def test_subnav_item_seeded(self):
        from core.models import NavigationItem

        item = NavigationItem.objects.get(code="tasks-sub-da-gestire")
        assert item.section == "subnav" and item.parent_code == "tasks"
        assert item.route_name == "tasks:da_gestire"
        assert item.required_permission_code == "tasks.kickoff.view"


class KickoffDaGestireAclTests(TestCase):
    def test_route_is_bound(self):
        from core.models import RoutePermissionBinding
        from tasks.acl_bootstrap import bootstrap_tasks_acl_endpoints

        bootstrap_tasks_acl_endpoints(force=True)
        assert RoutePermissionBinding.objects.filter(
            route_name="tasks:da_gestire", permission_id="tasks.kickoff.view", is_active=True
        ).exists()


class ProjectPhaseDeriveTests(TestCase):
    """Board per fase — derivazione fase iniziale (backfill)."""

    def test_derivation(self):
        from tasks.phase import derive_initial_phase

        assert derive_initial_phase(3, 0, "PENDING") == "DONE"
        assert derive_initial_phase(3, 2, "PENDING") == "EXEC"
        assert derive_initial_phase(0, 0, "UPLOADED") == "VRF"
        assert derive_initial_phase(0, 0, "PENDING") == "BOZZA"
        assert derive_initial_phase(0, 0, "NOT_REQUIRED") == "BOZZA"


class ProjectSetPhaseTests(TestCase):
    """Board per fase — endpoint cambio fase."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser(
            username="ph_admin", email="p@x.local", password="x"
        )
        cls.outsider = User.objects.create_user(username="ph_out", password="x")

    def _make_project(self):

        return make_project(name="P", created_by=self.admin)

    def test_editor_can_set_phase(self):
        p = self._make_project()
        self.client.force_login(self.admin)
        r = self.client.post(reverse("tasks:project_set_phase", args=[p.id]), {"phase": "EXEC"})
        assert r.status_code == 200
        assert r.json()["phase"] == "EXEC"
        p.refresh_from_db()
        assert p.phase == "EXEC"

    def test_invalid_phase_400(self):
        p = self._make_project()
        self.client.force_login(self.admin)
        r = self.client.post(reverse("tasks:project_set_phase", args=[p.id]), {"phase": "NOPE"})
        assert r.status_code == 400

    def test_non_editor_denied(self):
        # Il non-editor viene negato dal gate ACL (302 redirect del middleware
        # strict-canonical) o dalla view (403/404): in ogni caso NON cambia la fase.
        p = self._make_project()
        self.client.force_login(self.outsider)
        r = self.client.post(reverse("tasks:project_set_phase", args=[p.id]), {"phase": "EXEC"})
        assert r.status_code in (302, 403, 404)
        p.refresh_from_db()
        assert p.phase == "BOZZA"


class ProjectBoardRenderTests(TestCase):
    """Board per fase — render toggle Card/Board."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser(
            username="brd_admin", email="b@x.local", password="x"
        )

    def setUp(self):

        self.client.force_login(self.admin)
        make_project(name="P1", created_by=self.admin)

    def test_board_view_renders_columns(self):
        r = self.client.get(reverse("tasks:project_list") + "?view=board")
        assert r.status_code == 200
        assert b"kbf-board" in r.content
        assert b"In esecuzione" in r.content
        assert b"data-set-phase-url" in r.content

    def test_cards_view_default(self):
        r = self.client.get(reverse("tasks:project_list"))
        assert r.status_code == 200
        assert b"pf-grid" in r.content
        assert b"kbf-board" not in r.content


class PortfolioTimelineTests(TestCase):
    """Timeline cross-commessa — geometria roadmap pura."""

    def test_geometry_and_window(self):
        from datetime import date
        from tasks.timeline import build_portfolio_timeline

        items = [
            {"project": "A", "start": date(2026, 3, 1), "end": date(2026, 4, 30), "readiness": None},
            {"project": "B", "start": date(2026, 6, 1), "end": date(2026, 6, 15), "readiness": None},
        ]
        tl = build_portfolio_timeline(items, today=date(2026, 5, 10))
        assert tl["empty"] is False
        assert tl["window_start"] == date(2026, 3, 1)
        assert tl["window_end"] == date(2026, 6, 30)
        assert len(tl["months"]) == 4
        assert tl["rows"][0]["left_pct"] == 0.0
        assert 0 < tl["today_pct"] < 100

    def test_empty(self):
        from datetime import date
        from tasks.timeline import build_portfolio_timeline

        tl = build_portfolio_timeline([{"project": "X", "start": None, "end": None}], today=date(2026, 5, 1))
        assert tl["empty"] is True and tl["rows"] == []

    def test_clamp_end_before_start(self):
        from datetime import date
        from tasks.timeline import build_portfolio_timeline

        tl = build_portfolio_timeline(
            [{"project": "C", "start": date(2026, 5, 10), "end": date(2026, 5, 1)}],
            today=date(2026, 5, 5),
        )
        assert tl["rows"][0]["end"] == date(2026, 5, 10)


class PortfolioTimelineRenderTests(TestCase):
    """Timeline cross-commessa — render vista portfolio."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser(
            username="tl_admin", email="t@x.local", password="x"
        )

    def setUp(self):
        from tasks.models import Task

        self.client.force_login(self.admin)
        p = make_project(name="P", created_by=self.admin)
        Task.objects.create(
            title="t", created_by=self.admin, project=p,
            next_step_due=timezone.localdate(),
            due_date=timezone.localdate() + timedelta(days=20),
        )

    def test_timeline_view_renders_bar(self):
        r = self.client.get(reverse("tasks:project_list") + "?view=timeline")
        assert r.status_code == 200
        assert b"ptl-timeline" in r.content
        assert b"ptl-bar" in r.content

    def test_cards_still_default(self):
        r = self.client.get(reverse("tasks:project_list"))
        assert r.status_code == 200
        assert b"ptl-timeline" not in r.content


class VrfXlsxFormulaInjectionTests(TestCase):
    """Il generatore VRF (MOD.073) scrive nel template i campi del progetto
    (P/N, descrizione, cliente): sono testo libero, quindi `ws[cell] = valore`
    li scriverebbe come formula viva. Le formule ORIGINALI del template (K x R,
    totali) devono invece restare intatte. Dati sintetici.
    """

    EVIL = '=HYPERLINK("http://evil.test","x")'

    def _build(self, **project_kwargs):
        from types import SimpleNamespace

        from openpyxl import load_workbook

        from tasks import vrf_catalog
        from tasks.vrf_generator import build_vrf_xlsx

        defaults = {
            "vrf_quote_number": "Q-001",
            "versione": "1.0",
            "part_number": "PN-001",
            "vrf_description": "Descrizione sintetica",
            "vrf_esp": "ESP-1",
            "client_name": "Cliente Demo SRL",
        }
        defaults.update(project_kwargs)
        data = build_vrf_xlsx(SimpleNamespace(**defaults), None)
        wb = load_workbook(io.BytesIO(data))
        return wb[vrf_catalog.TEMPLATE_SHEET], vrf_catalog

    def test_header_fields_are_not_live_formulas(self):
        ws, catalog = self._build(part_number=self.EVIL, vrf_description="@SUM(A1)",
                                  client_name="+1+1", vrf_esp="-2+3")
        for key in ("part_number", "vrf_description", "client_name", "vrf_esp"):
            cell = ws[catalog.HEADER_CELLS[key]]
            self.assertNotEqual(cell.data_type, "f", f"{key} scritto come formula")
            self.assertTrue(cell.quotePrefix, f"{key} senza quote prefix")
        self.assertEqual(ws[catalog.HEADER_CELLS["part_number"]].value, self.EVIL)

    def test_template_formulas_and_benign_values_survive(self):
        ws, catalog = self._build()
        self.assertEqual(ws[catalog.HEADER_CELLS["part_number"]].value, "PN-001")
        self.assertFalse(ws[catalog.HEADER_CELLS["part_number"]].quotePrefix)
        # Non-regressione: le formule del template (totali TR) restano vive.
        tr_cell = ws[catalog.TR_CELLS["p"]]
        self.assertEqual(tr_cell.data_type, "f")
