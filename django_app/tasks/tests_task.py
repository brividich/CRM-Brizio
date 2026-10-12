"""Test delle attivita': elenco, audit, attrezzature, assenze, promemoria, import."""
from __future__ import annotations

from datetime import datetime, timedelta
import io
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.db import connection
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from core.models import Notifica
from attrezzature.models import (
    Attrezzatura,
    AttrezzaturaKickoffLink,
    AttrezzaturaStato,
    AttrezzaturaTask as GestioneAttrezzaturaTask,
    AttrezzaturaTaskOrigine as GestioneAttrezzaturaTaskOrigine,
    AttrezzaturaTaskStato as GestioneAttrezzaturaTaskStato,
    AttrezzaturaTaskTipo as GestioneAttrezzaturaTaskTipo,
)

from .models import (
    SubTask,
    Task,
    TaskCalendarEvent,
    TaskComment,
    TaskEvent,
    TaskEventType,
    TaskImpostazioni,
    TaskPriority,
    TaskStatus,
)
from .views import _task_date_absence_conflicts
from .tests_utils import (
    TasksBaseTestCase,
    _clear_assenze_table,
    _create_user_with_legacy,
    _ensure_assenze_table,
    _ensure_role,
    _grant_role_actions,
    make_project,
)

User = get_user_model()


@override_settings(LEGACY_AUTH_ENABLED=False, SECURE_SSL_REDIRECT=False)
class TaskAttrezzaturaEmbeddedPanelTests(TasksBaseTestCase):
    def setUp(self):
        super().setUp()
        _ensure_role(2, "utente")
        _grant_role_actions(2, ["tasks_view", "tasks_create", "tasks_edit", "tasks_comment"])
        self._refresh_acl_cache()
        self.user = _create_user_with_legacy(
            username="attpanel",
            legacy_user_id=2101,
            role_id=2,
            role_name="utente",
        )
        self.project = make_project(
            name="Kickoff panel",
            created_by=self.user,
            part_number=" pn-77 ",
        )
        self.task = Task.objects.create(
            title="Verifica attrezzatura",
            description="Controllare attrezzatura per produzione",
            created_by=self.user,
            assigned_to=self.user,
            project=self.project,
        )

    def test_detail_embeds_attrezzatura_panel_for_project_part_number(self):
        Attrezzatura.objects.create(
            codice="AT-77",
            part_number="PN-77",
            descrizione="Staffa prova",
            stato=AttrezzaturaStato.IN_CORSO,
            created_by=self.user,
        )
        GestioneAttrezzaturaTask.objects.create(
            part_number="PN-77",
            tipo=GestioneAttrezzaturaTaskTipo.VERIFICA_DISPONIBILITA,
            titolo="Verificare disponibilita PN-77",
            origine=GestioneAttrezzaturaTaskOrigine.KICKOFF,
            external_kickoff_id=str(self.project.id),
            external_kickoff_activity_id=str(self.task.id),
            created_by=self.user,
        )

        self.client.force_login(self.user)
        response = self.client.get(reverse("tasks:detail", args=[self.task.id]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Attrezzatura collegata - P/N PN-77")
        self.assertContains(response, "AT-77")
        self.assertContains(response, "Verificare disponibilita PN-77")
        self.assertContains(response, reverse("tasks:attrezzature_action", args=[self.task.id]))

    def test_attrezzatura_action_creates_kickoff_availability_task(self):
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("tasks:attrezzature_action", args=[self.task.id]),
            {"action": "create_availability_task"},
        )

        self.assertRedirects(response, reverse("tasks:detail", args=[self.task.id]))
        linked = GestioneAttrezzaturaTask.objects.get(
            tipo=GestioneAttrezzaturaTaskTipo.VERIFICA_DISPONIBILITA,
            external_kickoff_activity_id=str(self.task.id),
        )
        self.assertEqual(linked.part_number, "PN-77")
        self.assertEqual(linked.external_kickoff_id, str(self.project.id))
        self.assertEqual(linked.origine, GestioneAttrezzaturaTaskOrigine.KICKOFF)

    def test_create_task_can_request_new_tooling_workflow(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("tasks:create"),
            {
                "task_scope": "project",
                "project_link_mode": "existing",
                "project_choice": str(self.project.id),
                "title": "Creare attrezzo P/N 77",
                "description": "Serve attrezzo dedicato",
                "status": TaskStatus.TODO,
                "priority": TaskPriority.MEDIUM,
                "tooling_mode": "request_new",
                "tooling_code": "AT-REQ",
                "tooling_description": "Da creare per produzione",
                "reminder_portal_enabled_field": "on",
            },
        )
        task = Task.objects.get(title="Creare attrezzo P/N 77")
        self.assertRedirects(response, reverse("tasks:detail", args=[task.id]))
        tool = Attrezzatura.objects.get(codice="AT-REQ")
        self.assertEqual(tool.part_number, "PN-77")
        self.assertTrue(AttrezzaturaKickoffLink.objects.filter(attrezzatura=tool, task=task).exists())

    def test_create_task_can_link_existing_tooling(self):
        tool = Attrezzatura.objects.create(codice="AT-LINK", part_number="PN-77")
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("tasks:create"),
            {
                "task_scope": "project",
                "project_link_mode": "existing",
                "project_choice": str(self.project.id),
                "title": "Verificare attrezzo esistente",
                "status": TaskStatus.TODO,
                "priority": TaskPriority.MEDIUM,
                "tooling_mode": "link_existing",
                "tooling_existing_attrezzatura": str(tool.id),
                "reminder_portal_enabled_field": "on",
            },
        )
        task = Task.objects.get(title="Verificare attrezzo esistente")
        self.assertRedirects(response, reverse("tasks:detail", args=[task.id]))
        self.assertTrue(AttrezzaturaKickoffLink.objects.filter(attrezzatura=tool, task=task).exists())

    def test_attrezzatura_action_creates_draft_tool_from_kickoff_context(self):
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("tasks:attrezzature_action", args=[self.task.id]),
            {
                "action": "create_draft_tool",
                "codice": "AT-DRAFT",
                "description": "Bozza da activity",
            },
        )

        self.assertRedirects(response, reverse("tasks:detail", args=[self.task.id]))
        tool = Attrezzatura.objects.get(codice="AT-DRAFT")
        self.assertEqual(tool.part_number, "PN-77")
        self.assertEqual(tool.descrizione, "Bozza da activity")
        self.assertEqual(tool.origine_import, "kickoff")

    def test_attrezzatura_action_updates_progress_and_completes_task(self):
        tool = Attrezzatura.objects.create(
            codice="AT-UPD",
            part_number="PN-77",
            stato=AttrezzaturaStato.IN_CORSO,
            created_by=self.user,
        )
        linked = GestioneAttrezzaturaTask.objects.create(
            attrezzatura=tool,
            part_number="PN-77",
            tipo=GestioneAttrezzaturaTaskTipo.AGGIORNA_AVANZAMENTO,
            titolo="Aggiornare AT-UPD",
            origine=GestioneAttrezzaturaTaskOrigine.KICKOFF,
            external_kickoff_id=str(self.project.id),
            external_kickoff_activity_id=str(self.task.id),
            created_by=self.user,
        )
        self.client.force_login(self.user)

        progress_response = self.client.post(
            reverse("tasks:attrezzature_action", args=[self.task.id]),
            {
                "action": "update_progress",
                "attrezzatura_id": str(tool.id),
                "percentuale": "80",
                "stato": AttrezzaturaStato.IN_ATTESA,
                "note": "Aggiornato da dettaglio activity",
            },
        )
        self.assertRedirects(progress_response, reverse("tasks:detail", args=[self.task.id]))
        tool.refresh_from_db()
        self.assertEqual(tool.avanzamento_percentuale, 80)
        self.assertEqual(tool.stato, AttrezzaturaStato.IN_ATTESA)
        self.assertEqual(tool.avanzamenti.count(), 1)

        complete_response = self.client.post(
            reverse("tasks:attrezzature_action", args=[self.task.id]),
            {
                "action": "complete_attrezzatura_task",
                "attrezzatura_task_id": str(linked.id),
                "note": "Completata da KICK-OFF",
            },
        )
        self.assertRedirects(complete_response, reverse("tasks:detail", args=[self.task.id]))
        linked.refresh_from_db()
        self.assertEqual(linked.stato, GestioneAttrezzaturaTaskStato.COMPLETATA)
        self.assertEqual(linked.completed_by, self.user)


@override_settings(LEGACY_AUTH_ENABLED=False, SECURE_SSL_REDIRECT=False)
class TaskAuditTrailTests(TasksBaseTestCase):
    def setUp(self):
        super().setUp()
        _ensure_role(2, "utente")
        _grant_role_actions(2, ["tasks_view", "tasks_edit", "tasks_comment"])
        self._refresh_acl_cache()

        self.user = _create_user_with_legacy(username="audituser", legacy_user_id=3001, role_id=2, role_name="utente")
        self.task = Task.objects.create(title="Audit task", created_by=self.user, status=TaskStatus.TODO)

    def test_status_change_creates_audit_event(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("tasks:change_status", args=[self.task.id]),
            {"status": TaskStatus.IN_PROGRESS},
        )
        self.assertEqual(response.status_code, 302)
        event = TaskEvent.objects.filter(task=self.task, type=TaskEventType.STATUS_CHANGE).first()
        self.assertIsNotNone(event)
        self.assertEqual(event.payload.get("from"), TaskStatus.TODO)
        self.assertEqual(event.payload.get("to"), TaskStatus.IN_PROGRESS)

    def test_add_comment_creates_audit_event(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("tasks:add_comment", args=[self.task.id]),
            {"body": "Commento di test"},
        )
        self.assertEqual(response.status_code, 302)
        comment = TaskComment.objects.filter(task=self.task).first()
        self.assertIsNotNone(comment)
        event = TaskEvent.objects.filter(task=self.task, type=TaskEventType.COMMENT_ADDED).first()
        self.assertIsNotNone(event)
        self.assertEqual(event.payload.get("comment_id"), comment.id)

    def test_subtask_events_are_created(self):
        self.client.force_login(self.user)
        response_add = self.client.post(
            reverse("tasks:add_subtask", args=[self.task.id]),
            {"title": "Sub 1", "order_index": 1},
        )
        self.assertEqual(response_add.status_code, 302)
        subtask = SubTask.objects.get(task=self.task, title="Sub 1")

        response_status = self.client.post(
            reverse("tasks:edit_subtask_status", args=[self.task.id, subtask.id]),
            {"status": TaskStatus.DONE},
        )
        self.assertEqual(response_status.status_code, 302)

        add_event = TaskEvent.objects.filter(task=self.task, type=TaskEventType.SUBTASK_ADDED).first()
        status_event = TaskEvent.objects.filter(task=self.task, type=TaskEventType.SUBTASK_STATUS_CHANGE).first()
        self.assertIsNotNone(add_event)
        self.assertIsNotNone(status_event)
        self.assertEqual(add_event.payload.get("subtask_id"), subtask.id)
        self.assertEqual(status_event.payload.get("subtask_id"), subtask.id)
        self.assertEqual(status_event.payload.get("from"), TaskStatus.TODO)
        self.assertEqual(status_event.payload.get("to"), TaskStatus.DONE)

    def test_subtask_rollup_updates_parent_task_status(self):
        self.client.force_login(self.user)
        self.client.post(
            reverse("tasks:add_subtask", args=[self.task.id]),
            {"title": "Sub rollup", "order_index": 1},
        )
        subtask = SubTask.objects.get(task=self.task, title="Sub rollup")
        response = self.client.post(
            reverse("tasks:edit_subtask_status", args=[self.task.id, subtask.id]),
            {"status": TaskStatus.DONE},
        )
        self.assertEqual(response.status_code, 302)
        self.task.refresh_from_db()
        self.assertEqual(self.task.status, TaskStatus.DONE)
        rollup_events = TaskEvent.objects.filter(task=self.task, type=TaskEventType.STATUS_CHANGE)
        self.assertTrue(any((event.payload or {}).get("source") == "subtask_rollup" for event in rollup_events))

    def test_detail_timeline_renders_semantic_event_story(self):
        TaskEvent.objects.create(
            task=self.task,
            actor=self.user,
            type=TaskEventType.EDIT,
            payload={
                "changes": {
                    "next_step_due": {"from": "2026-03-10", "to": "2026-03-12"},
                    "due_date": {"from": "2026-03-15", "to": "2026-03-18"},
                }
            },
        )

        self.client.force_login(self.user)
        response = self.client.get(reverse("tasks:detail", args=[self.task.id]))

        self.assertContains(response, "Scheda attivita modificata")
        self.assertContains(response, "Campi aggiornati: data inizio, data fine.")
        self.assertContains(response, "Data inizio")
        self.assertContains(response, "10-03-2026")
        self.assertContains(response, "12-03-2026")
        self.assertContains(response, "Payload tecnico")
        self.assertNotContains(response, "{&#x27;changes&#x27;:")


@override_settings(LEGACY_AUTH_ENABLED=False, SECURE_SSL_REDIRECT=False)
class TaskListFiltersTests(TasksBaseTestCase):
    def setUp(self):
        super().setUp()
        _ensure_role(2, "utente")
        _ensure_role(3, "manager")
        _grant_role_actions(2, ["tasks_view"])
        _grant_role_actions(3, ["tasks_view", "tasks_admin"])
        self._refresh_acl_cache()

        self.user = _create_user_with_legacy(username="filteruser", legacy_user_id=4001, role_id=2, role_name="utente")
        self.manager = _create_user_with_legacy(username="manager", legacy_user_id=4002, role_id=3, role_name="manager")
        self.other = _create_user_with_legacy(username="other", legacy_user_id=4003, role_id=2, role_name="utente")
        self.project_alpha = make_project(name="Project Alpha", created_by=self.manager)
        self.project_beta = make_project(name="Project Beta", created_by=self.manager)

        today = timezone.localdate()
        self.t_overdue = Task.objects.create(
            title="Overdue TODO",
            created_by=self.other,
            assigned_to=self.user,
            project=self.project_alpha,
            status=TaskStatus.TODO,
            priority=TaskPriority.HIGH,
            due_date=today - timedelta(days=2),
            tags="produzione, urgente",
        )
        self.t_done_past = Task.objects.create(
            title="Done old",
            created_by=self.other,
            assigned_to=self.user,
            status=TaskStatus.DONE,
            priority=TaskPriority.HIGH,
            due_date=today - timedelta(days=1),
        )
        self.t_future_medium = Task.objects.create(
            title="Future medium",
            created_by=self.other,
            assigned_to=self.other,
            project=self.project_beta,
            status=TaskStatus.IN_PROGRESS,
            priority=TaskPriority.MEDIUM,
            due_date=today + timedelta(days=4),
            tags="it, inventory",
        )
        self.t_future_low = Task.objects.create(
            title="Future low",
            created_by=self.other,
            assigned_to=self.user,
            status=TaskStatus.TODO,
            priority=TaskPriority.LOW,
            due_date=today + timedelta(days=10),
            tags="planning",
        )
        self.t_unassigned = Task.objects.create(
            title="Unassigned task",
            created_by=self.other,
            status=TaskStatus.TODO,
            priority=TaskPriority.MEDIUM,
            due_date=today + timedelta(days=2),
        )
        self.t_without_due = Task.objects.create(
            title="Task without due date",
            created_by=self.other,
            assigned_to=self.user,
            project=self.project_alpha,
            status=TaskStatus.TODO,
            priority=TaskPriority.MEDIUM,
            due_date=None,
        )

    def test_filter_overdue(self):
        self.client.force_login(self.manager)
        response = self.client.get(reverse("tasks:list"), {"mine": "0", "overdue": "on"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.t_overdue.title)
        self.assertNotContains(response, self.t_done_past.title)
        self.assertNotContains(response, self.t_future_medium.title)

    def test_filter_status_priority_and_assigned_to(self):
        self.client.force_login(self.manager)
        response = self.client.get(
            reverse("tasks:list"),
            {
                "mine": "0",
                "status": TaskStatus.TODO,
                "priority": TaskPriority.LOW,
                "assigned_to": str(self.user.id),
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.t_future_low.title)
        self.assertNotContains(response, self.t_overdue.title)
        self.assertNotContains(response, self.t_future_medium.title)

    def test_filter_due_date_range(self):
        self.client.force_login(self.manager)
        today = timezone.localdate()
        response = self.client.get(
            reverse("tasks:list"),
            {
                "mine": "0",
                "due_from": (today + timedelta(days=1)).isoformat(),
                "due_to": (today + timedelta(days=6)).isoformat(),
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.t_future_medium.title)
        self.assertNotContains(response, self.t_future_low.title)

    def test_filter_by_tag(self):
        self.client.force_login(self.manager)
        response = self.client.get(
            reverse("tasks:list"),
            {
                "mine": "0",
                "tag": "urgente",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.t_overdue.title)
        self.assertNotContains(response, self.t_future_medium.title)

    def test_filter_by_project(self):
        self.client.force_login(self.manager)
        response = self.client.get(
            reverse("tasks:list"),
            {
                "mine": "0",
                "project": str(self.project_alpha.id),
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.t_overdue.title)
        self.assertNotContains(response, self.t_future_medium.title)

    def test_filter_unassigned(self):
        self.client.force_login(self.manager)
        response = self.client.get(
            reverse("tasks:list"),
            {
                "mine": "0",
                "unassigned": "on",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.t_unassigned.title)
        self.assertNotContains(response, self.t_overdue.title)

    def test_filter_without_due_date(self):
        self.client.force_login(self.manager)
        response = self.client.get(
            reverse("tasks:list"),
            {
                "mine": "0",
                "without_due_date": "on",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.t_without_due.title)
        self.assertNotContains(response, self.t_future_low.title)


@override_settings(LEGACY_AUTH_ENABLED=False, SECURE_SSL_REDIRECT=False)
class TaskAbsenceConflictTests(TasksBaseTestCase):
    def setUp(self):
        super().setUp()
        _ensure_role(2, "utente")
        _grant_role_actions(2, ["tasks_view", "tasks_create"])
        self._refresh_acl_cache()

        _ensure_assenze_table()
        _clear_assenze_table()

        self.owner = _create_user_with_legacy(username="absence_owner", legacy_user_id=6501, role_id=2, role_name="utente")
        self.assignee = _create_user_with_legacy(
            username="absence_assignee",
            legacy_user_id=6502,
            role_id=2,
            role_name="utente",
        )

    def _insert_absence(
        self,
        *,
        person_name: str,
        person_email: str,
        date_value,
        tipo: str = "Ferie",
        moderation_status: int = 0,
    ) -> None:
        start_dt = datetime.combine(date_value, datetime.min.time())
        end_dt = datetime.combine(date_value, datetime.max.time())
        with connection.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO assenze (
                    copia_nome,
                    email_esterna,
                    data_inizio,
                    data_fine,
                    tipo_assenza,
                    moderation_status,
                    consenso
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s)
                """,
                [person_name, person_email, start_dt, end_dt, tipo, moderation_status, "Approvato"],
            )

    def test_absence_conflicts_detected_on_task_dates(self):
        target_day = timezone.localdate() + timedelta(days=3)
        self._insert_absence(
            person_name=self.assignee.username,
            person_email=self.assignee.email,
            date_value=target_day,
        )
        task = Task.objects.create(
            title="Task con assenza",
            created_by=self.owner,
            assigned_to=self.assignee,
            status=TaskStatus.TODO,
            priority=TaskPriority.MEDIUM,
            due_date=target_day,
            next_step_due=target_day,
        )
        conflicts = _task_date_absence_conflicts(task)
        self.assertIn("due_date", conflicts)
        self.assertIn("next_step_due", conflicts)

    def test_project_gantt_marks_absence_cells(self):
        self.client.force_login(self.owner)
        target_day = timezone.localdate() + timedelta(days=2)
        self._insert_absence(
            person_name=self.assignee.username,
            person_email=self.assignee.email,
            date_value=target_day,
        )
        project = make_project(name="Project assenze", created_by=self.owner)
        Task.objects.create(
            title="Task conflitto gantt",
            created_by=self.owner,
            assigned_to=self.assignee,
            project=project,
            status=TaskStatus.IN_PROGRESS,
            next_step_due=target_day - timedelta(days=1),
            due_date=target_day + timedelta(days=1),
        )
        response = self.client.get(reverse("tasks:project_gantt", args=[project.id]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "assenze 1gg")
        self.assertTrue(response.context["gantt_rows"][0]["has_absence_conflicts"])


@override_settings(LEGACY_AUTH_ENABLED=False, SECURE_SSL_REDIRECT=False)
class TaskOutlookReminderTests(TasksBaseTestCase):
    def setUp(self):
        super().setUp()
        _ensure_role(2, "utente")
        _grant_role_actions(2, ["tasks_view", "tasks_create", "tasks_edit"])
        self._refresh_acl_cache()
        self.user = _create_user_with_legacy(
            username="outlookuser", legacy_user_id=8001, role_id=2, role_name="utente",
        )
        self.project = make_project(name="KO Reminder", created_by=self.user)
        self.impostazioni = TaskImpostazioni.get_singleton()
        self.impostazioni.notifiche_scadenza_attive = True
        self.impostazioni.giorni_preavviso = 3
        self.impostazioni.save()

    def test_reminder_created_on_task_save_when_enabled(self):
        from .models import TaskReminder

        target_due = timezone.localdate() + timedelta(days=10)
        self.client.force_login(self.user)
        payload = {
            "title": "Task con reminder",
            "description": "",
            "status": TaskStatus.TODO,
            "priority": TaskPriority.MEDIUM,
            "task_scope": "project",
            "project_link_mode": "existing",
            "project_choice": str(self.project.id),
            "due_date": target_due.isoformat(),
            "assigned_to": str(self.user.id),
            "reminder_portal_enabled_field": "on",
        }
        response = self.client.post(reverse("tasks:create"), payload)
        self.assertEqual(response.status_code, 302)
        task = Task.objects.get(title="Task con reminder")
        self.assertTrue(task.reminder_portal_enabled)
        reminders = TaskReminder.objects.filter(task=task)
        self.assertEqual(reminders.count(), 1)
        r = reminders.first()
        self.assertEqual(r.fire_at, target_due - timedelta(days=3))
        self.assertEqual(r.legacy_user_id, 8001)
        self.assertFalse(r.fired)

    def test_reminder_not_created_when_disabled(self):
        from .models import TaskReminder

        target_due = timezone.localdate() + timedelta(days=10)
        self.client.force_login(self.user)
        payload = {
            "title": "Task senza reminder",
            "description": "",
            "status": TaskStatus.TODO,
            "priority": TaskPriority.MEDIUM,
            "task_scope": "project",
            "project_link_mode": "existing",
            "project_choice": str(self.project.id),
            "due_date": target_due.isoformat(),
            "assigned_to": str(self.user.id),
            # reminder_portal_enabled_field NOT sent => False
        }
        response = self.client.post(reverse("tasks:create"), payload)
        self.assertEqual(response.status_code, 302)
        task = Task.objects.get(title="Task senza reminder")
        self.assertFalse(task.reminder_portal_enabled)
        self.assertEqual(TaskReminder.objects.filter(task=task).count(), 0)

    def test_send_task_reminders_command_creates_notifica(self):
        from .models import TaskReminder
        from django.core.management import call_command
        from io import StringIO

        task = Task.objects.create(
            title="T reminder fire",
            project=self.project,
            assigned_to=self.user,
            due_date=timezone.localdate() + timedelta(days=1),
            created_by=self.user,
            reminder_portal_enabled=True,
        )
        TaskReminder.objects.create(
            task=task, legacy_user_id=8001,
            fire_at=timezone.localdate(),
        )
        out = StringIO()
        call_command("send_task_reminders", stdout=out)
        r = TaskReminder.objects.get(task=task)
        self.assertTrue(r.fired)
        self.assertIsNotNone(r.fired_at)
        self.assertEqual(
            Notifica.objects.filter(legacy_user_id=8001).count(),
            1,
        )
        notif = Notifica.objects.get(legacy_user_id=8001)
        self.assertIn("T reminder fire", notif.messaggio)
        self.assertTrue(notif.url_azione)

    def test_send_task_reminders_skips_closed_task(self):
        from .models import TaskReminder
        from django.core.management import call_command
        from io import StringIO

        task = Task.objects.create(
            title="T chiuso",
            project=self.project, assigned_to=self.user,
            due_date=timezone.localdate() + timedelta(days=1),
            status=TaskStatus.DONE,
            created_by=self.user, reminder_portal_enabled=True,
        )
        TaskReminder.objects.create(
            task=task, legacy_user_id=8001, fire_at=timezone.localdate(),
        )
        call_command("send_task_reminders", stdout=StringIO())
        self.assertTrue(TaskReminder.objects.get(task=task).fired)
        self.assertEqual(Notifica.objects.filter(legacy_user_id=8001).count(), 0)

    def test_edit_form_prefills_existing_outlook_target(self):
        from .forms import TaskForm

        task = Task.objects.create(
            title="Task con evento Outlook",
            project=self.project,
            assigned_to=self.user,
            due_date=timezone.localdate() + timedelta(days=2),
            created_by=self.user,
            reminder_portal_enabled=True,
        )
        TaskCalendarEvent.objects.create(
            task=task,
            source_key=f"tasks.task:{task.id}:due",
            target_email="planner@example.com",
            target_display_name="Planner",
            due_date=task.due_date,
            subject="Task con evento Outlook",
            transaction_id="tx-001",
            graph_event_id="evt-001",
            graph_event_web_link="https://outlook.office.com/calendar/item/evt-001",
            created_by=self.user,
        )

        form = TaskForm(instance=task)

        self.assertTrue(form.initial["add_to_outlook"])
        self.assertEqual(form.initial["outlook_target_email"], "planner@example.com")

    def test_sync_task_outlook_event_access_denied_message_is_actionable(self):
        from .outlook_reminder import sync_task_outlook_event

        self.user.email = "outlook.user@cnovicrom.local"
        self.user.save(update_fields=["email"])
        task = Task.objects.create(
            title="Task access denied",
            project=self.project,
            assigned_to=self.user,
            due_date=timezone.localdate() + timedelta(days=4),
            created_by=self.user,
            reminder_portal_enabled=True,
        )

        with patch("tasks.outlook_reminder.graph_ready", return_value=True), patch(
            "tasks.outlook_reminder.create_event",
            side_effect=RuntimeError("Access is denied. Check credentials and try again."),
        ):
            level, message = sync_task_outlook_event(
                request=None,
                task=task,
                requested=True,
                explicit_email="",
            )

        self.assertEqual(level, "warning")
        self.assertIn("Calendars.ReadWrite", message)
        self.assertIn("Email calendario Outlook", message)
        self.assertIn(".local", message)


class ImportTemplateXlsxTests(TestCase):
    """Il template di import kickoff resta scaricabile e con le celle testuali."""

    def setUp(self):
        self.user = get_user_model().objects.create_superuser(
            username="tasks-xlsx-template",
            email="tasks-xlsx-template@example.invalid",
            password="pwd12345",
        )
        self.client.force_login(self.user)

    def test_template_download_has_no_live_formula(self):
        from openpyxl import load_workbook

        response = self.client.get(reverse("tasks:download_excel_template"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        wb = load_workbook(io.BytesIO(response.content))
        try:
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        self.assertNotEqual(cell.data_type, "f", f"{ws.title}!{cell.coordinate}")
            ws = wb["Kickoff Import"]
            self.assertEqual(ws["A1"].value, "Kickoff / Riferimento")
            self.assertTrue(ws["A1"].font.bold)
            self.assertEqual(ws["A2"].value, "Rif. kickoff Alfa")
            self.assertEqual(wb["Istruzioni"]["A1"].value, "Istruzioni import attivita kickoff")
            self.assertEqual(wb["Istruzioni"]["A3"].value, "Campo")
            self.assertEqual(wb["Istruzioni"]["A4"].value, "Kickoff / Riferimento")
        finally:
            wb.close()
