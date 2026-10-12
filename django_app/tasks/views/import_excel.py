"""Import massivo attivita' da Excel."""
from __future__ import annotations

import logging

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.db import transaction
from django.shortcuts import redirect, render

from core.audit import log_action
from core.upload_mime import (
    UploadMimeValidationError,
    safe_filename,
    validate_extension_and_mime,
)

from ..models import (
    Project,
    Task,
    TaskPriority,
    TaskStatus,
)
from ..access import (
    task_permissions_required,
)
from ..services.absences import (
    _coerce_date,
)
from ._shared import (
    _tasks_shell_context,
)

logger = logging.getLogger(__name__)
User = get_user_model()


# ---------------------------------------------------------------------------
# Excel import / template download
# ---------------------------------------------------------------------------

_IMPORT_COLUMNS = [
    "Kickoff / Riferimento",
    "Cliente",
    "P/N",
    "Revisione",
    "Versione",
    "PM (username)",
    "Titolo attivita",
    "Descrizione attivita",
    "Stato",       # TODO / IN_PROGRESS / DONE / CANCELED
    "Priorita",    # LOW / MEDIUM / HIGH
    "Assegnato a (username)",
    "Data inizio (YYYY-MM-DD)",
    "Data fine (YYYY-MM-DD)",
    "Tag",
]


@task_permissions_required("tasks_create")
def download_excel_template(request):
    """Genera e scarica il template .xlsx per l'import attivita kickoff."""
    import io
    import openpyxl
    from django.http import HttpResponse

    from core.excel_export import write_cell, write_cell_at, write_row

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kickoff Import"

    # Header row
    for col_idx, col_name in enumerate(_IMPORT_COLUMNS, start=1):
        cell = write_cell(ws, 1, col_idx, col_name)
        cell.font = openpyxl.styles.Font(bold=True)

    # Example row
    example = [
        "Rif. kickoff Alfa",
        "Cliente SRL",
        "PN-001",
        "A",
        "1.0",
        "",
        "Titolo attivita di esempio",
        "Descrizione operativa",
        "TODO",
        "MEDIUM",
        "",
        "",
        "",
        "",
    ]
    write_row(ws, 2, example)

    # Istruzioni foglio
    ws_info = wb.create_sheet("Istruzioni")
    write_cell_at(ws_info, "A1", "Istruzioni import attivita kickoff")
    ws_info["A1"].font = openpyxl.styles.Font(bold=True, size=13)
    instructions = [
        ("Kickoff / Riferimento", "Riferimento libero per raggruppare piu righe nello stesso kickoff quando il P/N non e disponibile."),
        ("Cliente", "Ragione sociale cliente (testo libero)."),
        ("P/N", "Part Number (testo libero)."),
        ("Revisione", "Revisione del P/N (testo libero, es. A, B, 01)."),
        ("Versione", "Versione del P/N (testo libero, es. 1.0, 2.3)."),
        ("PM (username)", "Username del project manager. Deve esistere nel sistema."),
        ("Titolo attivita", "Titolo dell'attivita kickoff (obbligatorio)."),
        ("Descrizione attivita", "Descrizione operativa (opzionale)."),
        ("Stato", "TODO / IN_PROGRESS / DONE / CANCELED. Default: TODO."),
        ("Priorita", "LOW / MEDIUM / HIGH. Default: MEDIUM."),
        ("Assegnato a (username)", "Username dell'operatore. Deve esistere nel sistema."),
        ("Data inizio (YYYY-MM-DD)", "Data inizio (next_step_due). Formato: 2025-03-15."),
        ("Data fine (YYYY-MM-DD)", "Data fine prevista (due_date). Formato: 2025-04-30."),
        ("Tag", "Etichette separate da virgola."),
    ]
    write_cell_at(ws_info, "A3", "Campo")
    write_cell_at(ws_info, "B3", "Descrizione")
    ws_info["A3"].font = openpyxl.styles.Font(bold=True)
    ws_info["B3"].font = openpyxl.styles.Font(bold=True)
    for row_idx, (field, desc) in enumerate(instructions, start=4):
        write_row(ws_info, row_idx, [field, desc])
    ws_info.column_dimensions["A"].width = 28
    ws_info.column_dimensions["B"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    response = HttpResponse(
        buf.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="template_import_kickoff.xlsx"'
    return response


@task_permissions_required("tasks_create")
def import_excel(request):
    """Upload + anteprima + commit import attivita kickoff da Excel."""
    import io
    import openpyxl

    if request.method == "GET":
        return render(request, "tasks/import.html", {
            **_tasks_shell_context(request, active="import"),
            "page_title": "Import attivita kickoff da Excel",
            "columns": _IMPORT_COLUMNS,
        })

    # --- STEP 1: upload file e anteprima ---
    if "file" in request.FILES and "confirm" not in request.POST:
        uploaded = request.FILES["file"]
        try:
            validate_extension_and_mime(
                uploaded,
                allowed_extensions={".xlsx", ".xls"},
                allowed_mimes={
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    "application/vnd.ms-excel",
                    "application/octet-stream",
                    "application/zip",
                    "application/x-zip-compressed",
                },
                max_bytes=20 * 1024 * 1024,
                label=safe_filename(uploaded.name) or "Import Excel",
                allow_empty=False,
            )
        except UploadMimeValidationError as exc:
            messages.error(request, str(exc))
            return redirect("tasks:import_excel")
        try:
            wb = openpyxl.load_workbook(io.BytesIO(uploaded.read()), read_only=True, data_only=True)
        except Exception:
            messages.error(request, "File non valido. Carica un file .xlsx corretto.")
            return redirect("tasks:import_excel")

        ws = wb.active
        rows_raw = list(ws.iter_rows(min_row=2, values_only=True))
        wb.close()

        preview_rows = []
        errors = []
        for row_idx, row in enumerate(rows_raw, start=2):
            if not any(row):
                continue
            def _cell(n):
                try:
                    return str(row[n] or "").strip()
                except IndexError:
                    return ""

            kickoff_reference = _cell(0)
            part_number = _cell(2)
            revisione = _cell(3) if part_number else ""
            versione = _cell(4) if part_number else ""
            titolo = _cell(6)
            if not titolo:
                errors.append(f"Riga {row_idx}: campo Titolo attivita vuoto — riga ignorata.")
                continue

            stato_raw = _cell(8).upper() or "TODO"
            if stato_raw not in {s.value for s in TaskStatus}:
                stato_raw = "TODO"
            priorita_raw = _cell(9).upper() or "MEDIUM"
            if priorita_raw not in {p.value for p in TaskPriority}:
                priorita_raw = "MEDIUM"

            preview_rows.append({
                "progetto": kickoff_reference,
                "cliente": _cell(1),
                "part_number": part_number,
                "revisione": revisione,
                "versione": versione,
                "pm_username": _cell(5),
                "titolo": titolo,
                "descrizione": _cell(7),
                "stato": stato_raw,
                "priorita": priorita_raw,
                "assegnato_username": _cell(10),
                "data_inizio": _cell(11),
                "data_fine": _cell(12),
                "tag": _cell(13),
            })

        request.session["tasks_import_preview"] = preview_rows
        return render(request, "tasks/import.html", {
            **_tasks_shell_context(request, active="import"),
            "page_title": "Import attivita kickoff da Excel",
            "columns": _IMPORT_COLUMNS,
            "preview_rows": preview_rows,
            "errors": errors,
            "filename": uploaded.name,
        })

    # --- STEP 2: commit ---
    if "confirm" in request.POST:
        preview_rows = request.session.pop("tasks_import_preview", [])
        if not preview_rows:
            messages.error(request, "Sessione scaduta o nessun dato da importare. Ricarica il file.")
            return redirect("tasks:import_excel")

        created_projects = 0
        updated_projects = 0
        created_tasks = 0
        updated_project_ids: set[int] = set()
        cached_projects_by_identity: dict[tuple[str, str, str], Project] = {}
        cached_projects_by_reference: dict[str, Project] = {}

        try:
            with transaction.atomic():
                for row in preview_rows:
                    # Risolvi PM
                    pm_user = None
                    if row["pm_username"]:
                        pm_user = User.objects.filter(username__iexact=row["pm_username"]).first()

                    proj_defaults = {
                        "client_name": row["cliente"],
                    }
                    part_number = (row["part_number"] or "").strip()
                    revisione = (row["revisione"] or "").strip()
                    versione = (row["versione"] or "").strip()
                    kickoff_reference = (row["progetto"] or "").strip()

                    project = None
                    project_was_created = False
                    if part_number:
                        identity_key = (
                            part_number.lower(),
                            revisione.lower(),
                            versione.lower(),
                        )
                        project = cached_projects_by_identity.get(identity_key)
                        if project is None:
                            project = (
                                Project.objects.filter(
                                    part_number__iexact=part_number,
                                    revisione__iexact=revisione,
                                    versione__iexact=versione,
                                )
                                .order_by("id")
                                .first()
                            )
                            if project is None:
                                project = Project.objects.create(
                                    name="",
                                    client_name=row["cliente"],
                                    part_number=part_number,
                                    revisione=revisione,
                                    versione=versione,
                                    created_by=request.user,
                                )
                                project_was_created = True
                                created_projects += 1
                            cached_projects_by_identity[identity_key] = project
                    elif kickoff_reference:
                        reference_key = kickoff_reference.lower()
                        project = cached_projects_by_reference.get(reference_key)
                        if project is None:
                            project = Project.objects.create(
                                name="",
                                client_name=row["cliente"],
                                created_by=request.user,
                            )
                            project_was_created = True
                            cached_projects_by_reference[reference_key] = project
                            created_projects += 1
                    else:
                        project = Project.objects.create(
                            name="",
                            client_name=row["cliente"],
                            created_by=request.user,
                        )
                        project_was_created = True
                        created_projects += 1

                    updated_fields = []
                    for field_name, value in proj_defaults.items():
                        if getattr(project, field_name) != value:
                            setattr(project, field_name, value)
                            updated_fields.append(field_name)
                    if updated_fields:
                        project.save(update_fields=updated_fields + ["updated_at"])
                    if pm_user is not None and not project.project_managers.filter(pk=pm_user.pk).exists():
                        # Il PM del file si aggiunge all'elenco, non sostituisce chi c'e' gia'.
                        project.project_managers.add(pm_user)
                        updated_fields.append("project_managers")
                    if project.id not in updated_project_ids and updated_fields and not project_was_created:
                        updated_project_ids.add(project.id)
                        updated_projects += 1

                    # Risolvi assegnatario
                    assigned = None
                    if row["assegnato_username"]:
                        assigned = User.objects.filter(username__iexact=row["assegnato_username"]).first()

                    # Date
                    data_inizio = _coerce_date(row["data_inizio"])
                    data_fine = _coerce_date(row["data_fine"])
                    # Sanity: data_fine deve essere > data_inizio
                    if data_inizio and data_fine and data_fine <= data_inizio:
                        data_fine = None

                    Task.objects.create(
                        title=row["titolo"],
                        description=row["descrizione"],
                        status=row["stato"],
                        priority=row["priorita"],
                        project=project,
                        assigned_to=assigned,
                        next_step_due=data_inizio,
                        due_date=data_fine,
                        tags=row["tag"],
                        created_by=request.user,
                    )
                    created_tasks += 1

        except Exception as exc:
            messages.error(request, f"Errore durante l'import: {exc}")
            return redirect("tasks:import_excel")

        log_action(
            request,
            "import_excel",
            "tasks",
            {
                "message": (
                    f"Import Excel: {created_tasks} attivita kickoff create, "
                    f"{created_projects} kickoff nuovi, {updated_projects} kickoff aggiornati."
                ),
                "created_tasks": created_tasks,
                "created_projects": created_projects,
                "updated_projects": updated_projects,
            },
        )
        messages.success(
            request,
            (
                f"Import completato: {created_tasks} attivita kickoff create "
                f"({created_projects} kickoff nuovi, {updated_projects} aggiornati)."
            ),
        )
        return redirect("tasks:list")

    return redirect("tasks:import_excel")
