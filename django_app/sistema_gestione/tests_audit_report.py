"""Cruscotto, export rilievi, allegato del rapporto e «I miei rilievi» (prompt 08, fasi 6-7)."""
from datetime import date
from io import BytesIO
from unittest.mock import patch

import fitz
from django.contrib.auth import get_user_model
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.utils import timezone

from core.models import SiteConfig
from core.test_excel_export import assert_no_live_formula

from . import miei_rilievi_views, procedure_views, report_views
from .audit_exports import rapporto_audit_pdf
from .models import (Audit, AuditEsito, AuditPersona, AuditVerificaEfficacia, AzioneCorrettivaAudit, CellaProgramma,
                     ProgrammaAudit, RigaProgramma)
from .services import audit as service
from .services import cruscotto
from .tests_audit import AuditBase, _request

User = get_user_model()


class CruscottoTests(AuditBase):
    def setUp(self):
        super().setUp()
        self.anno = timezone.localdate().year
        self.processo.punti_9100 = "8.1, 8.5"
        self.processo.save()

    def _audit_chiuso(self, n, mese=1):
        audit = self.make_audit(numero=f"RAIS-{self.anno}-{n:02d}", stato=Audit.STATO_CHIUSO, en9100=True,
                                data_inizio=date(self.anno, mese, 10))
        audit.processi_catalogo.add(self.processo)
        return audit

    def _rilievo(self, audit, esito="NC", gravita="MINORE", punti="8.5"):
        return AuditEsito.objects.create(audit=audit, processo=self.processo, testo_aggiuntivo="x", punti_aggiuntivi=punti,
                                         esito=esito, gravita=gravita if esito == "NC" else "", evidenze="e")

    def test_indicatori_programma_rilievi_car_efficacia_copertura(self):
        programma = ProgrammaAudit.objects.create(anno=self.anno, stato=ProgrammaAudit.STATO_APPROVATO)
        riga = RigaProgramma.objects.create(programma=programma, processo=self.processo, area="QA")
        primo = self._audit_chiuso(1)
        CellaProgramma.objects.create(riga=riga, mese=1, audit=primo)
        CellaProgramma.objects.create(riga=riga, mese=2)
        nc = self._rilievo(primo, gravita="MAGGIORE")
        self._rilievo(primo, esito="OFI")
        AuditEsito.objects.create(audit=primo, processo=self.processo, testo_aggiuntivo="c", punti_aggiuntivi="8.1", esito="CONFORME", evidenze="ok")
        secondo = self._audit_chiuso(2, mese=1)
        self._rilievo(secondo)
        voce = service.sincronizza_ofi(nc)
        AzioneCorrettivaAudit.objects.create(registro=voce, data_richiesta=date(self.anno, 1, 1), responsabile=self.direzione,
                                             chiusa_il=timezone.make_aware(timezone.datetime(self.anno, 1, 31)))
        AuditVerificaEfficacia.objects.create(esito=nc, risultato="EFFICACE", metodo="m", evidenza="e", data_verifica=date(self.anno, 2, 1))
        c = cruscotto.cruscotto(self.anno, oggi=date(self.anno, 6, 1))
        self.assertEqual((c["avanzamento"]["pianificati"], c["avanzamento"]["eseguiti"], c["avanzamento"]["in_ritardo"]), (2, 1, 1))
        self.assertEqual((c["totali"]["NC_MAGGIORE"], c["totali"]["NC_MINORE"], c["totali"]["OFI"]), (1, 1, 1))
        self.assertEqual(c["giorni_medi_chiusura"], 30)
        self.assertEqual(c["efficacia"]["percentuale"], 100)
        self.assertEqual(c["ricorrenti"][0]["audit"], 2)
        en9100 = next(n for n in c["copertura"] if n["norma"] == "EN 9100")
        self.assertEqual((en9100["verificate"], en9100["richieste"]), (2, 2))

    def test_query_costanti_al_crescere_dei_rilievi(self):
        audit = self._audit_chiuso(1)
        self._rilievo(audit)
        with CaptureQueriesContext(connection) as pochi:
            cruscotto.cruscotto(self.anno)
        for _ in range(6):
            self._rilievo(audit, esito="OFI")
        with CaptureQueriesContext(connection) as molti:
            cruscotto.cruscotto(self.anno)
        self.assertEqual(len(pochi), len(molti))

    @patch("sistema_gestione.report_views._has_perm", return_value=True)
    def test_export_rilievi_senza_formule_vive(self, _perm):
        audit = self._audit_chiuso(1)
        AuditEsito.objects.create(audit=audit, processo=self.processo, testo_aggiuntivo="x", punti_aggiuntivi="8.5",
                                  esito="NC", gravita="MINORE", evidenze="e", scostamento="=SOMMA(A1:A9)")
        response = report_views.rilievi_export(_request(self.user, method="get", data={"audit": audit.pk}))
        self.assertEqual(response.status_code, 200)
        assert_no_live_formula(self, response.content)
        from openpyxl import load_workbook
        testo = " ".join(str(c.value) for row in load_workbook(BytesIO(response.content)).active.iter_rows() for c in row if c.value)
        self.assertIn("NC minore", testo)
        self.assertIn(audit.numero, testo)

    @patch("sistema_gestione.report_views._has_perm", return_value=True)
    def test_pagina_cruscotto(self, _perm):
        self.assertContains(report_views.cruscotto(_request(self.user, method="get")), "Copertura delle clausole")


class AllegatoRilieviPdfTests(AuditBase):
    def test_allegato_presente_e_disattivabile(self):
        audit = self.make_audit()
        AuditEsito.objects.create(audit=audit, testo_aggiuntivo="x", punti_aggiuntivi="7.5.3", esito="OSS", evidenze="Registrazione incompleta")
        testo = lambda: "".join(p.get_text() for p in fitz.open(stream=rapporto_audit_pdf(audit), filetype="pdf"))  # noqa: E731
        self.assertIn("RILIEVI E AZIONI", testo())
        SiteConfig.set("sistema_gestione.rapporto_allegato_rilievi", "0")
        self.assertNotIn("RILIEVI E AZIONI", testo())


class MieiRilieviTests(AuditBase):
    def setUp(self):
        super().setUp()
        self.auditato = User.objects.create_user(username="auditato-test", password="x", email="auditato@example.test")
        self.estraneo = User.objects.create_user(username="estraneo-test2", password="x", email="estraneo@example.test")
        self.audit = self.make_audit(rapporto_firmato_auditor_il=timezone.now(), rapporto_firmato_auditor_da=self.user)
        AuditPersona.objects.create(audit=self.audit, nome="Auditato", email="AUDITATO@example.test", ruolo=AuditPersona.RUOLO_AUDITATO)
        self.nc = AuditEsito.objects.create(audit=self.audit, testo_aggiuntivo="x", punti_aggiuntivi="8.5", esito="NC",
                                            gravita="MINORE", evidenze="e", responsabile_azione=self.estraneo)
        bozza = self.make_audit(numero="RAIS-2026-77")
        AuditPersona.objects.create(audit=bozza, nome="Auditato", email="auditato@example.test", ruolo=AuditPersona.RUOLO_AUDITATO)
        AuditEsito.objects.create(audit=bozza, testo_aggiuntivo="y", punti_aggiuntivi="9.1", esito="OFI", evidenze="in bozza")

    def test_persona_collegata_per_email_univoca(self):
        self.assertEqual(self.audit.persone.get().user, self.auditato)
        User.objects.create_user(username="doppio-1", password="x", email="doppio@example.test")
        User.objects.create_user(username="doppio-2", password="x", email="doppio@example.test")
        ambigua = AuditPersona.objects.create(audit=self.audit, nome="Ambiguo", email="doppio@example.test", ruolo=AuditPersona.RUOLO_AUDITATO)
        self.assertIsNone(ambigua.user)

    def test_vede_solo_i_propri_rilievi_a_rapporto_firmato(self):
        # Responsabile dell'azione: vede il rilievo. Auditato senza responsabilità sul rilievo: no.
        self.assertEqual(list(miei_rilievi_views.rilievi_utente(self.estraneo)), [self.nc])
        self.assertEqual(list(miei_rilievi_views.rilievi_utente(self.auditato)), [])
        # Responsabile del processo del rilievo, coinvolto nell'audit: lo vede.
        from .models import Processo
        AuditEsito.objects.filter(pk=self.nc.pk).update(processo=Processo.objects.create(codice="PX", nome="X", responsabile=self.auditato))
        self.assertEqual(list(miei_rilievi_views.rilievi_utente(self.auditato)), [self.nc])
        terzo = User.objects.create_user(username="terzo-test", password="x")
        self.assertEqual(list(miei_rilievi_views.rilievi_utente(terzo)), [])

    @patch("sistema_gestione.miei_rilievi_views._has_perm", side_effect=lambda r, p: p == "sistema_gestione.audit.miei_rilievi")
    def test_pagina_mostra_link_car_solo_al_responsabile(self, _perm):
        service.sincronizza_ofi(self.nc)
        self.assertNotContains(miei_rilievi_views.miei_rilievi(_request(self.auditato, method="get")), "Gestisci l'azione correttiva")
        self.assertContains(miei_rilievi_views.miei_rilievi(_request(self.estraneo, method="get")), "Gestisci l'azione correttiva")

    def test_car_aperta_dal_responsabile_senza_consultazione_audit(self):
        voce = service.sincronizza_ofi(self.nc)
        solo_miei = lambda r, p: p == "sistema_gestione.audit.miei_rilievi"  # noqa: E731
        with patch("sistema_gestione.procedure_views._has_perm", side_effect=solo_miei):
            self.assertEqual(procedure_views.car(_request(self.estraneo, method="get"), voce.pk).status_code, 200)
            self.assertEqual(procedure_views.car(_request(self.auditato, method="get"), voce.pk).status_code, 302)
