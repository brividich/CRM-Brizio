"""Import dell'inventario software (export WatchGuard EPDR/Panda, BusinessLog, altri CSV/XLSX).

Il formato reale degli export non è fissato nel codice: l'utente indica quale colonna è
hostname, vendor, prodotto, versione e data di rilevamento (mappatura salvabile come preset
per fonte). Flusso: upload → mappatura → anteprima (righe valide e scartate con motivo) →
import. L'import è idempotente: la chiave di un'installazione è (fonte, host, vendor,
prodotto, versione), quindi lo stesso file due volte non crea duplicati. Le installazioni
non più presenti per gli host del file si marcano «non più rilevate», mai cancellate.
"""
from __future__ import annotations

import csv
import hashlib
import io
import logging
import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta

from django.db import transaction
from django.utils import timezone

from security.models import (
    SecurityAsset,
    SoftwareAlias,
    SoftwareImportPreset,
    SoftwareInstallation,
    SoftwareInventoryImport,
)

logger = logging.getLogger(__name__)

MAX_UPLOAD_BYTES = 15 * 1024 * 1024
MAX_ROWS = 100_000
MAX_SKIPPED_DETAILS = 200
MAX_COLUMNS = 200
ALLOWED_EXTENSIONS = {".csv", ".xlsx"}
ALLOWED_MIMES = {
    "text/csv", "text/plain", "application/csv", "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/zip",
    "application/x-zip-compressed", "application/octet-stream",
}
STORAGE_DIR = "security_inventory"

FIELD_LABELS = {
    "hostname": "Hostname / computer",
    "vendor": "Vendor / editore",
    "product": "Prodotto / software",
    "version": "Versione",
    "detected_at": "Data rilevamento",
}
# Intestazioni tipiche (italiano e inglese) per proporre la mappatura: si conferma sempre a mano.
HEADER_HINTS = {
    "hostname": ("hostname", "computer", "nome computer", "computer name", "device", "dispositivo", "host", "endpoint", "pc", "macchina"),
    "vendor": ("vendor", "publisher", "editore", "produttore", "fabbricante", "manufacturer", "developer", "sviluppatore"),
    "product": ("software", "product", "prodotto", "application", "applicazione", "program", "programma", "name", "nome", "nome software"),
    "version": ("version", "versione", "ver", "release"),
    "detected_at": ("detected", "rilevato", "data rilevamento", "install date", "data installazione", "installed on", "last seen",
                    "ultimo rilevamento", "date", "data"),
}
DATE_FORMATS = ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d", "%m/%d/%Y", "%d.%m.%Y")

_VENDOR_SUFFIXES = re.compile(r"[,.]?\s+(inc|incorporated|corp|corporation|co|company|ltd|limited|llc|gmbh|srl|s\.r\.l|spa|s\.p\.a|ag|sa|bv|plc)\.?$", re.I)
_VERSIONISH = re.compile(r"^v?\d+([._-]\d+)*[a-z]?$", re.I)
_ARCH = re.compile(r"\((x64|x86|64-bit|32-bit|amd64|arm64)\)|\b(x64|x86|64-bit|32-bit|amd64)\b", re.I)


class InventoryFileError(ValueError):
    """File illeggibile o fuori dai limiti: messaggio per l'utente."""


@dataclass
class ParsedTable:
    headers: list
    rows: list
    truncated: bool = False


@dataclass
class RowPreview:
    valid: list = field(default_factory=list)
    skipped: list = field(default_factory=list)  # (numero riga, motivo)
    total: int = 0


# --- Lettura file ----------------------------------------------------------------------------

def read_table(data: bytes, filename: str) -> ParsedTable:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        return _read_xlsx(data)
    if name.endswith(".csv"):
        return _read_csv(data)
    raise InventoryFileError("Formato non supportato: usa CSV o XLSX.")


def _decode(data):
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise InventoryFileError("Codifica del CSV non riconosciuta (attese UTF-8 o Windows-1252).")


def _read_csv(data):
    text = _decode(data)
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    try:
        return _table_from_rows(reader)
    except csv.Error as exc:
        raise InventoryFileError(f"CSV non leggibile: {exc}.") from exc


def _read_xlsx(data):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - qualunque errore di openpyxl = file non leggibile per l'utente
        raise InventoryFileError(f"XLSX non leggibile: {exc.__class__.__name__}.") from exc
    try:
        sheet = workbook.worksheets[0]
        rows = (
            [("" if cell is None else cell) for cell in row]
            for row in sheet.iter_rows(values_only=True, max_col=MAX_COLUMNS)
        )
        return _table_from_rows(rows)
    finally:
        workbook.close()


def _table_from_rows(rows):
    headers, out, truncated = None, [], False
    for row in rows:
        cells = list(row)[:MAX_COLUMNS]
        if headers is None:
            # Intestazione = prima riga con almeno due celle compilate (gli export hanno spesso un titolo sopra).
            if sum(1 for cell in cells if str(cell).strip()) >= 2:
                headers = [str(cell).strip()[:120] for cell in cells]
                while headers and not headers[-1]:
                    headers.pop()  # colonne vuote in coda (openpyxl con max_col le riempie)
            continue
        cells = cells[: len(headers)]
        if not any(str(cell).strip() for cell in cells):
            continue
        if len(out) >= MAX_ROWS:
            truncated = True
            break
        out.append(cells)
    if headers is None:
        raise InventoryFileError("Nessuna riga di intestazione trovata nel file.")
    return ParsedTable(headers=headers, rows=out, truncated=truncated)


# --- Mappatura -------------------------------------------------------------------------------

def _header_key(text):
    return re.sub(r"[^a-z0-9]+", " ", str(text).lower()).strip()


def suggest_mapping(headers):
    """Proposta di mappatura dalle intestazioni (l'utente la conferma o la corregge)."""
    keys = {_header_key(h): h for h in headers if h}
    mapping, used = {}, set()
    for field_name, hints in HEADER_HINTS.items():
        for hint in hints:
            match = next((orig for key, orig in keys.items() if orig not in used and (key == hint or key.startswith(hint + " "))), None)
            if match:
                mapping[field_name] = match
                used.add(match)
                break
    return mapping


def validate_mapping(headers, column_map):
    errors = []
    for field_name in SoftwareImportPreset.REQUIRED_FIELDS:
        if not column_map.get(field_name):
            errors.append(f"Manca la colonna per «{FIELD_LABELS[field_name]}».")
    for field_name, header in column_map.items():
        if field_name not in FIELD_LABELS:
            errors.append(f"Campo sconosciuto: {field_name}.")
        elif header and header not in headers:
            errors.append(f"La colonna «{header}» non è nel file.")
    return errors


def _cell(row, headers, header):
    if not header:
        return ""
    try:
        value = row[headers.index(header)]
    except (ValueError, IndexError):
        return ""
    return value


def _parse_date(value, date_format=""):
    if isinstance(value, datetime):
        return timezone.make_aware(value) if timezone.is_naive(value) else value
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in ((date_format,) if date_format else ()) + DATE_FORMATS:
        try:
            return timezone.make_aware(datetime.strptime(text, fmt))
        except ValueError:
            continue
    return False  # presente ma non leggibile


def build_preview(table: ParsedTable, column_map, date_format="") -> RowPreview:
    preview = RowPreview(total=len(table.rows))
    seen = set()
    for index, row in enumerate(table.rows, start=1):
        host = str(_cell(row, table.headers, column_map.get("hostname")) or "").strip()
        product = str(_cell(row, table.headers, column_map.get("product")) or "").strip()
        version_cell = _cell(row, table.headers, column_map.get("version"))
        if isinstance(version_cell, float):
            # Excel ha trasformato la versione in numero («2.10» → 2.1): non è affidabile.
            preview.skipped.append((index, "versione salvata come numero nel foglio: formatta la colonna come testo"))
            continue
        version = str(version_cell or "").strip()
        vendor = str(_cell(row, table.headers, column_map.get("vendor")) or "").strip()
        if not host:
            preview.skipped.append((index, "hostname vuoto"))
            continue
        if not product:
            preview.skipped.append((index, "prodotto vuoto"))
            continue
        if not version:
            preview.skipped.append((index, "versione vuota (non confrontabile con i range CVE)"))
            continue
        detected = _parse_date(_cell(row, table.headers, column_map.get("detected_at")), date_format)
        if detected is False:
            preview.skipped.append((index, "data di rilevamento non leggibile"))
            continue
        key = (host.casefold(), vendor.casefold(), product.casefold(), version.casefold())
        if key in seen:
            preview.skipped.append((index, "riga duplicata nel file"))
            continue
        seen.add(key)
        preview.valid.append({"host": host[:255], "vendor": vendor[:255], "product": product[:255], "version": version[:120], "detected_at": detected})
    return preview


# --- Normalizzazione -------------------------------------------------------------------------

def _aliases(kind):
    return dict(SoftwareAlias.objects.filter(kind=kind).values_list("raw", "canonical"))


def normalize_vendor(raw, aliases=None):
    text = re.sub(r"\s+", " ", str(raw or "").strip().lower())
    aliases = _aliases(SoftwareAlias.KIND_VENDOR) if aliases is None else aliases
    if text in aliases:
        return aliases[text]
    # Una sola forma societaria finale: «Demo Corp Inc.» → «demo corp», non «demo».
    text = _VENDOR_SUFFIXES.sub("", text).strip(" ,.")
    return aliases.get(text, text)[:120]


def normalize_product(raw, aliases=None, version=""):
    text = _ARCH.sub(" ", str(raw or "").lower())
    text = re.sub(r"\s+", " ", text).strip()
    aliases = _aliases(SoftwareAlias.KIND_PRODUCT) if aliases is None else aliases
    if text in aliases:
        return aliases[text]
    # Gli export mettono spesso la versione nel nome («7-Zip 23.01»): via i token finali che
    # sono la versione della riga o hanno un separatore («23.01», «1.2.3»). «Windows 10»,
    # «Office 365», «Python 3» restano interi: il numero fa parte del nome del prodotto.
    version = str(version or "").strip().lower()
    tokens = text.split(" ")
    while len(tokens) > 1 and (tokens[-1] == version or (_VERSIONISH.match(tokens[-1]) and re.search(r"[._-]", tokens[-1]))):
        tokens.pop()
    text = " ".join(tokens).strip(" -")
    return aliases.get(text, text)[:160]


def normalize_host(raw):
    host = str(raw or "").strip().lower()
    return host.split(".")[0] if host and not re.match(r"^\d+\.\d+\.\d+\.\d+$", host) else host


def installation_key(source_kind, host, vendor, product, version):
    seed = "|".join([source_kind, host, vendor, product, str(version).strip().lower()])
    return hashlib.sha256(seed.encode()).hexdigest()


# --- Import ----------------------------------------------------------------------------------

def _link_assets(host_norm):
    sec = SecurityAsset.objects.filter(hostname__iexact=host_norm).select_related("hub_asset").first()
    if sec is None:
        sec = SecurityAsset.objects.filter(hostname__istartswith=f"{host_norm}.").select_related("hub_asset").first()
    hub = sec.hub_asset if sec else None
    if hub is None:
        from assets.models import Asset

        hub = Asset.objects.filter(name__iexact=host_norm).first()
    return sec, hub


CHUNK = 500  # SQL Server: massimo 2100 parametri per query


def _chunks(items, size=CHUNK):
    items = list(items)
    for start in range(0, len(items), size):
        yield items[start:start + size]


def run_import(inventory_import: SoftwareInventoryImport, preview: RowPreview):
    """Upsert idempotente delle righe valide; marca «non più rilevate» quelle sparite per gli host del file."""
    from security.services.versioning import normalize as normalize_version

    now = timezone.now()
    vendor_aliases, product_aliases = _aliases(SoftwareAlias.KIND_VENDOR), _aliases(SoftwareAlias.KIND_PRODUCT)
    source_kind = inventory_import.source_kind
    rows = {}
    for row in preview.valid:
        host = normalize_host(row["host"])
        vendor = normalize_vendor(row["vendor"], vendor_aliases)
        product = normalize_product(row["product"], product_aliases, row["version"])
        key = installation_key(source_kind, host, vendor, product, row["version"])
        rows.setdefault(key, {**row, "host_norm": host, "vendor_norm": vendor, "product_norm": product})
    hosts = {data["host_norm"] for data in rows.values()}
    links = {host: _link_assets(host) for host in hosts}
    with transaction.atomic():
        existing = {}
        for chunk in _chunks(rows):
            existing.update(SoftwareInstallation.objects.filter(dedup_key__in=chunk).values_list("dedup_key", "pk"))
        new_objects = []
        for key, data in rows.items():
            if key in existing:
                continue
            sec, hub = links[data["host_norm"]]
            new_objects.append(SoftwareInstallation(
                dedup_key=key, source_kind=source_kind, host=data["host_norm"], host_display=data["host"],
                vendor_raw=data["vendor"], product_raw=data["product"], vendor=data["vendor_norm"], product=data["product_norm"],
                version_raw=data["version"], version_norm=normalize_version(data["version"]), first_import=inventory_import,
                last_import=inventory_import, security_asset=sec, hub_asset=hub, detected_at=data["detected_at"],
                seen_at=now, last_seen_at=now, last_inventory_date=inventory_import.inventory_date, still_detected=True,
            ))
        SoftwareInstallation.objects.bulk_create(new_objects, batch_size=60)
        for chunk in _chunks(existing.values()):
            SoftwareInstallation.objects.filter(pk__in=chunk).update(
                last_import=inventory_import, last_seen_at=now, last_inventory_date=inventory_import.inventory_date, still_detected=True,
            )
        for host, (sec, hub) in links.items():
            if sec is not None:
                SoftwareInstallation.objects.filter(host=host, security_asset__isnull=True).update(security_asset=sec)
            if hub is not None:
                SoftwareInstallation.objects.filter(host=host, hub_asset__isnull=True).update(hub_asset=hub)
        missing = []
        for chunk in _chunks(hosts):
            for pk, key in SoftwareInstallation.objects.filter(source_kind=source_kind, host__in=chunk, still_detected=True).values_list("pk", "dedup_key"):
                if key not in rows:
                    missing.append(pk)
        for chunk in _chunks(missing):
            SoftwareInstallation.objects.filter(pk__in=chunk).update(still_detected=False)
        inventory_import.status = SoftwareInventoryImport.STATUS_IMPORTED
        inventory_import.rows_total = preview.total
        inventory_import.rows_imported = len(preview.valid)
        inventory_import.rows_skipped = len(preview.skipped)
        inventory_import.rows_new = len(new_objects)
        inventory_import.rows_marked_missing = len(missing)
        inventory_import.hosts_count = len(hosts)
        inventory_import.skipped_details = [{"row": n, "reason": r} for n, r in preview.skipped[:MAX_SKIPPED_DETAILS]]
        inventory_import.imported_at = now
        inventory_import.save()
    return inventory_import


def cleanup_stale_previews(days=2):
    """Elimina i file di import mai confermati (restano solo cifrati e solo per poco)."""
    cutoff = timezone.now() - timedelta(days=days)
    removed = 0
    stale = SoftwareInventoryImport.objects.filter(status=SoftwareInventoryImport.STATUS_PREVIEW, created_at__lt=cutoff)
    for item in stale:
        discard_stored(item)
        item.status = SoftwareInventoryImport.STATUS_FAILED
        item.save(update_fields=["status"])
        removed += 1
    return removed


def file_sha256(data: bytes):
    return hashlib.sha256(data).hexdigest()


def storage():
    from core.private_attachments import PrivateAttachmentStorage

    return PrivateAttachmentStorage(download_hint="Inventario software SOC: nessun download.")


def store_upload(data: bytes, sha: str, extension: str):
    from django.core.files.base import ContentFile

    return storage().save(f"{STORAGE_DIR}/{sha[:16]}{extension}", ContentFile(data))


def load_stored(inventory_import):
    if not inventory_import.stored_name:
        raise InventoryFileError("Il file non è più disponibile: ricaricalo.")
    with storage().open(inventory_import.stored_name, "rb") as handle:
        return handle.read()


def discard_stored(inventory_import):
    if not inventory_import.stored_name:
        return
    try:
        storage().delete(inventory_import.stored_name)
    except OSError:
        logger.warning("File inventario %s non eliminato", inventory_import.stored_name)
    inventory_import.stored_name = ""
    inventory_import.save(update_fields=["stored_name"])
