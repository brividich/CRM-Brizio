"""Eccezioni sulla singola persona (core.permessi.eccezioni). Dati sintetici."""
from __future__ import annotations

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import TestCase

from core.models import (
    AccessGroup,
    AuditLog,
    EccezionePermessoPersona,
    PermissionDefinition,
    RuoloPermessoModulo,
    UserPermissionGrant,
)
from core.permessi.eccezioni import NOTA_ECCEZIONE, sincronizza_pagine_persona
from core.permessi.motore import FONTE_ECCEZIONE, decisione, utenti_visibili
from core.permessi.preset import applica_preset
from core.tests_permessi import PersoneMixin

User = get_user_model()


class EccezioniMotoreTests(PersoneMixin, TestCase):
    def setUp(self):
        cache.clear()
        self.addCleanup(cache.clear)  # l'interruttore del sistema resta in cache oltre il rollback
        applica_preset()

    def eccezione(self, legacy, modulo, azione, limite=""):
        EccezionePermessoPersona.objects.create(
            legacy_user_id=legacy.id, modulo=modulo, azione=azione, ambito_massimo=limite, motivo="prova"
        )

    def test_negata_alla_persona_non_ai_colleghi_dello_stesso_ruolo(self):
        user, legacy = self.persona("Con eccezione", ruoli=("ruolo-responsabile",))
        collega, _ = self.persona("Collega", ruoli=("ruolo-responsabile",))
        self.eccezione(legacy, "diario_preposto", "vedere")
        esito = decisione(self.ricarica(user), "diario_preposto", "vedere")
        self.assertFalse(esito.consentito)
        self.assertEqual(esito.fonte, FONTE_ECCEZIONE)
        self.assertIn("prova", esito.motivo)
        self.assertTrue(decisione(collega, "diario_preposto", "vedere").consentito)
        # Le altre azioni dello stesso modulo restano come le dà il ruolo.
        user = self.ricarica(user)
        for azione in ("creare", "modificare", "esportare", "approvare"):
            self.assertEqual(decisione(user, "diario_preposto", azione), decisione(collega, "diario_preposto", azione))

    def test_limite_restringe_la_visibilita(self):
        user, legacy = self.persona("Limitato", ruoli=("ruolo-responsabile",))
        self.assertEqual(decisione(user, "diario_preposto", "vedere").ambito, "team")
        self.eccezione(legacy, "diario_preposto", "vedere", "propri")
        user = self.ricarica(user)
        esito = decisione(user, "diario_preposto", "vedere")
        self.assertEqual((esito.consentito, esito.ambito, esito.fonte), (True, "propri", FONTE_ECCEZIONE))
        self.assertEqual(utenti_visibili(user, "diario_preposto", "vedere"), frozenset({legacy.id}))

    def test_limite_piu_largo_del_ruolo_non_aggiunge_niente(self):
        user, legacy = self.persona("Base")
        self.eccezione(legacy, "assenze", "vedere", "team")
        esito = decisione(self.ricarica(user), "assenze", "vedere")
        self.assertEqual((esito.ambito, esito.fonte), ("propri", "ruolo"))

    def test_eccezione_non_concede_quello_che_il_ruolo_non_da(self):
        user, legacy = self.persona("Base")
        self.eccezione(legacy, "assenze", "approvare", "propri")
        self.assertFalse(decisione(self.ricarica(user), "assenze", "approvare").consentito)

    def test_superuser_non_toccato(self):
        user, legacy = self.persona("Super", superuser=True)
        self.eccezione(legacy, "assenze", "vedere")
        self.assertTrue(decisione(self.ricarica(user), "assenze", "vedere").consentito)

    def test_vale_anche_sul_ruolo_vede_tutto(self):
        user, legacy = self.persona("Direzione", ruoli=("ruolo-direzione",))
        if not AccessGroup.objects.get(code="ruolo-direzione").vede_tutto:
            self.skipTest("preset Direzione senza «vede tutto»")
        self.eccezione(legacy, "anagrafica", "vedere")
        self.assertFalse(decisione(self.ricarica(user), "anagrafica", "vedere").consentito)


class EccezioniPagineTests(PersoneMixin, TestCase):
    def setUp(self):
        cache.clear()
        self.addCleanup(cache.clear)  # l'interruttore del sistema resta in cache oltre il rollback
        from core.test_acl_v2 import _ensure_legacy_acl_tables

        _ensure_legacy_acl_tables()
        applica_preset()
        PermissionDefinition.objects.create(code="dpi.detail.view", label="Dettaglio DPI", module="dpi")
        base = AccessGroup.objects.get(code="ruolo-dipendente")
        # Solo «vedere» apre la pagina di lettura: la prova non dipende dal preset di «esportare».
        RuoloPermessoModulo.objects.filter(ruolo=base, modulo="dpi", azione="esportare").delete()
        from core.permessi.sync import sincronizza_tutti

        sincronizza_tutti()

    def test_pagina_chiusa_solo_alla_persona_e_riaperta_togliendo_eccezione(self):
        from core.acl_resolver import resolve_permission_decision

        self.attiva_sistema()  # il ruolo base vale per tutti solo a sistema acceso
        _user, legacy = self.persona("Senza DPI")
        _c, collega = self.persona("Collega")
        EccezionePermessoPersona.objects.create(
            legacy_user_id=legacy.id, modulo="dpi", azione="vedere", motivo="prova"
        )
        self.assertEqual(sincronizza_pagine_persona(legacy.id), (1, 0))
        grant = UserPermissionGrant.objects.get(legacy_user_id=legacy.id, permission_id="dpi.detail.view")
        self.assertEqual((grant.enabled, grant.note), (False, NOTA_ECCEZIONE))
        self.assertFalse(resolve_permission_decision(permission_code="dpi.detail.view", legacy_user=legacy).allowed)
        self.assertTrue(resolve_permission_decision(permission_code="dpi.detail.view", legacy_user=collega).allowed)

        EccezionePermessoPersona.objects.filter(legacy_user_id=legacy.id).delete()
        self.assertEqual(sincronizza_pagine_persona(legacy.id), (0, 1))
        self.assertFalse(UserPermissionGrant.objects.filter(legacy_user_id=legacy.id).exists())

    def test_pagina_che_serve_a_un_altra_azione_concessa_resta_aperta(self):
        _user, legacy = self.persona("Esporta ancora")
        base = AccessGroup.objects.get(code="ruolo-dipendente")
        RuoloPermessoModulo.objects.create(ruolo=base, modulo="dpi", azione="esportare", ambito="propri")
        EccezionePermessoPersona.objects.create(
            legacy_user_id=legacy.id, modulo="dpi", azione="vedere", motivo="prova"
        )
        self.assertEqual(sincronizza_pagine_persona(legacy.id), (0, 0))

    def test_override_messo_a_mano_non_viene_toccato(self):
        _user, legacy = self.persona("A mano")
        UserPermissionGrant.objects.create(
            legacy_user_id=legacy.id, permission_id="dpi.detail.view", enabled=True, note="a mano"
        )
        EccezionePermessoPersona.objects.create(
            legacy_user_id=legacy.id, modulo="dpi", azione="vedere", motivo="prova"
        )
        sincronizza_pagine_persona(legacy.id)
        self.assertTrue(UserPermissionGrant.objects.get(legacy_user_id=legacy.id).enabled)


class EccezioniPaginaRuoliTests(PersoneMixin, TestCase):
    def setUp(self):
        cache.clear()
        self.addCleanup(cache.clear)  # l'interruttore del sistema resta in cache oltre il rollback
        applica_preset()
        self.admin = User.objects.create_user(username="admin.eccezioni@example.test", password="x", is_superuser=True)
        self.client.force_login(self.admin)

    def url(self, legacy):
        return f"/admin-portale/ruoli-portale/persona/{legacy.id}/eccezione/"

    def test_aggiunta_e_rimozione_registrate(self):
        user, legacy = self.persona("Persona")
        risposta = self.client.post(self.url(legacy), {"voce": "assenze::vedere", "limite": "", "motivo": "riservatezza"})
        self.assertEqual(risposta.status_code, 302)
        eccezione = EccezionePermessoPersona.objects.get(legacy_user_id=legacy.id)
        self.assertEqual((eccezione.modulo, eccezione.azione, eccezione.ambito_massimo), ("assenze", "vedere", ""))
        voce = AuditLog.objects.filter(modulo="permessi", azione="permessi_eccezione_persona").latest("id")
        self.assertEqual((voce.dettaglio["dopo"], voce.dettaglio["motivo"]), ("negato", "riservatezza"))
        self.assertFalse(decisione(self.ricarica(user), "assenze", "vedere").consentito)

        pagina = self.client.get(f"/admin-portale/ruoli-portale/?tab=persone&persona={legacy.id}")
        self.assertContains(pagina, "Eccezioni sulla persona")
        self.assertContains(pagina, "riservatezza")
        anteprima = self.client.get(f"/admin-portale/ruoli-portale/?tab=anteprima&persona={legacy.id}")
        self.assertContains(anteprima, "eccezione sulla persona")
        registro = self.client.get("/admin-portale/ruoli-portale/?tab=registro")
        self.assertContains(registro, "Eccezione sulla persona")

        self.client.post(self.url(legacy), {"togli": eccezione.id})
        self.assertFalse(EccezionePermessoPersona.objects.filter(legacy_user_id=legacy.id).exists())
        self.assertTrue(decisione(self.ricarica(user), "assenze", "vedere").consentito)

    def test_motivo_obbligatorio_e_voce_valida(self):
        _user, legacy = self.persona("Persona")
        self.client.post(self.url(legacy), {"voce": "assenze::vedere", "limite": "", "motivo": " "})
        self.client.post(self.url(legacy), {"voce": "inesistente::vedere", "limite": "", "motivo": "x"})
        self.client.post(self.url(legacy), {"voce": "assenze::vedere", "limite": "tutti", "motivo": "x"})
        self.assertFalse(EccezionePermessoPersona.objects.exists())

    def test_non_amministratore_respinto(self):
        user, legacy = self.persona("Curioso")
        self.client.force_login(user)
        risposta = self.client.post(self.url(legacy), {"voce": "assenze::vedere", "limite": "", "motivo": "x"})
        self.assertIn(risposta.status_code, (302, 403))
        self.assertFalse(EccezionePermessoPersona.objects.exists())

    def test_riesame_ha_il_foglio_delle_eccezioni(self):
        from io import BytesIO

        from openpyxl import load_workbook

        _user, legacy = self.persona("Persona")
        EccezionePermessoPersona.objects.create(
            legacy_user_id=legacy.id, modulo="assenze", azione="vedere", motivo="riservatezza"
        )
        risposta = self.client.get("/admin-portale/ruoli-portale/riesame.xlsx")
        wb = load_workbook(BytesIO(risposta.content))
        ws = wb["Eccezioni sulla persona"]
        self.assertEqual(ws.cell(row=2, column=1).value, "Persona")
        self.assertEqual(ws.cell(row=2, column=5).value, "riservatezza")
