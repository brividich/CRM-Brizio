"""Account Microsoft 365 confrontati con l'anagrafica (core.account_m365). Graph sempre simulato."""
from __future__ import annotations

from datetime import date
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse

from core import account_m365 as svc
from core import m365_graph

OGGI = date(2026, 10, 19)


def _persona(pid, nome, chiavi, cessato=None):
    return {"id": pid, "nome": nome, "cessato_il": cessato, "chiavi": set(chiavi)}


def _utente(uid, upn, *, abilitato=True, licenze=1, ultimo="2026-10-15T08:00:00Z", mail=None):
    return {
        "id": uid, "displayName": upn.split("@")[0], "userPrincipalName": upn, "mail": mail,
        "accountEnabled": abilitato, "userType": "Member", "createdDateTime": "2024-01-10T08:00:00Z",
        "assignedLicenses": [{"skuId": "x"}] * licenze,
        "signInActivity": {"lastSignInDateTime": ultimo} if ultimo else {},
    }


class RiconciliazioneTests(SimpleTestCase):
    persone = [
        _persona(1, "ROSSI MARIO", {"mario.rossi@example.com"}),
        _persona(2, "BIANCHI ANNA", {"login:a.bianchi"}, cessato=date(2026, 9, 30)),
        _persona(3, "VERDI LUCA", {"l.verdi@example.com"}, cessato=date(2026, 12, 31)),
    ]

    def _esiti(self, utenti):
        return {r.upn: r for r in svc.riconcilia(utenti, self.persone, con_accessi=True, oggi=OGGI)}

    def test_collegamento_per_email_e_per_login(self):
        r = self._esiti([_utente("1", "mario.rossi@example.com"), _utente("2", "a.bianchi@example.com")])
        self.assertEqual(r["mario.rossi@example.com"].persona, "ROSSI MARIO")
        self.assertEqual(r["a.bianchi@example.com"].persona, "BIANCHI ANNA")

    def test_ex_dipendente_con_account_attivo(self):
        r = self._esiti([_utente("2", "a.bianchi@example.com"), _utente("3", "l.verdi@example.com")])
        self.assertEqual(r["a.bianchi@example.com"].esito, "cessato_attivo")
        self.assertEqual(r["a.bianchi@example.com"].cessato_il, "2026-09-30")
        # Cessazione futura: ancora in forza.
        self.assertEqual(r["l.verdi@example.com"].esito, "ok")

    def test_account_disattivato_di_ex_dipendente_in_regola(self):
        r = self._esiti([_utente("2", "a.bianchi@example.com", abilitato=False)])
        self.assertEqual(r["a.bianchi@example.com"].esito, "ok")

    def test_inattivo_e_mai_entrato(self):
        r = self._esiti([_utente("1", "mario.rossi@example.com", ultimo="2026-05-01T08:00:00Z"),
                         _utente("9", "mai@example.com", ultimo=None)])
        self.assertEqual(r["mario.rossi@example.com"].esito, "inattivo")
        self.assertEqual(r["mai@example.com"].esito, "inattivo")
        self.assertIn("mai entrato", r["mai@example.com"].note)

    def test_non_collegato_solo_se_con_licenza(self):
        r = self._esiti([_utente("8", "casella.condivisa@example.com", licenze=0),
                         _utente("7", "esterno@example.com")])
        self.assertEqual(r["casella.condivisa@example.com"].esito, "ok")
        self.assertEqual(r["esterno@example.com"].esito, "non_collegato")

    def test_ordine_dal_piu_grave(self):
        righe = svc.riconcilia(
            [_utente("1", "mario.rossi@example.com"), _utente("2", "a.bianchi@example.com")],
            self.persone, con_accessi=True, oggi=OGGI,
        )
        self.assertEqual(righe[0].esito, "cessato_attivo")

    def test_senza_ultimo_accesso_niente_inattivi(self):
        righe = svc.riconcilia([_utente("1", "mario.rossi@example.com", ultimo=None)], self.persone,
                               con_accessi=False, oggi=OGGI)
        self.assertEqual(righe[0].esito, "ok")
        self.assertIsNone(righe[0].giorni_senza_accesso)


class LetturaTests(SimpleTestCase):
    def test_senza_auditlog_rilegge_senza_ultimo_accesso(self):
        chiamate = []

        def finto(url, *, permesso, params=None, max_pages=50):
            chiamate.append(params["$select"])
            if "signInActivity" in params["$select"]:
                raise m365_graph.GraphError("403", 403)
            return [_utente("1", "a@example.com"), {**_utente("2", "ospite@example.com"), "userType": "Guest"}]

        with patch("core.m365_graph.get_all", side_effect=finto):
            utenti, con_accessi = svc.leggi_utenti()
        self.assertFalse(con_accessi)
        self.assertEqual([u["id"] for u in utenti], ["1"])  # gli ospiti restano fuori
        self.assertEqual(len(chiamate), 2)


class CredenzialiTests(SimpleTestCase):
    @override_settings(GRAPH_TENANT_ID="", GRAPH_CLIENT_ID="", GRAPH_CLIENT_SECRET="")
    def test_credenziali_dal_dotenv(self):
        env = {"GRAPH_TENANT_ID": "t-sintetico", "GRAPH_CLIENT_ID": "c-sintetico", "GRAPH_CLIENT_SECRET": "s-sintetico"}
        with patch.dict("os.environ", env):
            self.assertEqual(m365_graph.credenziali()["client_id"], "c-sintetico")

    def test_url_non_graph_rifiutato(self):
        with self.assertRaises(m365_graph.GraphError):
            m365_graph.get("https://esempio.it/x", permesso="x", tok="t")


@override_settings(LEGACY_AUTH_ENABLED=False)
class SchedaRuoliTests(TestCase):
    def setUp(self):
        from core.models import UserOnboarding

        cache.delete(svc.CACHE_KEY)
        self.admin = get_user_model().objects.create_superuser(username="adm", email="a@a.it", password="x")
        UserOnboarding.objects.create(user=self.admin, completed=True)
        self.client.force_login(self.admin)
        self.url = reverse("admin_portale:ruoli_portale") + "?tab=m365"

    def _dati(self):
        righe = svc.riconcilia([_utente("2", "a.bianchi@example.com")], RiconciliazioneTests.persone,
                               con_accessi=True, oggi=OGGI)
        from dataclasses import asdict

        return {"letto_alle": "2026-10-19T08:00:00+00:00", "con_accessi": True, "inattivo_giorni": 90,
                "conteggi": {"cessato_attivo": 1, "inattivo": 0, "non_collegato": 0, "ok": 0},
                "righe": [asdict(r) for r in righe]}

    def test_scheda_mostra_gli_ex_dipendenti(self):
        with patch("core.m365_graph.configurato", return_value=True), \
             patch("core.account_m365.esegui", return_value=self._dati()):
            resp = self.client.get(self.url)
        self.assertContains(resp, "Account Microsoft 365")
        self.assertContains(resp, "BIANCHI ANNA")
        self.assertContains(resp, "cessato il 30/09/2026")

    def test_scheda_senza_credenziali(self):
        with patch("core.m365_graph.configurato", return_value=False):
            resp = self.client.get(self.url)
        self.assertContains(resp, "Credenziali Microsoft Graph non configurate")

    def test_permesso_mancante_spiegato(self):
        with patch("core.m365_graph.configurato", return_value=True), \
             patch("core.account_m365.esegui", side_effect=m365_graph.GraphError("manca il permesso «User.Read.All»", 403)):
            resp = self.client.get(self.url)
        self.assertContains(resp, "User.Read.All")

    def test_riesame_ha_il_foglio_account(self):
        from openpyxl import load_workbook
        from io import BytesIO

        from admin_portale.riesame_accessi import genera_riesame_xlsx

        with patch("core.m365_graph.configurato", return_value=True), \
             patch("core.account_m365.esegui", return_value=self._dati()):
            wb = load_workbook(BytesIO(genera_riesame_xlsx(preparato_da="test", sistema_attivo=False)))
        ws = wb["Account Microsoft 365"]
        self.assertEqual(ws.cell(2, 2).value, "BIANCHI ANNA")
