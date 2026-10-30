"""Separazione dei compiti (SGI) e riesame degli accessi: dati sintetici.

- assenze e anomalie: con il sistema dei ruoli acceso nessuno decide sulla propria richiesta;
- CAPA: il responsabile non verifica la propria azione;
- riesame degli accessi: il file Excel contiene persone, provenienza e matrice.
"""
from __future__ import annotations

from io import BytesIO
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.db import connection
from django.test import RequestFactory, TestCase
from django.urls import reverse

from core.models import ActionItem, AccessGroup, AuditLog
from core.permessi.preset import applica_preset
from core.tests_permessi import PersoneMixin, ensure_anagrafica_table

User = get_user_model()


class AssenzePropriaRichiestaTests(PersoneMixin, TestCase):
    def setUp(self):
        cache.clear()
        applica_preset()
        self.user, self.legacy = self.persona("Capo con ferie", ruoli=("ruolo-responsabile",))
        self.request = RequestFactory().post("/assenze/")
        self.request.user = self.user
        self.request.legacy_user = self.legacy

    def test_con_i_ruoli_accesi_la_propria_assenza_non_si_decide(self):
        from assenze.views import _e_propria_assenza

        self.attiva_sistema(True)
        self.assertTrue(_e_propria_assenza(self.request, {"utente_id": self.legacy.id}))
        self.assertTrue(_e_propria_assenza(self.request, {"email_esterna": self.legacy.email.upper()}))
        self.assertFalse(_e_propria_assenza(self.request, {"utente_id": self.legacy.id + 999,
                                                            "email_esterna": "altro@example.test"}))

    def test_a_sistema_spento_vale_la_regola_storica(self):
        from assenze.views import _e_propria_assenza

        self.attiva_sistema(False)
        self.assertFalse(_e_propria_assenza(self.request, {"utente_id": self.legacy.id}))


class AnomaliePropriaTests(PersoneMixin, TestCase):
    COLS = {"id", "note_capocommessa", "segnalare_cliente", "aprire_rdc", "chiudere", "created_by_user_id"}

    def setUp(self):
        cache.clear()
        applica_preset()
        with connection.cursor() as cur:
            cur.execute(
                "CREATE TABLE IF NOT EXISTS anomalie (id INTEGER PRIMARY KEY, note_capocommessa TEXT, "
                "segnalare_cliente INTEGER, aprire_rdc INTEGER, chiudere INTEGER, created_by_user_id INTEGER)"
            )
            cur.execute("DELETE FROM anomalie")
        _u, self.autore = self.persona("Autore anomalia")
        _u, self.cc = self.persona("Capocommessa")
        with connection.cursor() as cur:
            cur.execute("INSERT INTO anomalie (id, note_capocommessa, segnalare_cliente, aprire_rdc, chiudere, "
                        "created_by_user_id) VALUES (1, '', 0, 0, 0, %s)", [self.autore.id])

    def test_l_autore_non_decide_ma_corregge(self):
        from anomalie.views import _decisione_su_propria_anomalia

        self.attiva_sistema(True)
        decisione = {"note_capocommessa": "Rilavorare", "segnalare_cliente": 0, "aprire_rdc": 0, "chiudere": 1}
        self.assertTrue(_decisione_su_propria_anomalia(1, self.COLS, decisione, self.autore))
        self.assertEqual(_decisione_su_propria_anomalia(1, self.COLS, decisione, self.cc), "")
        invariata = {"note_capocommessa": "", "segnalare_cliente": 0, "aprire_rdc": 0, "chiudere": 0}
        self.assertEqual(_decisione_su_propria_anomalia(1, self.COLS, invariata, self.autore), "")

    def test_a_sistema_spento_nessun_blocco(self):
        from anomalie.views import _decisione_su_propria_anomalia

        self.attiva_sistema(False)
        decisione = {"note_capocommessa": "Rilavorare", "chiudere": 1}
        self.assertEqual(_decisione_su_propria_anomalia(1, self.COLS, decisione, self.autore), "")


class CapaResponsabileNonVerificaTests(TestCase):
    def test_il_responsabile_non_verifica_la_propria_azione(self):
        responsabile = User.objects.create_user(username="resp.capa@example.test", password="x")
        chi_chiude = User.objects.create_user(username="chiude.capa@example.test", password="x")
        azione = ActionItem.objects.create(
            titolo="Azione di prova", responsabile=responsabile, stato=ActionItem.STATO_CHIUSA,
            chiusa_da=chi_chiude, source_code="manuale",
        )
        self.client.force_login(responsabile)
        with mock.patch("core.views_capa._can_manage_capa", return_value=True):
            self.client.post(reverse("capa_verify", args=[azione.pk]), {"note_verifica": ""})
        azione.refresh_from_db()
        self.assertEqual(azione.stato, ActionItem.STATO_CHIUSA)


class RiesameAccessiTests(PersoneMixin, TestCase):
    def setUp(self):
        cache.clear()
        ensure_anagrafica_table()
        applica_preset()
        self.admin = User.objects.create_user(username="admin.riesame@example.test", password="x", is_superuser=True)
        self.client.force_login(self.admin)

    def test_excel_con_persone_provenienza_e_matrice(self):
        from openpyxl import load_workbook

        from anagrafica.models import DipendenteRuoloOperativo, RuoloOperativo

        _u, esterno = self.persona("RSPP ESTERNO PROVA", ruoli=("ruolo-sicurezza",))
        _u, interno = self.persona("CAPO PROVA")
        ruolo_org = RuoloOperativo.objects.create(nome="Capo reparto RIESAME")
        ruolo_org.profili_permessi.add(AccessGroup.objects.get(code="ruolo-responsabile"))
        with connection.cursor() as c:
            c.execute("INSERT OR IGNORE INTO anagrafica_dipendenti (id, nome, utente_id) VALUES (%s, %s, %s)",
                      [9700 + interno.id, interno.nome, interno.id])
        DipendenteRuoloOperativo.objects.create(legacy_anagrafica_id=9700 + interno.id, ruolo=ruolo_org)

        risposta = self.client.get(reverse("admin_portale:ruoli_portale_riesame"))
        self.assertEqual(risposta.status_code, 200)
        wb = load_workbook(BytesIO(risposta.content))
        self.assertEqual(wb.sheetnames, ["Riesame", "Persone", "Matrice", "Organigramma", "Eccezioni sulla persona",
                                         "Variazioni da organigramma", "Account Microsoft 365"])
        persone = [tuple(r) for r in wb["Persone"].iter_rows(min_row=2, values_only=True)]
        self.assertTrue(any(r[0] == "RSPP ESTERNO PROVA" and r[4] == "a mano" for r in persone))
        self.assertTrue(any(r[0] == "CAPO PROVA" and "Capo reparto RIESAME" in r[4] for r in persone))
        self.assertGreater(wb["Matrice"].max_row, 10)
        self.assertTrue(AuditLog.objects.filter(azione="permessi_riesame_esportato").exists())

    def test_solo_amministratori(self):
        user, _ = self.persona("Curioso")
        self.client.force_login(user)
        risposta = self.client.get(reverse("admin_portale:ruoli_portale_riesame"))
        self.assertIn(risposta.status_code, (302, 403))
