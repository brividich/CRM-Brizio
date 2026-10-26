"""Modulo Registro OFI / NC (MOD.174, 4.2) — viste.

Modulo dedicato montato a ``/ofi-registro/`` (namespace ``registro_ofi``): dashboard,
registro fedele al MOD.174, «Le mie voci», scheda della voce con avanzamento PDCA,
evidenze e cronologia. Il modello ``RegistroOFI`` vive in ``gestione_specifiche``
(continuità con MOD.133 e ``AzioneOFI``); le voci arrivano anche da audit interni,
rilievi degli enti e azioni correttive, e si legano al catalogo processi.

Permessi (``registro_ofi.acl_bootstrap``): ``PERM_OFI_VIEW`` apre registro e dashboard
completa; ``PERM_OFI_MIE`` (a tutti) apre le sole voci di cui si è responsabili;
``PERM_OFI_ADD`` inserisce e modifica.
"""
from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO

import hashlib

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Q
from django.http import FileResponse, Http404, HttpResponse, HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from core.acl_v2 import request_has_permission_code
from core.audit import log_action
from gestione_specifiche import registro_ofi as reg
from gestione_specifiche.forms import RegistroOFIAllegatoForm, RegistroOFIForm
from gestione_specifiche.models import RegistroOFI, RegistroOFIAllegato, RegistroOFIEvento

from . import collegamenti
from .acl_bootstrap import PERM_OFI_ADD, PERM_OFI_VIEW




def _puo_gestire(request, voce) -> bool:
    """Chi gestisce il registro, oppure il responsabile della voce."""
    if request_has_permission_code(request, PERM_OFI_ADD):
        return True
    return voce.responsabile_id is not None and voce.responsabile_id == request.user.pk


def _puo_vedere(request, voce) -> bool:
    """Chi consulta il registro, oppure il responsabile della voce."""
    if request_has_permission_code(request, PERM_OFI_VIEW):
        return True
    return voce.responsabile_id is not None and voce.responsabile_id == request.user.pk


def _ctx(request, **extra) -> dict:
    """Contesto comune a tutte le pagine del modulo (menu interno)."""
    return {"ofi_vede_registro": request_has_permission_code(request, PERM_OFI_VIEW),
            "can_add": request_has_permission_code(request, PERM_OFI_ADD), **extra}


def _aperte(qs):
    return qs.exclude(fase=RegistroOFI.FASE_CHIUSO).filter(data_chiusura__isnull=True)


_FILTRO_ORIGINE = {
    "audit_interno": Q(content_type__model="auditesito"),
    "audit_esterno": Q(content_type__model="auditesterno"),
    "mod133": Q(content_type__model="rigamod133") | Q(modulo_origine="gestione_specifiche"),
    "manuale": (Q(content_type__isnull=True) & ~Q(modulo_origine="gestione_specifiche")),
}


def _voci_filtrate(request):
    """Applica i filtri GET (modulo/fase/priorità/scaduti) e ritorna (voci, ctx_filtri)."""
    fase = (request.GET.get("fase") or "").strip().upper()
    priorita = (request.GET.get("priorita") or "").strip().upper()
    modulo = (request.GET.get("modulo") or "").strip()
    tipo = (request.GET.get("tipo") or "").strip().upper()
    q = (request.GET.get("q") or "").strip()
    origine = (request.GET.get("origine") or "").strip()
    processo = (request.GET.get("processo") or "").strip()
    solo_scaduti = request.GET.get("scaduti") == "1"
    solo_mie = request.GET.get("mie") == "1"
    solo_aperte = request.GET.get("aperte") == "1"
    in_scadenza = request.GET.get("in_scadenza") == "1"
    senza_responsabile = request.GET.get("senza_responsabile") == "1"

    qs = RegistroOFI.objects.all().select_related("content_type", "responsabile", "processo_sgi")
    if solo_mie:
        qs = qs.filter(responsabile=request.user)
    if solo_aperte or in_scadenza or senza_responsabile:
        qs = _aperte(qs)
    if in_scadenza:
        oggi = timezone.localdate()
        qs = qs.filter(data_richiesta__gte=oggi, data_richiesta__lte=oggi + timedelta(days=30))
    if senza_responsabile:
        qs = qs.filter(responsabile__isnull=True)
    if origine in _FILTRO_ORIGINE:
        qs = qs.filter(_FILTRO_ORIGINE[origine])
    processo_obj = None
    if processo == "0":
        qs = qs.filter(processo_sgi__isnull=True)
    elif processo.isdigit():
        qs = qs.filter(processo_sgi_id=int(processo))
        from sistema_gestione.models import Processo
        processo_obj = Processo.objects.filter(pk=int(processo)).first()
    if modulo:
        qs = qs.filter(modulo_origine=modulo)
    if tipo in dict(RegistroOFI.TIPO_CHOICES):
        qs = qs.filter(tipo=tipo)
    if q:
        filtro = (Q(processo__icontains=q) | Q(opportunita__icontains=q) | Q(plan__icontains=q)
                  | Q(ref__icontains=q) | Q(rif_norma__icontains=q) | Q(fonte__icontains=q)
                  | Q(owner_processo__icontains=q) | Q(proprietario__icontains=q))
        if q.isdigit():
            filtro |= Q(numero=int(q))
        qs = qs.filter(filtro)
    if fase in dict(RegistroOFI.FASE_CHOICES):
        qs = qs.filter(fase=fase)
    if priorita in dict(RegistroOFI.PRIORITA_CHOICES):
        qs = qs.filter(priorita=priorita)
    voci = list(qs)
    if solo_scaduti:
        voci = [v for v in voci if v.is_scaduto]
    return voci, {"fase": fase, "priorita": priorita, "modulo": modulo, "tipo": tipo, "q": q,
                  "solo_scaduti": solo_scaduti, "solo_mie": solo_mie, "solo_aperte": solo_aperte,
                  "origine": origine if origine in _FILTRO_ORIGINE else "", "processo": processo,
                  "processo_obj": processo_obj, "in_scadenza": in_scadenza,
                  "senza_responsabile": senza_responsabile}


# Intestazioni del MOD.174 Rev.0, nell'ordine del modulo (colonna «27001» = fonte del rilievo).
_MOD174_HEADER = ["REF", "DATA", "OFI/NC", "FONTE", "NORME", "REF NORMA", "PROCESSO", "OPPORTUNITY",
                  "PLAN", "Allegato/Link", "DO", "CHECK", "ACT", "DATA REQUIRED", "DATA CLOSED",
                  "OWNER", "P", "D", "C", "A", "TOT", "MODULO"]


def _export_xlsx(voci, counters):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "MOD.174"
    ws.append(["Costruzioni Novicrom — MOD.174 - SGI Registro OFI_NC Rev.0"])
    ws.append([f"P {counters['p']}", f"D {counters['d']}", f"C {counters['c']}", f"A {counters['a']}",
               f"A/P {counters['pct']}% ({'≥90%' if counters['over90'] else '<90%'})"])
    ws.append(_MOD174_HEADER)
    for cella in ws[3]:
        cella.font = Font(bold=True, color="FFFFFF")
        cella.fill = PatternFill("solid", fgColor="1F3A5F")
    ws["A1"].font = Font(bold=True, size=13)
    for v in voci:
        m = v.pdca_marks
        ws.append([
            v.numero, v.data_apertura, v.tipo, v.fonte, ", ".join(v.norme_sigle), v.rif_norma, v.processo,
            v.opportunita, v.plan, v.allegato_link, v.do, v.verifica, v.act, v.data_richiesta,
            v.data_chiusura, v.owner_processo or v.proprietario,
            *("X" if m[k] else "" for k in ("p", "d", "c", "a")), v.tot_letter, v.modulo_display,
        ])
    for riga in ws.iter_rows(min_row=4):
        for cella in riga:
            cella.alignment = Alignment(vertical="top", wrap_text=True)
            if hasattr(cella.value, "year"):
                cella.number_format = "DD.MM.YYYY"
    larghezze = [7, 11, 7, 12, 14, 12, 18, 40, 40, 20, 30, 30, 30, 12, 12, 18, 4, 4, 4, 4, 5, 18]
    for i, w in enumerate(larghezze, start=1):
        ws.column_dimensions[ws.cell(row=3, column=i).column_letter].width = w
    ws.freeze_panes = "A4"
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="MOD174_registro_OFI_{timezone.localdate():%Y%m%d}.xlsx"'
    return resp


@login_required
def lista(request):
    """Registro OFI/NC (MOD.174) — elenco fedele all'Excel, a tutta larghezza.

    Nessun filtro modulo = registro UNICO (tutti i moduli); ``?modulo=<chiave>``
    = registro del singolo modulo. Contatori cumulativi P/D/C/A + KPI ≥90% (A/P)
    come la riga 2 del MOD.174.
    """
    voci, f = _voci_filtrate(request)
    modulo = f["modulo"]
    # Contatori/KPI coerenti col filtro modulo (il "mio" registro vs unico).
    base = RegistroOFI.objects.filter(modulo_origine=modulo) if modulo else RegistroOFI.objects.all()
    counters = reg.conta_pdca_cumulativo(base)
    if request.GET.get("export") == "xlsx":
        filtri = {k: v for k, v in f.items() if k != "processo_obj"}
        log_action(request, "registro_ofi_export", "gestione_specifiche", {"voci": len(voci), "filtri": filtri})
        return _export_xlsx(voci, counters)

    from sistema_gestione.models import Processo

    return render(request, "registro_ofi/lista.html", _ctx(
        request,
        voci=voci,
        counters=counters,
        n_scaduti=sum(1 for v in base if v.is_scaduto),
        fase=f["fase"], priorita=f["priorita"], solo_scaduti=f["solo_scaduti"],
        solo_mie=f["solo_mie"], solo_aperte=f["solo_aperte"],
        in_scadenza=f["in_scadenza"], senza_responsabile=f["senza_responsabile"],
        n_mie=RegistroOFI.objects.filter(responsabile=request.user).exclude(fase=RegistroOFI.FASE_CHIUSO).count(),
        tipo=f["tipo"], q=f["q"], TIPO_CHOICES=RegistroOFI.TIPO_CHOICES,
        filtri_query=request.GET.urlencode(),
        modulo=modulo, modulo_label=reg.modulo_label(modulo) if modulo else "",
        moduli_opts=reg.moduli_presenti(),
        origine=f["origine"], ORIGINI=collegamenti.ORIGINI,
        processo=f["processo"], processo_obj=f["processo_obj"],
        processi_opts=Processo.objects.filter(voci_ofi__isnull=False).distinct().order_by("codice"),
        FASE_CHOICES=RegistroOFI.FASE_CHOICES,
        PRIORITA_CHOICES=RegistroOFI.PRIORITA_CHOICES,
    ))


def _iniziali_da_processo(processo) -> dict:
    """Valori proposti per una voce aperta dalla scheda di un processo del catalogo."""
    iniziali = {"processo_sgi": processo.pk, "processo": str(processo)[:200], "fonte": "INTERNO"}
    if processo.responsabile_id and processo.responsabile.is_active:
        iniziali["responsabile"] = processo.responsabile_id
    norme = {"punti_9100": "norma_en9100", "punti_45001": "norma_iso45001", "punti_14001": "norma_iso14001",
             "punti_27001": "norma_iso27001", "punti_pdr125": "norma_pdr125"}
    for campo_processo, campo_voce in norme.items():
        if (getattr(processo, campo_processo, "") or "").strip():
            iniziali[campo_voce] = True
    return iniziali


@login_required
def nuovo(request):
    """Inserimento manuale di una voce OFI (ACL ``registroofi.add``).

    ``numero`` è assegnato server-side da ``prossimo_numero()`` (spazio numeri
    condiviso con l'OFI legacy). Voce del registro unico (``modulo_origine`` vuoto).
    Con ``?processo=<id>`` (dalla scheda del processo) propone processo, responsabile
    e norme del processo.
    """
    processo = None
    if (request.GET.get("processo") or "").isdigit():
        from sistema_gestione.models import Processo
        processo = Processo.objects.select_related("responsabile").filter(pk=int(request.GET["processo"])).first()
    if request.method == "POST":
        form = RegistroOFIForm(request.POST)
        if form.is_valid():
            voce = form.save(commit=False)
            from core.numbering_retry import save_with_next_number

            def assign():
                voce.numero = reg.prossimo_numero()

            save_with_next_number(assign, voce.save)  # audit S8
            reg.registra_evento(voce, "CREATA", utente=request.user, a=voce.get_fase_display(),
                                nota=f"Dalla scheda del processo {voce.processo_sgi}" if processo else "")
            log_action(request, "registro_ofi_creata", "gestione_specifiche", {"numero": voce.numero}, oggetto=voce)
            messages.success(request, f"Voce {voce.numero} creata.")
            return redirect("registro_ofi:dettaglio", pk=voce.pk)
    else:
        form = RegistroOFIForm(initial=_iniziali_da_processo(processo) if processo else None)
    return render(request, "registro_ofi/form.html", _ctx(
        request, form=form, voce=None, numero_preview=reg.prossimo_numero(),
        moduli_esistenti=reg.moduli_presenti(), processo_origine=processo,
    ))


@login_required
@transaction.atomic
def modifica(request, pk: int):
    """Modifica di una voce OFI (ACL ``registroofi.add``)."""
    voce = get_object_or_404(RegistroOFI.objects.select_for_update(), pk=pk)
    if request.method == "POST":
        prima = reg.fotografia(voce)
        form = RegistroOFIForm(request.POST, instance=voce)
        if form.is_valid():
            form.save()
            reg.registra_modifiche(prima, voce, request.user)
            messages.success(request, f"Voce {voce.numero} aggiornata.")
            return redirect("registro_ofi:dettaglio", pk=voce.pk)
    else:
        form = RegistroOFIForm(instance=voce)
    return render(request, "registro_ofi/form.html", _ctx(
        request, form=form, voce=voce, numero_preview=voce.numero,
        moduli_esistenti=reg.moduli_presenti(),
    ))


@login_required
def guida(request):
    """Guida d'uso del Registro OFI: stessa fonte della guida audit, solo i capitoli del registro."""
    from sistema_gestione import guida as g
    return g.risposta_guida(request, g.GUIDA_OFI)


@login_required
def dettaglio(request, pk: int):
    """Scheda della voce: PDCA, collegamenti (origine, processo, CAR), evidenze e cronologia."""
    voce = get_object_or_404(
        RegistroOFI.objects.select_related("content_type", "responsabile", "processo_sgi")
        .prefetch_related("azioni", "allegati__caricato_da", "eventi__utente"), pk=pk)
    if not _puo_vedere(request, voce):
        return HttpResponseForbidden("Questa voce è visibile a chi consulta il registro e al suo responsabile.")
    car = reg.car_collegata(voce)
    return render(request, "registro_ofi/dettaglio.html", _ctx(
        request,
        voce=voce,
        car=car,
        car_aperta=car is not None and car.chiusa_il is None,
        origine=collegamenti.origine(voce),
        processo_link=collegamenti.processo(voce),
        car_link=collegamenti.car(voce),
        puo_gestire=_puo_gestire(request, voce),
        testo_fase=getattr(voce, voce.campo_fase_corrente, "") if voce.campo_fase_corrente else "",
        allegato_form=RegistroOFIAllegatoForm(),
    ))


@login_required
@require_POST
@transaction.atomic
def avanza(request, pk: int):
    """Chiude la fase corrente del PDCA e passa alla successiva (gestore o responsabile)."""
    voce = get_object_or_404(RegistroOFI.objects.select_for_update(), pk=pk)
    if not _puo_gestire(request, voce):
        return HttpResponseForbidden("Solo chi gestisce il registro o il responsabile della voce.")
    try:
        nuova = reg.avanza_fase(voce, utente=request.user, testo=request.POST.get("testo", ""))
    except reg.AvanzamentoNonValido as exc:
        messages.error(request, str(exc))
    else:
        log_action(request, "registro_ofi_avanza", "gestione_specifiche", {"numero": voce.numero, "fase": nuova},
                   oggetto=voce)
        messages.success(request, f"Voce {voce.numero}: ora in fase «{voce.get_fase_display()}».")
    return redirect("registro_ofi:dettaglio", pk=voce.pk)


@login_required
@require_POST
@transaction.atomic
def riapri(request, pk: int):
    """Verifica non superata: la voce torna in PLAN con il motivo (gestore o responsabile)."""
    voce = get_object_or_404(RegistroOFI.objects.select_for_update(), pk=pk)
    if not _puo_gestire(request, voce):
        return HttpResponseForbidden("Solo chi gestisce il registro o il responsabile della voce.")
    try:
        reg.riapri_non_efficace(voce, utente=request.user, motivo=request.POST.get("motivo", ""))
    except reg.AvanzamentoNonValido as exc:
        messages.error(request, str(exc))
    else:
        log_action(request, "registro_ofi_riapri", "gestione_specifiche", {"numero": voce.numero}, oggetto=voce)
        messages.success(request, f"Voce {voce.numero} riaperta in PLAN.")
    return redirect("registro_ofi:dettaglio", pk=voce.pk)


@login_required
@require_POST
def allegato_carica(request, pk: int):
    """Carica un'evidenza nello storage privato cifrato (gestore o responsabile)."""
    voce = get_object_or_404(RegistroOFI, pk=pk)
    if not _puo_gestire(request, voce):
        return HttpResponseForbidden("Solo chi gestisce il registro o il responsabile della voce.")
    form = RegistroOFIAllegatoForm(request.POST, request.FILES)
    if not form.is_valid():
        for errori in form.errors.values():
            for e in errori:
                messages.error(request, e)
        return redirect("registro_ofi:dettaglio", pk=voce.pk)
    caricato = form.cleaned_data["file"]
    digest = hashlib.sha256()
    for blocco in caricato.chunks():
        digest.update(blocco)
    caricato.seek(0)
    allegato = RegistroOFIAllegato(
        voce=voce, nome=caricato.name[:200], descrizione=form.cleaned_data.get("descrizione", ""),
        sha256=digest.hexdigest(), dimensione=caricato.size, caricato_da=request.user)
    allegato.file.save(caricato.name, caricato, save=False)
    allegato.save()
    reg.registra_evento(voce, "ALLEGATO", utente=request.user, a=allegato.nome, nota=allegato.descrizione)
    log_action(request, "registro_ofi_allegato", "gestione_specifiche",
               {"numero": voce.numero, "allegato": allegato.pk, "sha256": allegato.sha256}, oggetto=voce)
    messages.success(request, f"Allegato «{allegato.nome}» caricato.")
    return redirect("registro_ofi:dettaglio", pk=voce.pk)


@login_required
def allegato_scarica(request, pk: int, allegato_id: int):
    """Download protetto di un'evidenza (chi può leggere il registro), con traccia."""
    allegato = get_object_or_404(RegistroOFIAllegato.objects.select_related("voce"), pk=allegato_id, voce_id=pk)
    if not _puo_vedere(request, allegato.voce):
        return HttpResponseForbidden("Questa voce è visibile a chi consulta il registro e al suo responsabile.")
    try:
        fh = allegato.file.open("rb")
    except FileNotFoundError as exc:
        raise Http404("File non trovato") from exc
    log_action(request, "registro_ofi_allegato_download", "gestione_specifiche",
               {"numero": allegato.voce.numero, "allegato": allegato.pk}, oggetto=allegato.voce)
    return FileResponse(fh, as_attachment=True, filename=allegato.nome)


@login_required
def cruscotto(request):
    """Il cruscotto è nella dashboard del modulo: il vecchio indirizzo ci porta, con i filtri."""
    query = request.GET.urlencode()
    return redirect(reverse("registro_ofi:home") + (f"?{query}" if query else ""))


@login_required
def home(request):
    """Dashboard del modulo.

    A tutti: «Il mio lavoro» (voci di cui si è responsabili). A chi consulta il registro:
    numeri chiave, cosa guardare adesso (scadute, in scadenza, senza responsabile),
    origine delle voci, processi del catalogo, andamento e attività recente.
    """
    oggi = timezone.localdate()
    limite = oggi + timedelta(days=30)
    mie = list(_aperte(RegistroOFI.objects.filter(responsabile=request.user))
               .select_related("processo_sgi").order_by("data_richiesta", "numero"))
    for v in mie:
        v.in_scadenza = bool(v.data_richiesta and oggi <= v.data_richiesta <= limite)
    ctx = _ctx(request, mie=mie, n_mie_scadute=sum(1 for v in mie if v.is_scaduto),
               n_mie_in_scadenza=sum(1 for v in mie if v.in_scadenza))
    if not ctx["ofi_vede_registro"]:
        return render(request, "registro_ofi/home.html", ctx)

    modulo = (request.GET.get("modulo") or "").strip()
    qs = RegistroOFI.objects.all().select_related("content_type", "responsabile", "processo_sgi")
    if modulo:
        qs = qs.filter(modulo_origine=modulo)
    voci = list(qs)
    aperte = [v for v in voci if not v.is_chiuso]
    scadute = sorted((v for v in aperte if v.is_scaduto), key=lambda v: (v.data_richiesta, v.numero))
    in_scadenza = sorted((v for v in aperte if v.data_richiesta and oggi <= v.data_richiesta <= limite),
                         key=lambda v: (v.data_richiesta, v.numero))
    senza_resp = sorted((v for v in aperte if v.responsabile_id is None), key=lambda v: v.numero)

    per_processo = {}
    for v in aperte:
        chiave = v.processo_sgi_id or 0
        riga = per_processo.setdefault(chiave, {"processo": v.processo_sgi, "aperte": 0, "nc": 0, "scadute": 0})
        riga["aperte"] += 1
        riga["nc"] += v.tipo == RegistroOFI.TIPO_NC
        riga["scadute"] += v.is_scaduto
    righe_processo = sorted((r for k, r in per_processo.items() if k),
                            key=lambda r: (-r["nc"], -r["aperte"], r["processo"].codice))
    origini = collegamenti.conta_per_origine(aperte)
    for r in origini:
        r["totale"] = 0
    tot_origine = {r["chiave"]: r for r in origini}
    for v in voci:
        tot_origine[collegamenti.tipo_origine(v)]["totale"] += 1

    eventi = (RegistroOFIEvento.objects.select_related("voce", "utente")
              .filter(voce__in=qs).order_by("-quando")[:12]) if voci else []
    ctx.update(
        c=reg.cruscotto(voci, oggi=oggi),
        counters=reg.conta_pdca_cumulativo(qs),
        scadute=scadute[:8], n_scadute=len(scadute),
        in_scadenza=in_scadenza[:8], n_in_scadenza=len(in_scadenza),
        senza_resp=senza_resp[:8], n_senza_resp=len(senza_resp),
        righe_processo=righe_processo[:12], n_senza_processo=per_processo.get(0, {}).get("aperte", 0),
        origini=origini, eventi=eventi,
        modulo=modulo, modulo_label=reg.modulo_label(modulo) if modulo else "",
        moduli_opts=reg.moduli_presenti(),
    )
    return render(request, "registro_ofi/home.html", ctx)


@login_required
def mie(request):
    """Le voci di cui l'utente è responsabile: aperte prima, poi le chiuse."""
    voci = list(RegistroOFI.objects.filter(responsabile=request.user)
                .select_related("content_type", "processo_sgi"))
    # Prima le scadenze più vicine; le voci senza scadenza in fondo.
    aperte = sorted((v for v in voci if not v.is_chiuso),
                    key=lambda v: (v.data_richiesta is None, v.data_richiesta or date.max, v.numero))
    chiuse = sorted((v for v in voci if v.is_chiuso), key=lambda v: v.data_chiusura or date.min, reverse=True)
    for v in voci:
        v.link_origine = collegamenti.origine(v)
    return render(request, "registro_ofi/mie.html", _ctx(request, aperte=aperte, chiuse=chiuse))
