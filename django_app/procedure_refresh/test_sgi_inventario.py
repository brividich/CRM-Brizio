"""Test del comando ``sgi_inventario`` (F0, sola lettura) su un albero sintetico."""

from __future__ import annotations

import io
import json
import tempfile
from pathlib import Path

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase, override_settings

from procedure_refresh.management.commands.sgi_inventario import analizza_pdf, build_inventario


def _pdf_nativo(path: Path, *, con_tabella: bool = True) -> None:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    y = 72
    for riga in ("1. Scopo", "Testo sintetico di prova per la sezione scopo del documento.",
                 "2. Campo di applicazione", "Altro testo sintetico, abbastanza lungo da superare la soglia."):
        page.insert_text((72, y), riga)
        y += 18
    if con_tabella:
        # Griglia 3x3 con bordi: find_tables la riconosce come tabella.
        x0, y0, w, h = 72, 200, 120, 24
        for r in range(4):
            page.draw_line((x0, y0 + r * h), (x0 + 3 * w, y0 + r * h))
        for c in range(4):
            page.draw_line((x0 + c * w, y0), (x0 + c * w, y0 + 3 * h))
        for r in range(3):
            for c in range(3):
                page.insert_text((x0 + c * w + 6, y0 + r * h + 16), f"C{r}{c}")
    page2 = doc.new_page()
    page2.insert_text((72, 72), "3. Responsabilita")
    page2.insert_text((72, 90), "Contenuto sintetico della terza sezione del documento di prova.")
    doc.save(str(path))
    doc.close()


def _pdf_scansione(path: Path) -> None:
    """PDF 'scansione': solo un'immagine, nessun testo estraibile."""
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 40, 40), False)
    pix.clear_with(200)
    page.insert_image(fitz.Rect(72, 72, 300, 300), pixmap=pix)
    doc.new_page()
    doc.save(str(path))
    doc.close()


def _build_share(root: Path) -> None:
    qual = root / "9100_Qualita"
    sic = root / "45001_Sicurezza"
    sup = qual / "SUPERATO"
    for d in (qual, sic, sup):
        d.mkdir(parents=True)
    _pdf_nativo(qual / "MT CN 06 Rev.21_Documento sintetico.pdf")
    _pdf_scansione(sic / "MOD.165 Rev.2_Modulo sintetico.pdf")
    _pdf_nativo(sic / "Norma esterna sintetica.pdf", con_tabella=False)  # fallback
    _pdf_nativo(sup / "MT CN 06 Rev.20_Documento sintetico.pdf")  # escluso
    # Editabili: stesso codice del PDF (stessa revisione) + un codice solo editabile.
    try:
        import docx

        d = docx.Document()
        d.add_heading("1. Scopo", level=1)
        d.add_paragraph("Testo sintetico.")
        d.save(str(qual / "MT CN 06 Rev.21_Documento sintetico.docx"))
    except ImportError:
        (qual / "MT CN 06 Rev.21_Documento sintetico.docx").write_bytes(b"PK\x03\x04")
    from openpyxl import Workbook

    wb = Workbook()
    wb.active["A1"] = "dato sintetico"
    wb.save(str(sic / "MOD.200 Rev.1_Registro sintetico.xlsx"))
    (sic / "MOD.165 Rev.2_Modulo sintetico.doc").write_bytes(b"\xd0\xcf\x11\xe0")
    (sic / "~$D.165 Rev.2_Modulo sintetico.doc").write_bytes(b"lock")  # lock Office
    (sic / "note.txt").write_text("sintetico", encoding="utf-8")


class SgiInventarioTests(TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)
        _build_share(self.root)

    def tearDown(self):
        self._tmp.cleanup()

    def test_conteggi_formati_e_aree(self):
        inv = build_inventario(self.root)
        f = inv["file"]
        self.assertEqual(f["per_formato"][".pdf"], 3)  # SUPERATO escluso
        self.assertEqual(f["per_formato"][".docx"], 1)
        self.assertEqual(f["per_formato"][".doc"], 1)
        self.assertEqual(f["per_formato"][".xlsx"], 1)
        self.assertEqual(f["per_formato"]["altro"], 1)
        self.assertEqual(f["altro_per_estensione"], {".txt": 1})
        self.assertEqual(f["esclusi_superato"], 1)
        self.assertEqual(f["esclusi_temporanei"], 1)
        self.assertEqual(inv["per_area"]["9100_Qualita"]["totale"], 2)
        self.assertEqual(inv["per_area"]["45001_Sicurezza"][".pdf"], 2)

    def test_nomi_e_confronto_pdf_editabile(self):
        inv = build_inventario(self.root)
        self.assertEqual(inv["nomi"]["pdf_non_riconosciuti_fallback"], 1)
        self.assertEqual(inv["nomi"]["pdf_documenti_unici"], 3)
        pe = inv["pdf_vs_editabile"]
        self.assertEqual(pe["codici_in_entrambi"], 2)  # MT CN 06 (.docx) + MOD.165 (.doc)
        self.assertEqual(pe["codici_in_entrambi_stessa_revisione"], 2)
        self.assertEqual(pe["codici_solo_editabile"], 1)  # MOD.200
        self.assertIn("MT CN 06", pe["esempi_in_entrambi"])

    def test_scansioni_tabelle_heading(self):
        inv = build_inventario(self.root)
        p = inv["pdf"]
        self.assertEqual(p["analizzati"], 3)
        self.assertEqual(p["scansioni"], 1)
        self.assertEqual(p["testo_nativo"], 2)
        self.assertEqual(p["pagine_totali"], 6)
        self.assertEqual(p["pagine_scansione"], 2)
        self.assertGreaterEqual(p["pagine_con_tabelle"], 1)
        self.assertEqual(p["heading_almeno_2_sezioni"], 2)

    def test_sample_e_senza_tabelle(self):
        inv = build_inventario(self.root, sample=1, tabelle=False)
        self.assertEqual(inv["pdf"]["analizzati"], 1)
        self.assertTrue(inv["pdf"]["campione"])
        self.assertIsNone(inv["pdf"]["pagine_con_tabelle"])
        # i conteggi dei file restano sull'intero albero
        self.assertEqual(inv["file"]["per_formato"][".pdf"], 3)

    def test_pdf_corrotto_fail_safe(self):
        bad = self.root / "rotto.pdf"
        bad.write_bytes(b"non un pdf")
        self.assertIn("errore", analizza_pdf(bad))
        inv = build_inventario(self.root)
        self.assertEqual(inv["pdf"]["illeggibili"], 1)

    @override_settings(OLLAMA_RAG_SGI_MAX_PROCS=2)
    def test_cap_rag(self):
        inv = build_inventario(self.root)
        self.assertEqual(inv["cap_rag"]["OLLAMA_RAG_SGI_MAX_PROCS"], 2)
        self.assertEqual(inv["cap_rag"]["revisioni_correnti_db"], 0)
        self.assertEqual(inv["cap_rag"]["eccedenza_share"], 1)

    def test_comando_json_sola_lettura(self):
        prima = sorted(str(p) for p in self.root.rglob("*"))
        out = io.StringIO()
        call_command("sgi_inventario", "--root", str(self.root), "--json", stdout=out)
        data = json.loads(out.getvalue())
        self.assertEqual(data["file"]["per_formato"][".pdf"], 3)
        self.assertEqual(sorted(str(p) for p in self.root.rglob("*")), prima)

    def test_comando_markdown_e_root_da_settings(self):
        out = io.StringIO()
        with override_settings(PROCEDURE_REFRESH_SGI_SHARE_ROOT=str(self.root)):
            call_command("sgi_inventario", stdout=out)
        self.assertIn("# Inventario SGI", out.getvalue())

    def test_output_sotto_root_rifiutato(self):
        with self.assertRaises(CommandError):
            call_command("sgi_inventario", "--root", str(self.root), "--output", str(self.root / "x.json"))

    def test_root_mancante(self):
        with override_settings(PROCEDURE_REFRESH_SGI_SHARE_ROOT=""):
            with self.assertRaises(CommandError):
                call_command("sgi_inventario")
        with self.assertRaises(CommandError):
            call_command("sgi_inventario", "--root", str(self.root / "non_esiste"))
