"""Parser del registro Excel storico dei contatori MFC (foglio trimestrale).

Formato atteso (un blocco per trimestre, colonne A-G; il resto e' ignorato):

    LUG-AGO-SET 2026 | A4 BN | A3 BN | A4 C | A3 C | MATRICOLA | TOTALE
    AMMINISTRAZIONE  | 66153 |  980  | 5027 |  435 | 2YM19959  | ...

L'etichetta del trimestre e' libera ("OTT-NOV-DIC-2022", "APR - MAG -GIU 2026"):
conta il primo mese riconosciuto e l'anno a quattro cifre. Nessuna scrittura qui:
il parser produce righe, `importa_storico_contatori` decide cosa salvare.
"""
from dataclasses import dataclass, field
from datetime import date
import re

from openpyxl import load_workbook

FOGLIO_DEFAULT = "CONTATORI CANON"

_MESI_TRIMESTRE = {
    "GEN": 1, "FEB": 1, "MAR": 1,
    "APR": 2, "MAG": 2, "GIU": 2,
    "LUG": 3, "AGO": 3, "SET": 3, "SETT": 3,
    "OTT": 4, "NOV": 4, "DIC": 4,
}
_FINE_TRIMESTRE = {1: (3, 31), 2: (6, 30), 3: (9, 30), 4: (12, 31)}
_ANNO_RE = re.compile(r"(20\d{2})")


@dataclass
class RigaStorico:
    riga_excel: int
    trimestre: str
    data: date
    reparto: str
    matricola: str
    a4_bn: int
    a3_bn: int
    a4_col: int
    a3_col: int

    @property
    def totale(self):
        return self.a4_bn + self.a3_bn + self.a4_col + self.a3_col


@dataclass
class EsitoParser:
    righe: list = field(default_factory=list)
    avvisi: list = field(default_factory=list)


def trimestre_da_etichetta(testo):
    """("2026-Q3", date(2026, 9, 30)) da "LUG - AGO - SET 2026"; None se non riconosciuta."""
    testo = str(testo or "").upper()
    anno = _ANNO_RE.search(testo)
    if not anno:
        return None
    for token in re.findall(r"[A-Z]+", testo):
        q = _MESI_TRIMESTRE.get(token)
        if q:
            mese, giorno = _FINE_TRIMESTRE[q]
            a = int(anno.group(1))
            return f"{a}-Q{q}", date(a, mese, giorno)
    return None


def _intero(valore):
    """Cella contatore -> intero >= 0; vuota = 0. Ritorna None se non numerica."""
    if valore is None or str(valore).strip() == "":
        return 0
    try:
        numero = float(str(valore).replace(".", "").replace(",", ".")) \
            if isinstance(valore, str) else float(valore)
    except ValueError:
        return None
    if numero < 0 or numero != int(numero):
        return None
    return int(numero)


def _e_intestazione(riga):
    return str(riga[1] or "").strip().upper().replace(" ", "") == "A4BN"


def leggi_righe(valori_righe):
    """Interpreta le righe (tuple di celle A..G) del foglio. Separato da openpyxl per i test."""
    esito = EsitoParser()
    corrente = None
    for n, riga in enumerate(valori_righe, start=1):
        riga = tuple(riga) + (None,) * (7 - len(riga))
        if _e_intestazione(riga):
            corrente = trimestre_da_etichetta(riga[0])
            if corrente is None:
                esito.avvisi.append(f"riga {n}: trimestre non riconosciuto «{riga[0]}», blocco ignorato")
            continue
        reparto = str(riga[0] or "").strip()
        if corrente is None or not reparto:
            continue
        contatori = [_intero(v) for v in riga[1:5]]
        if any(v is None for v in contatori):
            esito.avvisi.append(f"riga {n} {reparto} {corrente[0]}: contatore non numerico, saltata")
            continue
        if sum(contatori) == 0:
            esito.avvisi.append(f"riga {n} {reparto} {corrente[0]}: contatori vuoti, saltata")
            continue
        esito.righe.append(RigaStorico(
            riga_excel=n, trimestre=corrente[0], data=corrente[1], reparto=reparto,
            matricola=str(riga[5] or "").strip(), a4_bn=contatori[0], a3_bn=contatori[1],
            a4_col=contatori[2], a3_col=contatori[3],
        ))
    return esito


def leggi_file(percorso, foglio=FOGLIO_DEFAULT):
    wb = load_workbook(percorso, read_only=True, data_only=True)
    try:
        if foglio not in wb.sheetnames:
            raise ValueError(f"foglio «{foglio}» assente (presenti: {', '.join(wb.sheetnames)})")
        ws = wb[foglio]
        return leggi_righe(ws.iter_rows(min_col=1, max_col=7, values_only=True))
    finally:
        wb.close()
