"""Export PDF e XLSX dei report di conformita'.

Stesso :class:`ReportResult` della vista web: il file scaricato riporta gli
stessi numeri, piu' periodo, filtri, riferimenti normativi e data di
generazione (evidenza per l'auditor).
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from html import escape

from django.utils import timezone
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, Spacer

from core.pdf import PdfTheme, build_styles, data_table, header_footer_callback, make_document, section_heading

from .registry import TONE_DANGER, TONE_OK, TONE_WARN, Kpi, ReportDef, ReportParams, ReportResult

_TONE_BG = {
    TONE_DANGER: colors.HexColor("#fde8e8"),
    TONE_WARN: colors.HexColor("#fdf3d8"),
    TONE_OK: colors.HexColor("#e7f6ec"),
}
_TONE_FG = {
    TONE_DANGER: colors.HexColor("#b42318"),
    TONE_WARN: colors.HexColor("#8a5a00"),
    TONE_OK: colors.HexColor("#1f7a3f"),
}


@dataclass
class PdfSection:
    title: str
    clausole: str
    kpis: list[Kpi]
    columns: list[str] = field(default_factory=list)
    rows: list[list] = field(default_factory=list)
    row_tones: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


def _txt(value) -> str:
    return "" if value is None else str(value)


def _kpi_table(kpis: list[Kpi], theme: PdfTheme, styles, width: float):
    from reportlab.lib.styles import ParagraphStyle

    per_row = 4
    cells = []
    style = []
    kpi_style = ParagraphStyle("rc_kpi", parent=styles["cell"], leading=16)
    for index, kpi in enumerate(kpis):
        r, col = divmod(index, per_row)
        while len(cells) <= r:
            cells.append([""] * per_row)
        hint = f"<br/><font size=7 color='#64748b'>{escape(kpi.hint)}</font>" if kpi.hint else ""
        cells[r][col] = Paragraph(
            f"<font size=7 color='#64748b'>{escape(kpi.label.upper())}</font><br/>"
            f"<font size=13><b>{escape(_txt(kpi.value))}</b></font>{hint}",
            kpi_style,
        )
        if kpi.tone in _TONE_BG:
            style.append(("BACKGROUND", (col, r), (col, r), _TONE_BG[kpi.tone]))
    if not cells:
        return None
    return data_table(
        cells, theme, col_widths=[width / per_row] * per_row, header=False,
        extra_style=[("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white]), *style,
                     ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)],
    )


def _col_widths(columns: list[str], rows: list[list], width: float) -> list[float]:
    lengths = []
    for i, col in enumerate(columns):
        longest = max([len(_txt(col))] + [len(_txt(r[i])) for r in rows[:200] if i < len(r)])
        lengths.append(min(max(longest, 4), 60))
    total = sum(lengths) or 1
    return [width * n / total for n in lengths]


def _data_table(section: PdfSection, theme: PdfTheme, styles, width: float):
    rows = [[Paragraph(escape(_txt(h)), styles["table_header"]) for h in section.columns]]
    extra = [("FONTSIZE", (0, 0), (-1, -1), 7), ("VALIGN", (0, 0), (-1, -1), "TOP")]
    for index, values in enumerate(section.rows, start=1):
        rows.append([Paragraph(escape(_txt(v)).replace("\n", "<br/>"), styles["cell"]) for v in values])
        tone = section.row_tones[index - 1] if index - 1 < len(section.row_tones) else ""
        if tone in (TONE_DANGER, TONE_WARN):
            extra.append(("BACKGROUND", (0, index), (-1, index), _TONE_BG[tone]))
    return data_table(rows, theme, col_widths=_col_widths(section.columns, section.rows, width),
                      repeat_rows=1, extra_style=extra)


def render_pdf(*, title: str, subtitle: str, sections: list[PdfSection], intro: list[str] | None = None) -> bytes:
    theme = PdfTheme.from_branding()
    styles = build_styles(theme)
    buf = io.BytesIO()
    doc = make_document(buf, title=title, landscape=True)
    width = doc.pagesize[0] - doc.leftMargin - doc.rightMargin
    elements: list = []
    for line in intro or []:
        elements.append(Paragraph(escape(line), styles["body"]))
    if intro:
        elements.append(Spacer(1, 4 * mm))

    for section in sections:
        block = []
        if len(sections) > 1:
            block += section_heading(escape(section.title), theme, styles)
        if section.clausole:
            block.append(Paragraph(f"<b>Riferimenti:</b> {escape(section.clausole)}", styles["body"]))
            block.append(Spacer(1, 2 * mm))
        kpi_table = _kpi_table(section.kpis, theme, styles, width)
        if kpi_table is not None:
            block.append(kpi_table)
            block.append(Spacer(1, 4 * mm))
        elements.append(KeepTogether(block))
        if section.columns:
            if section.rows:
                elements.append(_data_table(section, theme, styles, width))
            else:
                elements.append(Paragraph("Nessuna riga da segnalare.", styles["body"]))
            elements.append(Spacer(1, 3 * mm))
        for note in section.notes:
            elements.append(Paragraph(f"<font size=7 color='#64748b'>• {escape(note)}</font>", styles["body"]))
        elements.append(Spacer(1, 6 * mm))

    draw = header_footer_callback(theme, title=title.upper()[:90], subtitle=subtitle[:140])
    doc.build(elements, onFirstPage=draw, onLaterPages=draw)
    return buf.getvalue()


def section_for(report: ReportDef, result: ReportResult) -> PdfSection:
    return PdfSection(
        title=report.title,
        clausole=" · ".join(str(c) for c in report.clausole),
        kpis=result.kpis,
        columns=result.columns,
        rows=result.rows,
        row_tones=result.row_tones,
        notes=result.notes,
    )


def subtitle_for(report: ReportDef, params: ReportParams, filtri_label: str = "") -> str:
    parti = []
    if report.usa_periodo:
        parti.append(f"Periodo {params.periodo_label}")
    else:
        parti.append(f"Situazione al {params.today:%d/%m/%Y}")
    if filtri_label:
        parti.append(filtri_label)
    parti.append(" · ".join(report.norme))
    return " | ".join(parti)


def report_pdf(report: ReportDef, params: ReportParams, result: ReportResult, filtri_label: str = "") -> bytes:
    return render_pdf(
        title=report.title,
        subtitle=subtitle_for(report, params, filtri_label),
        sections=[section_for(report, result)],
    )


def report_xlsx(report: ReportDef, params: ReportParams, result: ReportResult, filtri_label: str = "") -> bytes:
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    from core.excel_export import append_row, build_xlsx_bytes

    raw = build_xlsx_bytes(
        columns=result.columns or ["Nessun dato"],
        rows=result.rows,
        sheet_title="Dati",
        title=report.title,
        subtitle=subtitle_for(report, params, filtri_label),
        filters_label=f"Generato il {timezone.localtime():%d/%m/%Y %H:%M}",
    )
    wb = load_workbook(io.BytesIO(raw))
    ws = wb.create_sheet("Indicatori", 0)
    ws.append([report.title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([subtitle_for(report, params, filtri_label)])
    ws.append([])
    ws.append(["Riferimenti normativi"])
    ws["A4"].font = Font(bold=True)
    for clausola in report.clausole:
        append_row(ws, [clausola.norma, clausola.punto])
    ws.append([])
    ws.append(["Indicatore", "Valore", "Nota"])
    header = ws.max_row
    for cell in ws[header]:
        cell.font = Font(bold=True)
    for kpi in result.kpis:
        append_row(ws, [kpi.label, _txt(kpi.value), kpi.hint])
    if result.notes:
        ws.append([])
        ws.append(["Note"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for note in result.notes:
            append_row(ws, [note])
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 30
    wb.active = 0
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
